"""The legacy onboarding world and the import pipeline of legacy template files (PRD J-01.2 to
J-01.9, §2.12 WLD-F-01, WLD-F-02; 04 T-IMP-01 ``LEGACY_V1``, T-MIG-01 seeded parity templates;
POLICIES §6.3; BUILD_SPEC DIN-4, DIN-5).

``legacy_world`` builds the WLD-T-20 preconditions of J-01 as a test world: Maya (Revenue
Accountant, SSP Analyst) uploads; Priya (Revenue Reviewer, SSP Approver, MFA) approves imports;
Marcus (Controller, Tenant Admin, MFA) approves configuration. Entities ``Mock Entity 1`` and
``Mock Entity 2`` (USD, America/New_York) share a January calendar whose FY2023 periods are open;
the setup accounts 21001, 21002, 15001 and 15002 exist; mapping ``LEGACY-MAP-2023-01`` covers the
roles a parity computation reaches; the four seeded parity templates are PUBLISHED from 01 Jan 2023;
and, unless ``preset = False``, the TENANT accounting policy set carries ``LEGACY_PARITY``.

The seeded templates and the preset are written as the configuration lifecycle leaves them
(``rows.publish_registry_version`` walks registry versions the same way): 04 T-MIG-01 seeds them,
and no command seeds them outside migration (L5-1-Q-13).
"""

from __future__ import annotations

import secrets
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Final
from uuid import UUID, uuid4

import openpyxl
from erev_api.auth.keyring import KeyRing
from erev_api.clock import FrozenClock
from erev_api.db import new_id
from erev_api.db.session import DbContext, tenant_session
from erev_api.db.tables import job, pob_template, pob_template_version
from erev_api.enums import (
    ApprovalRequestStatus,
    ApprovalSubjectType,
    ConfigStatus,
    PrincipalKind,
    RegistryCategory,
    RegistryScope,
)
from erev_api.files.store import LocalFileStore
from erev_api.jobs.context import JobRuntime
from erev_api.registry import presets
from fastapi import FastAPI
from sqlalchemy import insert, select, update
from support.factories import (
    IMPORTS_PATH,
    LEGACY_UAT,
    ImportWorld,
    Workspace,
    imported,
    open_periods,
    run_import_job,
    stamp_test_release,
)
from support.factories import workspace as workspace_of
from support.http import call
from support.principals import (
    Actor,
    colleague,
    cookie_headers,
    enrolled,
    member,
    sign_in,
    workspace,
)
from support.reference import (
    approve,
    assign,
    calendar,
    entity,
    get,
    gl_account,
    mapping_published,
    put,
)
from support.rows import INVITED_AT, insert_approval_request, publish_registry_version

__all__ = [
    "ENTITIES",
    "SETUP_2023",
    "SKU_SSP",
    "LegacyWorld",
    "committed",
    "diffed",
    "job_of",
    "legacy_world",
    "replayed",
    "shown",
    "submit",
    "workbook_rows",
]

SKU_SSP: Final = LEGACY_UAT / "01-ssp-upload" / "SKU SSP Template.xlsx"  # WLD-F-01
SETUP_2023: Final = (
    LEGACY_UAT / "02-contract-setup-2023-01-01" / "Contract Setup Template 1.1.2023.xlsx"
)  # WLD-F-02
ENTITIES: Final = ("Mock Entity 1", "Mock Entity 2")
PARITY_FROM: Final = datetime(2023, 1, 1, tzinfo=UTC)
MAPPING_FROM: Final = "2023-01-01T00:00:00Z"
# 04 T-MIG-01: (code, obligation kind, distinctness), all UNITS_DELIVERED (S03-R-18).
PARITY_TEMPLATES: Final = (
    ("LEGACY-DISTINCT", "STANDARD", "distinct"),
    ("LEGACY-MATERIAL-RIGHT", "MATERIAL_RIGHT", "distinct"),
    ("LEGACY-NONDISTINCT", "STANDARD", "nondistinct"),
    ("LEGACY-VC", "VC_LINE", "distinct"),
)
# The setup files' deferred revenue and unbilled accounts (LM-CL-11, LM-CL-12).
SETUP_ACCOUNTS: Final = (
    ("21001", "Deferred revenue - Mock Entity 1", "LIABILITY", "C"),
    ("21002", "Deferred revenue - Mock Entity 2", "LIABILITY", "C"),
    ("15001", "Unbilled receivable - Mock Entity 1", "ASSET", "D"),
    ("15002", "Unbilled receivable - Mock Entity 2", "ASSET", "D"),
)
# (code, name, type, normal balance, role) of the mapping a parity computation reaches.
PARITY_CHART: Final = (
    ("1100", "Accounts receivable", "ASSET", "D", "ACCOUNTS_RECEIVABLE"),
    ("1105", "Unbilled receivable", "ASSET", "D", "UNBILLED_RECEIVABLE"),
    ("1200", "Contract asset", "ASSET", "D", "CONTRACT_ASSET"),
    ("2060", "Clearing - billing", "LIABILITY", "C", "BILLING_CLEARING:BILLING"),
    ("2100", "Contract liability", "LIABILITY", "C", "CONTRACT_LIABILITY"),
    ("4000", "Revenue - products", "REVENUE", "C", "REVENUE"),
    ("4900", "Pre-standard revenue", "REVENUE", "C", "PRE_STANDARD_REVENUE"),
    ("7900", "Rounding", "EXPENSE", "D", "ROUNDING"),
)


@dataclass(frozen=True, slots=True)
class LegacyWorld:
    """The J-01 world: Maya's import world, the approvers and the entities."""

    imports: ImportWorld
    priya: Actor
    marcus: Actor
    entity_ids: Mapping[str, UUID]

    @property
    def app(self) -> FastAPI:
        return self.imports.app

    @property
    def maya(self) -> Actor:
        return self.imports.actor

    @property
    def tenant_id(self) -> UUID:
        return self.imports.tenant_id

    def place(self) -> Workspace:
        """Maya's workspace for domain units of work (``support.factories.Workspace``)."""
        runtime = self.imports.runtime
        assert runtime.keyring is not None and runtime.files is not None
        assert isinstance(runtime.files, LocalFileStore)
        return workspace_of(self.app, self.imports.clock, runtime.keyring, runtime.files, self.maya)


def _publish_parity_templates(tenant_id: UUID) -> None:
    """The four seeded parity templates, each with version 1 PUBLISHED from 01 Jan 2023."""
    context = DbContext(tenant_id=tenant_id, user_id=None, entity_scope="*")
    system = {"created_by_kind": PrincipalKind.SYSTEM.value}
    stamps = {**system, "updated_by_kind": PrincipalKind.SYSTEM.value}
    with tenant_session(context) as session:
        for code, kind, flag in PARITY_TEMPLATES:
            template_id = new_id()
            session.execute(
                insert(pob_template).values(
                    tenant_id=tenant_id,
                    id=template_id,
                    code=code,
                    name=f"Legacy parity {kind.lower()} ({flag})",
                    **stamps,
                )
            )
            version_id = new_id()
            session.execute(
                insert(pob_template_version).values(
                    tenant_id=tenant_id,
                    id=version_id,
                    pob_template_id=template_id,
                    obligation_kind=kind,
                    distinctness=flag,
                    satisfaction_pattern="POINT_IN_TIME",
                    over_time_criterion="NOT_APPLICABLE",
                    recognition_method="UNITS_DELIVERED",
                    principal_agent="PRINCIPAL",
                    version_no=1,
                    status=ConfigStatus.DRAFT.value,
                    effective_from=PARITY_FROM,
                    **stamps,
                )
            )
            where = pob_template_version.c.id == version_id
            session.execute(
                update(pob_template_version)
                .where(where)
                .values(status=ConfigStatus.TESTED.value, content_sha256=secrets.token_hex(32))
            )
            session.execute(
                update(pob_template_version)
                .where(where)
                .values(status=ConfigStatus.SUBMITTED.value)
            )
            request_id = insert_approval_request(
                session,
                tenant_id=tenant_id,
                status=ApprovalRequestStatus.APPROVED,
                subject_type=ApprovalSubjectType.POB_TEMPLATE_VERSION.value,
                subject_id=version_id,
            )
            session.execute(
                update(pob_template_version)
                .where(where)
                .values(status=ConfigStatus.APPROVED.value, approval_request_id=request_id)
            )
            session.execute(
                update(pob_template_version)
                .where(where)
                .values(status=ConfigStatus.PUBLISHED.value, published_at=INVITED_AT)
            )


def _publish_preset(tenant_id: UUID) -> None:
    """The TENANT accounting policy set with preset ``LEGACY_PARITY`` (J-01.4, J-01.5)."""
    context = DbContext(tenant_id=tenant_id, user_id=None, entity_scope="*")
    with tenant_session(context) as session:
        publish_registry_version(
            session,
            tenant_id=tenant_id,
            category=RegistryCategory.ACCOUNTING_POLICY,
            values=presets.legacy_parity_values(scope=RegistryScope.TENANT, book_code=None),
            preset_code=presets.LEGACY_PARITY,
        )


def legacy_world(
    app: FastAPI,
    keyring: KeyRing,
    clock: FrozenClock,
    files: LocalFileStore,
    *,
    preset: bool = True,
) -> LegacyWorld:
    """PRD J-01.2 to J-01.5 as a test world (module docstring)."""
    maya_member = member(keyring, clock)
    assign(maya_member, "revenue_accountant")
    assign(maya_member, "ssp_analyst")
    maya = workspace(app, maya_member, sign_in(app, maya_member.email))
    approvers: dict[str, Actor] = {}
    for name, roles in (
        ("priya", ("revenue_reviewer", "ssp_approver")),
        ("marcus", ("controller", "ssp_approver", "tenant_admin")),
    ):
        someone = colleague(maya_member.tenant_id, name)
        for code in roles:
            assign(someone, code)
        approvers[name] = enrolled(app, clock, someone)
    marcus = approvers["marcus"]
    entity_ids = legacy_reference(app, maya, marcus)
    if preset:
        _publish_preset(maya_member.tenant_id)
    stamp_test_release()
    return LegacyWorld(
        imports=ImportWorld(
            app=app,
            actor=maya,
            runtime=JobRuntime(clock=clock, keyring=keyring, files=files),
            clock=clock,
        ),
        priya=approvers["priya"],
        marcus=marcus,
        entity_ids=entity_ids,
    )


def legacy_reference(app: FastAPI, maya: Actor, marcus: Actor) -> dict[str, UUID]:
    """The WLD-T-20 reference data of the J-01 world, created by ``maya`` and approved by
    ``marcus`` (Tenant Admin, Controller): USD enabled; a January calendar with FY2023 and FY2024;
    both entities with FY2023 open; the setup accounts; mapping ``LEGACY-MAP-2023-01``; the four
    seeded parity templates. The entity ids by code (BUILD_SPEC GPA-1 shares it with DIN-4 to
    DIN-6)."""
    enabled = put(app, "/api/v1/tenant-currencies", marcus, {"currency_codes": ["USD"]})
    assert enabled.status_code == 200, enabled.text
    calendar_id = calendar(app, maya, years=(2023, 2024))
    entity_ids: dict[str, UUID] = {}
    for code in ENTITIES:
        created = entity(app, maya, code=code, calendar_id=calendar_id)
        entity_ids[code] = UUID(str(created["id"]))
        open_periods(
            app, maya, entity_code=code, keys=[f"FY2023-P{month:02d}" for month in range(1, 13)]
        )
    for code, name, kind, normal in SETUP_ACCOUNTS:
        gl_account(app, maya, code=code, name=name, account_type=kind, normal_balance=normal)
    rules: list[dict[str, Any]] = []
    for code, name, kind, normal, named in PARITY_CHART:
        role, _, purpose = named.partition(":")
        rule: dict[str, Any] = {
            "account_role": role,
            "gl_account_id": gl_account(
                app, maya, code=code, name=name, account_type=kind, normal_balance=normal
            ),
        }
        if purpose:
            rule["clearing_purpose"] = purpose
        rules.append(rule)
    mapping_published(
        app, maya, marcus, name="LEGACY-MAP-2023-01", effective_from=MAPPING_FROM, rules=rules
    )
    _publish_parity_templates(maya.member.tenant_id)
    return entity_ids


# --- the pipeline -------------------------------------------------------------------------------


def job_of(world: ImportWorld, import_id: UUID, kind: str) -> UUID:
    """The latest job of ``kind`` for an import."""
    rows = world.rows(
        select(job.c.id)
        .where(job.c.subject_id == import_id, job.c.kind == kind)
        .order_by(job.c.created_at)
    )
    assert rows, f"no {kind} job for {import_id}"
    return UUID(str(rows[-1]["id"]))


def shown(world: ImportWorld, import_id: str) -> dict[str, Any]:
    """API-S-Import."""
    response = get(world.app, f"{IMPORTS_PATH}/{import_id}", world.actor)
    assert response.status_code == 200, response.text
    return dict(response.json())


def diffed(
    world: ImportWorld,
    name: str,
    content: bytes,
    template_code: str,
    parameters: Mapping[str, Any] | None = None,
) -> str:
    """Upload, validate and diff a file; the id of the DIFF_READY import."""
    import_id, validated = imported(world, name, content, template_code, parameters)
    assert validated["status"] == "VALIDATED", validated
    run_import_job(world, job_of(world, UUID(import_id), "IMPORT_DIFF"))
    ready = shown(world, import_id)
    assert ready["status"] == "DIFF_READY", ready
    return import_id


def submit(world: ImportWorld, import_id: str) -> Any:
    """``POST /imports/{id}/submit``."""
    headers = cookie_headers(
        world.actor.token,
        world.actor.csrf_token,
        key=False,
        **{"Idempotency-Key": f"k-{uuid4()}"},
    )
    return call(
        world.app,
        "POST",
        f"{IMPORTS_PATH}/{import_id}/submit",
        json={"comment": "Legacy template upload"},
        headers=headers,
    )


def committed(
    world: LegacyWorld,
    name: str,
    content: bytes,
    template_code: str,
    parameters: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Upload, validate, diff, submit, approve by Priya and commit; API-S-Import of the COMMITTED
    import."""
    import_id = diffed(world.imports, name, content, template_code, parameters)
    submitted = submit(world.imports, import_id)
    assert submitted.status_code == 200, submitted.text
    decided = approve(world.app, str(submitted.json()["approval_request_id"]), world.priya)
    assert decided.status_code == 200, decided.text
    run_import_job(world.imports, job_of(world.imports, UUID(import_id), "IMPORT_COMMIT"))
    done = shown(world.imports, import_id)
    assert done["status"] == "COMMITTED", done
    return done


def workbook_rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    """The first sheet of a fixture workbook: its headers and its non-blank data rows."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        values = book.worksheets[0].iter_rows(values_only=True)
        headers = [str(cell) for cell in next(values) if cell is not None]
        rows = [
            list(row[: len(headers)])
            for row in values
            if any(cell is not None and cell != "" for cell in row)
        ]
    finally:
        book.close()
    return headers, rows


def assigned(someone: Actor, role_code: str, entity_ids: Sequence[UUID] = ()) -> None:
    """Assign one more role to an actor's membership."""
    assign(someone.member, role_code, entity_ids=entity_ids)


def replayed(world: LegacyWorld, through: str) -> dict[str, dict[str, Any]]:
    """Golden steps 01 to ``through`` replayed through the legacy v1 import pipeline in DG-PAR-04
    order: the step handler gives the template, ``step.json`` ``date_input`` the upload parameter
    ``effective_date`` and the handler the ``mode``; Maya prepares (``ak-preparer``) and Priya
    approves (``ak-approver``). API-S-Import of each committed step, by step number (BUILD_SPEC
    DIN-5)."""
    from support import golden_streams

    done: dict[str, dict[str, Any]] = {}
    for step in golden_streams.steps(through):
        parameters: dict[str, Any] = {}
        if step.date_input is not None:
            parameters["effective_date"] = step.date_input.isoformat()
        if step.mode is not None:
            parameters["mode"] = step.mode
        done[step.number] = committed(
            world,
            step.workbook.name,
            step.workbook.read_bytes(),
            step.template_code,
            parameters or None,
        )
    return done
