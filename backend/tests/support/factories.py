"""Test data factories that create data through domain commands (dev-guide DG-TST-13, DG-TST-16).

Besides tenants, the module builds the reference world of the contract computations of BUILD_SPEC
CTR-2 through the RFD routes and the contract commands: a calendar with open periods, the entity
AVM-US, customers, published obligation templates, approved SSP book versions and a published
account mapping (PRD §2.6); a booked and activated contract; and one computation. ``k11_world``
builds the AVM-DE world of PRD WLD-K-11 for CTR-3, and ``activated_contract`` activates a booked
contract as the SYSTEM principal (BUILD_SPEC BS3-D-19).

``engine()`` answers ``erev_engine.compute`` once the package has it (BUILD_SPEC END-9, lane L3-2)
and ``fake_compute`` until then (SPRINT-vabc V-C, R-03): the number-asserting CTR-2 and CTR-3 tests
run through the fake in the lane and through the real engine in the L3 merge gate. ``fake_compute``
follows the ENGINE_SPEC §0.5 output contract for single-entity groups of booked lines: relative
SSP over the approved point or range entries (S05-R-07, CV-34), daily time-elapsed schedules
cumulatively rounded (CV-35), revenue on delivered units less returns (ALG-01), awaiting-trigger
amounts for every other method, a trace whose nodes re-evaluate (DG-KRN-EXP-04), and, for an ACTIVE
contract only (S02-R-02), JET-02 posting intents of the revenue delta against the bundle's posted
amounts in the period of the latest delivery or return (RCP-04, RCP-06; S14-R-12, S14-R-15). Lines
flagged out of scope follow S03-R-11 and S04-R-03: their ``out_of_scope_amount`` is the version's
``out_of_scope_amount``, a ``LEASE_842`` line stays an allocation target without schedule or
revenue, and any other flag leaves the obligations. Revenue schedules exist only for activated
contracts (REQ-CON-002). ``j03_world`` builds the AVM-US world of PRD WLD-X-23 (``SF-ORD-20417``)
for CTR-4. ``k02_world`` builds the AVM-US world of PRD WLD-K-02 (``SF-ORD-10002``) for CTR-15.
"""

from __future__ import annotations

import importlib
import secrets
from collections.abc import Callable, Iterator, Mapping, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, replace
from datetime import date
from fractions import Fraction
from pathlib import Path
from types import MappingProxyType
from typing import TYPE_CHECKING, Any, Final, Literal, cast
from uuid import UUID

import erev_engine
from erev_api.auth.keyring import KeyRing
from erev_api.auth.permissions import DEFAULT_ROLES
from erev_api.auth.principal import Principal, RequestContext, system_principal
from erev_api.clock import Clock, FrozenClock
from erev_api.config import REPO_ROOT, Environment
from erev_api.controls.release import stamp_release
from erev_api.db.session import DbContext, tenant_session
from erev_api.domain.contracts import activation, bundles, computation
from erev_api.domain.contracts.commands import BookedContract, book_contract
from erev_api.domain.platform.provisioning import (
    OperatorActor,
    TenantProvisionRequest,
    TenantProvisionResult,
    provision_tenant,
)
from erev_api.enums import ComputationTrigger, ContractEventType, PrincipalKind, TenantKind
from erev_api.events.payloads import (
    ChecklistItemV1,
    ContractActivatedV1,
    ContractBookedV1,
    DeliveryRecordedV1,
)
from erev_api.events.stream import EventIn, append_events
from erev_api.files.store import LocalFileStore
from erev_api.uow import UnitOfWork, unit_of_work
from erev_engine.bundle import (
    BookInput,
    BookOutput,
    ContractVersionOut,
    EntityInput,
    EventInput,
    InputBundle,
    IntentLine,
    ObligationVersionOut,
    OutputBundle,
    PostingIntent,
    SspEntryInput,
    SspVersionInput,
    TemplateInput,
)
from erev_engine.canonical import sha256_hex
from erev_engine.errors import EngineError
from erev_engine.money import (
    cumulative_posted,
    format_exact,
    format_money,
    largest_remainder,
    round_half_up,
    to_fraction,
)
from erev_engine.stages.s01_canonicalize import obligation_subject_key
from erev_engine.stages.state import ScheduleKind, ScheduleLineOut, ScheduleLineType
from erev_engine.trace import SourceRef, TraceBuilder
from fastapi import FastAPI
from sqlalchemy import Select
from support import upload_fixtures
from support.clock import frozen_clock
from support.http import HttpResponse, call

if TYPE_CHECKING:
    from erev_api.jobs.context import JobRuntime
    from support.principals import Actor, Member


def _through(module: str, name: str) -> Callable[..., Any]:
    """``module.name`` imported on use: ``support.principals`` (and ``support.reference``, which
    imports it) import this module for ``tenant_factory``."""

    def call_through(*args: Any, **kwargs: Any) -> Any:
        return getattr(importlib.import_module(module), name)(*args, **kwargs)

    call_through.__name__ = name
    return call_through


approve = _through("support.reference", "approve")
assign = _through("support.reference", "assign")
calendar = _through("support.reference", "calendar")
entity = _through("support.reference", "entity")
get = _through("support.reference", "get")
gl_account = _through("support.reference", "gl_account")
holding = _through("support.reference", "holding")
mapping_published = _through("support.reference", "mapping_published")
new_product = _through("support.reference", "new_product")
patch = _through("support.reference", "patch")
periods = _through("support.reference", "periods")
post = _through("support.reference", "post")
put = _through("support.reference", "put")
colleague = _through("support.principals", "colleague")
cookie_headers = _through("support.principals", "cookie_headers")
enrolled = _through("support.principals", "enrolled")
member = _through("support.principals", "member")

# 04 §14.3 example tenant: the platform domain tests share one ``acme-test`` per session.
ACME = TenantProvisionRequest(
    code="acme-test",
    display_name="Acme Test",
    reporting_currency="USD",
    is_demo=False,
    admin_email="Admin@Acme.test",
)
ACME_ACTOR = OperatorActor(
    channel="CLI", operator_user_id=None, os_user="builder", request_id="r-12345678"
)


def tenant_factory(
    *,
    keyring: KeyRing,
    clock: Clock | None = None,
    code: str | None = None,
    admin_email: str | None = None,
    is_demo: bool = False,
) -> TenantProvisionResult:
    """A fresh production tenant with a random code, provisioned by ``provision_tenant``; with
    ``is_demo`` it carries the demo marker."""
    tenant_code = code or f"t-{secrets.token_hex(6)}"
    request = TenantProvisionRequest(
        code=tenant_code,
        display_name=f"Tenant {tenant_code}",
        reporting_currency="USD",
        is_demo=is_demo,
        admin_email=admin_email or f"admin@{tenant_code}.test",
    )
    actor = OperatorActor(
        channel="CLI", operator_user_id=None, os_user="tests", request_id=f"r-factory-{tenant_code}"
    )
    return provision_tenant(request, actor=actor, clock=clock or frozen_clock(), keyring=keyring)


def tenant_id_of(result: TenantProvisionResult) -> UUID:
    tenant_id = result.tenant["id"]
    assert isinstance(tenant_id, UUID)
    return tenant_id


# --- the engine of computation tests (V-C) --------------------------------------------------------


def engine() -> computation.Engine:
    """``erev_engine.compute`` when END-9 has landed, else ``fake_compute``."""
    found = getattr(erev_engine, "compute", None)
    return fake_compute if found is None else cast(computation.Engine, found)


def _group_policy(book_input: BookInput, code: str, default: str) -> str:
    for policy in book_input.policies:
        if policy.code == code and policy.scope == "GROUP" and isinstance(policy.value, str):
            return policy.value
    return default


def _template(bundle: InputBundle, code: str | None, at: date) -> TemplateInput:
    eligible = [
        item
        for item in bundle.pob_template_versions
        if item.template_code == code
        and item.effective_from <= at
        and (item.effective_to is None or at < item.effective_to)
    ]
    found = max(eligible, key=lambda item: item.version_no, default=None)
    if found is None:
        raise EngineError("PRODUCT_UNMAPPED", "no template resolves for the line", detail={})
    return found


def _ssp(
    bundle: InputBundle, product_code: str, currency: str, at: date
) -> tuple[SspVersionInput, SspEntryInput]:
    candidates: list[tuple[int, SspVersionInput, SspEntryInput]] = []
    for version in bundle.ssp_versions:
        if version.effective_from_date is not None and version.effective_from_date > at:
            continue
        if version.effective_to_date is not None and at > version.effective_to_date:
            continue
        for item in version.entries:
            if item.product_code == product_code and item.currency == currency and item.ranges:
                candidates.append((version.version_no, version, item))
    if not candidates:
        raise EngineError("SSP_KEY_NOT_FOUND", "no approved SSP entry", detail={})
    _, version, item = max(candidates, key=lambda found: found[0])
    return version, item


@dataclass(frozen=True, slots=True)
class _Line:
    contract_key: str
    event: EventInput
    line: Mapping[str, Any]
    inception: date
    entity_code: str

    @property
    def subject_key(self) -> str:
        return obligation_subject_key(self.contract_key, str(self.line["obligation_key"]))


def _lines(bundle: InputBundle) -> list[_Line]:
    bookings: dict[str, EventInput] = {}
    for event in bundle.events:
        if event.event_type == "CONTRACT_BOOKED":
            bookings[event.contract_key] = event
    headers = {item.external_id: item for item in bundle.contracts}
    found: list[_Line] = []
    for contract_key, event in sorted(bookings.items()):
        header = headers[contract_key]
        members = event.payload.get("lines")
        for line in members if isinstance(members, list | tuple) else ():
            found.append(
                _Line(
                    contract_key,
                    event,
                    cast(Mapping[str, Any], line),
                    header.inception_date,
                    header.contracting_entity_code,
                )
            )
    return sorted(found, key=lambda item: item.subject_key)


DELIVERY_RECORDED: Final = "DELIVERY_RECORDED"
RETURN_RECORDED: Final = "RETURN_RECORDED"
REVENUE_RECOGNITION: Final = "REVENUE_RECOGNITION"
UNIT_METHODS: Final = frozenset({"UNITS_DELIVERED", "POINT_IN_TIME"})
POSTABLE_STATES: Final = frozenset({"open", "closing", "reopened"})
CLOSED_STATES: Final = frozenset({"closed", "permanently_locked"})


def _deliveries(bundle: InputBundle) -> dict[str, tuple[Fraction, Fraction, EventInput]]:
    """(delivered, returned, latest event) of each obligation with deliveries or returns."""
    found: dict[str, tuple[Fraction, Fraction, EventInput]] = {}
    for event in bundle.events:
        if event.event_type not in (DELIVERY_RECORDED, RETURN_RECORDED):
            continue
        subject = obligation_subject_key(event.contract_key, str(event.payload["obligation_key"]))
        delivered, returned, _ = found.get(subject, (Fraction(0), Fraction(0), event))
        quantity = to_fraction(str(event.payload["quantity"]))
        if event.event_type == DELIVERY_RECORDED:
            delivered += quantity
        else:
            returned += quantity
        found[subject] = (delivered, returned, event)
    return found


def _posting_period(
    entity_input: EntityInput, book_code: str, effective: date
) -> tuple[str, str | None, str | None] | None:
    """RCP-04: (posting period, origin period, reason code) of an event-driven amount, or None
    when its period is future (nothing posts)."""
    ordered = sorted(entity_input.periods, key=lambda item: item.start_date)
    for index, found in enumerate(ordered):
        if not found.start_date <= effective <= found.end_date:
            continue
        state = dict(found.states).get(book_code)
        if state in POSTABLE_STATES:
            return found.period_key, None, None
        if state in CLOSED_STATES:
            later = next(
                (
                    item
                    for item in ordered[index + 1 :]
                    if dict(item.states).get(book_code) in POSTABLE_STATES
                ),
                None,
            )
            return None if later is None else (later.period_key, found.period_key, "LATE_EVENT")
        return None
    return None


def _account_code(book_input: BookInput, role: str, entity_code: str, product_code: str) -> str:
    """T-REF-15 step 3 in its simplest form: the first rule of the role that names no other entity,
    book or product."""
    for rule in book_input.account_mapping.rules:
        if (
            rule.account_role == role
            and rule.clearing_purpose is None
            and rule.entity_code in (None, entity_code)
            and rule.book_code in (None, book_input.book_code)
            and rule.product_code in (None, product_code)
        ):
            return rule.account_code
    raise EngineError(
        "ACCOUNT_MAPPING_MISSING",
        "no mapping rule resolves the role",
        detail={"role": role, "clearing_purpose": ""},
    )


def _source(tb: TraceBuilder, measure: str, subject: str, period: str | None, **node: Any) -> str:
    """A node valued by its source (``alloc.relative_ssp.v1`` total over one source input)."""
    return tb.node(
        measure=measure,
        subject_key=subject,
        period_key=period,
        formula_id="alloc.relative_ssp.v1",
        params={"role": "total"},
        narrative_key="alloc.relative_ssp",
        **node,
    )


IN_SCOPE: Final = "IN_SCOPE_606"
# S03-R-11 (PT-09): LEASE_842 lines stay allocation targets without schedules or revenue targets.
ALLOCATION_TARGET_FLAGS: Final = frozenset({IN_SCOPE, "LEASE_842"})


def _scope_flag(item: _Line) -> str:
    return str(item.line.get("scope_flag") or IN_SCOPE)


def _out_of_scope(lines: Sequence[_Line], minor: int) -> int:
    """S04-R-03: Σ ``out_of_scope_amount`` (else the stated price) of the lines flagged out of
    Topic 606 scope, in minor units (REQ-CON-016)."""
    return sum(
        round_half_up(
            to_fraction(str(item.line.get("out_of_scope_amount") or item.line["total_price"])),
            minor,
        )
        for item in lines
        if _scope_flag(item) != IN_SCOPE
    )


def _fake_book(bundle: InputBundle, book_input: BookInput) -> BookOutput:
    currency = bundle.group.transaction_currency
    minor = bundle.currencies[currency].minor_unit
    group = bundle.group.group_key
    tb = TraceBuilder(engine_version=bundle.engine_version)
    booked = _lines(bundle)
    out_of_scope = _out_of_scope(booked, minor)
    lines = [item for item in booked if _scope_flag(item) in ALLOCATION_TARGET_FLAGS]
    products = {item.code: item for item in bundle.group.products}
    inside = _group_policy(book_input, "ssp.inside_range_point", "CONTRACT_PRICE")
    as_of = max(event.effective_date for event in bundle.events)
    deliveries = _deliveries(bundle)
    targets: dict[str, tuple[int, EventInput]] = {}
    prices: dict[str, Fraction] = {}
    price_nodes: list[str] = []
    selected: dict[str, Fraction] = {}
    selected_nodes: list[str] = []
    lineage: dict[str, dict[str, object]] = {}
    for item in lines:
        subject = item.subject_key
        price = to_fraction(str(item.line["total_price"]))
        prices[subject] = price
        posted = round_half_up(price, minor)
        price_nodes.append(
            _source(
                tb,
                "stated_price",
                subject,
                None,
                value=posted,
                currency=currency,
                minor_unit=minor,
                inputs=[
                    SourceRef(
                        "contract_event",
                        item.event.event_key,
                        {"value": format_money(posted, minor)},
                    )
                ],
            )
        )
        quantity = to_fraction(str(item.line["quantity"]))
        version, entry = _ssp(bundle, str(item.line["product_code"]), currency, item.inception)
        band = entry.ranges[0]
        low = mid = high = None
        in_range: bool | None = None
        if band.point_value is not None:
            chosen = quantity * to_fraction(band.point_value)
        else:
            assert band.low_value is not None and band.mid_value is not None
            assert band.high_value is not None
            low, mid, high = (
                quantity * to_fraction(value)
                for value in (band.low_value, band.mid_value, band.high_value)
            )
            in_range = low <= price <= high
            if in_range:
                chosen = price if inside == "CONTRACT_PRICE" else mid
            else:
                chosen = low if price < low else high
        selected[subject] = chosen
        selected_nodes.append(
            _source(
                tb,
                "original_ssp_selected",
                subject,
                None,
                value=chosen,
                currency=None,
                minor_unit=None,
                inputs=[SourceRef("ssp_entry", entry.entry_key, {"value": format_exact(chosen)})],
            )
        )
        lineage[subject] = {
            "ssp_book_version_key": version.version_key,
            "ssp_entry_key": entry.entry_key,
            "ssp_method": entry.method,
            "ssp_version_label": version.legacy_version_label,
            "original_ssp_low": low,
            "original_ssp_mid": mid,
            "original_ssp_high": high,
            "original_ssp_in_range": in_range,
            "original_ssp_selected": chosen,
        }
    total_price = round_half_up(sum(prices.values(), Fraction(0)), minor)
    price_node = tb.node(
        measure="transaction_price",
        subject_key=group,
        period_key=None,
        value=total_price,
        currency=currency,
        minor_unit=minor,
        formula_id="alloc.relative_ssp.v1",
        inputs=price_nodes,
        params={"role": "total"},
        narrative_key="alloc.relative_ssp",
    )
    total_ssp = sum(selected.values(), Fraction(0))
    total_node = tb.node(
        measure="total_ssp",
        subject_key=group,
        period_key=None,
        value=total_ssp,
        currency=None,
        minor_unit=None,
        formula_id="alloc.relative_ssp.v1",
        inputs=selected_nodes,
        params={"role": "total"},
        narrative_key="alloc.relative_ssp",
    )
    subjects = [item.subject_key for item in lines]
    shares = largest_remainder(total_price, [selected[key] for key in subjects], subjects)
    versions: list[ObligationVersionOut] = []
    schedule_lines: list[ScheduleLineOut] = []
    entities = {item.code: item for item in bundle.entities}
    activated = {
        event.contract_key for event in bundle.events if event.event_type == "CONTRACT_ACTIVATED"
    }
    for index, item in enumerate(lines):
        subject = item.subject_key
        product = products[str(item.line["product_code"])]
        template = _template(bundle, product.default_template_code, item.inception)
        weight = selected[subject] / total_ssp
        weight_node = tb.node(
            measure="allocation_weight",
            subject_key=subject,
            period_key=None,
            value=weight,
            currency=None,
            minor_unit=None,
            formula_id="alloc.relative_ssp.v1",
            inputs=[selected_nodes[index], total_node],
            params={"role": "weight"},
            narrative_key="alloc.relative_ssp",
        )
        allocated = shares[index]
        exact = Fraction(total_price, 10**minor) * weight
        allocated_node = tb.node(
            measure="allocated_amount",
            subject_key=subject,
            period_key=None,
            value=allocated,
            currency=currency,
            minor_unit=minor,
            formula_id="alloc.largest_remainder.v1",
            inputs=[price_node, *selected_nodes],
            params={"keys": ",".join(subjects), "index": str(index)},
            exact=exact,
            narrative_key="alloc.largest_remainder",
        )
        exact_node = tb.node(
            measure="allocated_exact",
            subject_key=subject,
            period_key=None,
            value=exact,
            currency=None,
            minor_unit=None,
            formula_id="alloc.relative_ssp.v1",
            inputs=[price_node, weight_node],
            params={"role": "exact"},
            narrative_key="alloc.relative_ssp",
        )
        start = item.line.get("start_date")
        end = item.line.get("end_date")
        start_date = None if start is None else date.fromisoformat(str(start))
        end_date = None if end is None else date.fromisoformat(str(end))
        scheduled = 0
        in_scope = _scope_flag(item) == IN_SCOPE
        # REQ-CON-002: revenue schedules exist once the contract is active.
        schedulable = in_scope and item.contract_key in activated
        if (
            template.recognition_method == "TIME_ELAPSED"
            and start_date
            and end_date
            and schedulable
        ):
            scheduled = allocated
            total_days = (end_date - start_date).days + 1
            previous = 0
            for period in entities[item.entity_code].periods:
                if period.end_date < start_date or period.start_date > end_date:
                    continue
                through = min(period.end_date, end_date)
                progress = Fraction((through - start_date).days + 1, total_days)
                cumulative = cumulative_posted(exact, allocated, progress, minor)
                amount = cumulative - previous
                previous = cumulative
                node = _source(
                    tb,
                    "revenue",
                    subject,
                    period.period_key,
                    value=amount,
                    currency=currency,
                    minor_unit=minor,
                    inputs=[
                        SourceRef(
                            "contract_event",
                            item.event.event_key,
                            {"value": format_money(amount, minor)},
                        )
                    ],
                )
                schedule_lines.append(
                    ScheduleLineOut(
                        schedule_kind=ScheduleKind.REVENUE,
                        subject_type="obligation",
                        subject_key=subject,
                        entity=item.entity_code,
                        period_key=period.period_key,
                        line_type=ScheduleLineType.NORMAL,
                        amount=amount,
                        cumulative_amount=cumulative,
                        cumulative_exact=exact * progress,
                        quantity=None,
                        is_released_at_close=True,
                        trace_node_id=node,
                    )
                )
        quantity = to_fraction(str(item.line["quantity"]))
        stated = round_half_up(prices[subject], minor)
        delivered_cum, returned_cum, latest = deliveries.get(
            subject, (Fraction(0), Fraction(0), None)
        )
        progress = Fraction(0)
        revenue = 0
        if template.recognition_method in UNIT_METHODS and latest is not None and in_scope:
            progress = min(max((delivered_cum - returned_cum) / quantity, Fraction(0)), Fraction(1))
            revenue = cumulative_posted(exact, allocated, progress, minor)
            targets[subject] = (revenue, latest)
        columns: dict[str, object] = {
            "obligation_key": str(item.line["obligation_key"]),
            "product_code": product.code,
            "sku_number": product.sku_number,
            "stratification": str(item.line.get("stratification") or ""),
            "obligation_kind": template.obligation_kind,
            "distinctness": template.distinctness,
            "series_increment_unit": template.series_increment_unit,
            "pob_template_version_key": template.version_key,
            "scope_flag": str(item.line.get("scope_flag") or "IN_SCOPE_606"),
            "satisfaction_pattern": template.satisfaction_pattern,
            "over_time_criterion": template.over_time_criterion,
            "recognition_method": template.recognition_method,
            "ratable_convention": template.ratable_convention,
            "principal_agent": product.principal_agent,
            "licence_nature": template.licence_nature,
            "warranty_type": template.warranty_type,
            "start_date": start_date,
            "end_date": end_date,
            "contracting_entity_code": item.entity_code,
            "performing_entity_code": item.entity_code,
            "txn_currency": currency,
            "account_overrides": {},
            "original_quantity": quantity,
            "original_stated_price": stated,
            "quantity": quantity,
            "stated_price": stated,
            **lineage[subject],
            "original_total_contract_price": total_price,
            "original_total_contract_ssp": total_ssp,
            "original_allocated_amount": allocated,
            "original_allocated_exact": exact,
            "original_unit_ssp": selected[subject] / quantity,
            "allocation_weight": weight,
            "allocated_amount": allocated,
            "allocated_exact": exact,
            "allocation_adjustment": allocated - stated,
            "unit_ssp": selected[subject] / quantity,
            "remaining_ssp": selected[subject],
            "remaining_quantity": quantity,
            "remaining_allocation": allocated,
            "remaining_billing": stated,
            "delivered_quantity_cum": delivered_cum,
            "returned_quantity_cum": returned_cum,
            "progress_ratio": progress,
            "revenue_cum": revenue,
            "scheduled_amount": scheduled,
            "awaiting_trigger_amount": allocated - scheduled - revenue,
            "position_obligation": 0,
            "position_contract_entity": 0,
            "satisfaction_status": "UNSATISFIED",
            "hold_types": (),
            "effective_date": as_of,
        }
        versions.append(
            ObligationVersionOut(
                subject_key=subject,
                columns=MappingProxyType(columns),
                trace_nodes=MappingProxyType(
                    {
                        "allocated_amount": allocated_node,
                        "allocated_exact": exact_node,
                        "allocation_weight": weight_node,
                        "original_ssp_selected": selected_nodes[index],
                    }
                ),
            )
        )
    statuses = tuple(
        sorted(
            (key, "ACTIVE" if key in activated else "DRAFT")
            for key in bundle.group.member_contract_keys
        )
    )
    scheduled_total = sum(int(cast(int, v.columns["scheduled_amount"])) for v in versions)
    revenue_total = sum(int(cast(int, v.columns["revenue_cum"])) for v in versions)
    version = ContractVersionOut(
        subject_key=group,
        columns=MappingProxyType(
            {
                "book_code": book_input.book_code,
                "transaction_currency": currency,
                "status_in_book": statuses[0][1],
                "status_reason_in_book": None,
                "transaction_price": total_price,
                "fixed_consideration": total_price,
                "out_of_scope_amount": out_of_scope,
                "total_ssp": total_ssp,
                "revenue_cum": revenue_total,
                "billed_cum": 0,
                "net_position": 0,
                "rpo_amount": total_price - revenue_total,
                "scheduled_amount": scheduled_total,
                "awaiting_trigger_amount": total_price - scheduled_total - revenue_total,
                "modification_boundary_no": 0,
            }
        ),
        trace_nodes=MappingProxyType({"total_ssp": total_node, "transaction_price": price_node}),
    )
    starts = {
        (item.code, period.period_key): period.start_date
        for item in bundle.entities
        for period in item.periods
    }
    return BookOutput(
        book_code=book_input.book_code,
        contract_version=version,
        status_in_book=statuses,
        obligation_versions=tuple(versions),
        balances=(),
        schedules=tuple(
            sorted(
                schedule_lines,
                key=lambda line: (
                    str(line.schedule_kind),
                    line.subject_key,
                    str(line.line_type),
                    starts[(line.entity, line.period_key)],
                ),
            )
        ),
        cost_asset_versions=(),
        loss_provision_versions=(),
        fx_layer_movements=(),
        posting_intents=_fake_intents(bundle, book_input, tb, lines, targets, activated),
        proposals=(),
        time_triggers=(),
        trace=tb.build(root_measures={"transaction_price": price_node}),
    )


def _fake_intents(
    bundle: InputBundle,
    book_input: BookInput,
    tb: TraceBuilder,
    lines: Sequence[_Line],
    targets: Mapping[str, tuple[int, EventInput]],
    activated: set[str],
) -> tuple[PostingIntent, ...]:
    """JET-02 principal (Dr CONTRACT_LIABILITY, Cr REVENUE) of each revenue delta of an ACTIVE
    contract against the posted EVENT amounts (S02-R-02; RCP-04 to RCP-06; S14-R-12 to S14-R-15)."""
    currency = bundle.group.transaction_currency
    minor = bundle.currencies[currency].minor_unit
    entities = {item.code: item for item in bundle.entities}
    products = {item.code: item for item in bundle.group.products}
    book_code = book_input.book_code
    intents: list[PostingIntent] = []
    for item in lines:
        target = targets.get(item.subject_key)
        if target is None or item.contract_key not in activated:
            continue
        revenue, event = target
        posted = -sum(
            amount.amount_txn
            for amount in bundle.posted
            if amount.book_code == book_code
            and amount.subject_key == item.subject_key
            and amount.entry_kind == REVENUE_RECOGNITION
            and amount.account_role == "REVENUE"
            and amount.posting_class == "EVENT"
        )
        delta = revenue - posted
        if delta == 0:
            continue
        entity_input = entities[item.entity_code]
        if entity_input.functional_currency != currency:
            raise EngineError("FX_RATE_MISSING", "the fake posts single-currency groups", detail={})
        assigned = _posting_period(entity_input, book_code, event.effective_date)
        if assigned is None:
            continue
        period_key, origin_key, reason = assigned
        product = products[str(item.line["product_code"])]
        node = _source(
            tb,
            "revenue_target",
            item.subject_key,
            origin_key or period_key,
            value=revenue,
            currency=currency,
            minor_unit=minor,
            inputs=[
                SourceRef(
                    "contract_event", event.event_key, {"value": format_money(revenue, minor)}
                )
            ],
        )
        dimensions = {
            "contract": bundle.group.group_key,
            "contract_key": item.contract_key,
            "obligation_key": str(item.line["obligation_key"]),
            "product": product.code,
        }
        if product.revenue_category is not None:
            dimensions["revenue_category"] = product.revenue_category
        sides = (
            (("D", "CONTRACT_LIABILITY"), ("C", "REVENUE"))
            if delta > 0
            else (("D", "REVENUE"), ("C", "CONTRACT_LIABILITY"))
        )
        intent_lines = tuple(
            IntentLine(
                line_key=sha256_hex(
                    {
                        "account_role": role,
                        "book_code": book_code,
                        "clearing_purpose": None,
                        "counterparty_entity": None,
                        "entity": item.entity_code,
                        "entry_kind": REVENUE_RECOGNITION,
                        "origin_period_key": origin_key,
                        "posting_class": "EVENT",
                        "posting_period_key": period_key,
                        "subject_key": item.subject_key,
                    }
                ),
                side=side,
                account_role=role,
                clearing_purpose=None,
                counterparty_entity=None,
                account_code=_account_code(book_input, role, item.entity_code, product.code),
                amount_txn=abs(delta),
                amount_functional=abs(delta),
                txn_currency=currency,
                functional_currency=currency,
                dimensions=MappingProxyType(dict(sorted(dimensions.items()))),
                source_event_key=None,
                trace_node_id=node,
            )
            for side, role in sides
        )
        intents.append(
            PostingIntent(
                entry_key=sha256_hex(
                    {
                        "book_code": book_code,
                        "entity": item.entity_code,
                        "entry_kind": REVENUE_RECOGNITION,
                        "origin_period_key": origin_key,
                        "posting_class": "EVENT",
                        "posting_period_key": period_key,
                        "reason_code": reason,
                        "subject_key": item.subject_key,
                    }
                ),
                book_code=book_code,
                entity=item.entity_code,
                posting_period_key=period_key,
                origin_period_key=origin_key,
                entry_kind=REVENUE_RECOGNITION,
                posting_class="EVENT",
                subject_key=item.subject_key,
                reason_code=reason,
                lines=intent_lines,
            )
        )
    return tuple(sorted(intents, key=lambda intent: intent.entry_key))


def fake_compute(bundle: InputBundle) -> OutputBundle:
    """A stand-in for ``erev_engine.compute`` over the documented output contract (§0.5)."""
    return OutputBundle(
        engine_version=bundle.engine_version,
        input_sha256=bundle.sha256(),
        books=tuple(_fake_book(bundle, book_input) for book_input in bundle.books),
        diagnostics=(),
    )


# --- the reference world of computation tests (PRD §2.6) ------------------------------------------

POB_TEMPLATES: Final = "/api/v1/pob-templates"
POB_TEMPLATE_VERSIONS: Final = "/api/v1/pob-template-versions"
CONFIG_TEST_CASES: Final = "/api/v1/config-test-cases"
SSP_BOOKS: Final = "/api/v1/ssp-books"
SSP_BOOK_VERSIONS: Final = "/api/v1/ssp-book-versions"
PRODUCTS: Final = "/api/v1/products"
CUSTOMERS: Final = "/api/v1/customers"
JANUARY: Final = "2026-01-01T00:00:00Z"
OVER_TIME: Final = {"satisfaction_pattern": "OVER_TIME", "over_time_criterion": "OT_A"}
# PRD §2.6: TPL-SUB-DAILY series by day, time elapsed, daily; TPL-SVC-PCT distinct, output percent.
TPL_SUB_DAILY: Final = {
    "distinctness": "series",
    "series_increment_unit": "day",
    "recognition_method": "TIME_ELAPSED",
    "ratable_convention": "DAILY",
    **OVER_TIME,
}
TPL_SVC_PCT: Final = {
    "distinctness": "distinct",
    "recognition_method": "OUTPUT_PERCENT",
    **OVER_TIME,
}
RELEASE_BUILD: Final = "tests-ctr-2"


def stamp_test_release() -> None:
    """The engine release row a computation names (05 REL-03), as the app's startup stamps it."""
    # The factory pins its own build whatever release-manifest.json a `make release-manifest` or
    # `make docker-build` left at the repository root (REL-03 reads a present manifest in every
    # environment): the absent manifest path keeps the derived facts.
    stamp_release(
        Environment.TEST,
        request_id="tests-release",
        manifest_path=REPO_ROOT / ".run" / "tests-release-manifest-absent.json",
        build_sha=lambda: RELEASE_BUILD,
    )


def open_periods(app: FastAPI, actor: Actor, *, entity_code: str, keys: Sequence[str]) -> None:
    """Open the named ASC606 periods of an entity through ``POST /periods/{id}/open``."""
    states = {
        item["period"]["period_key"]: item for item in periods(app, actor, entity=entity_code)
    }
    for key in keys:
        state = states[key]
        opened = post(
            app,
            f"/api/v1/periods/{state['id']}/open",
            actor,
            {"comment": "Open for contracts"},
            if_match=f'"r{state["row_version"]}"',
        )
        assert opened.status_code == 200, opened.text


def published_template(
    app: FastAPI,
    author: Actor,
    approver: Actor,
    *,
    code: str,
    outputs: Mapping[str, Any],
    case_line: Mapping[str, Any],
) -> dict[str, str]:
    """A template whose version 1 (effective 2026-01-01) passes one case and is published."""
    created = post(app, POB_TEMPLATES, author, {"code": code, "name": f"Template {code}"})
    assert created.status_code == 201, created.text
    template_id = str(created.json()["id"])
    version = post(
        app,
        f"{POB_TEMPLATES}/{template_id}/versions",
        author,
        {**outputs, "effective_from": JANUARY},
    )
    assert version.status_code == 201, version.text
    version_id = str(version.json()["id"])
    case = post(
        app,
        CONFIG_TEST_CASES,
        author,
        {
            "subject_type": "pob_template_version",
            "subject_id": version_id,
            "name": f"{case_line['product_code']} booking",
            "input": {"booking_date": "2026-01-01", "lines": [dict(case_line)]},
            "expected_output": {"drafts": [{"obligation_key": case_line["obligation_key"]}]},
        },
    )
    assert case.status_code == 201, case.text
    tested = post(app, f"{POB_TEMPLATE_VERSIONS}/{version_id}/test", author, {})
    assert (tested.status_code, tested.json()["status"]) == (200, "TESTED"), tested.text
    sent = post(
        app, f"{POB_TEMPLATE_VERSIONS}/{version_id}/submit", author, {"comment": "Ready for review"}
    )
    assert sent.status_code == 200, sent.text
    decided = approve(app, str(sent.json()["pending_approval_request_id"]), approver)
    assert decided.status_code == 200, decided.text
    shown = get(app, f"{POB_TEMPLATE_VERSIONS}/{version_id}", author).json()
    assert shown["status"] == "PUBLISHED", shown
    return {"template_id": template_id, "version_id": version_id}


def product_with_template(
    app: FastAPI,
    actor: Actor,
    *,
    code: str,
    name: str,
    revenue_category: str,
    template_id: str | None = None,
) -> str:
    """``POST /products`` (principal), then ``PATCH`` its default template once published."""
    created = new_product(
        app,
        actor,
        code=code,
        name=name,
        revenue_category=revenue_category,
        principal_agent="PRINCIPAL",
    )
    product_id = str(created["id"])
    if template_id is not None:
        set_default_template(app, actor, product_id, template_id)
    return product_id


def set_default_template(app: FastAPI, actor: Actor, product_id: str, template_id: str) -> None:
    shown = get(app, f"{PRODUCTS}/{product_id}", actor)
    assert shown.status_code == 200, shown.text
    changed = patch(
        app,
        f"{PRODUCTS}/{product_id}",
        actor,
        {"default_pob_template_id": template_id},
        if_match=shown.headers["ETag"],
    )
    assert changed.status_code == 200, changed.text


def point_entry(
    product_code: str,
    point: str,
    method: str = "observable",
    *,
    currency: str = "USD",
    value_basis: str | None = None,
) -> dict[str, Any]:
    """A point-value SSP entry; ``value_basis`` is the E-49 basis a series product's entry must
    declare (D-97 (3a): the API refuses a series product's entry without one)."""
    entry: dict[str, Any] = {
        "product_code": product_code,
        "currency": currency,
        "method": method,
        "distinctness": "distinct",
        "ranges": [{"point_value": point}],
    }
    if value_basis is not None:
        entry["value_basis"] = value_basis
    return entry


def range_entry(
    product_code: str,
    low: str,
    mid: str,
    high: str,
    method: str = "observable",
    *,
    currency: str = "USD",
    value_basis: str | None = None,
    quantity_unit: str | None = None,
) -> dict[str, Any]:
    """An SSP entry with one low, mid and high band per unit (T-REF-31); ``value_basis`` is the
    E-49 pricing basis a series product's entry declares (D-93 (4)) and ``quantity_unit`` the E-125
    unit a PER_INCREMENT entry declares (D-97 (3))."""
    entry: dict[str, Any] = {
        "product_code": product_code,
        "currency": currency,
        "method": method,
        "distinctness": "distinct",
        "ranges": [{"low_value": low, "mid_value": mid, "high_value": high}],
    }
    if value_basis is not None:
        entry["value_basis"] = value_basis
    if quantity_unit is not None:
        entry["quantity_unit"] = quantity_unit
    return entry


def approved_ssp_version(
    app: FastAPI,
    author: Actor,
    approvers: Sequence[Actor],
    book_id: str,
    *,
    label: str,
    effective_from: str,
    entries: Sequence[Mapping[str, Any]],
) -> str:
    """An SSP book version with ``entries``, its study attached, submitted and approved."""
    body = {
        "legacy_version_label": label,
        "effective_from_date": effective_from,
        "methodology_label": "List-price study",
    }
    created = post(app, f"{SSP_BOOKS}/{book_id}/versions", author, body)
    assert created.status_code == 201, created.text
    version_id = str(created.json()["id"])
    stored = post(
        app, f"{SSP_BOOK_VERSIONS}/{version_id}/entries", author, {"entries": list(entries)}
    )
    assert stored.status_code == 200, stored.text
    uploaded = call(
        app,
        "POST",
        "/api/v1/files",
        data={"purpose": "SSP_STUDY"},
        files={"file": ("study.pdf", upload_fixtures.PDF, "application/pdf")},
        headers=cookie_headers(author.token, author.csrf_token),
    )
    assert uploaded.status_code == 201, uploaded.text
    attached = post(
        app,
        "/api/v1/attachments",
        author,
        {
            "file_object_id": uploaded.json()["id"],
            "subject_type": "ssp_book_version",
            "subject_id": version_id,
        },
    )
    assert attached.status_code == 201, attached.text
    current = get(app, f"{SSP_BOOK_VERSIONS}/{version_id}", author)
    sent = post(
        app,
        f"{SSP_BOOK_VERSIONS}/{version_id}/submit",
        author,
        {"comment": "Supported by the study."},
        if_match=current.headers["ETag"],
    )
    assert sent.status_code == 200, sent.text
    request_id = str(sent.json()["approval_request_id"])
    for approver in approvers:
        decided = approve(app, request_id, approver)
        assert decided.status_code == 200, decided.text
        if decided.json()["status"] == "APPROVED":
            break
    shown = get(app, f"{SSP_BOOK_VERSIONS}/{version_id}", author).json()
    assert shown["status"] == "APPROVED", shown
    return version_id


def ssp_book(app: FastAPI, author: Actor, *, code: str = "US-LIST", currency: str = "USD") -> str:
    created = post(
        app, SSP_BOOKS, author, {"code": code, "name": f"{code} list prices", "currency": currency}
    )
    assert created.status_code == 201, created.text
    return str(created.json()["id"])


# (code, name, account type, normal balance, role) of a minimal chart (PRD §2.6).
AVM_US_CHART: Final = (
    ("1200", "Accounts receivable", "ASSET", "D", "ACCOUNTS_RECEIVABLE"),
    ("1210", "Unbilled receivables", "ASSET", "D", "UNBILLED_RECEIVABLE"),
    ("2100", "Deferred revenue", "LIABILITY", "C", "CONTRACT_LIABILITY"),
    ("4010", "Revenue - services and subscriptions", "REVENUE", "C", "REVENUE"),
)
# CTR-7: the AVM-US chart with the accounts of a receipt while not a contract (JET-01b: Dr
# BILLING_CLEARING UNAPPLIED_CASH / Cr DEPOSIT_LIABILITY).
STEP1_CHART: Final = (
    *AVM_US_CHART,
    ("1900", "Clearing - unapplied cash", "ASSET", "D", "BILLING_CLEARING:UNAPPLIED_CASH"),
    ("2105", "Deposit liability", "LIABILITY", "C", "DEPOSIT_LIABILITY"),
)
# PRD §2.6 accounts of the roles a K-11 computation reaches.
K11_CHART: Final = (
    ("1100", "Accounts receivable", "ASSET", "D", "ACCOUNTS_RECEIVABLE"),
    ("1105", "Unbilled receivable", "ASSET", "D", "UNBILLED_RECEIVABLE"),
    ("1200", "Contract asset", "ASSET", "D", "CONTRACT_ASSET"),
    ("2100", "Contract liability", "LIABILITY", "C", "CONTRACT_LIABILITY"),
    ("4000", "Revenue - products", "REVENUE", "C", "REVENUE"),
)


def published_mapping(
    app: FastAPI,
    author: Actor,
    approver: Actor,
    *,
    chart: Sequence[tuple[str, str, str, str, str]] = AVM_US_CHART,
) -> str:
    """AVM-MAP-2026-01 (effective 2026-01-01) over ``chart`` (PRD §2.6). A role written
    ``BILLING_CLEARING:<purpose>`` names its clearing purpose."""
    rules = []
    for code, name, kind, normal, named in chart:
        role, _, purpose = named.partition(":")
        rule: dict[str, Any] = {
            "account_role": role,
            "gl_account_id": gl_account(
                app, author, code=code, name=name, account_type=kind, normal_balance=normal
            ),
        }
        if purpose:
            rule["clearing_purpose"] = purpose
        rules.append(rule)
    version = mapping_published(
        app, author, approver, name="AVM-MAP-2026-01", effective_from=JANUARY, rules=rules
    )
    return str(version["id"])


def customer_id(app: FastAPI, actor: Actor, *, code: str, name: str) -> UUID:
    created = post(app, CUSTOMERS, actor, {"code": code, "name": name})
    assert created.status_code == 201, created.text
    return UUID(str(created.json()["id"]))


def maya_principal(someone: Member) -> Principal:
    """Maya with the Revenue Accountant permissions for every entity (PRD §5.6)."""
    permissions = DEFAULT_ROLES["revenue_accountant"]
    scopes: dict[str, Literal["*"] | frozenset[UUID]] = {code: "*" for code in permissions}
    return Principal(
        kind=PrincipalKind.USER,
        id=someone.user_id,
        tenant_id=someone.tenant_id,
        membership_id=someone.membership_id,
        display_name="Maya Okafor",
        roles=("revenue_accountant",),
        permissions=permissions,
        permission_scopes=MappingProxyType(scopes),
        entity_scope="*",
        auth_method="password",
        mfa_verified_at=None,
        session_id=None,
        support_grant_id=None,
        on_behalf_of_id=None,
    )


@dataclass(frozen=True, slots=True)
class Workspace:
    """A signed-in author and the units of work of the domain commands."""

    app: FastAPI
    clock: FrozenClock
    keyring: KeyRing
    files: LocalFileStore
    author: Actor
    principal: Principal

    @property
    def tenant_id(self) -> UUID:
        return self.principal.tenant_id

    @contextmanager
    def uow(self, principal: Principal | None = None) -> Iterator[UnitOfWork]:
        ctx = RequestContext(
            principal=self.principal if principal is None else principal,
            tenant_kind=TenantKind.PRODUCTION,
            request_id="tests-contract-computation",
            source_ip=None,
            user_agent=None,
            idempotency_key=None,
            if_match=None,
            now=self.clock.now(),
            format_locale="en-US",
        )
        with unit_of_work(ctx, clock=self.clock, keyring=self.keyring, files=self.files) as uow:
            yield uow

    def rows(self, statement: Select[Any]) -> list[dict[str, Any]]:
        context = DbContext(tenant_id=self.tenant_id, user_id=None, entity_scope="*")
        with tenant_session(context, read_only=True) as session:
            return [dict(row) for row in session.execute(statement).mappings()]

    def scalar(self, statement: Select[Any]) -> Any:
        context = DbContext(tenant_id=self.tenant_id, user_id=None, entity_scope="*")
        with tenant_session(context, read_only=True) as session:
            return session.execute(statement).scalar_one()


def workspace(
    app: FastAPI, clock: FrozenClock, keyring: KeyRing, files: LocalFileStore, author: Actor
) -> Workspace:
    return Workspace(
        app=app,
        clock=clock,
        keyring=keyring,
        files=files,
        author=author,
        principal=maya_principal(author.member),
    )


def booked_contract(place: Workspace, body: Mapping[str, Any], *, activate: bool) -> BookedContract:
    """``book_contract``, then ``CONTRACT_ACTIVATED`` at the inception date when ``activate``."""
    with place.uow() as uow:
        booked = book_contract(uow, body=ContractBookedV1.model_validate(dict(body)), origin="UI")
        if activate:
            append_events(
                uow,
                contract_id=booked.contract["id"],
                expected_stream_version=1,
                events=[
                    EventIn(
                        event_type=ContractEventType.CONTRACT_ACTIVATED,
                        effective_date=booked.contract["inception_date"],
                        payload=ContractActivatedV1(
                            checklist=(ChecklistItemV1(code="SOURCE_REFERENCE", passed=True),)
                        ),
                    )
                ],
                origin="UI",
            )
        uow.commit()
    if activate:
        booked = replace(booked, contract={**booked.contract, "status": "ACTIVE"})
    return booked


def activated_contract(
    place: Workspace,
    booked: BookedContract,
    *,
    run: computation.Engine | None = None,
    compute: bool = True,
) -> BookedContract:
    """BS3-D-19 as CTR-9 builds it: the approved activation ``activation.activate`` as the SYSTEM
    principal, which appends ``CONTRACT_ACTIVATED`` at the inception date with the evaluated
    checklist, computes and persists the group and fails closed. The fixtures record no Step 1
    review, so the checklist is stored without the gate and no request is routed; ``compute =
    False`` leaves the computation to the caller (L4-1-Q-21)."""
    with place.uow(system_principal(place.tenant_id)) as uow:
        activation.activate(
            uow,
            contract_id=UUID(str(booked.contract["id"])),
            approval_request_id=None,
            on_behalf_of=None,
            engine=run,
            gate=False,
            compute=compute,
        )
        uow.commit()
    return replace(booked, contract={**booked.contract, "status": "ACTIVE"})


def appended(place: Workspace, contract_id: UUID, expected: int, events: Sequence[EventIn]) -> None:
    """Append ``events`` to a contract after head ``expected`` in one committed unit of work."""
    with place.uow() as uow:
        append_events(
            uow,
            contract_id=contract_id,
            expected_stream_version=expected,
            events=list(events),
            origin="UI",
        )
        uow.commit()


def drafted_override(
    place: Workspace,
    contract_id: UUID,
    policy_key: str,
    value: Any,
    *,
    obligation_key: str | None = None,
    rationale: str = "Stated by the order form (a test's own DRAFT override).",
) -> UUID:
    """A DRAFT policy override of ``contract_id`` written as the workspace's author, by the
    fixture and not by the product: release 1.0 creates none (04 T-CON-23 "Not offered in release
    1.0", rev 1.322; ``support.rows.insert_policy_override``). A test of the kept code submits it
    through ``POST /policy-overrides/{id}/submit`` and has it approved."""
    # Imported here, as in import_world: support.principals imports this module.
    from support.rows import insert_policy_override

    context = DbContext(tenant_id=place.tenant_id, user_id=None, entity_scope="*")
    with tenant_session(context) as session:
        return insert_policy_override(
            session,
            place.tenant_id,
            contract_id=contract_id,
            obligation_key=obligation_key,
            created_by=place.author.member.user_id,
            policy_key=policy_key,
            value=value,
            rationale=rationale,
        )


def computed(
    place: Workspace,
    group_id: UUID,
    *,
    run: computation.Engine | None = None,
    trigger: ComputationTrigger = ComputationTrigger.COMMAND,
) -> tuple[InputBundle, OutputBundle, Mapping[str, Any]]:
    """Build the group's bundle at the clock, compute it once and persist the output."""
    compute = run if run is not None else engine()
    with place.uow() as uow:
        bundle = bundles.build(uow.session, group_id, uow.now, (), trigger)
        output = compute(bundle)
        stored = computation.persist(uow, bundle, output)
        uow.commit()
    return bundle, output, stored


def world_calendar(
    app: FastAPI,
    author: Actor,
    *,
    entity_code: str = "AVM-US",
    functional_currency: str = "USD",
    time_zone: str = "America/New_York",
) -> tuple[str, UUID]:
    """A January calendar for 2026 and 2027 and an entity whose FY2026 P01 to P09 are open."""
    calendar_id = calendar(app, author, years=(2026, 2027))
    created = entity(
        app,
        author,
        code=entity_code,
        calendar_id=calendar_id,
        functional_currency=functional_currency,
        time_zone=time_zone,
    )
    open_periods(
        app,
        author,
        entity_code=entity_code,
        keys=[f"FY2026-P{month:02d}" for month in range(1, 10)],
    )
    return calendar_id, UUID(str(created["id"]))


def random_suffix() -> str:
    return secrets.token_hex(3).upper()


# --- the K-11 world of CTR-3 (PRD §2.6, WLD-K-11, WLD-X-22) ---------------------------------------

GATEWAY: Final = "AVM-GW"
SUPPORT: Final = "AVM-SUP-12"
K11_EXTERNAL_ID: Final = "NS-SO-DE-5004"
# PRD §2.6 TPL-PROD-UNITS: point in time per unit, units delivered.
TPL_PROD_UNITS: Final = {
    "distinctness": "distinct",
    "satisfaction_pattern": "POINT_IN_TIME",
    "over_time_criterion": "NOT_APPLICABLE",
    "recognition_method": "UNITS_DELIVERED",
}
GATEWAY_CASE: Final = {
    "obligation_key": "POB-01",
    "product_code": GATEWAY,
    "quantity": "200",
    "total_price": "90000.00",
}
SUPPORT_CASE: Final = {
    "obligation_key": "POB-01",
    "product_code": SUPPORT,
    "quantity": "1",
    "total_price": "18000.00",
    "start_date": "2026-09-15",
    "end_date": "2027-09-14",
}


@dataclass(frozen=True, slots=True)
class K11World:
    place: Workspace
    priya: Actor
    marcus: Actor
    entity_id: UUID
    customer_id: UUID
    book_id: str
    version_id: str
    mapping_id: str

    @property
    def app(self) -> FastAPI:
        return self.place.app


def k11_world(
    app: FastAPI,
    keyring: KeyRing,
    clock: FrozenClock,
    files: LocalFileStore,
    *,
    chart: Sequence[tuple[str, str, str, str, str]] | None = None,
) -> K11World:
    """PRD §2.6 for K-11: Maya (Revenue Accountant, SSP Analyst) prepares; Priya (SSP Approver)
    approves DE-LIST 2026; Marcus (Controller, SSP Approver, Tenant Admin; MFA) enables EUR and
    approves the templates and the mapping. AVM-DE (EUR, Europe/Berlin) has FY2026-P01 to P09 open;
    customer C-11; AVM-GW (TPL-PROD-UNITS; observable 405.00 / 450.00 / 495.00 EUR per unit) and
    AVM-SUP-12 (TPL-SUB-DAILY; cost plus margin 19,000.00 / 20,000.00 / 21,000.00 EUR);
    AVM-MAP-2026-01 over ``chart``, by default ``K11_CHART``."""
    maya_member = member(keyring, clock)
    assign(maya_member, "revenue_accountant")
    maya = holding(app, maya_member, "ssp_analyst")
    approvers: dict[str, Actor] = {}
    for name, roles in (
        ("priya", ("ssp_approver",)),
        ("marcus", ("controller", "ssp_approver", "tenant_admin")),
    ):
        someone = colleague(maya_member.tenant_id, name)
        for code in roles:
            assign(someone, code)
        approvers[name] = enrolled(app, clock, someone)
    marcus = approvers["marcus"]
    enabled = put(app, "/api/v1/tenant-currencies", marcus, {"currency_codes": ["USD", "EUR"]})
    assert enabled.status_code == 200, enabled.text
    _, entity_id = world_calendar(
        app, maya, entity_code="AVM-DE", functional_currency="EUR", time_zone="Europe/Berlin"
    )
    buyer = customer_id(app, maya, code="C-11", name="Hollenbrand Medizintechnik GmbH (Demo)")
    products = {
        GATEWAY: product_with_template(
            app, maya, code=GATEWAY, name="Sensor gateway unit", revenue_category="PRODUCT"
        ),
        SUPPORT: product_with_template(
            app, maya, code=SUPPORT, name="Platform support, 12 months", revenue_category="SERVICES"
        ),
    }
    units = published_template(
        app, maya, marcus, code="TPL-PROD-UNITS", outputs=TPL_PROD_UNITS, case_line=GATEWAY_CASE
    )
    daily = published_template(
        app, maya, marcus, code="TPL-SUB-DAILY", outputs=TPL_SUB_DAILY, case_line=SUPPORT_CASE
    )
    set_default_template(app, maya, products[GATEWAY], units["template_id"])
    set_default_template(app, maya, products[SUPPORT], daily["template_id"])
    book_id = ssp_book(app, maya, code="DE-LIST", currency="EUR")
    version_id = approved_ssp_version(
        app,
        maya,
        [approvers["priya"]],
        book_id,
        label="2026",
        effective_from="2026-01-01",
        entries=[
            range_entry(GATEWAY, "405.00", "450.00", "495.00", currency="EUR"),
            # D-97 (3a): a series product's entry declares its basis (AMOUNT: the remaining
            # increments at d, the reading these worlds embody).
            range_entry(
                SUPPORT,
                "19000.00",
                "20000.00",
                "21000.00",
                "cost_plus_margin",
                currency="EUR",
                value_basis="AMOUNT",
            ),
        ],
    )
    mapping_id = published_mapping(app, maya, marcus, chart=chart or K11_CHART)
    stamp_test_release()
    return K11World(
        place=workspace(app, clock, keyring, files, maya),
        priya=approvers["priya"],
        marcus=marcus,
        entity_id=entity_id,
        customer_id=buyer,
        book_id=book_id,
        version_id=version_id,
        mapping_id=mapping_id,
    )


def k11_body(customer: UUID) -> dict[str, Any]:
    """PRD WLD-K-11 ``NS-SO-DE-5004``: O1 AVM-GW 200 units × 450.00 = 90,000.00 EUR; O2 AVM-SUP-12
    18,000.00 EUR from 15 Sep 2026 to 14 Sep 2027; inception 2026-09-01."""
    return {
        "external_id": K11_EXTERNAL_ID,
        "customer_id": str(customer),
        "contracting_entity_code": "AVM-DE",
        "transaction_currency": "EUR",
        "inception_date": "2026-09-01",
        "lines": [
            {
                "obligation_key": "O1",
                "product_code": GATEWAY,
                "quantity": "200",
                "total_price": {"amount": "90000.00", "currency": "EUR"},
            },
            {
                "obligation_key": "O2",
                "product_code": SUPPORT,
                "quantity": "1",
                "total_price": {"amount": "18000.00", "currency": "EUR"},
                "start_date": "2026-09-15",
                "end_date": "2027-09-14",
            },
        ],
    }


def delivered_k11(world: K11World, *, activate: bool = True) -> BookedContract:
    """K-11 booked, activated as SYSTEM when ``activate`` (BS3-D-19), then ``DELIVERY_RECORDED``
    of 120 O1 units effective 2026-09-12 (PRD J-04.1)."""
    booked = booked_contract(world.place, k11_body(world.customer_id), activate=False)
    head = 1
    if activate:
        # L4-1-Q-21: the CTR-3 subledger figures count the postings of the caller's computation
        # after the delivery, so the activation computes nothing here.
        booked = activated_contract(world.place, booked, compute=False)
        head = 2
    appended(
        world.place,
        booked.contract["id"],
        head,
        [
            EventIn(
                event_type=ContractEventType.DELIVERY_RECORDED,
                effective_date=date(2026, 9, 12),
                payload=DeliveryRecordedV1(obligation_key="O1", quantity="120", trigger="DELIVERY"),
            )
        ],
    )
    return booked


# --- the K-02 world of CTR-15 (PRD §2.6, WLD-C-02, WLD-K-02) --------------------------------------

SEAT_MONTH: Final = "AVM-SEAT-MO"
K02_EXTERNAL_ID: Final = "SF-ORD-10002"
# [J] L4-2-Q-8: 100 seats × 24 months as 2,400 seat-months, the unit of the SSP entry.
SEAT_MONTH_CASE: Final = {
    "obligation_key": "POB-01",
    "product_code": SEAT_MONTH,
    "quantity": "2400",
    "total_price": "240000.00",
    "start_date": "2026-01-01",
    "end_date": "2027-12-31",
}


@dataclass(frozen=True, slots=True)
class K02World:
    place: Workspace
    priya: Actor
    marcus: Actor
    entity_id: UUID
    customer_id: UUID
    book_id: str
    version_id: str
    mapping_id: str

    @property
    def app(self) -> FastAPI:
        return self.place.app


def k02_world(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> K02World:
    """PRD §2.6 for K-02: Maya (Revenue Accountant, SSP Analyst) prepares; Priya (SSP Approver)
    approves US-LIST 2026-H1; Marcus (Controller, SSP Approver, Tenant Admin; MFA) enables USD and
    approves the template and the mapping. AVM-US (USD, America/New_York) has FY2026-P01 to P09
    open; customer C-02; AVM-SEAT-MO (TPL-SUB-DAILY; observable 90.00 / 100.00 / 110.00 USD per
    seat-month); AVM-MAP-2026-01 over ``AVM_US_CHART``."""
    maya_member = member(keyring, clock)
    assign(maya_member, "revenue_accountant")
    maya = holding(app, maya_member, "ssp_analyst")
    approvers: dict[str, Actor] = {}
    for name, roles in (
        ("priya", ("ssp_approver",)),
        ("marcus", ("controller", "ssp_approver", "tenant_admin")),
    ):
        someone = colleague(maya_member.tenant_id, name)
        for code in roles:
            assign(someone, code)
        approvers[name] = enrolled(app, clock, someone)
    marcus = approvers["marcus"]
    enabled = put(app, "/api/v1/tenant-currencies", marcus, {"currency_codes": ["USD"]})
    assert enabled.status_code == 200, enabled.text
    _, entity_id = world_calendar(app, maya)
    buyer = customer_id(app, maya, code="C-02", name="Marrowby Health Partners LLC (Demo)")
    seats = product_with_template(
        app,
        maya,
        code=SEAT_MONTH,
        name="Platform seat, per seat per month",
        revenue_category="SUBSCRIPTION",
    )
    daily = published_template(
        app, maya, marcus, code="TPL-SUB-DAILY", outputs=TPL_SUB_DAILY, case_line=SEAT_MONTH_CASE
    )
    set_default_template(app, maya, seats, daily["template_id"])
    book_id = ssp_book(app, maya)
    version_id = approved_ssp_version(
        app,
        maya,
        [approvers["priya"]],
        book_id,
        label="2026-H1",
        effective_from="2026-01-01",
        # D-93 (4) / D-97 (3): the seat-month entry prices one increment and the line's quantity
        # counts increments (PRD K-02; 04 E-49 PER_INCREMENT, E-125 INCREMENTS).
        entries=[
            range_entry(
                SEAT_MONTH,
                "90.00",
                "100.00",
                "110.00",
                value_basis="PER_INCREMENT",
                quantity_unit="INCREMENTS",
            )
        ],
    )
    mapping_id = published_mapping(app, maya, marcus)
    stamp_test_release()
    return K02World(
        place=workspace(app, clock, keyring, files, maya),
        priya=approvers["priya"],
        marcus=marcus,
        entity_id=entity_id,
        customer_id=buyer,
        book_id=book_id,
        version_id=version_id,
        mapping_id=mapping_id,
    )


def k02_seat_month_body(customer: UUID) -> dict[str, Any]:
    """PRD WLD-K-02 ``SF-ORD-10002``: O1 AVM-SEAT-MO, 100 seats × 24 months for 240,000.00 USD from
    01 Jan 2026 to 31 Dec 2027; inception 2026-01-01."""
    return {
        "external_id": K02_EXTERNAL_ID,
        "customer_id": str(customer),
        "contracting_entity_code": "AVM-US",
        "transaction_currency": "USD",
        "inception_date": "2026-01-01",
        "lines": [
            {
                "obligation_key": "O1",
                "product_code": SEAT_MONTH,
                "quantity": "2400",
                "total_price": {"amount": "240000.00", "currency": "USD"},
                "start_date": "2026-01-01",
                "end_date": "2027-12-31",
            }
        ],
    }


# --- the SF-ORD-20417 world of CTR-4 (PRD §2.6, WLD-C-12, WLD-F-20, WLD-X-23) ---------------------

PLATFORM_100: Final = "AVM-PLAT-100"
IMPLEMENTATION_PLUS: Final = "AVM-IMPL-PLUS"
SF_ORD_20417: Final = "SF-ORD-20417"
# PRD §2.6 TPL-SVC-HOURS: distinct, over time, labour hours.
TPL_SVC_HOURS: Final = {
    "distinctness": "distinct",
    "recognition_method": "LABOUR_HOURS",
    **OVER_TIME,
}
PLATFORM_100_CASE: Final = {
    "obligation_key": "POB-01",
    "product_code": PLATFORM_100,
    "quantity": "1",
    "total_price": "96000.00",
    "start_date": "2026-09-01",
    "end_date": "2027-08-31",
}
IMPLEMENTATION_PLUS_CASE: Final = {
    "obligation_key": "POB-01",
    "product_code": IMPLEMENTATION_PLUS,
    "quantity": "1",
    "total_price": "24000.00",
}


@dataclass(frozen=True, slots=True)
class J03World:
    place: Workspace
    priya: Actor
    marcus: Actor
    calendar_id: str
    entity_id: UUID
    uk_entity_id: UUID
    customer_id: UUID
    book_id: str
    version_id: str
    mapping_id: str

    @property
    def app(self) -> FastAPI:
        return self.place.app


def j03_world(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> J03World:
    """PRD §2.6 for J-03: Maya (Revenue Accountant, SSP Analyst) prepares; Priya (SSP Approver)
    approves US-LIST 2026-H1; Marcus (Controller, SSP Approver, Tenant Admin; MFA) enables USD, EUR
    and GBP and approves the templates and the mapping. AVM-US (USD, America/New_York) and AVM-UK
    (GBP, Europe/London) share a January calendar with FY2026-P01 to P09 open; customer C-12;
    AVM-PLAT-100 (TPL-SUB-DAILY; observable 85,000.00 / 100,000.00 / 115,000.00 USD) and
    AVM-IMPL-PLUS (TPL-SVC-HOURS; cost plus margin 18,000.00 / 20,000.00 / 22,000.00 USD);
    AVM-MAP-2026-01 over ``AVM_US_CHART``."""
    maya_member = member(keyring, clock)
    assign(maya_member, "revenue_accountant")
    maya = holding(app, maya_member, "ssp_analyst")
    approvers: dict[str, Actor] = {}
    for name, roles in (
        ("priya", ("ssp_approver",)),
        ("marcus", ("controller", "ssp_approver", "tenant_admin")),
    ):
        someone = colleague(maya_member.tenant_id, name)
        for code in roles:
            assign(someone, code)
        approvers[name] = enrolled(app, clock, someone)
    marcus = approvers["marcus"]
    enabled = put(
        app, "/api/v1/tenant-currencies", marcus, {"currency_codes": ["USD", "EUR", "GBP"]}
    )
    assert enabled.status_code == 200, enabled.text
    calendar_id, entity_id = world_calendar(app, maya)
    uk = entity(
        app,
        maya,
        code="AVM-UK",
        calendar_id=calendar_id,
        functional_currency="GBP",
        time_zone="Europe/London",
    )
    open_periods(
        app, maya, entity_code="AVM-UK", keys=[f"FY2026-P{month:02d}" for month in range(1, 10)]
    )
    buyer = customer_id(app, maya, code="C-12", name="Kinsley Marrow Foods Inc. (Demo)")
    products = {
        PLATFORM_100: product_with_template(
            app,
            maya,
            code=PLATFORM_100,
            name="Platform, 100 seats, 12 months",
            revenue_category="SUBSCRIPTION",
        ),
        IMPLEMENTATION_PLUS: product_with_template(
            app,
            maya,
            code=IMPLEMENTATION_PLUS,
            name="Implementation, extended scope",
            revenue_category="SERVICES",
        ),
    }
    daily = published_template(
        app, maya, marcus, code="TPL-SUB-DAILY", outputs=TPL_SUB_DAILY, case_line=PLATFORM_100_CASE
    )
    hours = published_template(
        app,
        maya,
        marcus,
        code="TPL-SVC-HOURS",
        outputs=TPL_SVC_HOURS,
        case_line=IMPLEMENTATION_PLUS_CASE,
    )
    set_default_template(app, maya, products[PLATFORM_100], daily["template_id"])
    set_default_template(app, maya, products[IMPLEMENTATION_PLUS], hours["template_id"])
    book_id = ssp_book(app, maya)
    version_id = approved_ssp_version(
        app,
        maya,
        [approvers["priya"]],
        book_id,
        label="2026-H1",
        effective_from="2026-01-01",
        entries=[
            range_entry(PLATFORM_100, "85000.00", "100000.00", "115000.00", value_basis="AMOUNT"),
            range_entry(
                IMPLEMENTATION_PLUS, "18000.00", "20000.00", "22000.00", "cost_plus_margin"
            ),
        ],
    )
    mapping_id = published_mapping(app, maya, marcus)
    stamp_test_release()
    return J03World(
        place=workspace(app, clock, keyring, files, maya),
        priya=approvers["priya"],
        marcus=marcus,
        calendar_id=calendar_id,
        entity_id=entity_id,
        uk_entity_id=UUID(str(uk["id"])),
        customer_id=buyer,
        book_id=book_id,
        version_id=version_id,
        mapping_id=mapping_id,
    )


def sf_ord_20417_body(
    customer: UUID,
    *,
    external_id: str = SF_ORD_20417,
    implementation_price: str = "24000.00",
) -> dict[str, Any]:
    """PRD WLD-F-20 ``SF-ORD-20417`` for C-12 (AVM-US, USD, inception 2026-09-01, ``document_ref``
    the order number): O1 AVM-PLAT-100 96,000.00 from 01 Sep 2026 to 31 Aug 2027; O2 AVM-IMPL-PLUS
    ``implementation_price`` performed by AVM-UK."""
    return {
        "external_id": external_id,
        "customer_id": str(customer),
        "contracting_entity_code": "AVM-US",
        "transaction_currency": "USD",
        "inception_date": "2026-09-01",
        "document_ref": external_id,
        "lines": [
            {
                "obligation_key": "O1",
                "product_code": PLATFORM_100,
                "quantity": "1",
                "total_price": {"amount": "96000.00", "currency": "USD"},
                "start_date": "2026-09-01",
                "end_date": "2027-08-31",
            },
            {
                "obligation_key": "O2",
                "product_code": IMPLEMENTATION_PLUS,
                "quantity": "1",
                "total_price": {"amount": implementation_price, "currency": "USD"},
                "performing_entity_code": "AVM-UK",
            },
        ],
    }


# --- the AVM-US subscription world of CTR-5 (PRD §2.6, WLD-K-02, WLD-K-09, WLD-X-05) -------------

SEAT_MO: Final = "AVM-SEAT-MO"
K09_EXTERNAL_ID: Final = "SF-ORD-10417"
SEAT_MO_CASE: Final = {
    "obligation_key": "POB-01",
    "product_code": SEAT_MO,
    "quantity": "100",
    "total_price": "240000.00",
    "start_date": "2026-01-01",
    "end_date": "2027-12-31",
}


@dataclass(frozen=True, slots=True)
class SeatWorld:
    place: Workspace
    priya: Actor
    marcus: Actor
    entity_id: UUID
    customers: Mapping[str, UUID]
    book_id: str
    version_id: str
    mapping_id: str

    @property
    def app(self) -> FastAPI:
        return self.place.app


def seat_world(
    app: FastAPI,
    keyring: KeyRing,
    clock: FrozenClock,
    files: LocalFileStore,
    *,
    unpriced: Sequence[str] = (),
    chart: Sequence[tuple[str, str, str, str, str]] = AVM_US_CHART,
    zero_priced: Sequence[str] = (),
    untemplated: Sequence[str] = (),
) -> SeatWorld:
    """PRD §2.6 for K-02 and K-09: Maya (Revenue Accountant, SSP Analyst) prepares; Priya (SSP
    Approver) approves US-LIST 2026-H1; Marcus (Controller, SSP Approver, Tenant Admin; MFA)
    approves the template and the mapping. AVM-US (USD, America/New_York) has a January calendar
    for 2026 to 2029, which K-09's term needs, with FY2026-P01 to P09 open; customers C-02 and
    C-09; AVM-SEAT-MO (TPL-SUB-DAILY; observable 2,160.00 / 2,400.00 / 2,640.00 USD per seat);
    AVM-MAP-2026-01 over ``AVM_US_CHART``. Each ``unpriced`` product code gets the template and no
    SSP entry (CTR-5 quarantine). ``chart`` replaces the mapped accounts (CTR-7). Each
    ``zero_priced`` product code gets the template and an observable point of 0.00 USD (CTR-9
    ``TOTAL_SSP_ZERO``). Each ``untemplated`` product code gets no default template and an
    observable point of 2,400.00 USD (DIN-11 ``PRODUCT_UNMAPPED``)."""
    maya_member = member(keyring, clock)
    assign(maya_member, "revenue_accountant")
    maya = holding(app, maya_member, "ssp_analyst")
    approvers: dict[str, Actor] = {}
    for name, roles in (
        ("priya", ("ssp_approver",)),
        ("marcus", ("controller", "ssp_approver", "tenant_admin")),
    ):
        someone = colleague(maya_member.tenant_id, name)
        for code in roles:
            assign(someone, code)
        approvers[name] = enrolled(app, clock, someone)
    marcus = approvers["marcus"]
    enabled = put(app, "/api/v1/tenant-currencies", marcus, {"currency_codes": ["USD"]})
    assert enabled.status_code == 200, enabled.text
    calendar_id = calendar(app, maya, years=(2026, 2027, 2028, 2029))
    created = entity(app, maya, code="AVM-US", calendar_id=calendar_id)
    entity_id = UUID(str(created["id"]))
    open_periods(
        app, maya, entity_code="AVM-US", keys=[f"FY2026-P{month:02d}" for month in range(1, 10)]
    )
    customers = {
        "C-02": customer_id(app, maya, code="C-02", name="Marrowby Health Partners LLC (Demo)"),
        "C-09": customer_id(app, maya, code="C-09", name="Orrin Vale Architects LLP (Demo)"),
    }
    seat = product_with_template(
        app, maya, code=SEAT_MO, name="Platform seat, monthly", revenue_category="SUBSCRIPTION"
    )
    daily = published_template(
        app, maya, marcus, code="TPL-SUB-DAILY", outputs=TPL_SUB_DAILY, case_line=SEAT_MO_CASE
    )
    set_default_template(app, maya, seat, daily["template_id"])
    for code in (*unpriced, *zero_priced):
        extra = product_with_template(
            app, maya, code=code, name=f"Unpriced seat {code}", revenue_category="SUBSCRIPTION"
        )
        set_default_template(app, maya, extra, daily["template_id"])
    for code in untemplated:
        product_with_template(
            app,
            maya,
            code=code,
            name=f"Seat without template {code}",
            revenue_category="SUBSCRIPTION",
        )
    book_id = ssp_book(app, maya)
    version_id = approved_ssp_version(
        app,
        maya,
        [approvers["priya"]],
        book_id,
        label="2026-H1",
        effective_from="2026-01-01",
        entries=[
            range_entry(SEAT_MO, "2160.00", "2400.00", "2640.00", value_basis="AMOUNT"),
            # The zero-priced seats share the series template: D-97 (3a) wants their basis stated.
            *(point_entry(code, "0.00", value_basis="AMOUNT") for code in zero_priced),
            *(point_entry(code, "2400.00") for code in untemplated),
        ],
    )
    mapping_id = published_mapping(app, maya, marcus, chart=chart)
    stamp_test_release()
    return SeatWorld(
        place=workspace(app, clock, keyring, files, maya),
        priya=approvers["priya"],
        marcus=marcus,
        entity_id=entity_id,
        customers=customers,
        book_id=book_id,
        version_id=version_id,
        mapping_id=mapping_id,
    )


def seat_line(
    key: str, *, seats: str, price: str, start: str, end: str, product_code: str = SEAT_MO
) -> dict[str, Any]:
    return {
        "obligation_key": key,
        "product_code": product_code,
        "quantity": seats,
        "total_price": {"amount": price, "currency": "USD"},
        "start_date": start,
        "end_date": end,
    }


def seat_body(
    customer: UUID, *, external_id: str, inception: str, lines: Sequence[Mapping[str, Any]]
) -> dict[str, Any]:
    return {
        "external_id": external_id,
        "customer_id": str(customer),
        "contracting_entity_code": "AVM-US",
        "transaction_currency": "USD",
        "inception_date": inception,
        "document_ref": external_id,
        "lines": [dict(line) for line in lines],
    }


def k02_body(customer: UUID) -> dict[str, Any]:
    """PRD WLD-K-02 ``SF-ORD-10002``: O1 AVM-SEAT-MO, 100 seats × 24 months, 240,000.00 USD from 01
    Jan 2026 to 31 Dec 2027; inception 2026-01-01."""
    line = seat_line("O1", seats="100", price="240000.00", start="2026-01-01", end="2027-12-31")
    return seat_body(customer, external_id=K02_EXTERNAL_ID, inception="2026-01-01", lines=[line])


def step1_criteria(topic: str = "COLLECTIBILITY", **answers: str) -> dict[str, str]:
    """04 T-CON-19 ``questionnaire.criteria`` of a Step 1 review (rev 1.150; supervisor rulings
    R-113 (f), R-115 (f)): a record of topic COLLECTIBILITY or NOT_A_CONTRACT answers the five
    criteria of 606-10-25-1 when it is sent for review. Criteria (a) to (d) are met and (e)
    follows the topic; ``answers`` replace single keys."""
    met = "YES" if topic == "COLLECTIBILITY" else "NO"
    return {"a": "YES", "b": "YES", "c": "YES", "d": "YES", "e": met, **answers}


def k09_body(customer: UUID) -> dict[str, Any]:
    """PRD WLD-K-09 ``SF-ORD-10417``: O1 AVM-SEAT-MO, 30 seats × 36 months, 108,000.00 USD from 01
    Sep 2026 to 31 Aug 2029; inception 2026-09-01."""
    line = seat_line("O1", seats="30", price="108000.00", start="2026-09-01", end="2029-08-31")
    return seat_body(customer, external_id=K09_EXTERNAL_ID, inception="2026-09-01", lines=[line])


# --- imports (BUILD_SPEC DIN-1) -----------------------------------------------------------------

IMPORT_TEMPLATES_PATH: Final = "/api/v1/import-templates"
IMPORTS_PATH: Final = "/api/v1/imports"
IMPORT_ID_HEADER: Final = "X-Erev-Import-Id"
FILES_PATH: Final = "/api/v1/files"
LEGACY_UAT: Final = Path(__file__).resolve().parents[1] / "fixtures" / "legacy_uat"
SKU_SSP_FIXTURE: Final = LEGACY_UAT / "01-ssp-upload" / "SKU SSP Template.xlsx"
_TASK_FETCHED: Final = "UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"


@dataclass(frozen=True, slots=True)
class ImportWorld:
    """A Revenue Accountant (``import.upload``) signed in to a provisioned tenant, and the worker
    runtime that validation jobs run under."""

    app: FastAPI
    actor: Actor
    runtime: JobRuntime
    clock: FrozenClock

    @property
    def tenant_id(self) -> UUID:
        return self.actor.member.tenant_id

    def rows(self, statement: Select[Any]) -> list[dict[str, Any]]:
        context = DbContext(tenant_id=self.tenant_id, user_id=None, entity_scope="*")
        with tenant_session(context, read_only=True) as session:
            return [dict(row) for row in session.execute(statement).mappings()]

    def scalar(self, statement: Select[Any]) -> Any:
        context = DbContext(tenant_id=self.tenant_id, user_id=None, entity_scope="*")
        with tenant_session(context, read_only=True) as session:
            return session.execute(statement).scalar_one()


def import_world(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> ImportWorld:
    # Imported here: support.principals imports this module for tenant_factory.
    from erev_api.jobs.context import JobRuntime
    from support.principals import member, sign_in, workspace
    from support.rows import insert_role_assignment

    someone = member(keyring, clock)
    context = DbContext(tenant_id=someone.tenant_id, user_id=None, entity_scope="*")
    with tenant_session(context) as session:
        insert_role_assignment(
            session,
            tenant_id=someone.tenant_id,
            membership_id=someone.membership_id,
            role_code="revenue_accountant",
        )
    actor = workspace(app, someone, sign_in(app, someone.email))
    # 05 REL-03 (rev 1.15; D-98 60): the IMPORT_COMMIT job this world's runtime runs records
    # CTL-001 / CTL-044 evidence, whose producer stamps the process release — an unstamped
    # process fails closed (release-mismatch), so the world stamps as its siblings do
    # (P5-DOCTOR-R1).
    stamp_test_release()
    return ImportWorld(
        app=app,
        actor=actor,
        runtime=JobRuntime(clock=clock, keyring=keyring, files=files),
        clock=clock,
    )


def upload_import_source(world: ImportWorld, name: str, content: bytes) -> str:
    """``POST /files`` with purpose ``IMPORT_SOURCE``; the file id."""
    from support.principals import cookie_headers

    response = call(
        world.app,
        "POST",
        FILES_PATH,
        data={"purpose": "IMPORT_SOURCE"},
        files={"file": (name, content, "application/octet-stream")},
        headers=cookie_headers(world.actor.token, world.actor.csrf_token),
    )
    assert response.status_code == 201, response.text
    return str(response.json()["id"])


def create_import(
    world: ImportWorld,
    file_id: str,
    template_code: str,
    parameters: Mapping[str, Any] | None = None,
) -> HttpResponse:
    """``POST /imports``."""
    from support.reference import post

    body: dict[str, Any] = {"file_id": file_id, "template_code": template_code}
    if parameters is not None:
        body["parameters"] = dict(parameters)
    return post(world.app, IMPORTS_PATH, world.actor, body)


def run_import_job(world: ImportWorld, job_id: UUID) -> None:
    """The worker fetches the job's task and runs it."""
    from erev_api.db.tables import job
    from erev_api.jobs.registry import run_job
    from sqlalchemy import select, text

    context = DbContext(tenant_id=world.tenant_id, user_id=None, entity_scope="*")
    with tenant_session(context) as session:
        task_id = session.execute(
            select(job.c.procrastinate_job_id).where(job.c.id == job_id)
        ).scalar_one()
        session.execute(text(_TASK_FETCHED), {"id": task_id})
    run_job(job_id, world.tenant_id, attempt=1, runtime=world.runtime)


def imported(
    world: ImportWorld,
    name: str,
    content: bytes,
    template_code: str,
    parameters: Mapping[str, Any] | None = None,
) -> tuple[str, dict[str, Any]]:
    """Upload a file, create its import, run the validation; the import id and API-S-Import."""
    from support.reference import get

    file_id = upload_import_source(world, name, content)
    created = create_import(world, file_id, template_code, parameters)
    assert created.status_code == 202, created.text
    import_id = created.headers[IMPORT_ID_HEADER]
    run_import_job(world, UUID(str(created.json()["id"])))
    shown = get(world.app, f"{IMPORTS_PATH}/{import_id}", world.actor)
    assert shown.status_code == 200, shown.text
    return import_id, dict(shown.json())


def workbook_bytes(title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> bytes:
    """A one-sheet XLSX built with openpyxl; text starting with ``=`` is stored as a formula
    without a cached value."""
    import io

    import openpyxl

    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet(title)
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
