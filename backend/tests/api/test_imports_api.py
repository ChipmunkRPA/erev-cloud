"""API-R-43 imports (04 §15.3 API-R-43, §15.2 ``duplicate-import`` and ``upload-type-not-allowed``,
§16.6 error report and import commands; 05 IPL-01, UPL-02, UPL-10; PRD SM-05, IMP-05, ERR-37; 03
REQ-DAT-001, REQ-DAT-016; control CTL-001; BUILD_SPEC DIN-1).

World: ``support.factories.import_world`` (a Revenue Accountant of a provisioned tenant).
"""

from __future__ import annotations

import csv
import io
from datetime import date
from uuid import UUID, uuid4

import openpyxl
import pytest
from erev_api.auth.keyring import KeyRing
from erev_api.clock import FrozenClock
from erev_api.config import Settings
from erev_api.db.session import DbContext, tenant_session
from erev_api.db.tables import audit_event, file_object, import_row, import_upload, job
from erev_api.db.transitions import apply
from erev_api.domain.imports import upload
from erev_api.files import policy
from erev_api.files.store import LocalFileStore
from erev_api.main import create_app
from fastapi import FastAPI
from sqlalchemy import func, insert, select
from support import upload_fixtures
from support.db import TestDatabase
from support.factories import (
    IMPORT_TEMPLATES_PATH,
    IMPORTS_PATH,
    SKU_SSP_FIXTURE,
    ImportWorld,
    create_import,
    import_world,
    imported,
    run_import_job,
    upload_import_source,
    workbook_bytes,
)
from support.http import call
from support.principals import carrying, colleague, cookie_headers, sign_in, workspace
from support.reference import get, post, slug
from support.rows import import_row_values, insert_role_assignment

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PROGRESS_HEADERS = (
    "Contract Unique Name",
    "POB Unique ID",
    "SKU Name",
    "Current Delivery",
    "Current Billing",
    "Current Pre-ASC606 Revenue (Net Design Only)",
    "Memo 1",
    "Memo 2",
    "Memo 3",
)
SKU_SSP_HEADERS = (
    "SKU Unique ID",
    "SKU Name",
    "Distinct or Nondistinct",
    "SKU Unit List Price",
    "ASC 606 Stratification",
    "Midpoint Discount Percentage",
    "SSP Range Method (+-)",
    "SSP Version",
    "Revenue Account",
)
# E-40 along the pairs of PRD SM-05 up to the commit that DIN-3 builds.
COMMIT_PATH = (
    ("VALIDATED", "DIFFING"),
    ("DIFFING", "DIFF_READY"),
    ("DIFF_READY", "SUBMITTED"),
    ("SUBMITTED", "APPROVED"),
    ("APPROVED", "COMMITTING"),
    ("COMMITTING", "COMMITTED"),
)


@pytest.fixture
def app(committed_db: TestDatabase, app_settings: Settings, clock: FrozenClock) -> FastAPI:
    return create_app(app_settings, clock=clock)


@pytest.fixture
def world(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, app_settings: Settings
) -> ImportWorld:
    return import_world(app, keyring, clock, LocalFileStore(app_settings.file_root))


def test_rows_sort_severity(world: ImportWorld) -> None:
    """D-87 L6-4-Q-14: ``sort=severity`` lists ERROR, WARNING, VALID, AGGREGATED, then BLANK rows,
    rows of one status by row number, across keyset pages, and takes no ``-`` prefix."""
    content = workbook_bytes(
        "SKU Setup",
        SKU_SSP_HEADERS,
        [
            (1, "Hardware 1", "Distinct", "abc", "Hardware 1", None, 0.15, "2023-01-01", 5001),
            (2, None, "Distinct", 200, "Software 1", "x", 0.15, "2023-01-01", 5002),
            (3, "Consulting 1", "Nondistinct", 300, "Consulting 1", 0.5, 0, "2023-01-01", 5003),
        ],
    )
    import_id, shown = imported(world, "severity.xlsx", content, "legacy_sku_ssp")
    assert shown["status"] == "INVALID"
    upload_id = UUID(import_id)
    # Rows 5 to 9 of the other statuses; T-IMP-03 rows are insert-only, so the test stores them.
    context = DbContext(tenant_id=world.tenant_id, user_id=None, entity_scope="*")
    with tenant_session(context) as session:
        lead = session.execute(
            select(import_row.c.id).where(
                import_row.c.import_upload_id == upload_id, import_row.c.row_number == 4
            )
        ).scalar_one()
        for number, status in (
            (5, "BLANK"),
            (6, "AGGREGATED"),
            (7, "WARNING"),
            (8, "VALID"),
            (9, "ERROR"),
        ):
            session.execute(
                insert(import_row).values(
                    import_row_values(
                        world.tenant_id,
                        import_upload_id=upload_id,
                        row_number=number,
                        status=status,
                        normalized=None if status in ("BLANK", "ERROR") else {"SKU Name": "X"},
                        aggregated_into_row_id=lead if status == "AGGREGATED" else None,
                    )
                )
            )
    path = f"{IMPORTS_PATH}/{import_id}/rows"
    numbers: list[int] = []
    cursor: str | None = None
    pages = 0
    while True:
        query: dict[str, str | int] = {"sort": "severity", "limit": 3}
        if cursor is not None:
            query["cursor"] = cursor
        page = get(world.app, path, world.actor, query)
        assert page.status_code == 200, page.text
        numbers += [int(item["row_number"]) for item in page.json()["items"]]
        pages += 1
        cursor = page.json()["next_cursor"]
        if cursor is None:
            break
    assert (numbers, pages) == ([2, 3, 9, 7, 4, 8, 6, 5], 3)
    errors = get(world.app, path, world.actor, {"sort": "severity", "status": "ERROR"})
    assert [item["row_number"] for item in errors.json()["items"]] == [2, 3, 9]
    default = get(world.app, path, world.actor)
    assert [item["row_number"] for item in default.json()["items"]] == [2, 3, 4, 5, 6, 7, 8, 9]
    refused = get(world.app, path, world.actor, {"sort": "-severity"})
    assert (refused.status_code, slug(refused)) == (422, "validation-failed"), refused.text


def _committed(world: ImportWorld, import_id: UUID) -> None:
    """Walk a validated import to COMMITTED along DB-03 as SYSTEM (the commit is DIN-3)."""
    context = DbContext(tenant_id=world.tenant_id, user_id=None, entity_scope="*")
    stamps = {"updated_at": world.clock.now(), "updated_by": None, "updated_by_kind": "SYSTEM"}
    with tenant_session(context) as session:
        for source, target in COMMIT_PATH:
            extra = {"committed_at": world.clock.now()} if target == "COMMITTED" else {}
            apply(
                session,
                "import_upload",
                import_id,
                to_status=target,
                set_values={**stamps, **extra},
                expected_status=source,
            )


def test_template_download_example_rows_and_definitions(world: ImportWorld) -> None:
    path = f"{IMPORT_TEMPLATES_PATH}/legacy_progress_tracking/download"
    response = get(world.app, path, world.actor)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == (
        "attachment; filename*=UTF-8''legacy_progress_tracking.xlsx"
    )
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Contract Progress Tracking", "Definitions"]
    rows = list(book.worksheets[0].iter_rows(values_only=True))
    assert rows == [
        PROGRESS_HEADERS,
        (
            "Contract 1",
            "POB #1",
            "Hardware 1",
            "5",
            "500",
            "0",
            "Delivered",
            "Invoice 1001",
            "Region A",
        ),
    ]
    definitions = list(book["Definitions"].iter_rows(values_only=True))
    assert definitions[0] == ("name", "type", "required", "rule_ids")
    assert [row[0] for row in definitions[1:]] == list(PROGRESS_HEADERS)
    assert definitions[1] == ("Contract Unique Name", "identifier", True, "REQUIRED_VALUE_BLANK")
    assert definitions[4] == (
        "Current Delivery",
        "quantity",
        True,
        "REQUIRED_VALUE_BLANK, VALUE_NOT_NUMERIC",
    )
    assert definitions[7] == ("Memo 1", "text", False, "PROGRESS_MEMO_BLANK")
    missing = get(world.app, f"{IMPORT_TEMPLATES_PATH}/legacy_unknown/download", world.actor)
    assert (missing.status_code, slug(missing)) == (404, "not-found")


def test_upload_type_not_allowed(world: ImportWorld) -> None:
    response = call(
        world.app,
        "POST",
        "/api/v1/files",
        data={"purpose": "IMPORT_SOURCE"},
        files={"file": ("contract.pdf", upload_fixtures.PDF, "application/pdf")},
        headers=cookie_headers(world.actor.token, world.actor.csrf_token),
    )
    assert (response.status_code, slug(response)) == (422, "upload-type-not-allowed")
    assert response.json()["detail"] == policy.IMPORT_DETAIL
    assert world.scalar(select(func.count()).select_from(file_object)) == 0


def test_error_report_csv(world: ImportWorld) -> None:
    content = workbook_bytes(
        "SKU Setup",
        SKU_SSP_HEADERS,
        [
            (1, "Hardware 1", "Distinct", "abc", "Hardware 1", None, 0.15, "2023-01-01", 5001),
            (2, None, "Distinct", 200, "Software 1", "x", 0.15, "2023-01-01", 5002),
            (3, "Consulting 1", "Nondistinct", 300, "Consulting 1", 0.5, 0, "2023-01-01", 5003),
        ],
    )
    import_id, shown = imported(world, "errors.xlsx", content, "legacy_sku_ssp")
    assert shown["status"] == "INVALID"
    assert shown["counts"] == {
        "rows": 3,
        "valid": 1,
        "warnings": 0,
        "errors": 2,
        "aggregated": 0,
        "blank": 0,
    }
    report = get(world.app, f"{IMPORTS_PATH}/{import_id}/error-report", world.actor)
    assert report.status_code == 200, report.text
    assert report.headers["content-type"].startswith("text/csv")
    assert report.content.splitlines()[0] == b"row,column,rule_id,message"
    # BUILD_SPEC DIN-7: legacy row messages name worksheet, row, column, key and rule id.
    assert list(csv.reader(io.StringIO(report.text))) == [
        ["row", "column", "rule_id", "message"],
        [
            "2",
            "SKU Unit List Price",
            "VALUE_NOT_NUMERIC",
            "SKU Setup row 2, column SKU Unit List Price: SKU Unit List Price must be a number. "
            'Found "abc". Key Hardware 1 / Hardware 1 / 2023-01-01. (VALUE_NOT_NUMERIC)',
        ],
        [
            "2",
            "Midpoint Discount Percentage",
            "REQUIRED_VALUE_BLANK",
            "SKU Setup row 2, column Midpoint Discount Percentage: Midpoint Discount Percentage is "
            "required. Key Hardware 1 / Hardware 1 / 2023-01-01. (REQUIRED_VALUE_BLANK)",
        ],
        [
            "3",
            "SKU Name",
            "REQUIRED_VALUE_BLANK",
            "SKU Setup row 3, column SKU Name: SKU Name is required. Key Software 1 / 2023-01-01. "
            "(REQUIRED_VALUE_BLANK)",
        ],
        [
            "3",
            "Midpoint Discount Percentage",
            "VALUE_NOT_NUMERIC",
            "SKU Setup row 3, column Midpoint Discount Percentage: Midpoint Discount Percentage "
            'must be a number. Found "x". Key Software 1 / 2023-01-01. (VALUE_NOT_NUMERIC)',
        ],
    ]


def test_cancel_before_approval(world: ImportWorld) -> None:
    content = SKU_SSP_FIXTURE.read_bytes()
    first_id, shown = imported(world, "SKU SSP Template.xlsx", content, "legacy_sku_ssp")
    assert shown["status"] == "VALIDATED"
    cancelled = post(world.app, f"{IMPORTS_PATH}/{first_id}/cancel", world.actor, {})
    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()["status"] == "CANCELLED"
    audited = world.rows(
        select(audit_event.c.action).where(
            audit_event.c.object_id == UUID(first_id),
            audit_event.c.action == "import_upload.cancel",
        )
    )
    assert audited == [{"action": "import_upload.cancel"}]

    second_id, shown = imported(world, "SKU SSP Template.xlsx", content, "legacy_sku_ssp")
    assert shown["status"] == "VALIDATED"
    _committed(world, UUID(second_id))
    refused = post(world.app, f"{IMPORTS_PATH}/{second_id}/cancel", world.actor, {})
    assert (refused.status_code, slug(refused)) == (409, "invalid-transition")
    assert get(world.app, f"{IMPORTS_PATH}/{second_id}", world.actor).json()["status"] == (
        "COMMITTED"
    )


@pytest.mark.control("CTL-001")
def test_ctl_001_duplicate_upload_refused_before_validation(world: ImportWorld) -> None:
    """CTL-001, the upload's part of "duplicates rejected before commit": a second import of
    the same file content and kind is refused 409 ``duplicate-import`` before anything of it is
    stored or validated — no second upload row, no staged row, no second validation job — while
    the first is still to be validated and after it was; it is accepted again once the first is
    cancelled. Not witnessed here: a committed file refused again
    (``tests/domain/imports/test_commit.py``), a duplicate source version of an integration
    (``tests/domain/integrations/test_normalise.py``), and that a command's retry is safe (the
    idempotency store, REQ-PLT-026)."""
    content = SKU_SSP_FIXTURE.read_bytes()
    file_id = upload_import_source(world, "SKU SSP Template.xlsx", content)
    first = create_import(world, file_id, "legacy_sku_ssp")
    assert first.status_code == 202, first.text
    first_id = first.headers["X-Erev-Import-Id"]
    shown = get(world.app, f"{IMPORTS_PATH}/{first_id}", world.actor).json()
    expected = upload.duplicate_message(
        shown["import_no"], date.fromisoformat(shown["created_at"][:10])
    )

    copy_id = upload_import_source(world, "SKU SSP Template (copy).xlsx", content)
    assert copy_id == file_id  # identical content of one purpose is one file (DG-KRN-FILE-01)
    refused = create_import(world, copy_id, "legacy_sku_ssp")
    assert (refused.status_code, slug(refused)) == (409, "duplicate-import")
    body = refused.json()
    assert body["detail"] == expected
    assert [(error["rule_id"], error["message"]) for error in body["errors"]] == [
        ("IMPORT_FILE_DUPLICATE", expected)
    ]
    assert world.scalar(select(func.count()).select_from(import_upload)) == 1
    assert world.scalar(select(func.count()).select_from(import_row)) == 0
    validations = select(func.count()).where(job.c.kind == "IMPORT_VALIDATE")
    assert world.scalar(validations) == 1

    run_import_job(world, UUID(first.json()["id"]))
    again = create_import(world, copy_id, "legacy_sku_ssp")
    assert (again.status_code, slug(again)) == (409, "duplicate-import")
    assert world.scalar(select(func.count()).select_from(import_upload)) == 1
    assert world.scalar(select(func.count()).select_from(import_row)) == 7

    cancelled = post(world.app, f"{IMPORTS_PATH}/{first_id}/cancel", world.actor, {})
    assert cancelled.status_code == 200, cancelled.text
    accepted = create_import(world, copy_id, "legacy_sku_ssp")
    assert accepted.status_code == 202, accepted.text


def test_req_plt_012_an_import_names_only_a_source_its_creator_may_read(
    world: ImportWorld, app: FastAPI
) -> None:
    """Security review S2, its family (04 T-PLT-29 Read access, rev 1.151). ``POST
    /imports`` checked the purpose and the media type of the file it was given and nothing else,
    so a holder of ``import.upload`` named another member's upload; the import was then its own,
    and it read the file's rows back through ``GET /imports/{id}/rows``. The command asks the
    file-read question first and answers a source its caller may not read as a missing one."""
    someone = colleague(world.tenant_id, "ingrid")  # a second Revenue Accountant: import.upload
    context = DbContext(tenant_id=world.tenant_id, user_id=None, entity_scope="*")
    with tenant_session(context) as session:
        insert_role_assignment(
            session,
            tenant_id=world.tenant_id,
            membership_id=someone.membership_id,
            role_code="revenue_accountant",
        )
    ingrid = carrying(world, actor=workspace(app, someone, sign_in(app, someone.email)))
    file_id = upload_import_source(ingrid, "SKU SSP Template.xlsx", SKU_SSP_FIXTURE.read_bytes())

    # Ingrid's upload, which nothing owns yet, named by another holder of import.upload.
    refused = create_import(world, file_id, "legacy_sku_ssp")
    unknown = create_import(world, str(uuid4()), "legacy_sku_ssp")
    assert refused.status_code == 422, refused.text
    assert slug(refused) == "validation-failed"
    assert [(error["field"], error["rule_id"]) for error in refused.json()["errors"]] == [
        ("file_id", "API-R-43")
    ]
    assert refused.json()["errors"] == unknown.json()["errors"]
    assert world.scalar(select(func.count()).select_from(import_upload)) == 0

    # Positive control: the uploader's own source is accepted.
    created = create_import(ingrid, file_id, "legacy_sku_ssp")
    assert created.status_code == 202, created.text
    # ... and once an import owns the source, the readers of imports read it (T-PLT-29: here an
    # all-entities holder, the upload's entities not being resolved yet): the same request is
    # then the CTL-001 duplicate, not a missing file.
    duplicate = create_import(world, file_id, "legacy_sku_ssp")
    assert (duplicate.status_code, slug(duplicate)) == (409, "duplicate-import"), duplicate.text
