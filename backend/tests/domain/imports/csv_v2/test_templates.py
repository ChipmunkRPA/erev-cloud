"""DIN-9 modern CSV v2 templates (04 T-IMP-01 ``CSV_V2``, NC-19, §16.1 to §16.5, §16.3; PRD
WLD-F-21, WLD-F-27, J-04.1, J-12.1, J-12-ALT-1, CPY-06, IMP-20; SCREENS §12.2 sample world; 03
REQ-DAT-004, REQ-REF-006, REQ-CST-001, REQ-DAT-012; BUILD_SPEC DIN-9, BS3-D-07).

Worlds: ``support.factories.k11_world`` (K-11 ``NS-SO-DE-5004`` active), ``seat_world`` (K-09
``SF-ORD-10417`` active), ``j03_world`` (US-LIST, AVM-MAP-2026-01) and ``import_world`` with a
controller for FX rates. Maya uploads, Priya (Revenue Reviewer, MFA) approves imports, and the jobs
run as the worker runs them.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from decimal import Decimal
from typing import Any
from uuid import UUID

import openpyxl
import pytest
from erev_api.auth.keyring import KeyRing
from erev_api.clock import FrozenClock
from erev_api.config import Settings
from erev_api.db.tables import (
    account_mapping_rule,
    account_mapping_version,
    contract,
    contract_event,
    estimate,
    estimate_version,
    exception_item,
    fx_rate_set_version,
    gl_account,
    import_row,
    import_row_lineage,
    product,
    product_bundle_component,
    ssp_book_version,
    ssp_entry,
)
from erev_api.domain.imports.csv_v2 import TEMPLATES
from erev_api.domain.imports.csv_v2.account_mapping import AccountMappingRowsIn
from erev_api.domain.imports.csv_v2.bundles import BundleIn
from erev_api.domain.imports.csv_v2.cost_events import CostEventIn
from erev_api.domain.imports.csv_v2.estimates import EstimateParametersRowIn, EstimatesIn
from erev_api.domain.imports.csv_v2.framework import flatten
from erev_api.domain.imports.csv_v2.fx_rates import FxRatesIn
from erev_api.domain.imports.csv_v2.invoices import InvoiceIn
from erev_api.domain.imports.csv_v2.pre_standard_revenue import PreStandardRevenueIn
from erev_api.domain.imports.csv_v2.progress_events import ProgressEventIn
from erev_api.domain.imports.csv_v2.ssp_values import SspValuesIn
from erev_api.domain.imports.csv_v2.usage import UsageIn
from erev_api.files.store import LocalFileStore
from erev_api.jobs.context import JobRuntime
from erev_api.main import create_app
from erev_api.schemas.accounts import GlAccountIn
from erev_api.schemas.contracts import ContractCreateIn
from erev_api.schemas.customers import CustomerIn
from erev_api.schemas.db_json import ESTIMATE_PARAMETERS
from erev_api.schemas.products import ProductIn
from fastapi import FastAPI
from pydantic import BaseModel
from sqlalchemy import select
from support.db import TestDatabase
from support.factories import (
    IMPORT_TEMPLATES_PATH,
    K09_EXTERNAL_ID,
    K11_EXTERNAL_ID,
    ImportWorld,
    Workspace,
    activated_contract,
    booked_contract,
    create_import,
    import_world,
    imported,
    j03_world,
    k09_body,
    k11_body,
    k11_world,
    run_import_job,
    seat_world,
    upload_import_source,
)
from support.legacy_replay import diffed, job_of, shown, submit
from support.principals import Actor, colleague, enrolled
from support.reference import approve, assign, calendar, get, post, put

# NC-19: the API command request of each CSV v2 template with an emitter (L5-1-Q-19 to L5-1-Q-22).
COMMANDS: dict[str, type[BaseModel]] = {
    "customers": CustomerIn,
    "products": ProductIn,
    "bundles": BundleIn,
    "ssp_values": SspValuesIn,
    "contracts": ContractCreateIn,
    "invoices": InvoiceIn,
    "progress_events": ProgressEventIn,
    "usage": UsageIn,
    "fx_rates": FxRatesIn,
    "cost_events": CostEventIn,
    "pre_standard_revenue": PreStandardRevenueIn,
    "gl_accounts": GlAccountIn,
    "account_mapping": AccountMappingRowsIn,
    "estimates": EstimatesIn,
}
# L5-1-Q-23, D-86: modifications wait for CTR-17, which R-RC-1 moves post-rc.
WITHOUT_EMITTER = ("modifications",)


def csv_bytes(headers: Sequence[str], rows: Sequence[Sequence[str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def importer(place: Workspace) -> ImportWorld:
    return ImportWorld(
        app=place.app,
        actor=place.author,
        runtime=JobRuntime(clock=place.clock, keyring=place.keyring, files=place.files),
        clock=place.clock,
    )


def committed(
    imports: ImportWorld, approver: Actor, name: str, content: bytes, template_code: str
) -> dict[str, Any]:
    import_id = diffed(imports, name, content, template_code)
    submitted = submit(imports, import_id)
    assert submitted.status_code == 200, submitted.text
    decided = approve(imports.app, str(submitted.json()["approval_request_id"]), approver)
    assert decided.status_code == 200, decided.text
    run_import_job(imports, job_of(imports, UUID(import_id), "IMPORT_COMMIT"))
    done = shown(imports, import_id)
    assert done["status"] == "COMMITTED", done
    return done


@pytest.fixture
def app(committed_db: TestDatabase, app_settings: Settings, clock: FrozenClock) -> FastAPI:
    return create_app(app_settings, clock=clock)


@pytest.fixture
def files(app_settings: Settings) -> LocalFileStore:
    return LocalFileStore(app_settings.file_root)


def test_csv_headers_equal_command_fields(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    world = import_world(app, keyring, clock, files)
    response = get(app, IMPORT_TEMPLATES_PATH, world.actor)
    assert response.status_code == 200, response.text
    csv_items = {
        item["code"]: item for item in response.json()["items"] if item["family"] == "CSV_V2"
    }
    assert len(csv_items) == 15
    assert sorted(TEMPLATES) == sorted(COMMANDS)
    for code, command in COMMANDS.items():
        expected = [column.name for column in flatten(command)]
        assert [header["name"] for header in csv_items[code]["headers"]] == expected, code
        required = {column.name for column in flatten(command) if column.required}
        assert {
            header["name"] for header in csv_items[code]["headers"] if header["required"]
        } == required
    # Arrays of lines are one row per line with the header fields repeated.
    assert "lines.obligation_key" in [h["name"] for h in csv_items["contracts"]["headers"]]
    assert "lines.tax_lines.tax_type" in [h["name"] for h in csv_items["invoices"]["headers"]]
    for code in WITHOUT_EMITTER:
        assert csv_items[code]["headers"] == [], code
        file_id = upload_import_source(world, f"{code}.csv", b"external_id\nX\n")
        refused = create_import(world, file_id, code)
        assert refused.status_code == 422, refused.text
    download = get(app, f"{IMPORT_TEMPLATES_PATH}/progress_events/download", world.actor)
    assert download.status_code == 200, download.text
    book = openpyxl.load_workbook(io.BytesIO(download.content), read_only=True)
    first = next(book.worksheets[0].iter_rows(values_only=True))
    book.close()
    assert list(first) == [column.name for column in flatten(ProgressEventIn)]


def test_progress_events_k11(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    k11 = k11_world(app, keyring, clock, files)
    booked = booked_contract(k11.place, k11_body(k11.customer_id), activate=False)
    activated_contract(k11.place, booked)
    imports = importer(k11.place)
    headers = [column.name for column in flatten(ProgressEventIn)]
    row = dict.fromkeys(headers, "")
    row |= {
        "contract": K11_EXTERNAL_ID,
        "event_type": "DELIVERY_RECORDED",
        "effective_date": "2026-09-12",
        "obligation_key": "O1",
        "quantity": "120",
        "trigger": "DELIVERY",
    }
    import_id = diffed(
        imports,
        "avm-de-progress-2026-09.csv",
        csv_bytes(headers, [list(row.values())]),
        "progress_events",
    )
    ready = shown(imports, import_id)
    # SCREENS §12.2 J-04.1: revenue 53,504.59 on 12 Sep 2026.
    assert ready["diff_summary"]["revenue_by_period_delta"] == [
        {"period_key": "FY2026-P09", "amount": "53504.59"}
    ]
    assert ready["diff_summary"]["contracts_affected"] == 1
    listed = get(app, f"/api/v1/imports/{import_id}/diff", imports.actor)
    assert listed.status_code == 200, listed.text
    revenue = [item for item in listed.json()["items"] if item["measure"] == "revenue_cum"]
    assert [(item["obligation_key"], item["before"], item["after"]) for item in revenue] == [
        ("O1", "0.00", "53504.59")
    ]
    # Nothing is persisted before commit (BR-DAT-05).
    assert (
        imports.rows(
            select(contract_event.c.id).where(contract_event.c.event_type == "DELIVERY_RECORDED")
        )
        == []
    )


def test_progress_events_over_delivery_rejected(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    """PRD J-04-ALT-1 and WLD-B-04 (BUILD_SPEC DIN-15): the stage 01 bounds run at validation, in
    row order over the stored stream. K-11 O1 books 200 units: row 2 delivers 120, so row 3's 90
    breaks the bound and the upload is INVALID with one exception item; nothing is diffed."""
    k11 = k11_world(app, keyring, clock, files)
    booked = booked_contract(k11.place, k11_body(k11.customer_id), activate=False)
    activated_contract(k11.place, booked)
    imports = importer(k11.place)
    headers = [column.name for column in flatten(ProgressEventIn)]
    rows = []
    for quantity, day in (("120", "2026-09-12"), ("90", "2026-09-25")):
        row = dict.fromkeys(headers, "")
        row |= {
            "contract": K11_EXTERNAL_ID,
            "event_type": "DELIVERY_RECORDED",
            "effective_date": day,
            "obligation_key": "O1",
            "quantity": quantity,
            "trigger": "DELIVERY",
        }
        rows.append(list(row.values()))
    import_id, validated = imported(
        imports,
        "avm-de-progress-2026-09-overdelivery.csv",
        csv_bytes(headers, rows),
        "progress_events",
    )
    assert validated["status"] == "INVALID", validated
    assert (validated["counts"]["rows"], validated["counts"]["valid"]) == (2, 1)
    assert validated["counts"]["errors"] == 1
    assert validated["finding_counts"] == [
        {"code": "PROGRESS_OVER_DELIVERY", "severity": "ERROR", "rows": 1}
    ]
    listed = get(app, f"/api/v1/imports/{import_id}/rows", imports.actor)
    assert listed.status_code == 200, listed.text
    statuses = {item["row_number"]: item["status"] for item in listed.json()["items"]}
    assert statuses == {2: "VALID", 3: "ERROR"}
    [invalid] = [item for item in listed.json()["items"] if item["status"] == "ERROR"]
    assert [message["message"] for message in invalid["messages"]] == [
        "Row 3, column Quantity: NS-SO-DE-5004, obligation O1 (AVM-GW): requested 90, "
        "remaining 80. (PROGRESS_OVER_DELIVERY)"
    ]
    items = imports.rows(
        select(
            exception_item.c.code,
            exception_item.c.source,
            exception_item.c.severity,
            exception_item.c.status,
        ).where(exception_item.c.import_upload_id == UUID(import_id))
    )
    assert [
        (str(item["code"]), str(item["source"]), str(item["severity"]), str(item["status"]))
        for item in items
    ] == [("PROGRESS_OVER_DELIVERY", "IMPORT", "BLOCKING", "OPEN")]
    assert imports.rows(select(import_row.c.id).where(import_row.c.status == "AGGREGATED")) == []
    assert (
        imports.rows(
            select(contract_event.c.id).where(contract_event.c.event_type == "DELIVERY_RECORDED")
        )
        == []
    )


def test_usage_row_whose_period_ends_after_its_date_is_a_finding(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    """04 §16.3 "The usage period of a report" and table 15.4-B ``USAGE_PERIOD_NOT_ENDED`` (rev
    1.320; PRD IMP-148; item USAGE-REPORT-PERIOD-ENDED-1): a usage report states usage that has
    occurred. Row 2 reports 15 to 30 September on 30 September and is usable; row 3 reports the
    same period on 24 September — a period that has not ended — and is a finding of the row at
    ``usage_period_end`` when the file is validated: the upload is INVALID with one exception
    item, and nothing is diffed or appended.

    Before the rule both rows were valid and the commit appended both reports."""
    k11 = k11_world(app, keyring, clock, files)
    booked = booked_contract(k11.place, k11_body(k11.customer_id), activate=False)
    activated_contract(k11.place, booked)
    imports = importer(k11.place)
    headers = [column.name for column in flatten(UsageIn)]
    rows = []
    for day in ("2026-09-30", "2026-09-24"):
        row = dict.fromkeys(headers, "")
        row |= {
            "contract": K11_EXTERNAL_ID,
            "effective_date": day,
            "obligation_key": "O2",
            "usage_period_start": "2026-09-15",
            "usage_period_end": "2026-09-30",
            "metric": "SUPPORT_TICKETS",
            "quantity": "14",
        }
        rows.append(list(row.values()))
    import_id, validated = imported(
        imports, "avm-de-usage-2026-09-period.csv", csv_bytes(headers, rows), "usage"
    )
    # the status first: without the rule the upload is VALID and its commit appends both rows
    assert validated["status"] == "INVALID", validated
    assert (validated["counts"]["rows"], validated["counts"]["valid"]) == (2, 1)
    assert validated["counts"]["errors"] == 1
    assert validated["finding_counts"] == [
        {"code": "USAGE_PERIOD_NOT_ENDED", "severity": "ERROR", "rows": 1}
    ]
    listed = get(app, f"/api/v1/imports/{import_id}/rows", imports.actor)
    assert listed.status_code == 200, listed.text
    statuses = {item["row_number"]: item["status"] for item in listed.json()["items"]}
    assert statuses == {2: "VALID", 3: "ERROR"}
    [invalid] = [item for item in listed.json()["items"] if item["status"] == "ERROR"]
    assert [message["message"] for message in invalid["messages"]] == [
        "Row 3, column Usage period end: The usage period 15 Sep 2026 to 30 Sep 2026 ends after "
        "the report's date, 24 Sep 2026. A usage report states usage that has occurred: end the "
        "period on or before 24 Sep 2026. (USAGE_PERIOD_NOT_ENDED)"
    ]
    items = imports.rows(
        select(
            exception_item.c.code,
            exception_item.c.source,
            exception_item.c.severity,
            exception_item.c.status,
        ).where(exception_item.c.import_upload_id == UUID(import_id))
    )
    assert [
        (str(item["code"]), str(item["source"]), str(item["severity"]), str(item["status"]))
        for item in items
    ] == [("USAGE_PERIOD_NOT_ENDED", "IMPORT", "BLOCKING", "OPEN")]
    assert (
        imports.rows(
            select(contract_event.c.id).where(contract_event.c.event_type == "USAGE_REPORTED")
        )
        == []
    )


def test_contract_not_found_message(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    world = import_world(app, keyring, clock, files)
    headers = [column.name for column in flatten(CostEventIn)]
    row = dict.fromkeys(headers, "")
    row |= {
        "contract": "SF-ORD-99999",
        "effective_date": "2026-09-05",
        "purpose": "COST_TO_OBTAIN",
        "amount.amount": "5760.00",
        "amount.currency": "USD",
        "plan_code": "SALES-2026",
    }
    import_id, validated = imported(
        world,
        "avm-us-commissions-2026-09.csv",
        csv_bytes(headers, [list(row.values())]),
        "cost_events",
    )
    assert validated["status"] == "INVALID", validated
    (item,) = world.rows(
        select(exception_item.c.code, exception_item.c.message, import_row.c.row_number)
        .join(import_row, import_row.c.id == exception_item.c.import_row_id)
        .where(exception_item.c.import_upload_id == UUID(import_id))
    )
    assert (item["code"], item["row_number"]) == ("CONTRACT_NOT_FOUND", 2)
    assert item["message"] == (
        "Row 2, column Contract: Contract SF-ORD-99999 does not exist in this workspace. "
        "(CONTRACT_NOT_FOUND)"
    )


def test_cost_events_csv(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    seats = seat_world(app, keyring, clock, files)
    booked = booked_contract(seats.place, k09_body(seats.customers["C-09"]), activate=False)
    activated_contract(seats.place, booked)
    assign(seats.priya.member, "revenue_reviewer")
    headers = [column.name for column in flatten(CostEventIn)]
    row = dict.fromkeys(headers, "")
    row |= {
        "contract": K09_EXTERNAL_ID,
        "effective_date": "2026-09-05",
        "purpose": "COST_TO_OBTAIN",
        "amount.amount": "6480.00",
        "amount.currency": "USD",
        "payee": "Account executive",
        "plan_code": "SALES-2026",
        "is_incremental": "true",
        "has_clawback": "false",
    }
    done = committed(
        importer(seats.place),
        seats.priya,
        "avm-us-commissions-2026-09.csv",
        csv_bytes(headers, [list(row.values())]),
        "cost_events",
    )
    (event,) = importer(seats.place).rows(
        select(
            contract_event.c.event_type, contract_event.c.origin, contract_event.c.payload
        ).where(contract_event.c.import_upload_id == UUID(done["id"]))
    )
    assert (str(event["event_type"]), str(event["origin"])) == ("COST_INCURRED", "IMPORT")
    payload = event["payload"]
    assert (
        payload["purpose"],
        payload["amount"],
        payload["plan_code"],
        payload["is_incremental"],
    ) == (
        "COST_TO_OBTAIN",
        {"amount": "6480.00", "currency": "USD"},
        "SALES-2026",
        True,
    )


def test_fx_rates_csv_creates_version_for_approval(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    world = import_world(app, keyring, clock, files)
    calendar(app, world.actor)
    people: dict[str, Actor] = {}
    for name, roles in (
        ("priya", ("revenue_reviewer",)),
        ("carmen", ("controller", "tenant_admin")),
    ):
        someone = colleague(world.tenant_id, name)
        for code in roles:
            assign(someone, code)
        people[name] = enrolled(app, clock, someone)
    enabled = put(
        app, "/api/v1/tenant-currencies", people["carmen"], {"currency_codes": ["USD", "EUR"]}
    )
    assert enabled.status_code == 200, enabled.text
    created = post(
        app,
        "/api/v1/fx-rate-sets",
        world.actor,
        {"code": "AVM-RATES-SPOT", "name": "Spot rates", "rate_type": "spot"},
    )
    assert created.status_code == 201, created.text
    headers = [column.name for column in flatten(FxRatesIn)]
    row = dict.fromkeys(headers, "")
    row |= {
        "fx_rate_set_code": "AVM-RATES-SPOT",
        "coverage_from": "2026-09-01",
        "coverage_to": "2026-09-30",
        "lines.base_currency": "EUR",
        "lines.quote_currency": "USD",
        "lines.rate": "1.105",
        "lines.effective_date": "2026-09-12",
    }
    done = committed(
        world,
        people["priya"],
        "avm-rates-2026-09.csv",
        csv_bytes(headers, [list(row.values())]),
        "fx_rates",
    )
    (version,) = world.rows(
        select(
            fx_rate_set_version.c.id,
            fx_rate_set_version.c.status,
            fx_rate_set_version.c.import_upload_id,
            fx_rate_set_version.c.approval_request_id,
        ).where(fx_rate_set_version.c.import_upload_id == UUID(done["id"]))
    )
    assert str(version["status"]) == "SUBMITTED"
    query = {"rate_type": "spot", "base": "EUR", "quote": "USD", "date": "2026-09-12"}
    before = get(app, "/api/v1/fx-rates", world.actor, query)
    assert (before.status_code, before.json()["items"]) == (200, []), before.text
    shown_version = get(app, f"/api/v1/fx-rate-set-versions/{version['id']}", world.actor).json()
    decided = approve(app, str(shown_version["pending_approval_request_id"]), people["carmen"])
    assert decided.status_code == 200, decided.text
    after = get(app, "/api/v1/fx-rates", world.actor, query)
    assert after.status_code == 200, after.text
    assert [
        (item["base_currency"], item["quote_currency"], Decimal(item["rate"]))
        for item in after.json()["items"]
    ] == [("EUR", "USD", Decimal("1.105"))]


def test_master_data_and_configuration_csv(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    j03 = j03_world(app, keyring, clock, files)
    assign(j03.priya.member, "revenue_reviewer")
    imports = importer(j03.place)

    product_headers = [column.name for column in flatten(ProductIn)]
    products_rows = []
    for code, name, bundle in (
        ("AVM-KIT-10", "Sensor kit, 10 units", "false"),
        ("AVM-KIT-11", "Sensor mount", "false"),
        ("AVM-BUNDLE-01", "Kit and mount bundle", "true"),
    ):
        values = dict.fromkeys(product_headers, "")
        values |= {"code": code, "name": name, "principal_agent": "PRINCIPAL", "is_bundle": bundle}
        products_rows.append(list(values.values()))
    done = committed(
        imports, j03.priya, "products.csv", csv_bytes(product_headers, products_rows), "products"
    )
    ids = {
        row["code"]: row["id"]
        for row in imports.rows(
            select(product.c.code, product.c.id).where(
                product.c.code.in_(["AVM-KIT-10", "AVM-KIT-11", "AVM-BUNDLE-01"])
            )
        )
    }
    assert sorted(ids) == ["AVM-BUNDLE-01", "AVM-KIT-10", "AVM-KIT-11"]
    lineage = imports.rows(
        select(import_row_lineage.c.target_type).where(
            import_row_lineage.c.import_upload_id == UUID(done["id"])
        )
    )
    assert (
        sorted(str(row["target_type"]) for row in lineage)
        == ["product"] * 3 + ["source_record"] * 3
    )

    bundle_headers = [column.name for column in flatten(BundleIn)]
    bundle_rows = []
    for sequence, component in enumerate(("AVM-KIT-10", "AVM-KIT-11"), start=1):
        values = dict.fromkeys(bundle_headers, "")
        values |= {
            "product_code": "AVM-BUNDLE-01",
            "lines.component_product_id": str(ids[component]),
            "lines.quantity_per_bundle": "1",
            "lines.split_basis": "relative_ssp",
            "lines.sequence": str(sequence),
            "lines.valid_from": "2026-01-01",
        }
        bundle_rows.append(list(values.values()))
    committed(imports, j03.priya, "bundles.csv", csv_bytes(bundle_headers, bundle_rows), "bundles")
    components = imports.rows(
        select(product_bundle_component.c.component_product_id, product_bundle_component.c.sequence)
        .where(product_bundle_component.c.bundle_product_id == ids["AVM-BUNDLE-01"])
        .order_by(product_bundle_component.c.sequence)
    )
    assert [row["component_product_id"] for row in components] == [
        ids["AVM-KIT-10"],
        ids["AVM-KIT-11"],
    ]

    account_headers = [column.name for column in flatten(GlAccountIn)]
    account_rows = []
    for code, name, kind, normal in (
        ("4020", "Revenue - licences", "REVENUE", "C"),
        ("1215", "Contract asset - hardware", "ASSET", "D"),
    ):
        values = dict.fromkeys(account_headers, "")
        values |= {"code": code, "name": name, "account_type": kind, "normal_balance": normal}
        account_rows.append(list(values.values()))
    committed(
        imports,
        j03.priya,
        "gl-accounts.csv",
        csv_bytes(account_headers, account_rows),
        "gl_accounts",
    )
    accounts = imports.rows(
        select(gl_account.c.code, gl_account.c.id, gl_account.c.source_system)
        .where(gl_account.c.code.in_(["4020", "1215"]))
        .order_by(gl_account.c.code)
    )
    assert [(row["code"], str(row["source_system"])) for row in accounts] == [
        ("1215", "CSV_V2"),
        ("4020", "CSV_V2"),
    ]

    ssp_headers = [column.name for column in flatten(SspValuesIn)]
    values = dict.fromkeys(ssp_headers, "")
    values |= {
        "ssp_book_code": "US-LIST",
        "legacy_version_label": "2026-H2",
        "effective_from_date": "2026-07-01",
        "methodology_label": "Q3 list-price study",
        "lines.product_code": "AVM-KIT-10",
        "lines.currency": "USD",
        "lines.method": "observable",
        "lines.distinctness": "distinct",
        "lines.ranges.low_value": "900.00",
        "lines.ranges.mid_value": "1000.00",
        "lines.ranges.high_value": "1100.00",
    }
    committed(
        imports,
        j03.priya,
        "ssp-values.csv",
        csv_bytes(ssp_headers, [list(values.values())]),
        "ssp_values",
    )
    (drafted,) = imports.rows(
        select(
            ssp_book_version.c.id, ssp_book_version.c.status, ssp_book_version.c.entry_count
        ).where(ssp_book_version.c.legacy_version_label == "2026-H2")
    )
    assert (str(drafted["status"]), drafted["entry_count"]) == ("DRAFT", 1)
    assert (
        len(
            imports.rows(
                select(ssp_entry.c.id).where(ssp_entry.c.ssp_book_version_id == drafted["id"])
            )
        )
        == 1
    )

    mapping_headers = [column.name for column in flatten(AccountMappingRowsIn)]
    values = dict.fromkeys(mapping_headers, "")
    values |= {
        "name": "AVM-MAP-2026-07",
        "effective_from": "2026-07-01T00:00:00Z",
        "lines.account_role": "REVENUE",
        "lines.gl_account_id": str(next(row["id"] for row in accounts if row["code"] == "4020")),
        "lines.priority": "10",
    }
    committed(
        imports,
        j03.priya,
        "account-mapping.csv",
        csv_bytes(mapping_headers, [list(values.values())]),
        "account_mapping",
    )
    (mapping,) = imports.rows(
        select(account_mapping_version.c.id, account_mapping_version.c.status).where(
            account_mapping_version.c.name == "AVM-MAP-2026-07"
        )
    )
    assert str(mapping["status"]) == "DRAFT"
    rules = imports.rows(
        select(account_mapping_rule.c.account_role).where(
            account_mapping_rule.c.account_mapping_version_id == mapping["id"]
        )
    )
    assert [str(row["account_role"]) for row in rules] == ["REVENUE"]


def test_estimates_modifications_usage_pre_standard_csv(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    k11 = k11_world(app, keyring, clock, files)
    booked = booked_contract(k11.place, k11_body(k11.customer_id), activate=False)
    activated_contract(k11.place, booked)
    assign(k11.priya.member, "revenue_reviewer")
    imports = importer(k11.place)

    usage_headers = [column.name for column in flatten(UsageIn)]
    values = dict.fromkeys(usage_headers, "")
    values |= {
        "contract": K11_EXTERNAL_ID,
        "effective_date": "2026-09-30",
        "obligation_key": "O2",
        "usage_period_start": "2026-09-15",
        "usage_period_end": "2026-09-30",
        "metric": "SUPPORT_TICKETS",
        "quantity": "14",
    }
    usage_done = committed(
        imports, k11.priya, "usage.csv", csv_bytes(usage_headers, [list(values.values())]), "usage"
    )
    (usage,) = imports.rows(
        select(contract_event.c.event_type, contract_event.c.payload).where(
            contract_event.c.import_upload_id == UUID(usage_done["id"])
        )
    )
    assert (str(usage["event_type"]), usage["payload"]["metric"], usage["payload"]["quantity"]) == (
        "USAGE_REPORTED",
        "SUPPORT_TICKETS",
        "14",
    )

    pre_headers = [column.name for column in flatten(PreStandardRevenueIn)]
    rows = []
    for key, amount in (("O1", "250.00"), ("O2", "-120.00")):
        values = dict.fromkeys(pre_headers, "")
        values |= {
            "contract": K11_EXTERNAL_ID,
            "effective_date": "2026-09-30",
            "obligation_key": key,
            "amount.amount": amount,
            "amount.currency": "EUR",
        }
        rows.append(list(values.values()))
    pre_done = committed(
        imports, k11.priya, "pre-standard.csv", csv_bytes(pre_headers, rows), "pre_standard_revenue"
    )
    recorded = imports.rows(
        select(contract_event.c.event_type, contract_event.c.payload)
        .where(contract_event.c.import_upload_id == UUID(pre_done["id"]))
        .order_by(contract_event.c.record_seq)
    )
    assert [(str(row["event_type"]), row["payload"]["amount"]["amount"]) for row in recorded] == [
        ("PRE_STANDARD_REVENUE_RECORDED", "250.00"),
        ("PRE_STANDARD_REVENUE_RECORDED", "-120.00"),
    ]
    # D-86: `estimates` runs the CTR-12 commands (L7-4-Q-2). Rows 2 and 3 are the two outcomes of
    # version 1 of a new rebate element; row 4 prepares an implicit price concession without them.
    assert set(EstimateParametersRowIn.model_fields) == {
        name for model in ESTIMATE_PARAMETERS.values() for name in model.model_fields
    }
    estimate_headers = [column.name for column in flatten(EstimatesIn)]
    rebate = {
        "contract": K11_EXTERNAL_ID,
        "estimate_kind": "VARIABLE_CONSIDERATION",
        "element_code": "REBATE-DE-01",
        "vc_element_type": "REBATE",
        "method": "EXPECTED_VALUE",
        "effective_date": "2026-09-30",
        "parameters.refund_liability_target": "1620.00",
        "unconstrained_amount": "1620.00",
        "most_conservative_amount": "0.00",
        "constrained_amount": "1620.00",
        "rationale": "Hollenbrand expects to reach the 2026 volume threshold.",
    }
    rows = []
    for outcome, amount, probability in (
        ("Volume threshold reached", "2700.00", "0.6"),
        ("Volume threshold missed", "0.00", "0.4"),
    ):
        values = dict.fromkeys(estimate_headers, "") | rebate
        values |= {
            "lines.outcome": outcome,
            "lines.amount": amount,
            "lines.probability": probability,
        }
        rows.append(list(values.values()))
    concession = dict.fromkeys(estimate_headers, "") | {
        "contract": K11_EXTERNAL_ID,
        "estimate_kind": "IMPLICIT_PRICE_CONCESSION",
        "element_code": "IPC-DE-01",
        "method": "MOST_LIKELY_AMOUNT",
        "effective_date": "2026-09-30",
        "constrained_amount": "500.00",
        "rationale": "Support credits granted at past renewals.",
    }
    rows.append(list(concession.values()))
    estimates_done = committed(
        imports, k11.priya, "estimates.csv", csv_bytes(estimate_headers, rows), "estimates"
    )
    contract_id = UUID(str(booked.contract["id"]))
    elements = {
        str(row["element_code"]): row
        for row in imports.rows(
            select(
                estimate.c.id,
                estimate.c.element_code,
                estimate.c.estimate_kind,
                estimate.c.direction,
                estimate.c.method,
            ).where(estimate.c.contract_id == contract_id)
        )
    }
    assert {
        code: (str(row["estimate_kind"]), str(row["direction"]), str(row["method"]))
        for code, row in elements.items()
    } == {
        "IPC-DE-01": ("IMPLICIT_PRICE_CONCESSION", "DECREASE", "MOST_LIKELY_AMOUNT"),
        "REBATE-DE-01": ("VARIABLE_CONSIDERATION", "DECREASE", "EXPECTED_VALUE"),
    }
    versions = imports.rows(
        select(
            estimate_version.c.estimate_id,
            estimate_version.c.version_no,
            estimate_version.c.status,
            estimate_version.c.scenarios,
            estimate_version.c.parameters,
            estimate_version.c.constrained_amount,
            estimate_version.c.approval_request_id,
        ).order_by(estimate_version.c.constrained_amount)
    )
    assert [
        (
            row["estimate_id"],
            row["version_no"],
            str(row["status"]),
            Decimal(str(row["constrained_amount"])),
            row["approval_request_id"],
        )
        for row in versions
    ] == [
        (elements["IPC-DE-01"]["id"], 1, "DRAFT", Decimal("500.00"), None),
        (elements["REBATE-DE-01"]["id"], 1, "DRAFT", Decimal("1620.00"), None),
    ]
    assert (versions[0]["scenarios"], versions[0]["parameters"]) == ([], {})
    assert versions[1]["parameters"] == {"refund_liability_target": "1620.00"}
    assert versions[1]["scenarios"] == [
        {"outcome": "Volume threshold reached", "amount": "2700.00", "probability": "0.6"},
        {"outcome": "Volume threshold missed", "amount": "0.00", "probability": "0.4"},
    ]
    lineage = imports.rows(
        select(import_row_lineage.c.target_type).where(
            import_row_lineage.c.import_upload_id == UUID(estimates_done["id"])
        )
    )
    assert (
        sorted(str(row["target_type"]) for row in lineage)
        == ["estimate_version"] * 3 + ["source_record"] * 3
    )
    # A DRAFT version appends no event; ESTIMATE_CHANGED follows its approval (E-12).
    assert (
        imports.rows(
            select(contract_event.c.id).where(
                contract_event.c.import_upload_id == UUID(estimates_done["id"])
            )
        )
        == []
    )
    # A file states ONE version of an element (item EST-ONE-OPEN-VERSION-1; 04 T-CON-13, §16.14
    # rev 1.241; PRD ERR-93): one version of an element is prepared at a time and an imported
    # version stays DRAFT. A later file that names the element while its version 1 is open is
    # refused as the API refuses it — a row finding of the dry run — and so is the second
    # version a single file states for one element. Before, both were stored as drafts.
    revised = dict.fromkeys(estimate_headers, "") | rebate
    revised |= {
        "parameters.refund_liability_target": "2000.00",
        "unconstrained_amount": "2000.00",
        "constrained_amount": "2000.00",
        "rationale": "Tessen volumes confirm the threshold.",
        "lines.outcome": "Volume threshold reached",
        "lines.amount": "2000.00",
        "lines.probability": "1",
    }

    def refusals(import_id: str) -> list[tuple[str, str]]:
        return [
            (str(row["code"]), str(row["message"]))
            for row in imports.rows(
                select(exception_item.c.code, exception_item.c.message)
                .where(exception_item.c.import_upload_id == UUID(import_id))
                .order_by(exception_item.c.message)
            )
        ]

    def one_open(row_number: int, code: str) -> tuple[str, str]:
        return (
            "IMPORT_PROCESSING_FAILED",
            f"Row {row_number}, column Contract: {K11_EXTERNAL_ID} / {code} cannot be applied: "
            f"invalid-transition; row: Estimate version {code} v1 is open. One version of an "
            "element is prepared at a time: approve or discard it first. "
            "(IMPORT_PROCESSING_FAILED)",
        )

    early_id = diffed(
        imports,
        "estimates-revised.csv",
        csv_bytes(estimate_headers, [list(revised.values())]),
        "estimates",
    )
    assert refusals(early_id) == [one_open(2, "REBATE-DE-01")]
    # sent for approval and approved all the same, the file is not applied: its commit stops at
    # the refused command and stores nothing (PRD SM-05 ``COMMITTING`` → ``FAILED``)
    sent = submit(imports, early_id)
    assert sent.status_code == 200, sent.text
    decided = approve(imports.app, str(sent.json()["approval_request_id"]), k11.priya)
    assert decided.status_code == 200, decided.text
    run_import_job(imports, job_of(imports, UUID(early_id), "IMPORT_COMMIT"))
    assert shown(imports, early_id)["status"] == "FAILED"
    twice = [
        list((revised | {"element_code": "REBATE-DE-03", "effective_date": day}).values())
        for day in ("2026-09-30", "2026-10-31")
    ]
    twice_id = diffed(
        imports, "estimates-two-versions.csv", csv_bytes(estimate_headers, twice), "estimates"
    )
    assert refusals(twice_id) == [one_open(3, "REBATE-DE-03")]
    stored_versions = select(estimate_version.c.version_no, estimate_version.c.status).order_by(
        estimate_version.c.version_no
    )
    rebate_versions = imports.rows(
        stored_versions.where(estimate_version.c.estimate_id == elements["REBATE-DE-01"]["id"])
    )
    assert [(row["version_no"], str(row["status"])) for row in rebate_versions] == [(1, "DRAFT")]
    assert "REBATE-DE-03" not in {
        str(row["element_code"])
        for row in imports.rows(
            select(estimate.c.element_code).where(estimate.c.contract_id == contract_id)
        )
    }

    # Once version 1 is discarded the later file is applied: its version is numbered after the
    # latest, and the number of the discarded draft is not given out again.
    (first_version,) = imports.rows(
        select(estimate_version.c.id).where(
            estimate_version.c.estimate_id == elements["REBATE-DE-01"]["id"]
        )
    )
    discarded = post(
        k11.app, f"/api/v1/estimate-versions/{first_version['id']}/discard", k11.place.author, {}
    )
    assert (discarded.status_code, discarded.json()["status"]) == (200, "VOIDED"), discarded.text
    resent = revised | {"rationale": "Tessen volumes confirm the threshold (after the discard)."}
    committed(
        imports,
        k11.priya,
        "estimates-revised-2.csv",
        csv_bytes(estimate_headers, [list(resent.values())]),
        "estimates",
    )
    rebate_versions = imports.rows(
        stored_versions.where(estimate_version.c.estimate_id == elements["REBATE-DE-01"]["id"])
    )
    assert [(row["version_no"], str(row["status"])) for row in rebate_versions] == [
        (1, "VOIDED"),
        (2, "DRAFT"),
    ]
    # A refused command is a row finding in the CSV form of CPY-06 (D-87 L6-1-Q-2; L7-4-Q-4).
    unknown = dict.fromkeys(estimate_headers, "") | rebate
    unknown |= {"element_code": "REBATE-DE-02", "obligation_key": "O9"}
    failed_id = diffed(
        imports,
        "estimates-unknown-obligation.csv",
        csv_bytes(estimate_headers, [list(unknown.values())]),
        "estimates",
    )
    (item,) = imports.rows(
        select(
            exception_item.c.code,
            exception_item.c.message,
            exception_item.c.contract_id,
            exception_item.c.entity_id,
        ).where(exception_item.c.import_upload_id == UUID(failed_id))
    )
    assert item["code"] == "IMPORT_PROCESSING_FAILED"
    # 04 T-IMP-05 rev 1.218 (item EXC-IMPORT-SCOPE-1): the finding of a plan the dry run refuses
    # names the existing contract of its row and that contract's entity — the plan's key is
    # "<contract> / <element>" here, the contract is the one the row's Contract column names —
    # so it is an item of that entity and no longer one every member reads
    (booked,) = imports.rows(
        select(contract.c.id, contract.c.contracting_entity_id).where(
            contract.c.external_id == K11_EXTERNAL_ID
        )
    )
    assert (item["contract_id"], item["entity_id"]) == (
        booked["id"],
        booked["contracting_entity_id"],
    )
    assert item["message"].startswith(
        f"Row 2, column Contract: {K11_EXTERNAL_ID} / REBATE-DE-02 cannot be applied: "
    ), item["message"]
    assert item["message"].endswith(
        "obligation_key: Choose an obligation of the contract. (IMPORT_PROCESSING_FAILED)"
    ), item["message"]

    # L5-1-Q-23, D-86: `modifications` has no emitter, so the upload is refused before validation.
    for code in WITHOUT_EMITTER:
        file_id = upload_import_source(
            imports, f"{code}.csv", f"contract\n{K11_EXTERNAL_ID}\n".encode()
        )
        refused = create_import(imports, file_id, code)
        assert refused.status_code == 422, refused.text


def test_csv_v2_row_findings_located(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> None:
    """D-87 L6-1-Q-2: every CSV v2 row finding takes the CPY-06 CSV form "Row <n>, column <name>:"
    and its code, in the rows and the error report; copy that locates itself (CONTRACT_NOT_FOUND)
    is not located twice."""
    world = import_world(app, keyring, clock, files)
    headers = [column.name for column in flatten(CostEventIn)]
    base = dict.fromkeys(headers, "") | {
        "contract": "SF-ORD-99999",
        "effective_date": "2026-09-05",
        "purpose": "COST_TO_OBTAIN",
        "amount.amount": "5760.00",
        "amount.currency": "USD",
        "plan_code": "SALES-2026",
    }
    rows = [
        list((base | {"effective_date": "", "amount.amount": "abc"}).values()),
        list((base | {"effective_date": "05/09/2026"}).values()),
        list(base.values()),
    ]
    import_id, validated = imported(
        world, "avm-us-commissions-located.csv", csv_bytes(headers, rows), "cost_events"
    )
    assert validated["status"] == "INVALID", validated
    listed = get(app, f"/api/v1/imports/{import_id}/rows", world.actor)
    assert listed.status_code == 200, listed.text
    messages = {
        item["row_number"]: [message["message"] for message in item["messages"]]
        for item in listed.json()["items"]
    }
    assert messages == {
        2: [
            "Row 2, column Effective date: effective_date is required. (REQUIRED_VALUE_BLANK)",
            'Row 2, column Amount amount: amount.amount must be a number. Found "abc". '
            "(VALUE_NOT_NUMERIC)",
        ],
        3: [
            "Row 3, column Effective date: effective_date is not a date. Use YYYY-MM-DD or an "
            "Excel date. (DATE_INVALID)"
        ],
        4: [
            "Row 4, column Contract: Contract SF-ORD-99999 does not exist in this workspace. "
            "(CONTRACT_NOT_FOUND)"
        ],
    }
    report = get(app, f"/api/v1/imports/{import_id}/error-report", world.actor)
    assert report.status_code == 200, report.text
    reported = [line[3] for line in csv.reader(io.StringIO(report.text))][1:]
    assert reported == [message for number in (2, 3, 4) for message in messages[number]]
