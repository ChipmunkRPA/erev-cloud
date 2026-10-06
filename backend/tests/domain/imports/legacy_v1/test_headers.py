"""DIN-6 legacy v1 header sets, type-stage codes and the first-sheet rule (04 T-IMP-01 headers and
``sheet_rule``, table 15.4-A; PRD IMP-01 to IMP-03, IMP-09, IMP-17, IMP-20, WLD-F-17, J-01-ALT-2;
legacy 01 §7.3 TC-setup-21, 02 §7.3 TC-delivery-11, 03 §7.3 TC-18, 05 §7.3 TC-pob-vc-18; BUILD_SPEC
DIN-6).

World: ``support.legacy_replay.legacy_world``; the modification tests replay golden steps 01 and 02
(SKU SSP and setup 1.1.2023) first. Row findings are read through ``GET /imports/{id}/rows`` (sheet,
Excel row, rule and column); file findings are the upload's exception items without a row. The date
prompt of the legacy handlers is the upload parameter ``effective_date``, refused at
``POST /imports``.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from typing import Any
from uuid import UUID

import openpyxl
import pytest
from erev_api.auth.keyring import KeyRing
from erev_api.clock import FrozenClock
from erev_api.config import Settings
from erev_api.db.tables import exception_item, import_row
from erev_api.files.store import LocalFileStore
from erev_api.main import create_app
from fastapi import FastAPI
from sqlalchemy import select
from support.db import TestDatabase
from support.factories import (
    IMPORTS_PATH,
    create_import,
    imported,
    upload_import_source,
    workbook_bytes,
)
from support.legacy_replay import LegacyWorld, legacy_world, replayed
from support.reference import get

SETUP = "legacy_contract_setup"
PROGRESS = "legacy_progress_tracking"
MODIFICATION = "legacy_contract_modification"
SHEET = "Sheet1"
SETUP_HEADERS = (
    "Contract Unique Name",
    "POB Unique ID",
    "SKU Name",
    "POB Start Date",
    "POB End Date",
    "ASC 606 Stratification",
    "Original POB Total Selling Price",
    "Original POB Total Qty",
    "Selling Entity",
    "SSP Version",
    "Deferred Revenue Account",
    "Unbilled A/R Account",
    "Current Period",
    "Memo 1",
    "Memo 2",
    "Memo 3",
)
PROGRESS_HEADERS = (
    "Contract Unique Name",
    "POB Unique ID",
    "SKU Name",
    "Current Delivery",
    "Current Billing",
    "Current Pre-ASC606 Revenue (Net Design Only)",
    "Memo 1",
    "Memo 2",
    "Memo 3",
)
MODIFICATION_HEADERS = (
    "Contract Unique Name",
    "POB Unique ID",
    "SKU Name",
    "Mod Start Date",
    "Mod End Date",
    "ASC 606 Stratification",
    "Mod Billing",
    "Mod Qty",
    "Selling Entity",
    "Deferred Revenue Account",
    "Unbilled A/R Account",
    "SSP Version",
    "Memo 1",
    "Memo 2",
    "Memo 3",
)
MEMOS = ("Mod 1", "Mod 2", "Mod 3")


@pytest.fixture
def app(committed_db: TestDatabase, app_settings: Settings, clock: FrozenClock) -> FastAPI:
    return create_app(app_settings, clock=clock)


@pytest.fixture
def files(app_settings: Settings) -> LocalFileStore:
    return LocalFileStore(app_settings.file_root)


@pytest.fixture
def legacy(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, files: LocalFileStore
) -> LegacyWorld:
    return legacy_world(app, keyring, clock, files)


def without_memo_3(headers: Sequence[str], extra: str) -> list[str]:
    return [name for name in headers if name != "Memo 3"] + [extra]


def mod_row(
    contract_name: str,
    pob: str,
    sku: str,
    billing: Any,
    quantity: Any,
    *,
    start: Any = "2023-01-01",
    entity: str = "Mock Entity 1",
    accounts: tuple[int, int] = (21001, 15001),
    ssp_version: Any = "2023-01-01",
) -> list[Any]:
    return [
        contract_name,
        pob,
        sku,
        start,
        "2024-05-31",
        sku,
        billing,
        quantity,
        entity,
        accounts[0],
        accounts[1],
        ssp_version,
        *MEMOS,
    ]


def file_findings(world: LegacyWorld, import_id: str) -> list[tuple[str, str]]:
    """The upload's findings without a row (IPL-03)."""
    return [
        (str(item["code"]), str(item["message"]))
        for item in world.imports.rows(
            select(exception_item.c.code, exception_item.c.message)
            .where(
                exception_item.c.import_upload_id == UUID(import_id),
                exception_item.c.import_row_id.is_(None),
            )
            .order_by(exception_item.c.code)
        )
    ]


def row_messages(world: LegacyWorld, import_id: str) -> list[tuple[str, int, str, str | None, str]]:
    """(sheet, Excel row, rule, column, message) of every row message (API-S-ImportRow)."""
    response = get(world.app, f"{IMPORTS_PATH}/{import_id}/rows", world.maya)
    assert response.status_code == 200, response.text
    return sorted(
        (
            str(item["sheet_name"]),
            int(item["row_number"]),
            str(message["rule_id"]),
            message["field"],
            str(message["message"]),
        )
        for item in response.json()["items"]
        for message in item["messages"]
    )


def refused_date(world: LegacyWorld, rows: Sequence[Sequence[Any]], mode: str, value: str) -> None:
    """The legacy date prompt (VR-03, VR-pob-vc-01) as the upload parameter: 422 naming it."""
    file_id = upload_import_source(
        world.imports,
        f"date {value.replace('/', '.')}.xlsx",
        workbook_bytes(SHEET, MODIFICATION_HEADERS, rows),
    )
    response = create_import(
        world.imports, file_id, MODIFICATION, {"effective_date": value, "mode": mode}
    )
    assert response.status_code == 422, response.text
    assert [(error["field"], error["message"]) for error in response.json()["errors"]] == [
        ("parameters.effective_date", "effective_date must be a date in the form YYYY-MM-DD.")
    ]


def test_tc_setup_21_setup_header_mismatch_lists_missing_and_extra(legacy: LegacyWorld) -> None:
    content = workbook_bytes(
        SHEET,
        without_memo_3(SETUP_HEADERS, "Notes"),
        [
            [
                "Contract 1",
                "POB #1",
                "Hardware 1",
                "2023-01-01",
                "2023-12-31",
                "Hardware 1",
                1000,
                10,
                "Mock Entity 1",
                "2023-01-01",
                21001,
                15001,
                "2023-01-01",
                "Setup 1",
                "Setup 2",
                "Region A",
            ]
        ],
    )
    import_id, validated = imported(
        legacy.imports, "legacy-setup-header-mismatch.xlsx", content, SETUP
    )
    assert validated["status"] == "INVALID"
    ((code, message),) = file_findings(legacy, import_id)
    assert code == "TEMPLATE_HEADER_MISMATCH"
    assert message.startswith("The file does not match the ")
    assert message.endswith(" template. Missing: Memo 3. Extra: Notes.")
    assert row_messages(legacy, import_id) == []


def test_tc_delivery_11_progress_header_and_type_errors(legacy: LegacyWorld) -> None:
    parameters = {"effective_date": "2023-01-31"}
    mismatch_id, mismatch = imported(
        legacy.imports,
        "progress header 1.31.2023.xlsx",
        workbook_bytes(
            SHEET,
            without_memo_3(PROGRESS_HEADERS, "Notes"),
            [["Contract 1", "POB #1", "Hardware 1", 1, 100, 0, "Delivered", "Invoice", "Region"]],
        ),
        PROGRESS,
        parameters,
    )
    assert mismatch["status"] == "INVALID"
    ((code, message),) = file_findings(legacy, mismatch_id)
    assert code == "TEMPLATE_HEADER_MISMATCH"
    assert message.endswith(" template. Missing: Memo 3. Extra: Notes.")
    typed_id, typed = imported(
        legacy.imports,
        "progress types 1.31.2023.xlsx",
        workbook_bytes(
            SHEET,
            PROGRESS_HEADERS,
            [
                ["Contract 1", "POB #1", "Hardware 1", 1, None, 0, "Delivered", "Invoice", "A"],
                ["Contract 1", "POB #2", "Software 1", "abc", 50, 0, "Delivered", "Invoice", "A"],
            ],
        ),
        PROGRESS,
        parameters,
    )
    assert typed["status"] == "INVALID"
    assert file_findings(legacy, typed_id) == []
    # BUILD_SPEC DIN-7: row messages name worksheet, row, column, key and rule id (REQ-DAT-006).
    assert row_messages(legacy, typed_id) == [
        (
            SHEET,
            2,
            "REQUIRED_VALUE_BLANK",
            "Current Billing",
            "Sheet1 row 2, column Current Billing: Current Billing is required. "
            "Key Contract 1 / POB #1 / Hardware 1. (REQUIRED_VALUE_BLANK)",
        ),
        (
            SHEET,
            3,
            "VALUE_NOT_NUMERIC",
            "Current Delivery",
            "Sheet1 row 3, column Current Delivery: Current Delivery must be a number. "
            'Found "abc". Key Contract 1 / POB #2 / Software 1. (VALUE_NOT_NUMERIC)',
        ),
    ]


def test_tc_18_modification_validation_codes(legacy: LegacyWorld) -> None:
    replayed(legacy, "02")
    rows = [
        mod_row("Contract 9", "POB #1", "Hardware 1", 100, 1),
        mod_row("Contract 1", "POB #1", "Hardware 1", 100, None),
        mod_row("Contract 1", "POB #2", "Software 1", "abc", 1),
        mod_row("Contract 1", "POB #3", "Consulting 1", 100, 1, start="2023-5-15"),
    ]
    import_id, validated = imported(
        legacy.imports,
        "Contract Modification Template 05.15.2023 - defects.xlsx",
        workbook_bytes(SHEET, MODIFICATION_HEADERS, rows),
        MODIFICATION,
        {"effective_date": "2023-05-15", "mode": "prospective"},
    )
    assert validated["status"] == "INVALID"
    assert row_messages(legacy, import_id) == [
        (
            SHEET,
            2,
            "CONTRACT_NOT_FOUND",
            "Contract Unique Name",
            "Sheet1 row 2, column Contract Unique Name: Contract Contract 9 does not exist in this "
            "workspace. Key Contract 9 / POB #1 / Hardware 1. (CONTRACT_NOT_FOUND)",
        ),
        (
            SHEET,
            3,
            "REQUIRED_VALUE_BLANK",
            "Mod Qty",
            "Sheet1 row 3, column Mod Qty: Mod Qty is required. Key Contract 1 / POB #1 / Hardware "
            "1. (REQUIRED_VALUE_BLANK)",
        ),
        (
            SHEET,
            4,
            "VALUE_NOT_NUMERIC",
            "Mod Billing",
            'Sheet1 row 4, column Mod Billing: Mod Billing must be a number. Found "abc". Key '
            "Contract 1 / POB #2 / Software 1. (VALUE_NOT_NUMERIC)",
        ),
        (
            SHEET,
            5,
            "DATE_INVALID",
            "Mod Start Date",
            "Sheet1 row 5, column Mod Start Date: Mod Start Date is not a date. Use YYYY-MM-DD or "
            "an Excel date. Key Contract 1 / POB #3 / Consulting 1. (DATE_INVALID)",
        ),
    ]
    refused_date(legacy, rows[:1], "prospective", "2023-5-15")


def test_tc_pob_vc_18_price_change_validation_codes(legacy: LegacyWorld) -> None:
    replayed(legacy, "02")
    parameters = {"effective_date": "2023-05-31", "mode": "pob_price_change"}
    mismatch_id, mismatch = imported(
        legacy.imports,
        "Contract Modification Template 05.31.2023 - headers.xlsx",
        workbook_bytes(
            SHEET,
            without_memo_3(MODIFICATION_HEADERS, "Foo"),
            [mod_row("Contract 2", "POB #1", "Hardware 1", -200, 0)[:14] + ["Foo"]],
        ),
        MODIFICATION,
        parameters,
    )
    assert mismatch["status"] == "INVALID"
    ((code, message),) = file_findings(legacy, mismatch_id)
    assert code == "TEMPLATE_HEADER_MISMATCH"
    assert message == (
        "The file does not match the Contract Modification template. Missing: Memo 3. Extra: Foo."
    )
    rows = [
        mod_row("Contract 1", "POB #1", "Hardware 1", "abc", 0),
        mod_row("Contract 1", "POB #2", "Software 1", -20, None),
        mod_row("Contract 1", "POB #3", "Consulting 1", -10, 0, ssp_version="2024-01-01"),
        mod_row(
            "Contract 2",
            "POB #1",
            "Hardware 1",
            -200,
            0,
            start="2023/05/31",
            entity="Mock Entity 2",
            accounts=(21002, 15002),
        ),
    ]
    import_id, validated = imported(
        legacy.imports,
        "Contract Modification Template 05.31.2023 - defects.xlsx",
        workbook_bytes(SHEET, MODIFICATION_HEADERS, rows),
        MODIFICATION,
        parameters,
    )
    assert validated["status"] == "INVALID"
    assert row_messages(legacy, import_id) == [
        (
            SHEET,
            2,
            "VALUE_NOT_NUMERIC",
            "Mod Billing",
            'Sheet1 row 2, column Mod Billing: Mod Billing must be a number. Found "abc". Key '
            "Contract 1 / POB #1 / Hardware 1. (VALUE_NOT_NUMERIC)",
        ),
        (
            SHEET,
            3,
            "REQUIRED_VALUE_BLANK",
            "Mod Qty",
            "Sheet1 row 3, column Mod Qty: Mod Qty is required. Key Contract 1 / POB #2 / Software "
            "1. (REQUIRED_VALUE_BLANK)",
        ),
        (
            SHEET,
            4,
            "SSP_KEY_NOT_FOUND",
            "SSP Version",
            "Sheet1 row 4, column SSP Version: No approved SSP for Consulting 1 / Consulting 1 / "
            "2024-01-01. Key Contract 1 / POB #3 / Consulting 1. (SSP_KEY_NOT_FOUND)",
        ),
        (
            SHEET,
            5,
            "DATE_INVALID",
            "Mod Start Date",
            "Sheet1 row 5, column Mod Start Date: Mod Start Date is not a date. Use YYYY-MM-DD or "
            "an Excel date. Key Contract 2 / POB #1 / Hardware 1. (DATE_INVALID)",
        ),
    ]
    refused_date(legacy, rows[3:], "pob_price_change", "2023/05/31")


def test_first_sheet_only(legacy: LegacyWorld) -> None:
    book = openpyxl.Workbook(write_only=True)
    first = book.create_sheet(SHEET)
    first.append(list(MODIFICATION_HEADERS))
    first.append(mod_row("Contract 9", "POB #1", "Hardware 1", 100, 1))
    second = book.create_sheet("Notes")
    second.append(list(MODIFICATION_HEADERS))
    for number in (1, 2, 3):
        second.append(mod_row("Contract 1", f"POB #{number}", "Hardware 1", "abc", "abc"))
    buffer = io.BytesIO()
    book.save(buffer)
    import_id, validated = imported(
        legacy.imports,
        "two sheets.xlsx",
        buffer.getvalue(),
        MODIFICATION,
        {"effective_date": "2023-05-15", "mode": "prospective"},
    )
    assert validated["counts"]["rows"] == 1
    stored = legacy.imports.rows(
        select(import_row.c.sheet_name, import_row.c.row_number).where(
            import_row.c.import_upload_id == UUID(import_id)
        )
    )
    assert [(item["sheet_name"], item["row_number"]) for item in stored] == [(SHEET, 2)]
    assert [
        (sheet, number, rule) for sheet, number, rule, _, _ in row_messages(legacy, import_id)
    ] == [(SHEET, 2, "CONTRACT_NOT_FOUND")]
