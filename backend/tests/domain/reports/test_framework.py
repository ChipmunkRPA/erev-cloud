"""Report run framework and stamped outputs (04 §16.9 API-R-41, T-RPT-02; SCREENS_B RV-02, RV-03,
RV-07, RV-14, RPT-27; DESIGN_SYSTEM DS-FMT-25, DS-FMT-26, DS-VER-03; 03 REQ-RPT-002, REQ-RPT-024,
REQ-SEC-006, REQ-SEC-011; CTL-029; BUILD_SPEC RPS-2).

World (SCREENS_B RPT-27 sample world, WLD-T-01): Maya is a Revenue Accountant (``report.run``,
``report.export``) who adds the entity AVM-US; Tomas, a Tenant Admin enrolled in MFA, creates the
API clients ``svc-salesforce``, ``svc-netsuite`` and ``svc-metering`` for all entities. The tests
play the worker as ``tests/api/test_jobs_api.py`` does: they mark the Procrastinate task fetched and
call ``run_job``. The output writers are also exercised on a sample dataset with money, date and
timestamp columns, which the API client inventory does not have.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from typing import Any, Final
from uuid import UUID

import openpyxl
import pytest
from erev_api.auth.keyring import KeyRing
from erev_api.auth.principal import system_principal
from erev_api.clock import FrozenClock
from erev_api.config import Settings
from erev_api.controls import release as release_module
from erev_api.db.session import DbContext, tenant_session
from erev_api.db.tables import api_client, audit_event, engine_release, job, report_run
from erev_api.domain.journals import views
from erev_api.domain.reports.builders import ReportParams, api_client_inventory
from erev_api.domain.reports.outputs import Column, ReportData, RunStamp
from erev_api.domain.reports.outputs import csv as csv_output
from erev_api.domain.reports.outputs import manifest as manifest_output
from erev_api.domain.reports.outputs import pdf as pdf_output
from erev_api.domain.reports.outputs import xlsx as xlsx_output
from erev_api.enums import JobKind, PrincipalKind
from erev_api.files.store import LocalFileStore
from erev_api.jobs import registry
from erev_api.jobs.context import JobRuntime, system_unit_of_work
from erev_api.main import create_app
from erev_engine import ENGINE_VERSION
from fastapi import FastAPI
from sqlalchemy import Engine, event, insert, select, text
from support.api_clients import access_approver, granted_client
from support.db import TestDatabase
from support.factories import RELEASE_BUILD, stamp_test_release
from support.principals import Actor, colleague, enrolled, member
from support.reference import assign, calendar, entity, get, holding, post
from support.rows import api_client_values, report_run_values

RUNS: Final = "/api/v1/report-runs"
JOBS: Final = "/api/v1/jobs"
CLIENTS: Final = "/api/v1/api-clients"
PROBLEM_BASE: Final = "https://erev.dev/problems/"
INVENTORY: Final = "api_client_inventory"
CLIENT_NAMES: Final = ("svc-salesforce", "svc-netsuite", "svc-metering")
KNOWN_AT: Final = "2026-09-12T12:00:00Z"
CLIENTS_CREATED: Final = datetime(2026, 9, 1, 9, tzinfo=UTC)
HOSTILE: Final = 'x" OR 1=1 OR "x'
LEDGER_HEADS: Final = {
    book: {"chain_seq": 0, "seal_sha256": None} for book in ("ASC606", "IFRS15", "LEGACY")
}
_FETCHED: Final = text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id")


@pytest.fixture
def app(committed_db: TestDatabase, app_settings: Settings, clock: FrozenClock) -> FastAPI:
    return create_app(app_settings, clock=clock)


@pytest.fixture
def runtime(app_settings: Settings, keyring: KeyRing, clock: FrozenClock) -> JobRuntime:
    return JobRuntime(clock=clock, keyring=keyring, files=LocalFileStore(app_settings.file_root))


def _db(tenant_id: UUID) -> DbContext:
    return DbContext(tenant_id=tenant_id, user_id=None, entity_scope="*")


def run_deferred(tenant_id: UUID, job_id: UUID, runtime: JobRuntime, *, attempt: int = 1) -> None:
    """The worker fetches the job's current task and runs it."""
    with tenant_session(_db(tenant_id)) as session:
        task_id = session.execute(
            select(job.c.procrastinate_job_id).where(job.c.id == job_id)
        ).scalar_one()
        session.execute(_FETCHED, {"id": task_id})
    registry.run_job(job_id, tenant_id, attempt=attempt, runtime=runtime)


def insert_clients(tenant_id: UUID, names: Sequence[str]) -> None:
    """All-entities clients created before the frozen clock (T-PLT-15)."""
    with tenant_session(_db(tenant_id)) as session:
        session.execute(
            insert(api_client),
            [
                api_client_values(
                    tenant_id, name=name, created_at=CLIENTS_CREATED, updated_at=CLIENTS_CREATED
                )
                for name in names
            ],
        )


@dataclass(frozen=True, slots=True)
class InventoryWorld:
    app: FastAPI
    maya: Actor
    tenant_id: UUID
    entity_id: UUID


def inventory_world(app: FastAPI, keyring: KeyRing, clock: FrozenClock) -> InventoryWorld:
    stamp_test_release()
    maya_member = member(keyring, clock)
    maya = holding(app, maya_member, "revenue_accountant")
    created = entity(app, maya, code="AVM-US", calendar_id=calendar(app, maya))
    tomas_member = colleague(maya_member.tenant_id, "tomas")
    assign(tomas_member, "tenant_admin")
    tomas = enrolled(app, clock, tomas_member)
    # A client's scopes are an access grant (supervisor ruling R-38 (iii)): Tomas requests
    # each client and Grace, a second Tenant Admin, approves its grant.
    grace = access_approver(app, clock, maya_member, "grace")
    for name in CLIENT_NAMES:
        granted_client(app, tomas, {"name": name, "scopes": ["contract.read"]}, approver=grace)
    return InventoryWorld(
        app=app, maya=maya, tenant_id=maya_member.tenant_id, entity_id=UUID(str(created["id"]))
    )


def start_run(
    app: FastAPI, actor: Actor, *, output_format: str, parameters: Mapping[str, Any] | None = None
) -> tuple[UUID, UUID]:
    """``POST /report-runs`` for the inventory; returns the run and job ids."""
    started = post(
        app,
        RUNS,
        actor,
        {
            "report_code": INVENTORY,
            "parameters": dict(parameters or {}),
            "output_format": output_format,
        },
    )
    assert started.status_code == 202, started.text
    body = started.json()
    assert (body["kind"], body["state"]) == ("REPORT_RUN", "QUEUED")
    assert started.headers["location"] == f"{JOBS}/{body['id']}"
    return UUID(started.headers["x-erev-report-run-id"]), UUID(body["id"])


@pytest.mark.control("CTL-029")
def test_ctl_029_run_record_and_rerun(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, runtime: JobRuntime
) -> None:
    world = inventory_world(app, keyring, clock)
    run_id, job_id = start_run(app, world.maya, output_format="JSON")
    run_deferred(world.tenant_id, job_id, runtime)

    shown = get(app, f"{RUNS}/{run_id}", world.maya)
    assert shown.status_code == 200, shown.text
    record = shown.json()
    assert record["status"] == "SUCCEEDED"
    assert record["report"] == {"code": INVENTORY, "version": 1, "name": "API client inventory"}
    # Every parameter including the defaults: the entities in scope, now and the flag's false.
    assert record["parameters"] == {
        "entity_codes": ["AVM-US"],
        "include_revoked": False,
        "known_at": KNOWN_AT,
        "known_at_basis": "record",  # CUTOFF-R1 (04 §16.9 rev 1.54): the read basis stored too
    }
    assert [ref["code"] for ref in record["entity_scope"]] == ["AVM-US"]
    assert (record["book"], record["as_of"], record["period_lock_id"]) == (None, None, None)
    assert record["known_at"] == KNOWN_AT
    # 05 REL-03 (rev 1.15, D-96 consumers): the run names the release THIS process stamped —
    # here the world's `stamp_test_release()` row (build RELEASE_BUILD) — never "the latest row
    # of the running engine version", which another module's stamp may have deployed later in
    # the shared database (the CTL-029 failure of T1's chain on main 4ed5f423).
    process = release_module.current_release()
    assert process is not None and process.build_sha == RELEASE_BUILD
    assert record["engine_release"] == {
        "engine_version": process.engine_version,
        "build_sha": process.build_sha,
    }
    assert record["engine_release"]["engine_version"] == ENGINE_VERSION
    release = {"id": process.id, "build_sha": process.build_sha}
    assert record["row_count"] == 3
    assert record["control_totals"] == {"client_count": 3}
    assert record["tie_out_results"] == []
    # Book → {chain_seq, seal_sha256} at run time: provisioning's heads, before any posting.
    assert record["ledger_heads"] == LEDGER_HEADS
    assert record["problem"] is None
    assert (record["run_by"]["id"], record["run_by"]["kind"]) == (
        str(world.maya.member.user_id),
        "USER",
    )
    assert (record["started_at"], record["finished_at"]) == (KNOWN_AT, KNOWN_AT)
    output = record["output"]
    assert output["format"] == "JSON"
    assert re.fullmatch(r"[0-9a-f]{64}", output["sha256"])
    assert output["href"] == f"{RUNS}/{run_id}/output"
    assert output["manifest_href"] is None
    assert re.fullmatch(r"RPT-[0-9]{6}", record["report_run_no"])

    with tenant_session(_db(world.tenant_id)) as session:
        stored = (
            session.execute(select(report_run).where(report_run.c.id == run_id)).mappings().one()
        )
    assert stored["entity_ids"] == [world.entity_id]
    assert stored["engine_release_id"] == release["id"]
    assert (stored["output_sha256"], stored["row_count"]) == (output["sha256"], 3)
    assert stored["ledger_heads"] == LEDGER_HEADS
    assert stored["job_id"] == job_id

    downloaded = get(app, output["href"], world.maya)
    assert downloaded.status_code == 200, downloaded.text
    assert hashlib.sha256(downloaded.content).hexdigest() == output["sha256"]
    rows = get(app, f"{RUNS}/{run_id}/data", world.maya, {"limit": "200"}).json()["items"]
    assert [row["name"] for row in rows] == ["svc-metering", "svc-netsuite", "svc-salesforce"]
    first = rows[0]
    assert first["row_key"] == "client:svc-metering"
    assert (first["status"], first["scopes"], first["entity_scope"]) == (
        "ACTIVE",
        ["contract.read"],
        [],
    )
    assert (first["rate_limit_per_minute"], first["last_used_at"]) == (600, None)
    assert first["created_by"]["display_name"] == "Tomas"
    assert "secret_hash" not in first

    rerun = post(app, f"{RUNS}/{run_id}/rerun", world.maya, {})
    assert rerun.status_code == 202, rerun.text
    second_id = UUID(rerun.headers["x-erev-report-run-id"])
    run_deferred(world.tenant_id, UUID(rerun.json()["id"]), runtime)
    finished = get(app, f"{JOBS}/{rerun.json()['id']}", world.maya).json()
    assert finished["state"] == "SUCCEEDED"
    assert finished["result"]["output_sha256_equal"] is True
    assert finished["result"]["control_totals_equal"] is True
    assert finished["result"]["href"] == f"{RUNS}/{second_id}"
    assert finished["result"]["rerun_of"] == str(run_id)
    second = get(app, f"{RUNS}/{second_id}", world.maya).json()
    assert second["parameters"] == record["parameters"]
    assert second["output"]["sha256"] == output["sha256"]
    assert second["report_run_no"] != record["report_run_no"]


def sample_data(rows: int = 2) -> ReportData:
    """A register with money, date, timestamp, integer, flag and code-list columns."""
    columns = (
        Column("contract", "Contract", "text"),
        Column("status", "Status", "code"),
        Column("units", "Units", "integer"),
        Column("amount", "Amount", "money"),
        Column("fee", "Fee", "money"),
        Column("effective_date", "Effective", "date"),
        Column("recorded_at", "Recorded", "timestamp"),
        Column("billed", "Billed", "boolean"),
        Column("scopes", "Scopes", "codes", empty_text="All entities"),
    )
    items = [
        {
            "row_key": "contract:SF-ORD-10001",
            "contract": "SF-ORD-10001",
            "status": "ACTIVE",
            "units": 1200,
            "amount": {"amount": "-4000.00", "currency": "USD"},
            "fee": {"amount": "1200", "currency": "JPY"},
            "effective_date": date(2026, 9, 7),
            "recorded_at": datetime(2026, 9, 12, 12, tzinfo=UTC),
            "billed": True,
            "scopes": ("contract.read", "event.record"),
        },
        {
            "row_key": "contract:SF-ORD-10002",
            "contract": "SF-ORD-10002",
            "status": "DRAFT",
            "units": 3,
            "amount": {"amount": "1250.50", "currency": "USD"},
            "fee": None,
            "effective_date": date(2026, 9, 30),
            "recorded_at": datetime(2026, 9, 14, 8, 30, tzinfo=UTC),
            "billed": False,
            "scopes": (),
        },
    ]
    extra = [
        {**items[1], "row_key": f"contract:SF-ORD-{10003 + index}", "contract": f"SF-{index}"}
        for index in range(max(rows - 2, 0))
    ]
    return ReportData(
        columns=columns,
        rows=tuple([*items, *extra][:rows]),
        control_totals={
            "contract_count": rows,
            "amount_total": {"amount": "-2749.50", "currency": "USD"},
        },
    )


def sample_stamp(style: str = "PARENTHESES") -> RunStamp:
    return RunStamp(
        report_code="sample_register",
        report_name="Sample register",
        report_version=3,
        report_run_no="RPT-000042",
        parameters={
            "entity_codes": ["AVM-US", "AVM-DE"],
            "include_revoked": False,
            "known_at": KNOWN_AT,
        },
        entity_codes=("AVM-US", "AVM-DE"),
        book="ASC606",
        as_of="2026-09-30",
        source="Current, known at 12 Sep 2026 12:00 UTC",
        engine_version=ENGINE_VERSION,
        build_sha="abc1234",
        run_by="Maya Chen",
        run_at=datetime(2026, 9, 12, 12, tzinfo=UTC),
        output_sha256="f" * 64,
        negative_number_style=style,
    )


def csv_rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content.decode("utf-8"), newline="")))


def test_csv_manifest(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, runtime: JobRuntime
) -> None:
    data = sample_data()
    content = csv_output.render_csv(data)
    assert b"\r\n" in content and not content.startswith(b"\xef\xbb\xbf")
    parsed = csv_rows(content)
    assert parsed[0] == [
        "Contract",
        "Status",
        "Units",
        "Amount (USD)",
        "Fee (JPY)",
        "Effective",
        "Recorded",
        "Billed",
        "Scopes",
    ]
    # DS-FMT-25: signed decimal strings with a hyphen-minus and no grouping, ISO dates, UTC stamps.
    assert parsed[1] == [
        "SF-ORD-10001",
        "ACTIVE",
        "1200",
        "-4000.00",
        "1200",
        "2026-09-07",
        "2026-09-12T12:00:00Z",
        "true",
        "contract.read, event.record",
    ]
    assert parsed[2][3:8] == ["1250.50", "", "2026-09-30", "2026-09-14T08:30:00Z", "false"]
    manifest = json.loads(
        manifest_output.manifest_bytes(
            stamp=sample_stamp(),
            data=data,
            file_name="sample_register-RPT-000042.csv",
            media_type=csv_output.MEDIA_TYPE,
            content=content,
        )
    )
    assert manifest["row_count"] == 2
    assert manifest["control_totals"] == {
        "amount_total": {"amount": "-2749.50", "currency": "USD"},
        "contract_count": 2,
    }
    assert manifest["sha256"] == hashlib.sha256(content).hexdigest()
    assert manifest["file"] == {
        "name": "sample_register-RPT-000042.csv",
        "media_type": "text/csv",
        "bytes": len(content),
    }

    # A CSV run stores the output and its manifest; the manifest names the SHA-256 of the CSV bytes.
    stamp_test_release()
    maya_member = member(keyring, clock)
    maya = holding(app, maya_member, "revenue_accountant")
    insert_clients(maya_member.tenant_id, ["svc-billing", "svc-crm"])
    run_id, job_id = start_run(app, maya, output_format="CSV")
    run_deferred(maya_member.tenant_id, job_id, runtime)
    record = get(app, f"{RUNS}/{run_id}", maya).json()
    assert record["output"]["manifest_href"] == f"{RUNS}/{run_id}/output?part=manifest"
    output = get(app, record["output"]["href"], maya)
    assert output.status_code == 200, output.text
    assert output.headers["content-type"].startswith("text/csv")
    stored = json.loads(get(app, f"{RUNS}/{run_id}/output", maya, {"part": "manifest"}).content)
    assert (
        stored["sha256"] == hashlib.sha256(output.content).hexdigest() == record["output"]["sha256"]
    )
    assert (stored["row_count"], stored["control_totals"]) == (2, {"client_count": 2})
    assert stored["report_run_no"] == record["report_run_no"]
    assert [row[0] for row in csv_rows(output.content)] == ["Name", "svc-billing", "svc-crm"]


def _grid_start(sheet: Any) -> int:
    return next(
        row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "Contract"
    )


def test_xlsx_numbers_and_stamp() -> None:
    data = sample_data()
    for style, usd, jpy in (
        ("PARENTHESES", "#,##0.00_);(#,##0.00)", "#,##0_);(#,##0)"),
        ("MINUS", "#,##0.00;-#,##0.00", "#,##0;-#,##0"),
    ):
        stamp = sample_stamp(style)
        book = openpyxl.load_workbook(io.BytesIO(xlsx_output.render_xlsx(data, stamp)))
        sheet = book.worksheets[0]
        assert sheet.title == "Sample register"
        grid = _grid_start(sheet)
        block = {sheet.cell(row, 1).value: sheet.cell(row, 2).value for row in range(1, grid)}
        assert block["Report"] == "Sample register v3"
        assert block["Run"] == "RPT-000042"
        assert block["Entity"] == "AVM-US, AVM-DE"
        assert (block["Book"], block["As of"]) == ("ASC606", "2026-09-30")
        assert block["Engine"] == f"{ENGINE_VERSION} (abc1234)"
        assert block["Run by"] == "Maya Chen"
        assert block["Run at"] == "12 Sep 2026 12:00 UTC"
        assert block["Rows"] == 2
        assert block["Output SHA-256"] == "f" * 64
        # Every parameter, then the control totals.
        assert block["entity_codes"] == "AVM-US, AVM-DE"
        assert block["include_revoked"] == "false"
        assert block["known_at"] == KNOWN_AT
        assert block["contract_count"] == 2
        totals = next(row for row in range(1, grid) if sheet.cell(row, 1).value == "amount_total")
        assert sheet.cell(totals, 2).value == -2749.5
        assert sheet.cell(totals, 2).number_format == usd
        assert sheet.cell(totals, 3).value == "USD"
        headers = [sheet.cell(grid, col).value for col in range(1, 10)]
        assert headers[3:5] == ["Amount (USD)", "Fee (JPY)"]
        amount, fee = sheet.cell(grid + 1, 4), sheet.cell(grid + 1, 5)
        assert (amount.data_type, amount.value, amount.number_format) == ("n", -4000, usd)
        assert (fee.data_type, fee.value, fee.number_format) == ("n", 1200, jpy)
        units = sheet.cell(grid + 1, 3)
        assert (units.data_type, units.value) == ("n", 1200)
        effective = sheet.cell(grid + 1, 6)
        assert (effective.value, effective.number_format) == (datetime(2026, 9, 7), "dd mmm yyyy")
        assert sheet.cell(grid + 2, 9).value == "All entities"
        assert sheet.freeze_panes == f"A{grid + 1}"


def pdf_pages(content: bytes) -> list[list[str]]:
    """The text strings of each page's content stream, in drawing order."""
    assert content.startswith(b"%PDF-1.4")
    offset = int(re.findall(rb"startxref\n([0-9]+)\n%%EOF", content)[0])
    assert content[offset : offset + 4] == b"xref"
    streams = re.findall(rb"stream\n(.*?)endstream", content, re.S)
    pages: list[list[str]] = []
    for stream in streams:
        texts = []
        for raw in re.findall(rb"\(((?:\\.|[^\\)])*)\) Tj", stream):
            unescaped = re.sub(rb"\\(.)", rb"\1", raw)
            texts.append(unescaped.decode("cp1252"))
        pages.append(texts)
    assert len(pages) == len(re.findall(rb"/Type /Page /Parent", content))
    return pages


def test_pdf_footer_stamp() -> None:
    data = sample_data(rows=130)
    pages = pdf_pages(pdf_output.render_pdf(data, sample_stamp()))
    count = len(pages)
    assert count == 4
    assert [page[-1] for page in pages] == [
        f"Run RPT-000042 · page {number} of {count}" for number in range(1, count + 1)
    ]
    first = pages[0]
    assert first[0] == "Sample register v3"
    assert "Output SHA-256: " + "f" * 64 in first
    assert "include_revoked: false" in first
    assert "amount_total: USD (2,749.50)" in first
    # A continuation page repeats the grid header.
    assert pages[1][:3] == ["Contract", "Status", "Units"]


def test_formula_injection_prefix() -> None:
    hostile = ("=1+1", "+CMD", "-2", "@SUM(A1)", "\tTAB", "\rCR")
    data = ReportData(
        columns=(
            Column("memo", "Memo", "text"),
            Column("amount", "Amount", "money"),
            Column("units", "Units", "integer"),
        ),
        rows=tuple(
            {
                "row_key": f"row:{index}",
                "memo": memo,
                "amount": {"amount": "-2.00", "currency": "USD"},
                "units": -2,
            }
            for index, memo in enumerate(hostile)
        ),
        control_totals={"row_count": len(hostile)},
    )
    parsed = csv_rows(csv_output.render_csv(data))[1:]
    assert [row[0] for row in parsed] == ["'=1+1", "'+CMD", "'-2", "'@SUM(A1)", "'\tTAB", "'\rCR"]
    assert {(row[1], row[2]) for row in parsed} == {("-2.00", "-2")}

    sheet = openpyxl.load_workbook(
        io.BytesIO(xlsx_output.render_xlsx(data, sample_stamp()))
    ).worksheets[0]
    grid = next(row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "Memo")
    memos = [sheet.cell(grid + offset, 1) for offset in range(1, len(hostile) + 1)]
    # openpyxl leaves the OOXML escape of a carriage return in the value.
    assert [(cell.data_type, str(cell.value).replace("_x000D_", "\r")) for cell in memos] == [
        ("s", "'=1+1"),
        ("s", "'+CMD"),
        ("s", "'-2"),
        ("s", "'@SUM(A1)"),
        ("s", "'\tTAB"),
        ("s", "'\rCR"),
    ]
    for offset in range(1, len(hostile) + 1):
        assert (sheet.cell(grid + offset, 2).data_type, sheet.cell(grid + offset, 2).value) == (
            "n",
            -2,
        )
        assert (sheet.cell(grid + offset, 3).data_type, sheet.cell(grid + offset, 3).value) == (
            "n",
            -2,
        )


def _values(parameters: Any) -> Iterator[Any]:
    if isinstance(parameters, Mapping):
        yield from parameters.values()
    elif isinstance(parameters, Sequence) and not isinstance(parameters, str):
        for item in parameters:
            yield from _values(item) if isinstance(item, Mapping | list | tuple) else (item,)


def test_bound_parameters_only(keyring: KeyRing, clock: FrozenClock, runtime: JobRuntime) -> None:
    someone = member(keyring, clock)
    insert_clients(someone.tenant_id, CLIENT_NAMES)
    statements: list[tuple[str, Any]] = []

    def capture(
        conn: Any, cursor: Any, statement: str, parameters: Any, context: Any, many: bool
    ) -> None:
        statements.append((statement, parameters))

    def params(name: str) -> ReportParams:
        return ReportParams(
            report_code=INVENTORY,
            report_version=1,
            parameters={"include_revoked": False},
            entity_ids=(),
            known_at=clock.now(),
            filters={"name": name},
        )

    event.listen(Engine, "before_cursor_execute", capture)
    try:
        with system_unit_of_work(
            runtime, system_principal(someone.tenant_id), request_id="tests-bound-parameters"
        ) as uow:
            injected = api_client_inventory.build(uow, params(HOSTILE))
            control = api_client_inventory.build(uow, params("svc-netsuite"))
    finally:
        event.remove(Engine, "before_cursor_execute", capture)
    assert (injected.rows, dict(injected.control_totals)) == ((), {"client_count": 0})
    assert [row["name"] for row in control.rows] == ["svc-netsuite"]
    selects = [(sql, bound) for sql, bound in statements if "api_client" in sql]
    assert len(selects) == 2
    assert all(HOSTILE not in sql for sql, _ in selects)
    assert HOSTILE in list(_values(selects[0][1]))


def test_failed_run_keeps_record(
    app: FastAPI, keyring: KeyRing, clock: FrozenClock, runtime: JobRuntime
) -> None:
    """RV-14: after the last of its 2 attempts a run that cannot build ends FAILED with the job's
    problem; the first attempt leaves it RUNNING for the retry."""
    stamp_test_release()
    someone = member(keyring, clock)
    tenant_id = someone.tenant_id
    # A real admitted entity (Codex 0727 (3)): with `entity_ids` empty the legacy view refuses
    # `entity_codes` first (views.NO_ENTITY) and the DELTA-of-LEGACY branch is never reached.
    maya = holding(app, someone, "revenue_accountant")
    created = entity(app, maya, code="AVM-US", calendar_id=calendar(app, maya))
    with tenant_session(_db(tenant_id)) as session:
        release_id = session.execute(
            select(engine_release.c.id).where(engine_release.c.build_sha == RELEASE_BUILD)
        ).scalar_one()
        values = report_run_values(tenant_id, engine_release_id=release_id)
        # A run that cannot build: DELTA of the LEGACY book is refused by name by the legacy view
        # (journals/views.py RULE_DELTA) once the entity is admitted, on every attempt. (RPT-05
        # gained its builder with ENG-C4 — the former "no builder" probe now builds.)
        values.update(
            report_code="legacy_je_summary",
            entity_ids=[UUID(str(created["id"]))],
            parameters={
                "entity_codes": ["AVM-US"],
                "known_at": KNOWN_AT,
                "book": "LEGACY",
                "mode": "DELTA",
            },
        )
        session.execute(insert(report_run).values(**values))
        row = registry.insert_job(
            session,
            JobKind.REPORT_RUN,
            {"report_run_id": str(values["id"])},
            tenant_id=tenant_id,
            now=clock.now(),
            created_by=someone.user_id,
            created_by_kind=PrincipalKind.USER,
        )
        registry.dispatch(
            session, job_id=row["id"], tenant_id=tenant_id, queue=str(row["queue"]), now=clock.now()
        )
    job_id = UUID(str(row["id"]))
    run_id = UUID(str(values["id"]))

    def state() -> tuple[str, str, Mapping[str, Any] | None]:
        with tenant_session(_db(tenant_id)) as session:
            current = session.execute(select(job.c.state).where(job.c.id == job_id)).scalar_one()
            stored = (
                session.execute(
                    select(report_run.c.status, report_run.c.problem).where(
                        report_run.c.id == run_id
                    )
                )
                .mappings()
                .one()
            )
        return str(current), str(stored["status"]), stored["problem"]

    run_deferred(tenant_id, job_id, runtime)
    assert state() == ("QUEUED", "RUNNING", None)
    run_deferred(tenant_id, job_id, runtime, attempt=2)
    job_state, run_status, problem = state()
    assert (job_state, run_status) == ("FAILED", "FAILED")
    assert problem is not None
    assert problem["type"] == PROBLEM_BASE + "validation-failed"
    # the refusal that made the run fail: DELTA of the LEGACY book (journals/views RULE_DELTA;
    # a validation-failed problem whose detail counts the fields and whose error names the rule)
    assert problem["detail"] == "1 field needs attention."
    assert [(e["field"], e["rule_id"], e["message"]) for e in problem["errors"]] == [
        ("book", views.RULE_DELTA, views.DELTA_OF_LEGACY)
    ]
    with tenant_session(_db(tenant_id)) as session:
        actions = session.scalars(
            select(audit_event.c.action)
            .where(audit_event.c.object_id == run_id)
            .order_by(audit_event.c.chain_seq)
        ).all()
    assert list(actions) == ["report_run.start", "report_run.fail"]
