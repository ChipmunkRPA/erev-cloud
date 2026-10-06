"""Period balances of contracts whose identifiers hold CV-21 delimiters — DB witnesses of supervisor
ruling R-16 (2026-09-29; ENGINE_SPEC_B S15-R-07a, S15-R-07; ENGINE_SPEC CV-21, CV-50; 03
REQ-SEC-011). DB-bound.

The defect (measured in lane FIX-D2): ``tie_outs.balances_at`` looked a member contract's balance
nodes up under the RAW external id while the engine keys them CV-21-encoded
(``contract_asset:A%3AB@AVM-US:FY2026-P08``), found nothing for an id holding ``%`` ``/`` ``@``
``#`` or ``:``, and silently answered the stored ``contract_version_balance`` column — the
version's LATEST period. ``A:B`` showed its September balance 89,753.42 at August; the September
rollforward opened it there, showed OTHER −9,863.01 and failed ``TO_ROLLFORWARD_BALANCES``.

The world: WLD-K-01 (``SF-ORD-10001``) plus, for the K-01 customer in AVM-US,

- EIGHT identical contracts — one 2026 PLATFORM line of 120,000.00 USD, ratable daily, no billing
  — under the ids ``A``, ``=A``, ``'=A`` and one id per CV-21 delimiter (``A:B`` ``A/B`` ``A@B``
  ``A#B`` ``A%B``). Revenue to date is 120,000.00 × days ÷ 365: 79,890.41 at 31 Aug 2026 (243
  days) and 89,753.42 at 30 Sep 2026 (273 days); unbilled, that is the contract asset, and
  September adds 9,863.01 (PRD J-05.3: "schedule Sep 2026 9,863.01");
- ONE contract of the same line billed 120,000.00 on 2026-01-01 (the shape of WLD-K-02) under the
  id ``SO:2026/0042``: contract liability 40,109.59 at August and 30,246.58 at September — with
  K-01's 39,708.49 and 29,944.11 (WLD-X-03) the 79,818.08 and 60,190.69 that
  ``test_home_consistency_k01`` pins for K-01 and K-02.

The clock is moved onto the record-time clock (``support.record_clock``): a version is known at
its events' server stamp, so the reads and freezes below take the cutoff after it.
"""

from __future__ import annotations

import csv
import dataclasses
import io
from collections.abc import Callable, Iterator, Mapping
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Final
from uuid import UUID

import openpyxl
import pytest
from erev_api.auth.keyring import KeyRing
from erev_api.clock import FrozenClock
from erev_api.config import Settings
from erev_api.db.session import DbContext, tenant_session
from erev_api.db.tables import file_object, job, legal_entity, lock_snapshot, period
from erev_api.db.tables import report_run as report_run_table
from erev_api.domain.close import snapshots as close_snapshots
from erev_api.domain.reports import locked, snapshots, tie_outs
from erev_api.domain.reports.outputs import utc_text
from erev_api.enums import ContractEventType
from erev_api.events.payloads import BillingRecordedV1
from erev_api.events.stream import EventIn
from erev_api.explain import store
from erev_api.files.store import LocalFileStore
from erev_api.jobs.registry import run_job
from erev_api.main import create_app
from erev_api.money import MoneyIn
from erev_engine.stages.s01_canonicalize import contract_entity_subject_key
from erev_engine.trace import Trace
from fastapi import FastAPI
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session
from support.db import TestDatabase
from support.factories import appended, booked_contract, computed, open_periods
from support.record_clock import on_record_time
from support.reference import entity, get, post
from support.worlds import (
    AUGUST_2026,
    AVM_US,
    JOBS,
    K01,
    REPORT_RUN_ID_HEADER,
    REPORT_RUNS,
    SEPTEMBER_2026,
    ReportWorld,
    k01_body,
    k01_pellworth,
    run_now,
)
from support.worlds import report_run as run_report

BOOK: Final = "ASC606"
HOME: Final = "/api/v1/dashboard/home"
# the formula-trigger pair, a plain id, and one id per CV-21 delimiter (`:` `/` `@` `#` `%`)
UNBILLED: Final = ("A", "=A", "'=A", "A:B", "A/B", "A@B", "A#B", "A%B")
DELIMITED: Final = ("A:B", "A/B", "A@B", "A#B", "A%B")
BILLED: Final = "SO:2026/0042"
AUGUST_ASSET: Final = "79890.41"  # 120,000.00 × 243 ÷ 365
SEPTEMBER_ASSET: Final = "89753.42"  # 120,000.00 × 273 ÷ 365
SEPTEMBER_REVENUE: Final = "9863.01"  # PRD J-05.3
BILLED_AUGUST: Final = "40109.59"  # 120,000.00 − 79,890.41
BILLED_SEPTEMBER: Final = "30246.58"  # 120,000.00 − 89,753.42
K01_AUGUST: Final = "39708.49"  # PRD WLD-X-03
K01_SEPTEMBER: Final = "29944.11"  # PRD WLD-X-03
K01_SEPTEMBER_REVENUE: Final = "9764.38"  # PRD WLD-X-02
ZERO: Final = "0.00"
PASS: Final = "PASS"
RULE: Final = "S15-R-07a"
_FETCHED: Final = text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id")
D = Decimal


@pytest.fixture
def app(committed_db: TestDatabase, app_settings: Settings, clock: FrozenClock) -> FastAPI:
    return create_app(app_settings, clock=clock)


@dataclass(frozen=True)
class Scene:
    world: ReportWorld
    cutoff: datetime  # the record-time clock: every version of the world is known by it


def _body(world: ReportWorld, external_id: str) -> dict[str, Any]:
    """A contract of the K-01 customer with ONE 2026 PLATFORM line of 120,000.00 USD."""
    customer = world.contracts[K01].contract["customer_id"]
    body = k01_body(UUID(str(customer)))
    body["external_id"] = external_id
    body["lines"] = [body["lines"][0]]
    return body


@pytest.fixture
def scene(app: FastAPI, keyring: KeyRing, clock: FrozenClock, app_settings: Settings) -> Scene:
    world = k01_pellworth(app, keyring, clock, LocalFileStore(app_settings.file_root))
    for external_id in (*UNBILLED, BILLED):
        found = booked_contract(world.place, _body(world, external_id), activate=True)
        if external_id == BILLED:
            invoiced = date(2026, 1, 1)
            billing = BillingRecordedV1(
                invoice_number="INV-US-2042",
                line_external_id="INV-US-2042-1",
                obligation_key="O1",
                amount=MoneyIn(amount="120000.00", currency="USD"),
                issue_date=invoiced,
            )
            event = EventIn(
                event_type=ContractEventType.BILLING_RECORDED,
                effective_date=invoiced,
                payload=billing,
            )
            appended(world.place, UUID(str(found.contract["id"])), 2, [event])
        computed(world.place, UUID(str(found.combination_group["id"])))
    world = on_record_time(world)
    return Scene(world=world, cutoff=world.place.clock.now())


def _prerequisite(condition: bool, what: str) -> None:
    if not condition:
        raise RuntimeError(f"prerequisite not met (not the balance read under test): {what}")


def usd(amount: str) -> dict[str, str]:
    return {"amount": amount, "currency": "USD"}


def _request(scene: Scene, selectors: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "entity_codes": [AVM_US],
        "book": BOOK,
        **selectors,
        "known_at": utc_text(scene.cutoff),  # the cutoff = the clock: never future
    }


def _live(
    scene: Scene, code: str, selectors: Mapping[str, Any], *, output_format: str = "JSON"
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    run, rows = run_report(
        scene.world, code, _request(scene, selectors), output_format=output_format
    )
    _prerequisite(run["status"] == "SUCCEEDED", f"live {code} run: {run}")
    return run, rows


def _scope(scene: Scene, period_key: str) -> snapshots.SnapshotScope:
    world = scene.world
    period_id = world.place.scalar(
        select(period.c.id)
        .join(legal_entity, legal_entity.c.calendar_id == period.c.calendar_id)
        .where(legal_entity.c.id == world.entity_id, period.c.period_key == period_key)
    )
    return snapshots.SnapshotScope(
        entity_id=world.entity_id,
        book_code=BOOK,
        period_id=UUID(str(period_id)),
        known_at=scene.cutoff,
    )


def _frozen_rows(scene: Scene, kind: str, period_key: str) -> list[dict[str, str]]:
    """The rows the lock of ``period_key`` would freeze for ``kind`` at the cutoff (the registry's
    own builder, S15-R-18), as their canonical texts."""
    with scene.world.place.uow() as uow:
        encoded = snapshots.SNAPSHOT_DATASETS[kind](uow, _scope(scene, period_key))
        uow.commit()  # the freeze writes nothing itself
    _, rows = locked._rows(encoded.content)
    assert len(rows) == encoded.row_count
    return rows


def _by_contract(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Rows keyed by their contract id; line rows and totals (no contract) are left out."""
    found = {str(row["contract_external_id"]): row for row in rows if row["contract_external_id"]}
    assert len(found) == sum(1 for row in rows if row["contract_external_id"])  # ids stay apart
    return found


def _ties(run: Mapping[str, Any]) -> dict[str, str]:
    return {str(item["code"]): str(item["result"]) for item in run["tie_out_results"]}


# --- an earlier period's balances are each contract's own -----------------------------------------


def test_r16_august_balances_are_each_contracts_own_whatever_its_id_holds(scene: Scene) -> None:
    """``contract_balances`` for FY2026-P08, asked in September: the eight identical contracts
    each report the contract asset 79,890.41 — live and in the dataset an August lock freezes —
    and the billed one its August liability. Fail-first: the five delimiter-bearing ids and the
    billed contract answered their SEPTEMBER figures (89,753.42; 30,246.58)."""
    run, rows = _live(scene, "contract_balances", {"period_key": AUGUST_2026})
    live = _by_contract(rows)
    assert set(live) == {K01, BILLED, *UNBILLED}
    for external_id in UNBILLED:
        assert live[external_id]["contract_asset"] == usd(AUGUST_ASSET), external_id
        assert live[external_id]["contract_liability"] == usd(ZERO), external_id
        assert live[external_id]["unbilled_receivable"] == usd(ZERO), external_id
    assert live[BILLED]["contract_liability"] == usd(BILLED_AUGUST)
    assert live[BILLED]["contract_asset"] == usd(ZERO)
    assert live[K01]["contract_liability"] == usd(K01_AUGUST)
    assert run["control_totals"]["contract_asset"] == {"USD": f"{8 * D(AUGUST_ASSET)}"}
    assert run["control_totals"]["contract_liability"] == {
        "USD": f"{D(K01_AUGUST) + D(BILLED_AUGUST)}"
    }
    assert _ties(run)[tie_outs.TO_BALANCES_EQ_ROLLFORWARD] == PASS

    frozen = _by_contract(_frozen_rows(scene, "CONTRACT_BALANCES", AUGUST_2026))
    assert set(frozen) == set(live)
    assert {key: frozen[key]["contract_asset"] for key in UNBILLED} == dict.fromkeys(
        UNBILLED, AUGUST_ASSET
    )
    assert frozen[BILLED]["contract_liability"] == BILLED_AUGUST
    assert frozen[K01]["contract_liability"] == K01_AUGUST

    # the period the versions were computed through is unchanged by the fix
    _, september = _live(scene, "contract_balances", {"period_key": SEPTEMBER_2026})
    latest = _by_contract(september)
    assert {key: latest[key]["contract_asset"] for key in UNBILLED} == dict.fromkeys(
        UNBILLED, usd(SEPTEMBER_ASSET)
    )
    assert latest[BILLED]["contract_liability"] == usd(BILLED_SEPTEMBER)
    assert latest[K01]["contract_liability"] == usd(K01_SEPTEMBER)


# --- the September rollforward opens at the true August closing -----------------------------------


def _line_rows(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {str(row["line_code"]): row for row in rows if row["line_code"]}


def test_r16_september_rollforward_opens_every_contract_at_its_august_closing(
    scene: Scene,
) -> None:
    """S15-R-07: opening + lines = closing and nothing is ``OTHER`` — live for the asset role and
    the liability role, and in the dataset a September lock freezes. Fail-first: each
    delimiter-bearing contract opened at its September figure, ``OTHER`` was −9,863.01 per asset
    contract (+9,863.01 on the billed liability) and ``TO_ROLLFORWARD_BALANCES`` failed."""
    september = {"from_period_key": SEPTEMBER_2026, "to_period_key": SEPTEMBER_2026}
    run, rows = _live(
        scene, "contract_balance_rollforward", {**september, "balance_role": "CONTRACT_ASSET"}
    )
    assets = _by_contract(rows)
    assert set(assets) == {K01, BILLED, *UNBILLED}
    for external_id in UNBILLED:
        row = assets[external_id]
        assert row["opening"] == usd(AUGUST_ASSET), external_id
        assert row["revenue_from_period_billings"] == usd(SEPTEMBER_REVENUE), external_id
        assert row["other"] == usd(ZERO), external_id
        assert row["closing"] == usd(SEPTEMBER_ASSET), external_id
        for column in ("billings", "revenue_from_opening", "reclassifications"):
            assert row[column] == usd(ZERO), (external_id, column)
    assert _ties(run) == {
        tie_outs.TO_ROLLFORWARD_BALANCES: PASS,
        tie_outs.TO_BALANCES_EQ_ROLLFORWARD: PASS,
        tie_outs.TO_ROLLFORWARD_EQ_GL: tie_outs.NOT_APPLICABLE,
    }

    default_run, default_rows = _live(scene, "contract_balance_rollforward", september)
    liabilities = _by_contract(default_rows)  # the default role: the contract liability
    assert liabilities[BILLED]["opening"] == usd(BILLED_AUGUST)
    assert liabilities[BILLED]["revenue_from_opening"] == usd(f"-{SEPTEMBER_REVENUE}")
    assert liabilities[BILLED]["other"] == usd(ZERO)
    assert liabilities[BILLED]["closing"] == usd(BILLED_SEPTEMBER)
    assert liabilities[K01]["opening"] == usd(K01_AUGUST)
    assert liabilities[K01]["closing"] == usd(K01_SEPTEMBER)
    assert _ties(default_run)[tie_outs.TO_ROLLFORWARD_BALANCES] == PASS

    opening_assets = f"{8 * D(AUGUST_ASSET)}"  # 639,123.28
    closing_assets = f"{8 * D(SEPTEMBER_ASSET)}"  # 718,027.36
    opening_liability = f"{D(K01_AUGUST) + D(BILLED_AUGUST)}"  # 79,818.08
    closing_liability = f"{D(K01_SEPTEMBER) + D(BILLED_SEPTEMBER)}"  # 60,190.69
    expected_lines = {
        "OPENING": (opening_liability, opening_assets),
        "BILLINGS": (ZERO, ZERO),
        "REVENUE_FROM_OPENING": (f"-{D(K01_SEPTEMBER_REVENUE) + D(SEPTEMBER_REVENUE)}", ZERO),
        "REVENUE_FROM_PERIOD_BILLINGS": (ZERO, f"{8 * D(SEPTEMBER_REVENUE)}"),
        "RECLASSIFICATIONS": (ZERO, ZERO),
        "FX_REMEASUREMENT": (ZERO, ZERO),
        "BUSINESS_COMBINATIONS": (ZERO, ZERO),
        "OTHER": (ZERO, ZERO),
        "CLOSING": (closing_liability, closing_assets),
    }
    live_lines = _line_rows(default_rows)
    assert {
        code: (row["contract_liability"]["amount"], row["contract_asset"]["amount"])
        for code, row in live_lines.items()
    } == expected_lines

    frozen = _frozen_rows(scene, "CONTRACT_BALANCE_ROLLFORWARD", SEPTEMBER_2026)
    frozen_lines = _line_rows(frozen)
    assert {
        code: (row["contract_liability"], row["contract_asset"])
        for code, row in frozen_lines.items()
    } == expected_lines
    assert {row["unbilled_receivable"] for row in frozen_lines.values()} == {ZERO}
    # S15-R-07 restated over the frozen rows (a frozen dataset stores no tie-out result):
    # opening + the explained lines = closing for every balance, and OTHER is 0
    for balance in tie_outs.ROLLFORWARD_BALANCES:
        explained = sum(
            (D(row[balance]) for code, row in frozen_lines.items() if code != "CLOSING"), D(0)
        )
        assert explained == D(frozen_lines["CLOSING"][balance]), balance
        assert D(frozen_lines["OTHER"][balance]) == 0, balance
    by_contract = _by_contract(frozen)  # frozen under the default role: the contract liability
    assert set(by_contract) == {K01, BILLED, *UNBILLED}
    assert (by_contract[BILLED]["opening"], by_contract[BILLED]["closing"]) == (
        BILLED_AUGUST,
        BILLED_SEPTEMBER,
    )
    assert by_contract[BILLED]["other"] == ZERO
    # the August closings a lock of August freezes are the openings September reads
    august = _by_contract(_frozen_rows(scene, "CONTRACT_BALANCES", AUGUST_2026))
    assert f"{sum((D(row['contract_asset']) for row in august.values()), D(0))}" == opening_assets
    assert (
        f"{sum((D(row['contract_liability']) for row in august.values()), D(0))}"
        == opening_liability
    )

    # the dashboard's liability movement reads the same two period ends
    shown = get(
        scene.world.app,
        HOME,
        scene.world.maya,
        {"entity": AVM_US, "period": SEPTEMBER_2026, "book": BOOK},
    )
    assert shown.status_code == 200, shown.text
    assert shown.json()["contract_liability"] == {
        "closing": usd(closing_liability),
        "opening": usd(opening_liability),
    }


# --- the entity code is a key component too -------------------------------------------------------


def test_r16_an_entity_code_holding_a_delimiter_is_read_under_the_engines_key(
    scene: Scene,
) -> None:
    """The engine encodes BOTH components of ``<contract>@<entity>`` (CV-21;
    ``test_l9_run_q3_position_subjects_are_cv_21_encoded``), and 04 TY-06 admits ``:`` ``/`` and
    ``#`` in an entity code: a PLAIN contract id in such an entity is read under
    ``PLAIN-01@AVM%3AUK%2F1%23``. Fail-first: the raw prefix ``PLAIN-01@AVM:UK/1#`` found no node
    and every contract of the entity answered its latest balance for an earlier period."""
    world = scene.world
    code = "AVM:UK/1#"
    calendar_id = world.place.scalar(
        select(legal_entity.c.calendar_id).where(legal_entity.c.id == world.entity_id)
    )
    entity(world.app, world.maya, code=code, calendar_id=str(calendar_id))  # AVM-US's calendar
    open_periods(
        world.app,
        world.maya,
        entity_code=code,
        keys=[f"FY2026-P{month:02d}" for month in range(1, 10)],
    )
    body = _body(world, "PLAIN-01")
    body["contracting_entity_code"] = code
    found = booked_contract(world.place, body, activate=True)
    computed(world.place, UUID(str(found.combination_group["id"])))
    there = Scene(world=on_record_time(world), cutoff=world.place.clock.now())
    _prerequisite(
        contract_entity_subject_key("PLAIN-01", code) == "PLAIN-01@AVM%3AUK%2F1%23",
        "the engine's subject key of the member",
    )

    def request(selectors: Mapping[str, Any]) -> dict[str, Any]:
        return {**_request(there, selectors), "entity_codes": [code]}

    run, rows = run_report(there.world, "contract_balances", request({"period_key": AUGUST_2026}))
    assert run["status"] == "SUCCEEDED", run
    (row,) = [item for item in rows if item["contract_external_id"]]
    assert (row["contract_external_id"], row["entity_code"]) == ("PLAIN-01", code)
    assert row["contract_asset"] == usd(AUGUST_ASSET)
    rolled, rolled_rows = run_report(
        there.world,
        "contract_balance_rollforward",
        request(
            {
                "from_period_key": SEPTEMBER_2026,
                "to_period_key": SEPTEMBER_2026,
                "balance_role": "CONTRACT_ASSET",
            }
        ),
    )
    assert rolled["status"] == "SUCCEEDED", rolled
    member = _by_contract(rolled_rows)["PLAIN-01"]
    assert (member["opening"], member["other"], member["closing"]) == (
        usd(AUGUST_ASSET),
        usd(ZERO),
        usd(SEPTEMBER_ASSET),
    )
    assert _ties(rolled)[tie_outs.TO_ROLLFORWARD_BALANCES] == PASS


# --- a balance the trace cannot answer is refused by name -----------------------------------------


def _settled_run(scene: Scene, code: str, selectors: Mapping[str, Any]) -> dict[str, Any]:
    """``POST /report-runs`` and the worker's attempts until the job settles (RV-14: a refused
    build ends FAILED after the second of REPORT_RUN_RETRY's attempts); returns API-S-ReportRun."""
    world = scene.world
    started = post(
        world.app,
        REPORT_RUNS,
        world.maya,
        {"report_code": code, "parameters": _request(scene, selectors), "output_format": "JSON"},
    )
    _prerequisite(started.status_code == 202, started.text)
    job_id = UUID(str(started.json()["id"]))
    state = run_now(world, job_id)
    if state["state"] not in ("SUCCEEDED", "FAILED"):
        context = DbContext(tenant_id=world.tenant_id, user_id=None, entity_scope="*")
        with tenant_session(context) as session:
            task_id = session.execute(
                select(job.c.procrastinate_job_id).where(job.c.id == job_id)
            ).scalar_one()
            session.execute(_FETCHED, {"id": task_id})
        run_job(job_id, world.tenant_id, attempt=2, runtime=world.runtime)
        state = dict(get(world.app, f"{JOBS}/{job_id}", world.maya).json())
    _prerequisite(state["state"] in ("SUCCEEDED", "FAILED"), f"settled job: {state}")
    shown = get(world.app, f"{REPORT_RUNS}/{started.headers[REPORT_RUN_ID_HEADER]}", world.maya)
    _prerequisite(shown.status_code == 200, shown.text)
    return dict(shown.json())


def _refused(scene: Scene, run: Mapping[str, Any], *named: str) -> str:
    """The run ended FAILED with the S15-R-07a refusal naming ``A:B`` and ``named``, and produced
    no output; returns the refusal's message."""
    assert run["status"] == "FAILED", run
    problem = run["problem"]
    assert problem["type"].endswith("/validation-failed") and problem["status"] == 422
    (error,) = problem["errors"]
    assert (error["field"], error["rule_id"]) == ("balances[A:B@AVM-US]", RULE)
    for fact in ("Contract A:B", "entity AVM-US", *named):
        assert fact in error["message"], fact
    stored = scene.world.place.rows(
        select(report_run_table.c.output_file_id, report_run_table.c.row_count).where(
            report_run_table.c.id == UUID(str(run["id"]))
        )
    )
    assert [(row["output_file_id"], row["row_count"]) for row in stored] == [(None, None)]
    return str(error["message"])


def _balances_at(scene: Scene, period_key: str) -> tuple[tie_outs.BalanceRow, ...]:
    """``tie_outs.balances_at`` of AVM-US at the end of ``period_key`` under the scene's cutoff."""
    world = scene.world
    with world.place.uow() as uow:
        found = tie_outs.entities(uow.session, (world.entity_id,))
        calendar = tie_outs.calendars(uow.session, found)[world.entity_id]
        return tie_outs.balances_at(
            uow.session,
            entity_ids=(world.entity_id,),
            book_code=BOOK,
            period_keys={world.entity_id: calendar.named(period_key, "period_key")},
            cutoff=scene.cutoff,
        )


@pytest.fixture
def stored_traces(monkeypatch: pytest.MonkeyPatch) -> Iterator[Callable[..., None]]:
    """Substitute the traces the readers load: ``alter(keep)`` drops every stored node whose id
    ``keep`` rejects. The engine writes every member node (DG-KRN-EXP-01), so a trace that cannot
    answer a period exists only as a fault — a writer and a reader disagreeing on the key, a
    damaged row — which is what the refusal is for."""
    load = store.load_trace

    def alter(keep: Callable[[str], bool]) -> None:
        def loaded(session: Session, contract_version_id: UUID) -> Trace | None:
            trace = load(session, contract_version_id)
            if trace is None:
                return None
            return dataclasses.replace(
                trace, nodes=tuple(node for node in trace.nodes if keep(node.id))
            )

        monkeypatch.setattr(store, "load_trace", loaded)

    yield alter
    monkeypatch.setattr(store, "load_trace", load)


def test_r16_a_balance_the_trace_cannot_answer_is_refused_by_name_on_every_path(
    scene: Scene, stored_traces: Callable[..., None]
) -> None:
    """Fail closed: nothing is answered from the version's latest stored figures.

    (1) ``A:B``'s trace holds NO node under its subject key: the August balances run, the
    September rollforward run and the August freeze are refused by name — contract, entity,
    contract version, period — with no output, no file and no snapshot row.
    (2) ``A:B``'s trace holds its other balance nodes but no ``contract_asset`` node: the stored
    89,753.42 is the September figure — the read answers September with it and refuses August by
    name (the September REPORT also reads the August opening for its tie-out, so it is the read
    itself that is asked for September).
    """
    world = scene.world
    subject = contract_entity_subject_key("A:B", AVM_US)
    _prerequisite(subject == "A%3AB@AVM-US", "the engine's subject key of A:B in AVM-US")
    september = {"from_period_key": SEPTEMBER_2026, "to_period_key": SEPTEMBER_2026}

    stored_traces(lambda node_id: f":{subject}:" not in node_id)
    message = _refused(
        scene,
        _settled_run(scene, "contract_balances", {"period_key": AUGUST_2026}),
        AUGUST_2026,
        subject,
    )
    assert SEPTEMBER_ASSET not in message  # no figure is offered in place of the balance
    _refused(
        scene,
        _settled_run(scene, "contract_balance_rollforward", september),
        AUGUST_2026,  # the opening: the period before the range
        subject,
    )
    files_before = int(world.place.scalar(select(func.count()).select_from(file_object)))
    rows_before = int(world.place.scalar(select(func.count()).select_from(lock_snapshot)))
    scope = _scope(scene, AUGUST_2026)
    with world.place.uow() as uow, pytest.raises(tie_outs.BalanceUnreadable) as refused:
        close_snapshots.freeze_datasets(
            uow, scope.entity_id, scope.book_code, scope.period_id, scope.known_at
        )
    (error,) = refused.value.errors
    assert (error.field, error.rule_id) == ("balances[A:B@AVM-US]", RULE)
    assert AUGUST_2026 in error.message
    assert int(world.place.scalar(select(func.count()).select_from(file_object))) == files_before
    assert int(world.place.scalar(select(func.count()).select_from(lock_snapshot))) == rows_before
    denied = get(
        world.app, HOME, world.maya, {"entity": AVM_US, "period": SEPTEMBER_2026, "book": BOOK}
    )
    assert denied.status_code == 422, denied.text
    assert [item["rule_id"] for item in denied.json()["errors"]] == [RULE]

    stored_traces(lambda node_id: not node_id.startswith(f"contract_asset:{subject}:"))
    at_september = {row.external_id: row for row in _balances_at(scene, SEPTEMBER_2026)}
    assert at_september["A:B"].value("contract_asset") == D(SEPTEMBER_ASSET)  # its own period
    with pytest.raises(tie_outs.BalanceUnreadable):
        _balances_at(scene, AUGUST_2026)
    message = _refused(
        scene,
        _settled_run(scene, "contract_balances", {"period_key": AUGUST_2026}),
        f"contract_asset {SEPTEMBER_ASSET} USD",
        f"latest period {SEPTEMBER_2026}",
        f"earlier period {AUGUST_2026}",
    )
    assert "not answered for an earlier period" in message


# --- the export boundary is unchanged -------------------------------------------------------------


def _output(scene: Scene, run: Mapping[str, Any]) -> bytes:
    shown = get(scene.world.app, f"{REPORT_RUNS}/{run['id']}/output", scene.world.maya)
    _prerequisite(shown.status_code == 200, shown.text)
    return bytes(shown.content)


def test_r16_the_export_guard_writes_the_formula_ids_as_before(scene: Scene) -> None:
    """REQ-SEC-011 is untouched by the balance read: the CSV and XLSX exports of the August
    balances write the contract ``=A`` with the leading apostrophe and ``'=A`` as it is — two rows
    reading ``'=A``, each 79,890.41 — and an id whose delimiter is not a leading formula trigger
    (``A@B``) unguarded; the JSON rows keep the two admitted ids apart."""
    selectors = {"period_key": AUGUST_2026}
    _, rows = _live(scene, "contract_balances", selectors)
    assert {"=A", "'=A"} <= set(_by_contract(rows))

    exported, _ = _live(scene, "contract_balances", selectors, output_format="CSV")
    table = list(csv.reader(io.StringIO(_output(scene, exported).decode("utf-8"))))
    header = next(index for index, row in enumerate(table) if row and row[0] == "Contract")
    asset = table[header].index("Contract asset (USD)")  # RPT-R-03: one currency in the column
    written = [(row[0], row[asset]) for row in table[header + 1 :] if row and row[0]]
    assert [item for item in written if item[0].lstrip("'").startswith("=")] == [
        ("'=A", AUGUST_ASSET),
        ("'=A", AUGUST_ASSET),
    ]
    assert [cell for cell, _ in written if cell.startswith("A")] == sorted(
        key for key in UNBILLED if key.startswith("A")
    )

    workbook, _ = _live(scene, "contract_balances", selectors, output_format="XLSX")
    sheet = openpyxl.load_workbook(io.BytesIO(_output(scene, workbook))).worksheets[0]
    grid = next(
        index for index in range(1, sheet.max_row + 1) if sheet.cell(index, 1).value == "Contract"
    )
    columns = [sheet.cell(grid, index).value for index in range(1, sheet.max_column + 1)]
    money = columns.index("Contract asset (USD)") + 1
    cells = [
        (sheet.cell(index, 1).data_type, sheet.cell(index, 1).value, sheet.cell(index, money).value)
        for index in range(grid + 1, sheet.max_row + 1)
        if str(sheet.cell(index, 1).value or "").lstrip("'").startswith("=")
    ]
    assert [(kind, value, D(str(amount))) for kind, value, amount in cells] == [
        ("s", "'=A", D(AUGUST_ASSET)),
        ("s", "'=A", D(AUGUST_ASSET)),
    ]
