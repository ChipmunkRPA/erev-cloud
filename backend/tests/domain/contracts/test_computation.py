"""Computed versions, schedules and calculation traces (04 T-CON-07 to T-CON-09, T-CON-11, T-ENG-01
to T-ENG-03, §14.1 DB-17; dev-guide §6.1 DG-CMD-04, DG-CMD-10, §5.15 DG-KRN-REG-02, 03, §5.16
DG-KRN-EXP-05; 05 RCP-15, RCP-21; PRD WLD-K-01, WLD-X-01, WLD-X-02; 03 REQ-REF-015, REQ-SSP-011,
REQ-SSP-014, REQ-REC-020, REQ-CON-014; CTL-011; BUILD_SPEC CTR-2).

Maya holds Revenue Accountant and SSP Analyst and prepares every configuration; Priya (SSP Approver)
approves the SSP book and Marcus (Controller) the templates and the mapping, both enrolled in MFA.
The reference data of PRD §2.6 is created through the RFD routes: a January calendar for 2026 and
2027 whose AVM-US periods FY2026-P01 to P09 are open, customer C-01, templates TPL-SUB-DAILY and
TPL-SVC-PCT (effective 2026-01-01) as the defaults of AVM-PLAT-ENT and AVM-IMPL-STD, US-LIST 2026-H1
(AVM-PLAT-ENT observable point 132,000.00; AVM-IMPL-STD cost plus margin point 18,000.00) and
AVM-MAP-2026-01. K-01 (PRD WLD-K-01, ``SF-ORD-10001``) is booked and activated through the contract
commands. ``support.factories.engine()`` computes: the ENGINE_SPEC §0.5 fake until END-9 lands in
this checkout, ``erev_engine.compute`` in the L3 merge gate (V-C). The frozen clock reads
2026-09-12T12:00:00Z.
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Mapping
from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal
from types import MappingProxyType
from typing import Any
from uuid import UUID

import pytest
from erev_api.auth.keyring import KeyRing
from erev_api.auth.principal import system_principal
from erev_api.clock import FrozenClock
from erev_api.config import Settings
from erev_api.db.session import DbContext, tenant_session
from erev_api.db.tables import (
    audit_event,
    calc_trace,
    combination_group,
    computation_evidence,
    contract,
    contract_computation,
    contract_event,
    contract_version,
    file_object,
    ledger_chain_head,
    obligation_version,
    period,
    product,
    registry_version,
    schedule,
    schedule_line,
    ssp_entry,
    ssp_range,
    subledger_posting,
)
from erev_api.domain.contracts import bundles, computation
from erev_api.domain.contracts import computation_evidence as evidence_writer
from erev_api.domain.contracts.replay_verify import Trigger, VerificationRow
from erev_api.domain.platform import file_evidence
from erev_api.domain.reports import computation_sources
from erev_api.enums import ComputationTrigger, ContractEventType
from erev_api.events.payloads import SignificantChangeFlaggedV1
from erev_api.events.stream import EventIn, append_events
from erev_api.explain import store
from erev_api.files.store import LocalFileStore, open_file, store_file
from erev_api.jobs.context import JobContext
from erev_api.main import create_app
from erev_api.problems import Problem
from erev_api.registry.policies import POLICY_PARAMETERS
from erev_engine import ENGINE_VERSION, upgrade
from erev_engine.bundle import OutputBundle
from erev_engine.errors import EngineError
from erev_engine.trace import reevaluate
from fastapi import FastAPI
from sqlalchemy import exc, func, insert, select, update
from support.db import TestDatabase
from support.factories import (
    TPL_SUB_DAILY,
    TPL_SVC_PCT,
    Workspace,
    approved_ssp_version,
    booked_contract,
    computed,
    customer_id,
    engine,
    point_entry,
    product_with_template,
    published_mapping,
    published_template,
    set_default_template,
    ssp_book,
    stamp_test_release,
    workspace,
    world_calendar,
)
from support.principals import Actor, colleague, enrolled, member
from support.reference import assign, get, holding

PLATFORM = "AVM-PLAT-ENT"
IMPLEMENTATION = "AVM-IMPL-STD"
PLATFORM_LINE = {
    "obligation_key": "POB-01",
    "product_code": PLATFORM,
    "quantity": "1",
    "total_price": "120000.00",
    "start_date": "2026-01-01",
    "end_date": "2026-12-31",
}
IMPLEMENTATION_LINE = {
    "obligation_key": "POB-01",
    "product_code": IMPLEMENTATION,
    "quantity": "1",
    "total_price": "15000.00",
}


@dataclass(frozen=True, slots=True)
class World:
    place: Workspace
    priya: Actor
    marcus: Actor
    calendar_id: str
    entity_id: UUID
    customer_id: UUID
    templates: Mapping[str, Mapping[str, str]]
    book_id: str
    h1_id: str
    mapping_id: str

    @property
    def app(self) -> FastAPI:
        return self.place.app


@pytest.fixture
def app(committed_db: TestDatabase, app_settings: Settings, clock: FrozenClock) -> FastAPI:
    return create_app(app_settings, clock=clock)


@pytest.fixture
def world(app: FastAPI, keyring: KeyRing, clock: FrozenClock, app_settings: Settings) -> World:
    maya_member = member(keyring, clock)
    assign(maya_member, "revenue_accountant")
    maya = holding(app, maya_member, "ssp_analyst")
    approvers: dict[str, Actor] = {}
    for name, roles in (("priya", ("ssp_approver",)), ("marcus", ("controller", "ssp_approver"))):
        someone = colleague(maya_member.tenant_id, name)
        for role in roles:
            assign(someone, role)
        approvers[name] = enrolled(app, clock, someone)
    calendar_id, entity_id = world_calendar(app, maya)
    buyer = customer_id(app, maya, code="C-01", name="Pellworth Logistics Inc. (Demo)")
    products = {
        PLATFORM: product_with_template(
            app,
            maya,
            code=PLATFORM,
            name="Platform, enterprise tier",
            revenue_category="SUBSCRIPTION",
        ),
        IMPLEMENTATION: product_with_template(
            app, maya, code=IMPLEMENTATION, name="Implementation", revenue_category="SERVICES"
        ),
    }
    templates = {
        "TPL-SUB-DAILY": published_template(
            app,
            maya,
            approvers["marcus"],
            code="TPL-SUB-DAILY",
            outputs=TPL_SUB_DAILY,
            case_line=PLATFORM_LINE,
        ),
        "TPL-SVC-PCT": published_template(
            app,
            maya,
            approvers["marcus"],
            code="TPL-SVC-PCT",
            outputs=TPL_SVC_PCT,
            case_line=IMPLEMENTATION_LINE,
        ),
    }
    set_default_template(app, maya, products[PLATFORM], templates["TPL-SUB-DAILY"]["template_id"])
    set_default_template(
        app, maya, products[IMPLEMENTATION], templates["TPL-SVC-PCT"]["template_id"]
    )
    book_id = ssp_book(app, maya)
    h1_id = approved_ssp_version(
        app,
        maya,
        [approvers["priya"]],
        book_id,
        label="2026-H1",
        effective_from="2026-01-01",
        entries=[
            point_entry(PLATFORM, "132000.00", value_basis="AMOUNT"),  # series: D-97 (3a)
            point_entry(IMPLEMENTATION, "18000.00", "cost_plus_margin"),
        ],
    )
    mapping_id = published_mapping(app, maya, approvers["marcus"])
    stamp_test_release()
    place = workspace(app, clock, keyring, LocalFileStore(app_settings.file_root), maya)
    return World(
        place=place,
        priya=approvers["priya"],
        marcus=approvers["marcus"],
        calendar_id=calendar_id,
        entity_id=entity_id,
        customer_id=buyer,
        templates=templates,
        book_id=book_id,
        h1_id=h1_id,
        mapping_id=mapping_id,
    )


def k01(world: World) -> Any:
    """PRD WLD-K-01 booked and activated."""
    body = {
        "external_id": "SF-ORD-10001",
        "customer_id": str(world.customer_id),
        "contracting_entity_code": "AVM-US",
        "transaction_currency": "USD",
        "inception_date": "2026-01-01",
        "lines": [
            {
                "obligation_key": "O1",
                "product_code": PLATFORM,
                "quantity": "1",
                "total_price": {"amount": "120000.00", "currency": "USD"},
                "start_date": "2026-01-01",
                "end_date": "2026-12-31",
            },
            {
                "obligation_key": "O2",
                "product_code": IMPLEMENTATION,
                "quantity": "1",
                "total_price": {"amount": "15000.00", "currency": "USD"},
            },
        ],
    }
    return booked_contract(world.place, body, activate=True)


def flag_change(world: World, contract_id: UUID, expected: int) -> None:
    """``SIGNIFICANT_CHANGE_FLAGGED`` effective 2026-02-01, a new fact for the next computation."""
    with world.place.uow() as uow:
        append_events(
            uow,
            contract_id=contract_id,
            expected_stream_version=expected,
            events=[
                EventIn(
                    event_type=ContractEventType.SIGNIFICANT_CHANGE_FLAGGED,
                    effective_date=date(2026, 2, 1),
                    payload=SignificantChangeFlaggedV1(description="Customer asked for a review"),
                )
            ],
            origin="UI",
        )
        uow.commit()


def count(world: World, table: Any) -> int:
    return int(world.place.scalar(select(func.count()).select_from(table)))


def test_bundle_pins_reference_versions(world: World) -> None:
    booked = k01(world)
    group_id = booked.combination_group["id"]
    bundle, _, stored = computed(world.place, group_id)
    # RCP-15: the members, their stream, the pinned reference data and the calendar.
    assert bundle.group.group_key == "CG-CON-000001"
    assert bundle.group.member_contract_keys == ("SF-ORD-10001",)
    assert [(event.stream_version, event.event_type) for event in bundle.events] == [
        (1, "CONTRACT_BOOKED"),
        (2, "CONTRACT_ACTIVATED"),
    ]
    first_lines = bundle.events[0].payload["lines"]
    assert isinstance(first_lines, list)
    assert first_lines[0]["total_price"] == "120000.00"
    assert [version.version_key for version in bundle.ssp_versions] == ["US-LIST@v1"]
    assert [item.version_key for item in bundle.pob_template_versions] == [
        "TPL-SUB-DAILY@v1",
        "TPL-SVC-PCT@v1",
    ]
    assert [(item.code, item.default_template_code) for item in bundle.group.products] == [
        (IMPLEMENTATION, "TPL-SVC-PCT"),
        (PLATFORM, "TPL-SUB-DAILY"),
    ]
    (book_input,) = bundle.books
    assert (book_input.book_code, book_input.is_primary, book_input.entity_codes) == (
        "ASC606",
        True,
        ("AVM-US",),
    )
    assert book_input.account_mapping.version_key == "AVM-MAP-2026-01@v1"
    (avm_us,) = bundle.entities
    states = {period_input.period_key: period_input.states for period_input in avm_us.periods}
    assert (len(avm_us.periods), states["FY2026-P09"], states["FY2026-P10"]) == (
        24,
        (("ASC606", "open"),),
        (("ASC606", "future"),),
    )
    assert bundle.group.previous_stream_heads == ()
    # REQ-REF-015: the computation names every pinned version.
    (row,) = world.place.rows(select(contract_computation))
    published = world.place.rows(
        select(registry_version.c.id).where(registry_version.c.status == "PUBLISHED")
    )
    refs = row["pinned_refs"]
    # 04 T-CON-07 rev 1.110 (security ruling R-21): the computation of a contract past DRAFT also
    # records each product as it used it — the ``ProductInput`` members and the level-P values
    # of the product's obligations — and the next computation reads the product from here.
    unset: dict[str, Any] = {
        "sku_number": None,
        "product_family": None,
        "principal_agent": "PRINCIPAL",
        "distinctness_default": "distinct",
        "unit_of_measure": "EA",
        "is_bundle": False,
        "is_franchisor_preopening_service": False,
        "policy_values": {},
        "assurance_cost_per_unit": None,
        "components": [],
    }
    assert refs == {
        "ssp_book_version_ids": [world.h1_id],
        "registry_version_ids": sorted(str(item["id"]) for item in published),
        "rule_set_version_ids": [],
        "pob_template_version_ids": sorted(item["version_id"] for item in world.templates.values()),
        "account_mapping_version_id": world.mapping_id,
        "fx_rate_set_version_ids": [],
        "calendar_ids": [world.calendar_id],
        "estimate_version_ids": [],
        "products": {
            IMPLEMENTATION: {
                **unset,
                "revenue_category": "SERVICES",
                "default_template_code": "TPL-SVC-PCT",
                "obligation_policies": {},
            },
            PLATFORM: {
                **unset,
                "revenue_category": "SUBSCRIPTION",
                "default_template_code": "TPL-SUB-DAILY",
                "obligation_policies": {
                    "recognition.time_convention": {
                        "value": "DAILY",
                        "source": "TPL-SUB-DAILY@v1",
                    }
                },
            },
        },
    }
    assert published
    assert (row["id"], row["trigger"], row["status"], row["engine_version"]) == (
        stored["id"],
        "COMMAND",
        "SUCCEEDED",
        ENGINE_VERSION,
    )
    assert row["stream_heads"] == {str(booked.contract["id"]): 2}
    # DG-KRN-REG-02: pin K values are recorded on the version with their level and source.
    (version,) = world.place.rows(select(contract_version))
    assert version["pinned_policies"]["ssp.outside_range_point"] == {
        "value": "NEAREST_BOUND",
        "level": "DEFAULT",
        "source_id": POLICY_PARAMETERS["ssp.outside_range_point"].source_ref,
    }
    assert all(POLICY_PARAMETERS[code].pin == "K" for code in version["pinned_policies"]), (
        "pin P values resolve per posting period"
    )


def test_allocation_lineage_k01(world: World) -> None:
    booked = k01(world)
    computed(world.place, booked.combination_group["id"])
    entries = {
        item["code"]: str(item["id"])
        for item in world.place.rows(
            select(ssp_entry.c.id, product.c.code)
            .select_from(ssp_entry.join(product, product.c.id == ssp_entry.c.product_id))
            .where(ssp_entry.c.ssp_book_version_id == UUID(world.h1_id))
        )
    }
    rows = world.place.rows(
        select(obligation_version).order_by(obligation_version.c.obligation_key)
    )
    # PRD WLD-X-01: 135,000 × 132,000 ÷ 150,000 and the remainder to O2.
    assert [
        (
            item["obligation_key"],
            item["product_code"],
            item["stated_price"],
            item["original_ssp_selected"],
            item["allocated_amount"],
            item["allocation_weight"],
            item["allocated_exact"],
            item["allocation_adjustment"],
        )
        for item in rows
    ] == [
        (
            "O1",
            PLATFORM,
            Decimal("120000.00"),
            Decimal(132000),
            Decimal("118800.00"),
            Decimal("0.88"),
            Decimal(118800),
            Decimal("-1200.00"),
        ),
        (
            "O2",
            IMPLEMENTATION,
            Decimal("15000.00"),
            Decimal(18000),
            Decimal("16200.00"),
            Decimal("0.12"),
            Decimal(16200),
            Decimal("1200.00"),
        ),
    ]
    # REQ-SSP-011: the version, entry, bounds, in-range flag and selected SSP of each obligation; a
    # point entry has no low, mid or high and no in-range flag (S05-R-07).
    assert [
        (
            str(item["ssp_book_version_id"]),
            str(item["ssp_entry_id"]),
            item["ssp_version_label"],
            item["original_ssp_low"],
            item["original_ssp_mid"],
            item["original_ssp_high"],
            item["original_ssp_in_range"],
        )
        for item in rows
    ] == [
        (world.h1_id, entries[PLATFORM], "2026-H1", None, None, None, None),
        (world.h1_id, entries[IMPLEMENTATION], "2026-H1", None, None, None, None),
    ]
    assert [item["ssp_method"] for item in rows] == ["observable", "cost_plus_margin"]
    (version,) = world.place.rows(select(contract_version))
    assert (version["transaction_price"], version["total_ssp"], version["book_code"]) == (
        Decimal("135000.00"),
        Decimal(150000),
        "ASC606",
    )
    assert sum(item["allocated_amount"] for item in rows) == version["transaction_price"]
    for item in rows:
        assert item["allocated_amount"] == (
            item["revenue_cum"] + item["scheduled_amount"] + item["awaiting_trigger_amount"]
        )


def snapshot(world: World) -> dict[int, list[tuple[Any, ...]]]:
    rows = world.place.rows(
        select(
            obligation_version.c.version_no,
            obligation_version.c.obligation_key,
            obligation_version.c.ssp_book_version_id,
            obligation_version.c.ssp_entry_id,
            obligation_version.c.original_ssp_selected,
            obligation_version.c.allocation_weight,
            obligation_version.c.allocated_amount,
        ).order_by(obligation_version.c.version_no, obligation_version.c.obligation_key)
    )
    found: dict[int, list[tuple[Any, ...]]] = {}
    for item in rows:
        found.setdefault(item["version_no"], []).append(tuple(list(item.values())[1:]))
    return found


def test_ssp_snapshot_survives_later_publication(world: World) -> None:
    booked = k01(world)
    group_id = booked.combination_group["id"]
    computed(world.place, group_id)
    h2_id = approved_ssp_version(
        world.app,
        world.place.author,
        [world.priya, world.marcus],
        world.book_id,
        label="2026-H2",
        effective_from="2026-10-01",
        entries=[
            point_entry(PLATFORM, "140000.00", value_basis="AMOUNT"),
            point_entry(IMPLEMENTATION, "20000.00", "cost_plus_margin"),
        ],
    )
    bundle, _, stored = computed(world.place, group_id)
    assert [version.version_key for version in bundle.ssp_versions] == ["US-LIST@v1", "US-LIST@v2"]
    assert stored["pinned_refs"]["ssp_book_version_ids"] == sorted([world.h1_id, h2_id])
    # REQ-SSP-014: version 2 keeps the inception snapshot of every obligation.
    found = snapshot(world)
    assert sorted(found) == [1, 2]
    assert found[2] == found[1]
    assert [str(item[1]) for item in found[2]] == [world.h1_id, world.h1_id]


def test_schedules_versioned_never_deleted(world: World) -> None:
    booked = k01(world)
    group_id = booked.combination_group["id"]
    computed(world.place, group_id)
    first = world.place.rows(select(schedule_line).order_by(schedule_line.c.period_end_date))
    flag_change(world, booked.contract["id"], 2)
    computed(world.place, group_id)
    headers = world.place.rows(
        select(schedule.c.contract_version_id, schedule.c.schedule_kind, schedule.c.line_count)
    )
    versions = {
        item["version_no"]: item["id"]
        for item in world.place.rows(select(contract_version.c.version_no, contract_version.c.id))
    }
    assert sorted(
        (item["contract_version_id"], item["schedule_kind"]) for item in headers
    ) == sorted([(versions[1], "REVENUE"), (versions[2], "REVENUE")])
    # REQ-REC-020: the lines of version 1 stay as they were.
    kept = world.place.rows(
        select(schedule_line).where(schedule_line.c.contract_version_id == versions[1])
    )
    assert sorted(first, key=lambda item: str(item["id"])) == sorted(
        kept, key=lambda item: str(item["id"])
    )
    later = world.place.rows(
        select(schedule_line).where(schedule_line.c.contract_version_id == versions[2])
    )
    assert len(later) == len(first) == 12
    # PRD WLD-X-02: 88,855.89 − 79,091.51.
    september = world.place.rows(
        select(
            contract_version.c.version_no,
            schedule_line.c.amount,
            schedule_line.c.cumulative_amount,
            schedule_line.c.line_type,
            schedule_line.c.is_released_at_close,
        )
        .select_from(
            schedule_line.join(period, period.c.id == schedule_line.c.period_id).join(
                contract_version, contract_version.c.id == schedule_line.c.contract_version_id
            )
        )
        .where(period.c.period_key == "FY2026-P09")
        .order_by(contract_version.c.version_no)
    )
    assert [tuple(item.values()) for item in september] == [
        (1, Decimal("9764.38"), Decimal("88855.89"), "NORMAL", True),
        (2, Decimal("9764.38"), Decimal("88855.89"), "NORMAL", True),
    ]
    assert all(item["line_count"] == 12 for item in headers)


def test_l8_e_rate_stamp_takes_the_latest_rate_ref() -> None:
    """D-88 L7-6-Q-1: T-SL-04 holds one ``fx_rate_id``, so a foreign-currency line is stamped with
    the RateRef of the latest effective date, ties to the greatest rate key; a RateRef the bundle
    does not pin, or whose version differs, raises (REQ-FX-006)."""
    from erev_engine.bundle import FxRateInput

    def rate(key: str, on: date) -> FxRateInput:
        return FxRateInput(key, "FX@v1", "spot", "GBP", "USD", on, None, Decimal("1.2700"))

    pinned = {
        item.rate_key: item
        for item in (
            rate("FX@v1/spot/GBP/USD/2026-04-01", date(2026, 4, 1)),
            rate("FX@v1/spot/GBP/USD/2026-05-31", date(2026, 5, 31)),
            rate("FX@v1/average/GBP/USD/2026-05-31", date(2026, 5, 31)),
        )
    }
    version = UUID(int=9)
    ids = {
        key: (UUID(int=index), version, Decimal("1.2700") + index)
        for index, key in enumerate(pinned)
    }
    april, may, may_average = (
        (("FX@v1/spot/GBP/USD/2026-04-01", "FX@v1"),),
        (("FX@v1/spot/GBP/USD/2026-05-31", "FX@v1"),),
        (("FX@v1/average/GBP/USD/2026-05-31", "FX@v1"),),
    )
    # L8-merge: the stamp carries the rate of that fx_rate row, so T-SL-04 fx_rate is never NULL
    # beside fx_rate_id (API-S-SubledgerLine and the T-SL-06 drill format it).
    assert computation._rate_stamp("L", april, pinned, ids) == (
        UUID(int=0),
        version,
        Decimal("1.2700"),
    )
    assert computation._rate_stamp("L", (*april, *may), pinned, ids) == (
        UUID(int=1),
        version,
        Decimal("2.2700"),
    )
    # Same effective date: the greatest rate key wins ("FX@v1/spot/..." > "FX@v1/average/...").
    assert computation._rate_stamp("L", (*may_average, *may), pinned, ids) == (
        UUID(int=1),
        version,
        Decimal("2.2700"),
    )
    with pytest.raises(ValueError, match="does not pin"):
        computation._rate_stamp("L", (("FX@v1/spot/GBP/USD/2026-06-30", "FX@v1"),), pinned, ids)
    with pytest.raises(ValueError, match="does not pin"):
        computation._rate_stamp("L", (("FX@v1/spot/GBP/USD/2026-04-01", "FX@v2"),), pinned, ids)


def test_persist_idempotent_on_input_hash(world: World) -> None:
    booked = k01(world)
    run = engine()
    with world.place.uow() as uow:
        bundle = bundles.build(uow.session, booked.combination_group["id"], uow.now)
        output = run(bundle)
        first = computation.persist(uow, bundle, output)
        uow.commit()
    tables = (
        contract_computation,
        contract_version,
        obligation_version,
        schedule,
        schedule_line,
        calc_trace,
        computation_evidence,
        ledger_chain_head,
        subledger_posting,
        file_object,
        audit_event,
    )
    before = [count(world, table) for table in tables]
    with world.place.uow() as uow:
        second = computation.persist(uow, bundle, output)
        uow.commit()
    # RCP-21: the same input hash and engine version answer the stored computation.
    assert (first["replayed"], second["replayed"], second["id"]) == (False, True, first["id"])
    assert second["contract_version_ids"] == first["contract_version_ids"]
    assert second["input_sha256"] == bundle.sha256()
    assert [count(world, table) for table in tables] == before


def test_calc_trace_stored_and_reevaluates(world: World) -> None:
    booked = k01(world)
    _, output, _ = computed(world.place, booked.combination_group["id"])
    (trace_row,) = world.place.rows(select(calc_trace))
    (version,) = world.place.rows(select(contract_version))
    assert (version["calc_trace_id"], trace_row["contract_version_id"]) == (
        trace_row["id"],
        version["id"],
    )
    context = DbContext(tenant_id=world.place.tenant_id, user_id=None, entity_scope="*")
    with tenant_session(context, read_only=True) as session:
        loaded = store.load_trace(session, version["id"])
    assert loaded is not None
    assert loaded == output.books[0].trace
    assert (
        trace_row["trace_sha256"],
        trace_row["node_count"],
        trace_row["book_code"],
        trace_row["format_version"],
        trace_row["engine_version"],
        trace_row["root_measures"],
    ) == (
        loaded.sha256(),
        len(loaded.nodes),
        "ASC606",
        1,
        ENGINE_VERSION,
        dict(loaded.root_measures),
    )
    assert loaded.nodes
    # DG-KRN-EXP-04, 05: every stored node re-evaluates to its stored value.
    assert reevaluate(loaded) == {node.id: node.value for node in loaded.nodes}
    node_ids = {node.id for node in loaded.nodes}
    for item in world.place.rows(select(obligation_version.c.trace_nodes)):
        assert set(item["trace_nodes"].values()) <= node_ids
    lines = world.place.rows(select(schedule_line.c.trace_node_id))
    assert {item["trace_node_id"] for item in lines} <= node_ids
    tampered = {**trace_row, "trace_sha256": "0" * 64}
    with pytest.raises(store.TraceIntegrityError):
        store.trace_from_row(tampered)


def test_versions_immutable_and_hashed(world: World) -> None:
    booked = k01(world)
    group_id = booked.combination_group["id"]
    bundle, output, stored = computed(world.place, group_id)
    events = world.place.rows(
        select(contract_event.c.id, contract_event.c.recorded_at).order_by(
            contract_event.c.stream_version
        )
    )
    (row,) = world.place.rows(select(contract_computation))
    (version,) = world.place.rows(select(contract_version))
    assert (row["engine_version"], row["input_sha256"], row["known_at"]) == (
        ENGINE_VERSION,
        bundle.sha256(),
        max(item["recorded_at"] for item in events),
    )
    assert (
        version["contract_computation_id"],
        version["cause_event_ids"],
        version["output_sha256"],
        version["version_no"],
        version["previous_version_id"],
        version["known_at"],
    ) == (row["id"], [item["id"] for item in events], output.sha256(), 1, None, row["known_at"])
    assert world.place.scalar(select(contract.c.latest_computation_id)) == stored["id"]
    group = world.place.rows(
        select(combination_group.c.head_computation_id, combination_group.c.dirty_since)
    )
    assert group == [{"head_computation_id": stored["id"], "dirty_since": None}]
    # REQ-CON-014: erev_app holds no UPDATE on a computed version.
    context = DbContext(tenant_id=world.place.tenant_id, user_id=None, entity_scope="*")
    with tenant_session(context) as session:
        savepoint = session.begin_nested()
        with pytest.raises(exc.DBAPIError) as excinfo:
            session.execute(
                update(contract_version)
                .where(contract_version.c.id == version["id"])
                .values(transaction_price=Decimal("1"))
            )
        savepoint.rollback()
    assert getattr(excinfo.value.orig, "sqlstate", None) == "42501"
    assert world.place.scalar(select(contract_version.c.transaction_price)) == Decimal("135000")
    # The next version names the events it first includes and its predecessor.
    flag_change(world, booked.contract["id"], 2)
    computed(world.place, group_id)
    latest = world.place.rows(select(contract_version).where(contract_version.c.version_no == 2))
    flagged = world.place.scalar(
        select(contract_event.c.id).where(contract_event.c.stream_version == 3)
    )
    assert (latest[0]["previous_version_id"], latest[0]["cause_event_ids"]) == (
        version["id"],
        [flagged],
    )


def without_lineage(output: OutputBundle) -> OutputBundle:
    book_output = output.books[0]
    first = book_output.obligation_versions[0]
    broken = replace(
        first,
        columns=MappingProxyType(
            {**first.columns, "ssp_book_version_key": None, "ssp_entry_key": None}
        ),
    )
    books = (
        replace(book_output, obligation_versions=(broken, *book_output.obligation_versions[1:])),
        *output.books[1:],
    )
    return replace(output, books=books)


@pytest.mark.control("CTL-011")
def test_ctl_011_allocation_lineage_recorded_per_obligation(world: World) -> None:
    booked = k01(world)
    run = engine()
    with world.place.uow() as uow:
        bundle = bundles.build(uow.session, booked.combination_group["id"], uow.now)
        output = run(bundle)
        with pytest.raises(EngineError) as excinfo:
            computation.persist(uow, bundle, without_lineage(output))
    error = excinfo.value
    assert (error.code, error.subject_key) == ("ENGINE_INVARIANT_VIOLATED", "SF-ORD-10001/O1")
    assert error.detail == {
        "control": "CTL-011",
        "rule": "REQ-SSP-011",
        "book_code": "ASC606",
        "columns": "ssp_book_version_key,ssp_entry_key",
        "count": "1",
    }
    assert computation.lineage_findings(bundle, output) == []
    for table in (contract_computation, contract_version, obligation_version, calc_trace):
        assert count(world, table) == 0
    # With lineage on every obligation the computation persists.
    with world.place.uow() as uow:
        stored = computation.persist(uow, bundle, output)
        uow.commit()
    assert count(world, obligation_version) == 2
    assert stored["replayed"] is False


@pytest.mark.control("CTL-011")
def test_computation_persists_selected_range_identity(world: World) -> None:
    booked = k01(world)
    with world.place.uow() as uow:
        bundle = bundles.build(uow.session, booked.combination_group["id"], uow.now)
        output = engine()(bundle)
        index = bundles.index(uow.session, bundle)
        keys = [item.columns.get("ssp_range_key") for item in output.books[0].obligation_versions]
        assert any(keys)
        computation.persist(uow, bundle, output)
        uow.commit()
    with world.place.uow() as uow:
        rows = uow.session.execute(
            select(
                obligation_version.c.ssp_range_id,
                obligation_version.c.ssp_entry_id,
            )
        ).all()
        actual = {row.ssp_range_id for row in rows if row.ssp_range_id is not None}
        assert actual == {index.ssp_ranges[str(key)] for key in keys if key is not None}
        for row in rows:
            if row.ssp_range_id is not None:
                assert (
                    uow.session.scalar(
                        select(ssp_range.c.ssp_entry_id).where(ssp_range.c.id == row.ssp_range_id)
                    )
                    == row.ssp_entry_id
                )


def test_computation_retains_typed_original_evidence(world: World) -> None:
    booked = k01(world)
    with world.place.uow() as uow:
        bundle = bundles.build(uow.session, booked.combination_group["id"], uow.now)
        output = engine()(bundle)
        cutoffs = dict(
            uow.session.execute(
                select(ledger_chain_head.c.book_code, ledger_chain_head.c.last_chain_seq)
            )
            .tuples()
            .all()
        )
        produced = computation.persist(uow, bundle, output)
        row = (
            uow.session.execute(
                select(computation_evidence).where(
                    computation_evidence.c.contract_computation_id == produced["id"]
                )
            )
            .mappings()
            .one()
        )
        assert row["captured"] == "AT_COMPUTATION"
        assert row["ledger_cutoff"] == {
            book.book_code: cutoffs[book.book_code] for book in bundle.books
        }
        assert row["input_sha256"] == bundle.sha256()
        assert row["output_sha256"] == output.sha256()
        assert row["book_codes"] == [
            book.book_code for book in output.books if book.contract_version is not None
        ]
        contents = []
        for name in ("input", "output"):
            meta, stream = open_file(
                uow.session, row[f"{name}_file_id"], files=uow.files, keyring=uow.keyring
            )
            with stream:
                content = stream.read()
            assert (
                hashlib.sha256(content).hexdigest() == row[f"{name}_file_sha256"] == meta["sha256"]
            )
            assert meta["purpose"] == f"COMPUTATION_{name.upper()}"
            assert uow.files.exists(str(meta["storage_key"]) + ".dek")
            holds = file_evidence.holds_in(uow.session, row[f"{name}_file_id"])
            assert any(hold.reference.table == "computation_evidence" for hold in holds)
            assert file_evidence.first_refusing(holds) is not None
            contents.append(content)
        restored = upgrade.decode_input(contents[0])
        assert restored == bundle
        assert restored.known_at == bundle.known_at
        assert engine()(restored).sha256() == output.sha256()
        assert contents[1] == upgrade.encode_output(output)
        uow.commit()
    # Evidence is append-only, just like the computed versions it proves.
    with (
        pytest.raises(
            exc.ProgrammingError, match="permission denied for table computation_evidence"
        ),
        world.place.uow() as uow,
    ):
        uow.session.execute(
            update(computation_evidence)
            .where(computation_evidence.c.contract_computation_id == produced["id"])
            .values(input_sha256="0" * 64)
        )
        uow.commit()


def test_computation_evidence_failure_rolls_back_before_posting(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    booked = k01(world)
    observed_tables = (
        contract_computation,
        computation_evidence,
        contract_version,
        subledger_posting,
        file_object,
    )
    before = [count(world, table) for table in observed_tables]
    store_original = store_file
    calls = 0

    def fail_output(*args: Any, **kwargs: Any) -> Any:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("simulated output storage failure")
        return store_original(*args, **kwargs)

    monkeypatch.setattr(evidence_writer, "store_file", fail_output)
    with pytest.raises(OSError, match="simulated output storage failure"), world.place.uow() as uow:
        bundle = bundles.build(uow.session, booked.combination_group["id"], uow.now)
        computation.persist(uow, bundle, engine()(bundle))
        uow.commit()
    assert calls == 2
    assert [count(world, table) for table in observed_tables] == before


def test_computation_evidence_links_predecessor_and_has_no_generic_download(world: World) -> None:
    booked = k01(world)
    group_id = booked.combination_group["id"]
    computed(world.place, group_id)
    (first,) = world.place.rows(select(computation_evidence))
    assert first["previous_computation_id"] is None
    flag_change(world, booked.contract["id"], 2)
    computed(world.place, group_id)
    rows = world.place.rows(select(computation_evidence))
    assert len(rows) == 2
    second = next(
        row for row in rows if row["contract_computation_id"] != first["contract_computation_id"]
    )
    assert second["previous_computation_id"] == first["contract_computation_id"]
    assert (
        next(
            row
            for row in rows
            if row["contract_computation_id"] == first["contract_computation_id"]
        )
        == first
    )
    for row in rows:
        for column in ("input_file_id", "output_file_id"):
            for suffix in ("", "/content"):
                response = get(
                    world.app, f"/api/v1/files/{row[column]}{suffix}", world.place.author
                )
                assert response.status_code == 404, response.text


def test_report_reads_original_computation_after_group_advances(world: World) -> None:
    booked = k01(world)
    group_id = booked.combination_group["id"]
    bundle, output, first = computed(world.place, group_id)
    flag_change(world, booked.contract["id"], 2)
    _, second_output, second = computed(world.place, group_id)
    report_only = replace(
        world.place.principal,
        permissions=world.place.principal.permissions - {"contract.read"},
    )
    with world.place.uow(report_only) as uow:
        retained = computation_sources.read_for_report(uow, first["id"])
        assert retained.bundle == bundle
        assert retained.engine_release_id == first["engine_release_id"]
        assert retained.output_sha256 == output.sha256()
        assert engine()(retained.bundle).sha256() == output.sha256()
        assert retained.previous_computation_id is None
        assert retained.bundle.events[-1].stream_version == 2
        assert retained.output == upgrade.output_mapping(output)

        latest = computation_sources.read_for_report(uow, second["id"])
        assert latest.previous_computation_id == first["id"]
        assert latest.output_sha256 == second_output.sha256()


def test_variance_candidate_selection_excludes_later_recompute_without_new_events(
    world: World,
) -> None:
    """Actual retained inputs; the frozen population shell isolates version selection."""
    import hashlib
    from uuid import uuid4

    from erev_api.db.tables import legal_entity
    from erev_api.domain.reports import locked, variance_execution, variance_sources
    from erev_api.domain.reports.evidence_close import FrozenClose, FrozenReport
    from erev_api.domain.reports.variance_projection import Projection

    booked = k01(world)
    group_id = booked.combination_group["id"]
    original, first_output, first = computed(world.place, group_id)
    with world.place.uow() as uow:
        cutoff = uow.session.execute(select(func.clock_timestamp())).scalar_one()
        entity_id = uow.session.execute(
            select(legal_entity.c.id).where(legal_entity.c.code == "AVM-US")
        ).scalar_one()
    newer, second_output, second = computed(world.place, group_id)
    assert original.events == newer.events
    assert original.known_at < newer.known_at
    lock_id = uuid4()
    content = b"row_key,contract_external_id\ncontract:SF-ORD-10001,SF-ORD-10001\n"
    dataset = locked.LockedDataset(
        lock_id, "CONTRACT_BALANCES", uuid4(), hashlib.sha256(content).hexdigest(), 1, {}, content
    )
    frozen = FrozenClose(
        locked.LockScope(lock_id, entity_id, "AVM-US", "ASC606", uuid4(), "FY2026-P09"),
        {"cutoff_known_at": cutoff},
        (FrozenReport("contract_balances", uuid4(), None, dataset),),
    )
    with world.place.uow() as uow:
        earlier = variance_sources._side(uow, frozen)
        later = variance_sources._side(
            uow,
            replace(
                frozen,
                record={
                    "cutoff_known_at": uow.session.execute(
                        select(func.clock_timestamp())
                    ).scalar_one()
                },
            ),
        )
        assert earlier.contract_computations == {"SF-ORD-10001": first["id"]}
        assert later.contract_computations == {"SF-ORD-10001": second["id"]}
        assert earlier.computations[first["id"]].bundle == original

    result = variance_execution.execute(
        earlier.computations[first["id"]],
        later.computations[second["id"]],
        entity_code="AVM-US",
        book_code="ASC606",
        period_key="FY2026-P09",
        earlier_metrics=Projection(original, "AVM-US", "ASC606", "FY2026-P09")(first_output),
        later_metrics=Projection(newer, "AVM-US", "ASC606", "FY2026-P09")(second_output),
    )
    assert "posted" in result.normalized_fields
    assert "group.previous_stream_heads" in result.normalized_fields
    assert "books.ASC606.ssp_readback.SF-ORD-10001/O1" in result.normalized_fields
    assert all(item.total == 0 and item.sums_exactly for item in result.replay.measures.values())
    final_bundle, final_output, third = computed(world.place, group_id)
    with world.place.uow() as uow:
        final_source = computation_sources.read_for_report(uow, third["id"])
    # Calculate after leaving the read transaction. Monetary endpoints here are
    # persisted computes, not a substitute for full frozen-close acceptance.
    result = variance_execution.execute(
        later.computations[second["id"]],
        final_source,
        entity_code="AVM-US",
        book_code="ASC606",
        period_key="FY2026-P09",
        earlier_metrics=Projection(newer, "AVM-US", "ASC606", "FY2026-P09")(second_output),
        later_metrics=Projection(final_bundle, "AVM-US", "ASC606", "FY2026-P09")(final_output),
    )
    assert result.earlier.computation_id == second["id"]
    assert result.later.computation_id == third["id"]
    assert all(item.total == 0 and item.sums_exactly for item in result.replay.measures.values())
    assert final_source.bundle.sha256() == final_bundle.sha256()
    assert result.replay.executions == ()  # identical input hash reuses the saved computation


@pytest.mark.parametrize("name", ["input", "output"])
@pytest.mark.parametrize(
    "fault", ["bytes", "purpose", "size", "digest", "forged metadata", "missing"]
)
def test_report_refuses_corrupt_computation_files(
    world: World, monkeypatch: pytest.MonkeyPatch, fault: str, name: str
) -> None:
    booked = k01(world)
    _, _, produced = computed(world.place, booked.combination_group["id"])
    streams: list[io.BytesIO] = []

    def altered(*args: Any, **kwargs: Any) -> Any:
        meta, stream = open_file(*args, **kwargs)
        with stream:
            content = stream.read()
        if str(meta["purpose"]) != f"COMPUTATION_{name.upper()}":
            return meta, io.BytesIO(content)
        if fault == "missing":
            raise FileNotFoundError("retained file absent")
        changed = dict(meta)
        if fault == "forged metadata":
            content += b" "
            changed["size_bytes"] = len(content)
            changed["sha256"] = hashlib.sha256(content).hexdigest()
        if fault == "bytes":
            content += b" "
        if fault == "purpose":
            changed["purpose"] = "REPORT_OUTPUT"
        if fault == "size":
            changed["size_bytes"] += 1
        if fault == "digest":
            changed["sha256"] = "0" * 64
        result = io.BytesIO(content)
        streams.append(result)
        return changed, result

    monkeypatch.setattr(computation_sources, "open_file", altered)
    with world.place.uow() as uow, pytest.raises(Problem) as caught:
        computation_sources.read_for_report(uow, produced["id"])
    assert caught.value.slug == "validation-failed"
    reason = "EVIDENCE_MISSING" if fault == "missing" else "EVIDENCE_HASH_MISMATCH"
    assert caught.value.errors[0].rule_id == f"{name.upper()}_{reason}"
    assert (streams or fault == "missing") and all(stream.closed for stream in streams)


@pytest.mark.parametrize(
    "scope",
    [
        "no report permission",
        "hidden entity",
        "hidden failed",
        "different tenant",
        "system",
    ],
)
def test_report_authorizes_computation_before_reading_files(
    world: World, monkeypatch: pytest.MonkeyPatch, scope: str
) -> None:
    booked = k01(world)
    _, _, produced = computed(world.place, booked.combination_group["id"])
    if scope == "hidden failed":
        (source,) = world.place.rows(
            select(contract_computation).where(contract_computation.c.id == produced["id"])
        )
        failed_id = UUID(int=7788)
        with world.place.uow() as uow:
            with uow.as_system():
                uow.session.execute(
                    insert(contract_computation).values(
                        **{
                            **source,
                            "id": failed_id,
                            "status": "FAILED",
                            "problem": {"detail": "failure"},
                        }
                    )
                )
            uow.commit()
        produced = {**produced, "id": failed_id}
    principal = world.place.principal
    if scope == "no report permission":
        principal = replace(principal, permissions=principal.permissions - {"report.run"})
    elif scope == "different tenant":
        principal = replace(principal, tenant_id=UUID(int=999))
    elif scope == "system":
        principal = system_principal(principal.tenant_id)
    else:
        principal = replace(
            principal,
            permission_scopes={
                **principal.permission_scopes,
                "report.run": frozenset({UUID(int=999)}),
            },
        )

    def unread(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("Evidence files must not be read for an unauthorized caller")

    monkeypatch.setattr(computation_sources, "open_file", unread)
    with world.place.uow(principal) as uow, pytest.raises(Problem) as caught:
        computation_sources.read_for_report(uow, produced["id"])
    assert caught.value.slug == (
        "not-found"
        if scope in {"hidden entity", "hidden failed", "different tenant"}
        else "forbidden"
    )


def test_report_missing_evidence_never_rebuilds_current_inputs(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    booked = k01(world)
    # Simulate a computation written before 0160: it has versions but no retained files.
    monkeypatch.setattr(evidence_writer, "retain", lambda *args, **kwargs: None)
    _, _, produced = computed(world.place, booked.combination_group["id"])

    def forbidden_rebuild(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("Missing historical evidence cannot use current inputs or files")

    monkeypatch.setattr(bundles, "build", forbidden_rebuild)
    monkeypatch.setattr(computation_sources, "open_file", forbidden_rebuild)
    with world.place.uow() as uow, pytest.raises(Problem) as caught:
        computation_sources.read_for_report(uow, produced["id"])
    assert caught.value.errors[0].rule_id == "INPUT_EVIDENCE_MISSING"


def test_verification_schema_guards_execution_identity_and_acceptance(world: World) -> None:
    from uuid import uuid4

    from erev_api.db.tables import (
        computation_verification,
        release_validation,
        release_validation_attempt,
        release_validation_group,
    )
    from erev_api.db.transitions import apply
    from erev_api.enums import JobKind
    from erev_api.jobs.registry import insert_job

    booked = k01(world)
    group_id = booked.combination_group["id"]
    _, output, source = computed(world.place, group_id)
    with world.place.uow() as uow:
        stamps = {
            "tenant_id": uow.principal.tenant_id,
            "created_at": uow.now,
            "created_by": uow.principal.id,
            "created_by_kind": uow.principal.kind.value,
        }
        mutable = {
            **stamps,
            "updated_at": uow.now,
            "updated_by": uow.principal.id,
            "updated_by_kind": uow.principal.kind.value,
        }
        release_id, attempt_id, selected_id = uuid4(), uuid4(), uuid4()
        uow.session.execute(
            insert(release_validation).values(
                **mutable,
                id=release_id,
                engine_release_id=source["engine_release_id"],
                effective_level="MINOR",
                status="RUNNING",
            )
        )
        uow.session.execute(
            insert(release_validation_attempt).values(
                **mutable,
                id=attempt_id,
                release_validation_id=release_id,
                attempt_no=1,
                seed=20260912,
                population={},
                required_entity_ids=[world.entity_id],
                status="RUNNING",
            )
        )
        uow.session.execute(
            insert(release_validation_group).values(
                **mutable,
                id=selected_id,
                release_validation_attempt_id=attempt_id,
                combination_group_id=group_id,
                source_computation_id=None if source is None else source["id"],
                selection="OPEN_ACTIVITY",
                strata=[],
                rank_sha256="a" * 64,
                orchestration_state="RUNNING",
            )
        )
        job_row = insert_job(
            uow.session,
            JobKind.REPLAY_VERIFY,
            {},
            tenant_id=uow.principal.tenant_id,
            now=uow.now,
            created_by=uow.principal.id,
            created_by_kind=uow.principal.kind,
            lock=f"verify:{group_id}",
        )
        row = {
            **stamps,
            "source_computation_id": source["id"],
            "combination_group_id": group_id,
            "trigger": "REPLAY_VERIFY",
            "source_engine_release_id": source["engine_release_id"],
            "source_engine_version": ENGINE_VERSION,
            "process_engine_release_id": source["engine_release_id"],
            "process_engine_version": ENGINE_VERSION,
            "engine_executed": True,
            "executed_engine_release_id": source["engine_release_id"],
            "executed_engine_version": ENGINE_VERSION,
            "actual_output_sha256": output.sha256(),
            "outcome": "MATCH",
            "release_validation_attempt_id": attempt_id,
            "job_id": job_row["id"],
            "task_id": 11,
            "attempt": 1,
            "duration_ms": 10,
        }
        bad_cases: list[dict[str, Any]] = [
            {"executed_engine_release_id": None},
            {"executed_engine_version": None},
            {"engine_executed": False},
            {
                "engine_executed": False,
                "executed_engine_release_id": None,
                "executed_engine_version": None,
            },
            {"source_engine_version": None},
            {"executed_engine_version": "9.9.9"},
            {"source_computation_id": None},
            {"trigger": "EVENT"},
            {"outcome": "ERROR"},
            {"attempt": 0},
            {"duration_ms": -1},
        ]
        for change in bad_cases:
            with pytest.raises(exc.DBAPIError), uow.session.begin_nested():
                uow.session.execute(
                    insert(computation_verification).values(**{**row, **change, "id": uuid4()})
                )
        error_id = uuid4()
        uow.session.execute(
            insert(computation_verification).values(
                **{
                    **row,
                    "id": error_id,
                    "actual_output_sha256": None,
                    "outcome": "ERROR",
                    "reason": "ENGINE_ERROR",
                }
            )
        )
        with pytest.raises(exc.DBAPIError), uow.session.begin_nested():
            apply(
                uow.session,
                "release_validation_group",
                selected_id,
                to_status="COMPLETE",
                expected_status="RUNNING",
                expected_row_version=1,
                set_values={"accepted_verification_id": error_id},
            )
        match_id = uuid4()
        good = {**row, "id": match_id, "task_id": 12, "attempt": 2}
        uow.session.execute(insert(computation_verification).values(**good))
        other_attempt, other_result = uuid4(), uuid4()
        uow.session.execute(
            insert(release_validation_attempt).values(
                **mutable,
                id=other_attempt,
                release_validation_id=release_id,
                attempt_no=2,
                seed=20260912,
                population={},
                required_entity_ids=[world.entity_id],
                status="RUNNING",
            )
        )
        uow.session.execute(
            insert(computation_verification).values(
                **{**good, "id": other_result, "release_validation_attempt_id": other_attempt}
            )
        )
        with pytest.raises(exc.DBAPIError), uow.session.begin_nested():
            apply(
                uow.session,
                "release_validation_group",
                selected_id,
                to_status="COMPLETE",
                expected_status="RUNNING",
                expected_row_version=1,
                set_values={"accepted_verification_id": other_result},
            )
        with pytest.raises(exc.DBAPIError), uow.session.begin_nested():
            uow.session.execute(insert(computation_verification).values(**{**good, "id": uuid4()}))
        with pytest.raises(exc.DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                update(computation_verification)
                .where(computation_verification.c.id == match_id)
                .values(duration_ms=20)
            )
        apply(
            uow.session,
            "release_validation_group",
            selected_id,
            to_status="COMPLETE",
            expected_status="RUNNING",
            expected_row_version=1,
            set_values={"accepted_verification_id": match_id},
        )
        with pytest.raises(exc.DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                update(release_validation_attempt)
                .where(release_validation_attempt.c.id == attempt_id)
                .values(population={"changed": True})
            )
        with pytest.raises(exc.DBAPIError), uow.session.begin_nested():
            apply(
                uow.session,
                "release_validation",
                release_id,
                to_status="PASSED",
                expected_status="RUNNING",
                expected_row_version=1,
                set_values={"enabled_at": uow.now, "accepted_attempt_id": attempt_id},
            )
        apply(
            uow.session,
            "release_validation_attempt",
            attempt_id,
            to_status="COMPLETE",
            expected_status="RUNNING",
            expected_row_version=1,
            set_values={},
        )
        apply(
            uow.session,
            "release_validation",
            release_id,
            to_status="PASSED",
            expected_status="RUNNING",
            expected_row_version=1,
            set_values={"enabled_at": uow.now, "accepted_attempt_id": attempt_id},
        )

        apply(
            uow.session,
            "release_validation_attempt",
            other_attempt,
            to_status="COMPLETE",
            expected_status="RUNNING",
            expected_row_version=1,
            set_values={},
        )
        with pytest.raises(exc.DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                update(release_validation)
                .where(release_validation.c.id == release_id)
                .values(accepted_attempt_id=other_attempt)
            )
        uow.commit()
    from alembic import command
    from erev_api.db.session import DbContext, tenant_session
    from support.db import alembic_config

    # Published evidence remains visible to its tenant and hidden from an unrelated tenant.
    for tenant_id, expected in ((world.place.principal.tenant_id, 3), (uuid4(), 0)):
        with tenant_session(
            DbContext(tenant_id=tenant_id, user_id=None, entity_scope="*")
        ) as session:
            assert (
                session.scalar(select(func.count()).select_from(computation_verification))
                == expected
            )
    # A populated evidence schema must not silently disappear on downgrade.
    try:
        with pytest.raises(exc.DBAPIError):
            command.downgrade(alembic_config(), "0160")
    finally:
        # Alembic commits each completed migration before the populated-schema refusal.
        # Restore newer revisions so subsequent computations see the current schema.
        command.upgrade(alembic_config(), "head")


@dataclass(frozen=True)
class VerificationCase:
    context: JobContext
    decision: VerificationRow
    selection_id: UUID
    group_id: UUID
    contract_id: UUID


def verification_case(
    world: World, *, with_source: bool = True, queued_only: bool = False
) -> VerificationCase:
    from uuid import uuid4

    from erev_api.controls.release import current_release
    from erev_api.db.tables import (
        release_validation,
        release_validation_attempt,
        release_validation_group,
    )
    from erev_api.db.transitions import apply
    from erev_api.domain.contracts.release_validation import ExecutionToken
    from erev_api.domain.contracts.replay_verify import (
        EvidenceFacts,
        ProcessFacts,
        SourceFacts,
        verify,
    )
    from erev_api.enums import JobKind
    from erev_api.jobs.context import JobAttempt, JobContext, JobRuntime
    from erev_api.jobs.registry import dispatch, insert_job

    booked = k01(world)
    group_id = booked.combination_group["id"]
    bundle, output, source = computed(world.place, group_id) if with_source else (None, None, None)
    release = current_release()
    assert release is not None
    assert source is None or release.id == source["engine_release_id"]
    with world.place.uow() as uow:
        stamps = {
            "tenant_id": uow.principal.tenant_id,
            "created_at": uow.now,
            "created_by": uow.principal.id,
            "created_by_kind": uow.principal.kind.value,
            "updated_at": uow.now,
            "updated_by": uow.principal.id,
            "updated_by_kind": uow.principal.kind.value,
        }
        validation_id, attempt_id, selection_id = uuid4(), uuid4(), uuid4()
        uow.session.execute(
            insert(release_validation).values(
                **stamps,
                id=validation_id,
                engine_release_id=release.id,
                effective_level="MINOR",
                status="RUNNING",
            )
        )
        uow.session.execute(
            insert(release_validation_attempt).values(
                **stamps,
                id=attempt_id,
                release_validation_id=validation_id,
                attempt_no=1,
                seed=20260912,
                population={},
                required_entity_ids=[world.entity_id],
                status="RUNNING",
            )
        )
        queued = insert_job(
            uow.session,
            JobKind.REPLAY_VERIFY,
            {
                "release_validation_attempt_id": str(attempt_id),
                "combination_group_id": str(group_id),
                "trigger": "REPLAY_VERIFY",
            },
            subject_type="combination_group",
            subject_id=group_id,
            tenant_id=uow.principal.tenant_id,
            now=uow.now,
            created_by=uow.principal.id,
            created_by_kind=uow.principal.kind,
            lock=f"verify:{group_id}",
        )
        task = dispatch(
            uow.session,
            job_id=queued["id"],
            tenant_id=uow.principal.tenant_id,
            queue="compute",
            now=uow.now,
        )
        if not queued_only:
            apply(
                uow.session,
                "job",
                queued["id"],
                to_status="RUNNING",
                expected_status="QUEUED",
                expected_row_version=1,
                set_values={"started_at": uow.now},
            )
        uow.session.execute(
            insert(release_validation_group).values(
                **stamps,
                id=selection_id,
                release_validation_attempt_id=attempt_id,
                combination_group_id=group_id,
                source_computation_id=None if source is None else source["id"],
                selection="OPEN_ACTIVITY",
                strata=[],
                rank_sha256="a" * 64,
                orchestration_state="PENDING" if queued_only else "RUNNING",
                child_job_id=queued["id"],
            )
        )
        runtime = JobRuntime(
            clock=uow.clock,
            keyring=uow.keyring,
            files=uow.files,
            engine_release=release,
            attempt=JobAttempt(queued["id"], uow.principal.tenant_id, task, 1),
        )
        context = JobContext(
            job_id=queued["id"],
            tenant_id=uow.principal.tenant_id,
            kind=JobKind.REPLAY_VERIFY,
            principal=uow.principal,
            runtime=runtime,
        )
        uow.commit()
    if source is None:
        from erev_api.domain.contracts.replay_verify import no_source

        decision = no_source(
            trigger="REPLAY_VERIFY",
            group_id=group_id,
            process=ProcessFacts(release.id, release.engine_version),
            token=ExecutionToken(queued["id"], task, 1),
            attempt_id=attempt_id,
        )
        return VerificationCase(context, decision, selection_id, group_id, booked.contract["id"])
    assert bundle is not None and output is not None
    input_bytes, output_bytes = upgrade.encode_input(bundle), upgrade.encode_output(output)
    decision = verify(
        trigger="REPLAY_VERIFY",
        source=SourceFacts(
            source["id"],
            group_id,
            release.id,
            release.engine_version,
            source["input_sha256"],
            {"ASC606": output.sha256()},
            source["cutoff_at"],
        ),
        evidence=EvidenceFacts(
            input_bytes,
            upgrade.raw_digest(input_bytes),
            output_bytes,
            upgrade.raw_digest(output_bytes),
            output.sha256(),
            ("ASC606",),
            {"format_version": bundle.format_version, "engine_version": bundle.engine_version},
        ),
        process=ProcessFacts(release.id, release.engine_version),
        engine=engine(),
        token=ExecutionToken(queued["id"], task, 1),
        attempt_id=attempt_id,
    )
    return VerificationCase(context, decision, selection_id, group_id, booked.contract["id"])


def test_verification_store_atomic_result_pointer_and_duplicate(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    from erev_api.db.tables import (
        computation_verification,
        release_validation_group,
        subledger_posting,
    )
    from erev_api.domain.contracts import verification_store

    case = verification_case(world)
    with world.place.uow() as uow:
        before = [
            uow.session.scalar(select(func.count()).select_from(table))
            for table in (contract_computation, contract_version, subledger_posting)
        ]
    original = verification_store._Store.set_orchestration

    def fail_pointer(*args: Any, **kwargs: Any) -> None:
        raise RuntimeError("pointer write interrupted")

    monkeypatch.setattr(verification_store._Store, "set_orchestration", fail_pointer)
    with pytest.raises(RuntimeError, match="interrupted"), world.place.uow() as uow:
        verification_store.persist(uow, case.decision, context=case.context, duration_ms=10)
        uow.commit()
    monkeypatch.setattr(verification_store._Store, "set_orchestration", original)
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(computation_verification)) == 0
        first = verification_store.persist(uow, case.decision, context=case.context, duration_ms=10)
        assert first.result == "inserted" and first.row_id == first.accepted_id
        (event,) = uow.audit_events
        assert event["action"] == "release_validation_group.verify"
        assert event["object_id"] == case.selection_id
        assert event["detail"]["contract_id"] == str(case.contract_id)
        uow.commit()
    with world.place.uow() as uow:
        again = verification_store.persist(uow, case.decision, context=case.context, duration_ms=20)
        assert again.result == "duplicate" and again.row_id == first.row_id
        assert not uow.audit_events
        selected = (
            uow.session.execute(
                select(release_validation_group).where(
                    release_validation_group.c.id == case.selection_id
                )
            )
            .mappings()
            .one()
        )
        assert selected["orchestration_state"] == "COMPLETE"
        assert selected["accepted_verification_id"] == first.row_id
        row = uow.session.execute(select(computation_verification)).mappings().one()
        assert row["actual_output_sha256"] == case.decision.actual_output_sha256
        assert row["duration_ms"] == 10
        assert before == [
            uow.session.scalar(select(func.count()).select_from(table))
            for table in (contract_computation, contract_version, subledger_posting)
        ]
        uow.commit()


@pytest.mark.parametrize("failure", ["stale task", "changed head", "wrong token", "wrong release"])
def test_verification_store_refuses_changed_execution_or_selection(
    world: World, failure: str
) -> None:
    from erev_api.db.tables import computation_verification, job
    from erev_api.domain.contracts import verification_store
    from erev_api.jobs.context import JobLost

    case = verification_case(world)
    decision = case.decision
    if failure == "stale task":
        with world.place.uow() as uow:
            uow.session.execute(
                update(job).where(job.c.id == case.context.job_id).values(procrastinate_job_id=None)
            )
            uow.commit()
    elif failure == "changed head":
        flag_change(world, case.contract_id, 2)
        computed(world.place, case.group_id)
    elif failure == "wrong token":
        decision = replace(decision, token=replace(decision.token, attempt=2))
    else:
        decision = replace(decision, process_engine_version="9.9.9")
    with pytest.raises((Problem, JobLost)), world.place.uow() as uow:
        verification_store.persist(uow, decision, context=case.context, duration_ms=10)
        uow.commit()
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(computation_verification)) == 0


def test_verification_store_keeps_error_then_accepts_retry(world: World) -> None:
    from erev_api.db.tables import computation_verification, release_validation_group
    from erev_api.db.transitions import apply
    from erev_api.domain.contracts import verification_store
    from erev_api.enums import VerificationOutcome
    from erev_api.jobs.registry import dispatch
    from sqlalchemy import text

    case = verification_case(world)
    error = replace(
        case.decision,
        outcome=VerificationOutcome.ERROR,
        reason="ENGINE_ERROR",
        actual_output_sha256=None,
    )
    with world.place.uow() as uow:
        first = verification_store.persist(uow, error, context=case.context, duration_ms=10)
        assert first.accepted_id is None
        selected = (
            uow.session.execute(
                select(release_validation_group).where(
                    release_validation_group.c.id == case.selection_id
                )
            )
            .mappings()
            .one()
        )
        assert selected["orchestration_state"] == "RUNNING"
        assert selected["accepted_verification_id"] is None
        uow.commit()
    with world.place.uow() as uow:
        uow.session.execute(
            text("UPDATE procrastinate_jobs SET status='failed' WHERE id=:id"),
            {"id": case.context.task_id},
        )
        apply(
            uow.session,
            "job",
            case.context.job_id,
            to_status="QUEUED",
            expected_status="RUNNING",
            expected_row_version=1,
            set_values={},
        )
        task = dispatch(
            uow.session,
            job_id=case.context.job_id,
            tenant_id=uow.principal.tenant_id,
            queue="compute",
            now=uow.now,
            attempt=2,
        )
        apply(
            uow.session,
            "job",
            case.context.job_id,
            to_status="RUNNING",
            expected_status="QUEUED",
            expected_row_version=1,
            set_values={},
        )
        for old, new, version in (("RUNNING", "RETRYABLE", 1), ("RETRYABLE", "RUNNING", 2)):
            apply(
                uow.session,
                "release_validation_group",
                case.selection_id,
                to_status=new,
                expected_status=old,
                expected_row_version=version,
                set_values={"row_version": version + 1},
            )
        uow.commit()
    previous = case.context.runtime.attempt
    assert previous is not None
    context = JobContext(
        job_id=case.context.job_id,
        tenant_id=case.context.tenant_id,
        kind=case.context.kind,
        principal=case.context.principal,
        runtime=replace(
            case.context.runtime, attempt=replace(previous, task_id=task, attempt_number=2)
        ),
    )
    retry = replace(case.decision, token=replace(case.decision.token, task_id=task, attempt=2))
    with world.place.uow() as uow:
        second = verification_store.persist(uow, retry, context=context, duration_ms=20)
        assert second.accepted_id == second.row_id and second.row_id != first.row_id
        rows = (
            uow.session.execute(
                select(computation_verification.c.outcome).order_by(
                    computation_verification.c.attempt
                )
            )
            .scalars()
            .all()
        )
        assert rows == ["ERROR", "MATCH"]
        uow.commit()


def test_validation_population_reads_heads_activity_and_methods(world: World) -> None:
    from erev_api.domain.contracts.validation_population import read

    booked = k01(world)
    group_id = booked.combination_group["id"]
    with world.place.uow() as uow:
        before = read(uow)
        group = next(item for item in before.groups if item.group_id == group_id)
        assert group.open_activity
        assert before.required_entity_ids == (world.entity_id,)
        assert group.head_computation_id is None and group.strata == frozenset()
    _, _, source = computed(world.place, group_id)
    with world.place.uow() as uow:
        population = read(uow)
        group = next(item for item in population.groups if item.group_id == group_id)
        assert group.head_computation_id == source["id"]
        assert group.open_activity
        assert group.strata == frozenset({("ASC606", "TIME_ELAPSED"), ("ASC606", "OUTPUT_PERCENT")})
        assert population.required_entity_ids == (world.entity_id,)


def test_validation_population_refuses_partial_entity_scope(world: World) -> None:
    from erev_api.domain.contracts.validation_population import read

    principal = replace(world.place.principal, entity_scope=(world.entity_id,))
    with world.place.uow(principal) as uow, pytest.raises(Problem, match="complete tenant"):
        read(uow)


def test_validation_population_empty_other_tenant(world: World) -> None:
    from erev_api.domain.contracts.validation_population import read

    k01(world)
    principal = replace(world.place.principal, tenant_id=UUID(int=999))
    with world.place.uow(principal) as uow:
        population = read(uow)
        assert population.groups == ()
        assert population.required_entity_ids == ()


def test_validation_population_future_activity_is_not_open(world: World) -> None:
    from erev_api.domain.contracts.validation_population import read

    booked = booked_contract(
        world.place,
        {
            "external_id": "FUTURE-POPULATION",
            "customer_id": str(world.customer_id),
            "contracting_entity_code": "AVM-US",
            "transaction_currency": "USD",
            "inception_date": "2026-10-01",
            "lines": [
                {
                    "obligation_key": "O1",
                    "product_code": PLATFORM,
                    "quantity": "1",
                    "total_price": {"amount": "120.00", "currency": "USD"},
                    "start_date": "2026-10-01",
                    "end_date": "2026-12-31",
                }
            ],
        },
        activate=False,
    )
    with world.place.uow() as uow:
        population = read(uow)
        group = next(
            item for item in population.groups if item.group_id == booked.combination_group["id"]
        )
        assert not group.open_activity
        assert group.head_computation_id is None
        assert group.strata == frozenset()
        assert population.required_entity_ids == (world.entity_id,)


def test_validation_population_requires_sealed_ledger_activity(world: World) -> None:
    from erev_api.db.tables import subledger_line, subledger_posting, subledger_posting_seal
    from erev_api.domain.contracts.validation_population import read
    from support.rows import (
        insert_ledger_parts,
        ledger_seal_values,
        subledger_line_values,
        subledger_posting_values,
    )

    with world.place.uow() as uow:
        tenant_id = world.place.tenant_id
        parts = insert_ledger_parts(uow.session, tenant_id)
        posting = subledger_posting_values(tenant_id, parts=parts)
        uow.session.execute(insert(subledger_posting).values(**posting))
        lines = [
            subledger_line_values(tenant_id, posting=posting, parts=parts, amount=amount)
            for amount in (Decimal("10.00"), Decimal("-10.00"))
        ]
        uow.session.execute(insert(subledger_line), lines)
        before = next(group for group in read(uow).groups if group.group_id == parts.chain.group_id)
        assert not before.open_activity  # this probe has no contract events or current members
        seal = ledger_seal_values(uow.session, tenant_id, posting=posting, lines=lines)
        uow.session.execute(insert(subledger_posting_seal).values(**seal))
        after = next(group for group in read(uow).groups if group.group_id == parts.chain.group_id)
        assert after.open_activity
        uow.commit()


def pending_validation(world: World, *, level: str = "MINOR") -> UUID:
    from erev_api.controls.release import current_release
    from erev_api.db import new_id
    from erev_api.db.tables import release_validation

    release = current_release()
    assert release is not None
    validation_id = new_id()
    with world.place.uow() as uow:
        uow.session.execute(
            insert(release_validation).values(
                tenant_id=world.place.tenant_id,
                id=validation_id,
                engine_release_id=release.id,
                effective_level=level,
                status="PENDING",
                created_at=uow.now,
                created_by=uow.principal.id,
                created_by_kind=uow.principal.kind.value,
                updated_at=uow.now,
                updated_by=uow.principal.id,
                updated_by_kind=uow.principal.kind.value,
            )
        )
        uow.commit()
    return validation_id


def test_validation_attempt_freeze_preserves_selection_on_retry(world: World) -> None:
    from erev_api.db.tables import job, release_validation_attempt, release_validation_group
    from erev_api.domain.contracts.validation_attempts import freeze

    booked = k01(world)
    group_id = booked.combination_group["id"]
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        jobs_before = uow.session.scalar(select(func.count()).select_from(job))
        first = freeze(uow, validation_id=validation_id, attempt_no=1, seed=20260912)
        assert first.created
        assert len(uow.audit_events) == 1
        uow.commit()
    # The frozen NO_SOURCE remains NO_SOURCE even after a computation becomes available.
    computed(world.place, group_id)
    with world.place.uow() as uow:
        again = freeze(uow, validation_id=validation_id, attempt_no=1, seed=20260912)
        assert again.id == first.id and not again.created
        assert not uow.audit_events
        attempt = uow.session.execute(select(release_validation_attempt)).mappings().one()
        assert attempt["status"] == "SELECTED"
        assert attempt["required_entity_ids"] == [world.entity_id]
        assert attempt["population"]["no_source"] == 1
        selected = uow.session.execute(select(release_validation_group)).mappings().one()
        assert selected["source_computation_id"] is None
        assert selected["combination_group_id"] == group_id
        assert selected["selection"] == "OPEN_ACTIVITY"
        assert selected["child_job_id"] is None and selected["orchestration_state"] == "PENDING"
        assert uow.session.scalar(select(func.count()).select_from(job)) == jobs_before
        uow.commit()


def test_validation_attempt_freeze_rolls_back_selection_failure(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    from erev_api.db.tables import release_validation_attempt, release_validation_group
    from erev_api.domain.contracts.validation_attempts import freeze

    k01(world)
    validation_id = pending_validation(world)
    with (
        pytest.raises(RuntimeError, match="selection insert interrupted"),
        world.place.uow() as uow,
    ):
        execute = uow.session.execute

        def fail_selection(statement: Any, *args: Any, **kwargs: Any) -> Any:
            if getattr(statement, "table", None) is release_validation_group:
                raise RuntimeError("selection insert interrupted")
            return execute(statement, *args, **kwargs)

        monkeypatch.setattr(uow.session, "execute", fail_selection)
        freeze(uow, validation_id=validation_id, attempt_no=1, seed=20260912)
        uow.commit()
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(release_validation_attempt)) == 0
        assert uow.session.scalar(select(func.count()).select_from(release_validation_group)) == 0
        assert freeze(uow, validation_id=validation_id, attempt_no=1, seed=20260912).created
        uow.commit()


@pytest.mark.parametrize(
    "case",
    ["different seed", "active predecessor", "skipped number", "partial scope", "other tenant"],
)
def test_validation_attempt_freeze_refuses_invalid_redraw(world: World, case: str) -> None:
    from erev_api.domain.contracts.validation_attempts import freeze

    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        freeze(uow, validation_id=validation_id, attempt_no=1, seed=20260912)
        uow.commit()
    principal = world.place.principal
    if case == "partial scope":
        principal = replace(principal, entity_scope=(world.entity_id,))
    elif case == "other tenant":
        principal = replace(principal, tenant_id=UUID(int=999))
    number = 2 if case == "active predecessor" else 3 if case == "skipped number" else 1
    with world.place.uow(principal) as uow, pytest.raises(Problem):
        freeze(
            uow,
            validation_id=validation_id,
            attempt_no=number,
            seed=99 if case == "different seed" else 20260912,
        )


def test_validation_attempt_freeze_serializes_duplicate_creators(world: World) -> None:
    import threading

    from erev_api.db.tables import release_validation_attempt
    from erev_api.domain.contracts.validation_attempts import FrozenAttempt, freeze
    from support.interleave import await_lock_wait, backend_pid, observing_checkouts

    validation_id = pending_validation(world)
    results: list[FrozenAttempt] = []
    errors: list[Exception] = []

    def duplicate() -> None:
        try:
            with world.place.uow() as uow:
                results.append(
                    freeze(uow, validation_id=validation_id, attempt_no=1, seed=20260912)
                )
                uow.commit()
        except Exception as error:
            errors.append(error)

    request = threading.Thread(target=duplicate, name="freeze-duplicate")
    try:
        with observing_checkouts() as backends, world.place.uow() as holder:
            first = freeze(holder, validation_id=validation_id, attempt_no=1, seed=20260912)
            holder_pid = backend_pid(holder.session)
            request.start()
            blocked_pid, _ = await_lock_wait(
                holder.session,
                holder_pid=holder_pid,
                backends=backends,
                timeout=15,
                expect="release_validation",
            )
            assert blocked_pid != holder_pid and request.is_alive()
            holder.commit()
    finally:
        if request.ident is not None:
            request.join(timeout=30)
    assert not request.is_alive() and not errors
    assert results == [FrozenAttempt(first.id, False)]
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(release_validation_attempt)) == 1


def test_validation_attempt_resolves_seed_once_and_keeps_it_on_retry(world: World) -> None:
    from erev_api.db.tables import release_validation_attempt, release_validation_group
    from erev_api.db.transitions import apply
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.enums import RegistryCategory
    from support.rows import publish_registry_version

    k01(world)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        first = freeze(uow, validation_id=validation_id, attempt_no=1)
        row = uow.session.execute(select(release_validation_attempt)).mappings().one()
        assert row["seed"] == 20260912
        original_rank = uow.session.scalar(select(release_validation_group.c.rank_sha256))
        uow.commit()
    with world.place.uow() as uow:
        publish_registry_version(
            uow.session,
            tenant_id=world.place.tenant_id,
            category=RegistryCategory.PLATFORM,
            values={"platform.replay_seed": 123},
            at=uow.now,
        )
        uow.commit()
    with world.place.uow() as uow:
        assert freeze(uow, validation_id=validation_id, attempt_no=1).id == first.id
        assert not uow.audit_events
        apply(
            uow.session,
            "release_validation_attempt",
            first.id,
            to_status="SUPERSEDED",
            expected_status="SELECTED",
            expected_row_version=1,
            set_values={"row_version": 2},
        )
        uow.commit()
    with world.place.uow() as uow:
        second = freeze(uow, validation_id=validation_id, attempt_no=2)
        assert second.created and second.id != first.id
        rows = (
            uow.session.execute(
                select(release_validation_attempt).order_by(release_validation_attempt.c.attempt_no)
            )
            .mappings()
            .all()
        )
        assert [row["seed"] for row in rows] == [20260912, 123]
        rank = uow.session.scalar(
            select(release_validation_group.c.rank_sha256).where(
                release_validation_group.c.release_validation_attempt_id == second.id
            )
        )
        assert rank != original_rank
        uow.commit()


def test_validation_attempt_start_is_atomic_and_idempotent(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    from erev_api.db.tables import release_validation, release_validation_attempt
    from erev_api.domain.contracts import validation_attempts

    k01(world)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        attempt = validation_attempts.freeze(uow, validation_id=validation_id, attempt_no=1)
        uow.commit()
    from erev_api.db.transitions import apply as apply_transition

    def fail_attempt(session: Any, table: str, *args: Any, **kwargs: Any) -> Any:
        if table == "release_validation_attempt":
            raise RuntimeError("attempt transition interrupted")
        return apply_transition(session, table, *args, **kwargs)

    monkeypatch.setattr(validation_attempts, "apply", fail_attempt)
    with pytest.raises(RuntimeError, match="interrupted"), world.place.uow() as uow:
        validation_attempts.start(uow, validation_id=validation_id, attempt_id=attempt.id)
        uow.commit()
    monkeypatch.setattr(validation_attempts, "apply", apply_transition)
    with world.place.uow() as uow:
        assert uow.session.scalar(select(release_validation.c.status)) == "PENDING"
        assert uow.session.scalar(select(release_validation_attempt.c.status)) == "SELECTED"
        assert validation_attempts.start(uow, validation_id=validation_id, attempt_id=attempt.id)
        assert [event["action"] for event in uow.audit_events] == [
            "release_validation_attempt.start"
        ]
        uow.commit()
    with world.place.uow() as uow:
        assert not validation_attempts.start(
            uow, validation_id=validation_id, attempt_id=attempt.id
        )
        assert not uow.audit_events
        assert uow.session.scalar(select(release_validation.c.status)) == "RUNNING"
        assert uow.session.scalar(select(release_validation_attempt.c.status)) == "RUNNING"
        assert uow.session.scalar(select(release_validation.c.row_version)) == 2
        assert uow.session.scalar(select(release_validation_attempt.c.row_version)) == 2
        uow.commit()


@pytest.mark.parametrize("change", ["head", "event", "group", "running event"])
def test_validation_attempt_start_refuses_population_drift(world: World, change: str) -> None:
    from erev_api.domain.contracts.validation_attempts import freeze, start
    from support.rows import insert_contract_rows

    booked = k01(world)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        attempt = freeze(uow, validation_id=validation_id, attempt_no=1)
        if change == "running event":
            assert start(uow, validation_id=validation_id, attempt_id=attempt.id)
        uow.commit()
    if change == "head":
        computed(world.place, booked.combination_group["id"])
    elif change in {"event", "running event"}:
        flag_change(world, booked.contract["id"], 2)
    else:
        with world.place.uow() as uow:
            insert_contract_rows(uow.session, world.place.tenant_id)
            uow.commit()
    with world.place.uow() as uow, pytest.raises(Problem, match="population changed"):
        start(uow, validation_id=validation_id, attempt_id=attempt.id)


def test_validation_attempt_start_refuses_legacy_unfingerprinted_attempt(world: World) -> None:
    from erev_api.db.tables import release_validation_attempt
    from erev_api.domain.contracts.validation_attempts import start

    case = verification_case(world)
    attempt_id = case.decision.release_validation_attempt_id
    assert attempt_id is not None
    with world.place.uow() as uow:
        parent_id = uow.session.scalar(
            select(release_validation_attempt.c.release_validation_id).where(
                release_validation_attempt.c.id == attempt_id
            )
        )
        assert parent_id is not None
        with pytest.raises(Problem, match="population changed"):
            start(uow, validation_id=parent_id, attempt_id=attempt_id)
        assert not uow.audit_events


@pytest.mark.parametrize("state", ["SELECTED", "RUNNING", "COMPLETE", "INCOMPLETE"])
def test_validation_attempt_supersede_and_replace_atomically(world: World, state: str) -> None:
    from erev_api.db.tables import release_validation_attempt, release_validation_group
    from erev_api.db.transitions import apply
    from erev_api.domain.contracts.validation_attempts import freeze, start, supersede

    booked = k01(world)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        first = freeze(uow, validation_id=validation_id, attempt_no=1)
        if state != "SELECTED":
            start(uow, validation_id=validation_id, attempt_id=first.id)
        if state in {"COMPLETE", "INCOMPLETE"}:
            apply(
                uow.session,
                "release_validation_attempt",
                first.id,
                to_status=state,
                expected_status="RUNNING",
                set_values={},
            )
        frozen = dict(uow.session.execute(select(release_validation_attempt)).mappings().one())
        groups = [
            dict(row) for row in uow.session.execute(select(release_validation_group)).mappings()
        ]
        uow.commit()
    flag_change(world, booked.contract["id"], 2)
    with pytest.raises(RuntimeError, match="replacement interrupted"), world.place.uow() as uow:
        assert supersede(uow, validation_id=validation_id, attempt_id=first.id)
        freeze(uow, validation_id=validation_id, attempt_no=2)
        raise RuntimeError("replacement interrupted")
    with world.place.uow() as uow:
        assert uow.session.scalar(select(release_validation_attempt.c.status)) == state
        assert supersede(uow, validation_id=validation_id, attempt_id=first.id)
        second = freeze(uow, validation_id=validation_id, attempt_no=2)
        assert start(uow, validation_id=validation_id, attempt_id=second.id)
        assert not supersede(uow, validation_id=validation_id, attempt_id=first.id)
        assert [event["action"] for event in uow.audit_events].count(
            "release_validation_attempt.supersede"
        ) == 1
        old = (
            uow.session.execute(
                select(release_validation_attempt).where(
                    release_validation_attempt.c.id == first.id
                )
            )
            .mappings()
            .one()
        )
        assert old["status"] == "SUPERSEDED"
        for key in (
            "population",
            "seed",
            "required_entity_ids",
            "dataset_file_id",
            "report_run_id",
        ):
            assert old[key] == frozen[key]
        assert old["row_version"] == frozen["row_version"] + 1
        assert groups == [
            dict(row)
            for row in uow.session.execute(
                select(release_validation_group).where(
                    release_validation_group.c.release_validation_attempt_id == first.id
                )
            ).mappings()
        ]
        uow.commit()


def test_validation_attempt_supersede_rejects_late_result(world: World) -> None:
    from erev_api.db.tables import computation_verification, release_validation_attempt
    from erev_api.domain.contracts.validation_attempts import supersede
    from erev_api.domain.contracts.verification_store import persist

    case = verification_case(world)
    attempt_id = case.decision.release_validation_attempt_id
    assert attempt_id is not None
    with world.place.uow() as uow:
        validation_id = uow.session.scalar(
            select(release_validation_attempt.c.release_validation_id)
        )
        assert validation_id is not None
        assert supersede(uow, validation_id=validation_id, attempt_id=attempt_id)
        uow.commit()
    with world.place.uow() as uow, pytest.raises(Problem, match="running frozen validation"):
        persist(uow, case.decision, context=case.context, duration_ms=10)
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(computation_verification)) == 0


@pytest.mark.parametrize(
    "case", ["partial scope", "other tenant", "wrong parent", "pending approval", "linked approval"]
)
def test_validation_attempt_supersede_refuses_unsafe_context(world: World, case: str) -> None:
    from erev_api.db.tables import approval_request, release_validation_attempt
    from erev_api.db.transitions import apply
    from erev_api.domain.contracts.validation_attempts import freeze, start, supersede
    from support.rows import approval_request_values

    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        attempt = freeze(uow, validation_id=validation_id, attempt_no=1)
        if case == "pending approval":
            start(uow, validation_id=validation_id, attempt_id=attempt.id)
            apply(
                uow.session,
                "release_validation",
                validation_id,
                to_status="PENDING_APPROVAL",
                expected_status="RUNNING",
                set_values={},
            )
        elif case == "linked approval":
            request = approval_request_values(world.place.tenant_id)
            uow.session.execute(insert(approval_request).values(**request))
            uow.session.execute(
                update(release_validation_attempt)
                .where(release_validation_attempt.c.id == attempt.id)
                .values(approval_request_id=request["id"])
            )
        uow.commit()
    principal = world.place.principal
    if case == "partial scope":
        principal = replace(principal, entity_scope=(world.entity_id,))
    elif case == "other tenant":
        principal = replace(principal, tenant_id=UUID(int=999))
    elif case == "wrong parent":
        validation_id = UUID(int=999)
    with world.place.uow(principal) as uow, pytest.raises(Problem):
        supersede(uow, validation_id=validation_id, attempt_id=attempt.id)
    with world.place.uow() as uow:
        assert uow.session.scalar(select(release_validation_attempt.c.status)) != "SUPERSEDED"


def test_replay_sources_loads_frozen_evidence_after_head_changes(world: World) -> None:
    from erev_api.domain.contracts.replay_sources import load
    from erev_api.domain.contracts.replay_verify import verify

    case = verification_case(world)
    attempt_id = case.decision.release_validation_attempt_id
    assert attempt_id is not None
    flag_change(world, case.contract_id, 2)
    _, _, newer = computed(world.place, case.group_id)
    assert newer["id"] != case.decision.source_computation_id
    with world.place.uow() as uow:
        loaded = load(
            uow,
            context=case.context,
            attempt_id=attempt_id,
            group_id=case.group_id,
            trigger="REPLAY_VERIFY",
        )
        assert loaded.source is not None and loaded.evidence is not None
        assert loaded.source.computation_id == case.decision.source_computation_id
        assert loaded.token == case.decision.token
        row = verify(
            trigger="REPLAY_VERIFY",
            source=loaded.source,
            evidence=loaded.evidence,
            process=loaded.process,
            engine=engine(),
            token=loaded.token,
            attempt_id=attempt_id,
        )
        assert row == case.decision
        uow.commit()


@pytest.mark.parametrize(
    "case_name",
    [
        "partial scope",
        "superseded",
        "runtime mismatch",
        "missing file",
        "missing output",
        "corrupt input and missing output",
        "storage failure",
        "input metadata",
        "output metadata",
    ],
)
def test_replay_sources_refusals_before_execution(
    world: World, monkeypatch: pytest.MonkeyPatch, case_name: str
) -> None:
    from erev_api.db.tables import release_validation_attempt
    from erev_api.domain.contracts import replay_sources
    from erev_api.domain.contracts.replay_verify import verify
    from erev_api.domain.contracts.validation_attempts import supersede

    case = verification_case(world)
    attempt_id = case.decision.release_validation_attempt_id
    assert attempt_id is not None
    principal = world.place.principal
    if case_name == "partial scope":
        principal = replace(principal, entity_scope=(world.entity_id,))
    if case_name == "superseded":
        with world.place.uow() as uow:
            validation_id = uow.session.scalar(
                select(release_validation_attempt.c.release_validation_id)
            )
            assert validation_id is not None
            supersede(uow, validation_id=validation_id, attempt_id=attempt_id)
            uow.commit()
    if case_name == "runtime mismatch":
        release = case.context.runtime.engine_release
        assert release is not None
        context = case.context
        case = replace(
            case,
            context=JobContext(
                job_id=context.job_id,
                tenant_id=context.tenant_id,
                kind=context.kind,
                principal=context.principal,
                runtime=replace(
                    context.runtime, engine_release=replace(release, engine_version="9.9.9")
                ),
            ),
        )
    opened: list[UUID] = []
    original = open_file

    def observe(*args: Any, **kwargs: Any) -> Any:
        opened.append(args[1])
        if case_name in {"partial scope", "superseded", "runtime mismatch"}:
            pytest.fail("a refused replay must not open a file")
        if case_name == "missing file" or (
            case_name in {"missing output", "corrupt input and missing output"} and len(opened) == 2
        ):
            raise FileNotFoundError("missing retained file")
        if case_name == "storage failure":
            raise PermissionError("storage unavailable")
        meta, stream = original(*args, **kwargs)
        if case_name == "corrupt input and missing output":
            with stream:
                content = stream.read()
            stream = io.BytesIO(content + b" ")
        if (case_name == "input metadata" and len(opened) == 1) or (
            case_name == "output metadata" and len(opened) == 2
        ):
            meta = {**meta, "media_type": "text/plain"}
        return meta, stream

    monkeypatch.setattr(replay_sources, "open_file", observe)
    with world.place.uow(principal) as uow:
        if case_name in {"partial scope", "superseded", "storage failure"}:
            with pytest.raises(PermissionError if case_name == "storage failure" else Problem):
                replay_sources.load(
                    uow,
                    context=case.context,
                    attempt_id=attempt_id,
                    group_id=case.group_id,
                    trigger="REPLAY_VERIFY",
                )
            return
        loaded = replay_sources.load(
            uow,
            context=case.context,
            attempt_id=attempt_id,
            group_id=case.group_id,
            trigger="REPLAY_VERIFY",
        )
        assert loaded.source is not None

        def never_execute(bundle: Any) -> Any:
            pytest.fail("invalid or missing evidence must not execute the engine")

        row = verify(
            trigger="REPLAY_VERIFY",
            source=loaded.source,
            evidence=loaded.evidence,
            process=loaded.process,
            engine=never_execute,
            token=loaded.token,
            attempt_id=attempt_id,
        )
        expected = {
            "runtime mismatch": "SOURCE_RUNTIME_UNAVAILABLE",
            "missing file": "INPUT_EVIDENCE_MISSING",
            "missing output": "OUTPUT_EVIDENCE_MISSING",
            "corrupt input and missing output": "INPUT_EVIDENCE_HASH_MISMATCH",
            "input metadata": "INPUT_EVIDENCE_HASH_MISMATCH",
            "output metadata": "OUTPUT_EVIDENCE_HASH_MISMATCH",
        }
        assert row.reason == expected[case_name] and not row.engine_executed
        if case_name == "runtime mismatch":
            assert not opened
        uow.commit()


@pytest.mark.parametrize("wrong", ["task", "tenant", "release"])
def test_replay_sources_rejects_foreign_execution_before_file_access(
    world: World, monkeypatch: pytest.MonkeyPatch, wrong: str
) -> None:
    from erev_api.domain.contracts import replay_sources
    from erev_api.jobs.context import JobLost

    case = verification_case(world)
    attempt_id = case.decision.release_validation_attempt_id
    assert attempt_id is not None
    context = case.context
    runtime = context.runtime
    assert runtime.attempt is not None and runtime.engine_release is not None
    if wrong == "task":
        assert runtime.attempt.task_id is not None
        runtime = replace(
            runtime, attempt=replace(runtime.attempt, task_id=runtime.attempt.task_id + 10000)
        )
    elif wrong == "release":
        assert runtime.engine_release is not None
        runtime = replace(runtime, engine_release=replace(runtime.engine_release, id=UUID(int=999)))
    changed = JobContext(
        job_id=context.job_id,
        tenant_id=UUID(int=999) if wrong == "tenant" else context.tenant_id,
        kind=context.kind,
        principal=context.principal,
        runtime=runtime,
    )

    def forbidden_file(*args: Any, **kwargs: Any) -> Any:
        pytest.fail("foreign execution opened retained evidence")

    monkeypatch.setattr(replay_sources, "open_file", forbidden_file)
    with world.place.uow() as uow, pytest.raises(JobLost if wrong == "task" else Problem):
        replay_sources.load(
            uow,
            context=changed,
            attempt_id=attempt_id,
            group_id=case.group_id,
            trigger="REPLAY_VERIFY",
        )


@pytest.mark.parametrize("trigger", ["REPLAY_VERIFY", "UPGRADE_VALIDATE"])
@pytest.mark.parametrize("outcome", ["MATCH", "ERROR", "NO_SOURCE"])
def test_replay_execution_commits_only_evidence_and_reuses_result(
    world: World, monkeypatch: pytest.MonkeyPatch, outcome: str, trigger: Trigger
) -> None:
    from collections.abc import Iterator
    from contextlib import contextmanager

    from erev_api.db.tables import computation_verification, release_validation_group
    from erev_api.domain.contracts import replay_sources
    from erev_api.domain.contracts.replay_execution import execute_running
    from erev_api.uow import UnitOfWork

    case = verification_case(world, with_source=outcome != "NO_SOURCE")
    attempt_id = case.decision.release_validation_attempt_id
    assert attempt_id is not None
    original = case.context.unit_of_work
    active = 0
    calls = 0

    @contextmanager
    def tracked_unit() -> Iterator[UnitOfWork]:
        nonlocal active
        with original() as uow:
            active += 1
            try:
                yield uow
            finally:
                active -= 1

    monkeypatch.setattr(case.context, "unit_of_work", tracked_unit)

    def compute_once(bundle: Any) -> Any:
        nonlocal calls
        assert active == 0, "engine ran inside a database unit of work"
        calls += 1
        assert calls == 1 and outcome != "NO_SOURCE"
        if outcome == "ERROR":
            raise EngineError("ENGINE_INVARIANT_VIOLATED", "retry probe")
        return engine()(bundle)

    with world.place.uow() as uow:
        before = [
            uow.session.scalar(select(func.count()).select_from(table))
            for table in (contract_computation, contract_version, subledger_posting)
        ]
    result = execute_running(
        case.context,
        attempt_id=attempt_id,
        group_id=case.group_id,
        trigger=trigger,
        engine=compute_once,
    )
    assert result.outcome.value == outcome and result.accepted == (outcome != "ERROR")

    def no_file(*args: Any, **kwargs: Any) -> Any:
        pytest.fail("committed result retry reopened evidence")

    monkeypatch.setattr(replay_sources, "open_file", no_file)
    repeated = execute_running(
        case.context,
        attempt_id=attempt_id,
        group_id=case.group_id,
        trigger=trigger,
        engine=compute_once,
    )
    assert repeated == result
    with pytest.raises(Problem, match="stored replay result"):
        execute_running(
            case.context,
            attempt_id=attempt_id,
            group_id=case.group_id,
            trigger="UPGRADE_VALIDATE" if trigger == "REPLAY_VERIFY" else "REPLAY_VERIFY",
            engine=compute_once,
        )
    assert calls == (0 if outcome == "NO_SOURCE" else 1)
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(computation_verification)) == 1
        row = uow.session.execute(select(computation_verification)).mappings().one()
        assert row["engine_executed"] == (outcome != "NO_SOURCE")
        assert (
            row["actual_output_sha256"] is None
            if outcome != "MATCH"
            else row["actual_output_sha256"] is not None
        )
        assert uow.session.scalar(select(release_validation_group.c.orchestration_state)) == (
            "RUNNING" if outcome == "ERROR" else "COMPLETE"
        )
        assert before == [
            uow.session.scalar(select(func.count()).select_from(table))
            for table in (contract_computation, contract_version, subledger_posting)
        ]


def test_replay_execution_refuses_head_changed_during_engine(world: World) -> None:
    from erev_api.db.tables import computation_verification
    from erev_api.domain.contracts.replay_execution import execute_running

    case = verification_case(world)
    attempt_id = case.decision.release_validation_attempt_id
    assert attempt_id is not None

    def change_head(bundle: Any) -> Any:
        flag_change(world, case.contract_id, 2)
        computed(world.place, case.group_id)
        return engine()(bundle)

    with pytest.raises(Problem, match="head changed"):
        execute_running(
            case.context,
            attempt_id=attempt_id,
            group_id=case.group_id,
            trigger="REPLAY_VERIFY",
            engine=change_head,
        )
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(computation_verification)) == 0


def test_replay_job_scheduler_completes_selected_group(world: World) -> None:
    from erev_api.db.tables import computation_verification, job, release_validation_group
    from erev_api.domain.contracts import replay_job  # registers the worker handler
    from erev_api.enums import JobKind
    from erev_api.jobs.registry import HANDLERS, run_job

    assert HANDLERS[JobKind.REPLAY_VERIFY].handler is replay_job.replay
    case = verification_case(world, queued_only=True)
    run_job(
        case.context.job_id,
        world.place.tenant_id,
        attempt=1,
        runtime=replace(case.context.runtime, attempt=None),
        delivered_task_id=case.context.task_id,
    )
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(select(job.c.state).where(job.c.id == case.context.job_id))
            == "SUCCEEDED"
        )
        assert (
            uow.session.scalar(select(release_validation_group.c.orchestration_state)) == "COMPLETE"
        )
        assert uow.session.scalar(select(computation_verification.c.outcome)) == "MATCH"


def test_replay_job_scheduler_retries_error_then_exhausts(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    from erev_api.db.tables import computation_verification, job, release_validation_group
    from erev_api.domain.contracts import replay_job
    from erev_api.domain.contracts.replay_execution import execute_running
    from erev_api.jobs.registry import run_job
    from sqlalchemy import text

    case = verification_case(world, queued_only=True)

    def fail_engine(bundle: Any) -> Any:
        raise EngineError("ENGINE_INVARIANT_VIOLATED", "retry probe")

    def execute(context: JobContext, **kwargs: Any) -> Any:
        return execute_running(context, **kwargs, engine=fail_engine)

    monkeypatch.setattr(replay_job, "execute_running", execute)
    for attempt in (1, 2, 3):
        with world.place.uow() as uow:
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == case.context.job_id)
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            case.context.job_id,
            world.place.tenant_id,
            attempt=attempt,
            runtime=replace(case.context.runtime, attempt=None),
            delivered_task_id=task_id,
        )
        with world.place.uow() as uow:
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'succeeded' WHERE id = :id"),
                {"id": task_id},
            )
            assert uow.session.scalar(
                select(job.c.state).where(job.c.id == case.context.job_id)
            ) == ("FAILED" if attempt == 3 else "QUEUED")
            assert uow.session.scalar(select(release_validation_group.c.orchestration_state)) == (
                "FAILED_EXHAUSTED" if attempt == 3 else "RETRYABLE"
            )
            assert (
                uow.session.scalar(select(release_validation_group.c.accepted_verification_id))
                is None
            )
            assert (
                uow.session.scalar(select(func.count()).select_from(computation_verification))
                == attempt
            )
            uow.commit()


@pytest.mark.parametrize("during_engine", [False, True])
def test_replay_job_cancellation_preserves_no_accepted_result(
    world: World, monkeypatch: pytest.MonkeyPatch, during_engine: bool
) -> None:
    from erev_api.db.tables import computation_verification, job, release_validation_group
    from erev_api.domain.contracts import replay_job
    from erev_api.domain.contracts.replay_execution import execute_running
    from erev_api.domain.platform.jobs import cancel_job
    from erev_api.jobs.registry import run_job

    case = verification_case(world, queued_only=True)

    def cancel() -> None:
        with world.place.uow() as uow:
            cancel_job(uow, case.context.job_id)
            uow.commit()

    def cancel_while_computing(bundle: Any) -> Any:
        cancel()
        return engine()(bundle)

    def execute(context: JobContext, **kwargs: Any) -> Any:
        return execute_running(context, **kwargs, engine=cancel_while_computing)

    if during_engine:
        monkeypatch.setattr(replay_job, "execute_running", execute)
        run_job(
            case.context.job_id,
            world.place.tenant_id,
            attempt=1,
            runtime=replace(case.context.runtime, attempt=None),
            delivered_task_id=case.context.task_id,
        )
    else:
        cancel()
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(select(job.c.state).where(job.c.id == case.context.job_id))
            == "CANCELLED"
        )
        assert (
            uow.session.scalar(select(release_validation_group.c.orchestration_state))
            == "CANCELLED"
        )
        assert (
            uow.session.scalar(select(release_validation_group.c.accepted_verification_id)) is None
        )
        assert uow.session.scalar(select(func.count()).select_from(computation_verification)) == 0


def test_replay_job_cancel_does_not_wait_on_worker_commit(world: World) -> None:
    from erev_api.db.tables import job
    from erev_api.domain.contracts import replay_job
    from erev_api.domain.platform.jobs import cancel_job

    assert replay_job.guard_cancel is not None
    case = verification_case(world)
    with case.context.unit_of_work() as worker:
        with world.place.uow() as caller, pytest.raises(Problem) as refused:
            cancel_job(caller, case.context.job_id)
        assert refused.value.slug == "lock-conflict"
        worker.commit()
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(
                select(job.c.cancel_requested_at).where(job.c.id == case.context.job_id)
            )
            is None
        )


def test_replay_job_refuses_cancel_after_acceptance(world: World) -> None:
    from erev_api.db.tables import job, release_validation_group
    from erev_api.domain.contracts import replay_job
    from erev_api.domain.platform.jobs import cancel_job

    case = verification_case(world)
    result = replay_job.replay(
        case.context,
        {
            "release_validation_attempt_id": str(case.decision.release_validation_attempt_id),
            "combination_group_id": str(case.group_id),
            "trigger": "REPLAY_VERIFY",
        },
    )
    assert result.result["outcome"] == "MATCH"
    with world.place.uow() as uow, pytest.raises(Problem) as refused:
        cancel_job(uow, case.context.job_id)
    assert refused.value.slug == "invalid-transition"
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(
                select(job.c.cancel_requested_at).where(job.c.id == case.context.job_id)
            )
            is None
        )
        assert (
            uow.session.scalar(select(release_validation_group.c.orchestration_state)) == "COMPLETE"
        )


def test_replay_job_failure_cleanup_rolls_back_while_worker_is_active(world: World) -> None:
    from erev_api.db.tables import job, release_validation_group
    from erev_api.domain.contracts import replay_job
    from erev_api.jobs.registry import fail_attempt

    assert replay_job.failed is not None
    case = verification_case(world)
    with case.context.unit_of_work() as worker:
        with pytest.raises(Problem) as refused:
            fail_attempt(
                case.context.job_id,
                world.place.tenant_id,
                attempt=3,
                error=RuntimeError("failure probe"),
                runtime=replace(case.context.runtime, attempt=None),
                task_id=case.context.task_id,
            )
        assert refused.value.slug == "lock-conflict"
        worker.commit()
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(select(job.c.state).where(job.c.id == case.context.job_id))
            == "RUNNING"
        )
        assert (
            uow.session.scalar(select(release_validation_group.c.orchestration_state)) == "RUNNING"
        )


@pytest.mark.parametrize("different", [False, True])
def test_validation_fanout_runs_frozen_match_and_no_source(
    world: World, monkeypatch: pytest.MonkeyPatch, different: bool
) -> None:
    import json

    from erev_api.controls.release import current_release
    from erev_api.db.tables import (
        computation_verification,
        job,
        release_validation,
        release_validation_attempt,
        release_validation_group,
    )
    from erev_api.domain.contracts import replay_job
    from erev_api.domain.contracts.replay_execution import execute_running
    from erev_api.domain.contracts.validation_aggregation import assess
    from erev_api.domain.contracts.validation_artifacts import retain
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.domain.contracts.validation_fanout import enqueue
    from erev_api.enums import JobKind
    from erev_api.jobs.context import JobRuntime
    from erev_api.jobs.registry import HANDLERS, run_job
    from erev_engine.bundle import Diagnostic
    from sqlalchemy import text
    from support.rows import insert_contract_rows

    assert HANDLERS[JobKind.REPLAY_VERIFY].handler is replay_job.replay
    booked = k01(world)
    computed(world.place, booked.combination_group["id"])
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        insert_contract_rows(uow.session, world.place.tenant_id)
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        first = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        assert first.created == 2
        runtime = JobRuntime(
            clock=uow.clock, keyring=uow.keyring, files=uow.files, engine_release=current_release()
        )
        uow.commit()

    def candidate(bundle: Any) -> Any:
        output = engine()(bundle)
        if different:
            return replace(
                output,
                diagnostics=(
                    *output.diagnostics,
                    Diagnostic("DATASET_PROBE", "INFO", None, None, None, {}),
                ),
            )
        return output

    def accepted_before_job_finishes(context: JobContext, **kwargs: Any) -> Any:
        result = execute_running(context, **kwargs, engine=candidate)
        with world.place.uow() as uow:
            assert assess(uow, validation_id=validation_id, attempt_id=frozen.id).status is None
        return result

    monkeypatch.setattr(replay_job, "execute_running", accepted_before_job_finishes)
    for job_id in first.job_ids:
        with world.place.uow() as uow:
            child = uow.session.execute(select(job).where(job.c.id == job_id)).mappings().one()
            task_id = child["procrastinate_job_id"]
            assert task_id is not None
            lock = uow.session.scalar(
                text("SELECT lock FROM procrastinate_jobs WHERE id = :id"), {"id": task_id}
            )
            assert lock == f"verify:{world.place.tenant_id}:{child['subject_id']}"
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            job_id, world.place.tenant_id, attempt=1, runtime=runtime, delivered_task_id=task_id
        )
    from erev_api.db.transitions import apply as original_apply
    from erev_api.domain.contracts import validation_artifacts

    def fail_dataset_pointer(*args: Any, **kwargs: Any) -> None:
        raise RuntimeError("dataset pointer interrupted")

    with world.place.uow() as uow:
        files_before = uow.session.scalar(select(func.count()).select_from(file_object))
    monkeypatch.setattr(validation_artifacts, "apply", fail_dataset_pointer)
    with pytest.raises(RuntimeError, match="dataset pointer"), world.place.uow() as uow:
        retain(uow, validation_id=validation_id, attempt_id=frozen.id)
        uow.commit()
    monkeypatch.setattr(validation_artifacts, "apply", original_apply)
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(file_object)) == files_before
        assert uow.session.scalar(select(release_validation_attempt.c.dataset_file_id)) is None
        assert uow.session.scalar(select(release_validation_attempt.c.status)) == "RUNNING"
        again = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        assert again.job_ids == first.job_ids and again.created == again.recovery_requested == 0
        decision = assess(uow, validation_id=validation_id, attempt_id=frozen.id)
        assert decision.status == ("FAILED" if different else "PASSED")
        assert decision.counts == {
            "COMPLETE": 2,
            "DIFFERENCES" if different else "MATCH": 1,
            "NO_SOURCE": 1,
        }
        assert not uow.audit_events
        assert set(uow.session.scalars(select(computation_verification.c.outcome))) == {
            "DIFFERENCES" if different else "MATCH",
            "NO_SOURCE",
        }
        assert set(uow.session.scalars(select(release_validation_group.c.orchestration_state))) == {
            "COMPLETE"
        }
        dataset = retain(uow, validation_id=validation_id, attempt_id=frozen.id)
        meta, stream = open_file(uow.session, dataset.file_id, files=uow.files, keyring=uow.keyring)
        with stream:
            payload = json.loads(stream.read())
        assert {row["outcome"] for row in payload["groups"]} == {
            "DIFFERENCES" if different else "MATCH",
            "NO_SOURCE",
        }
        assert any(row["differences"] for row in payload["groups"]) is different
        assert meta["sha256"] == dataset.sha256
        assert uow.session.scalar(select(release_validation_attempt.c.status)) == "COMPLETE"
        assert uow.session.scalar(select(release_validation.c.enabled_at)) is None
        uow.commit()
    with world.place.uow() as uow:
        count = uow.session.scalar(select(func.count()).select_from(file_object))
        assert retain(uow, validation_id=validation_id, attempt_id=frozen.id) == dataset
        assert assess(uow, validation_id=validation_id, attempt_id=frozen.id) == decision
        assert not uow.audit_events
        assert uow.session.scalar(select(func.count()).select_from(file_object)) == count

    from erev_api.db import new_id
    from erev_api.db.tables import release_group_processing
    from sqlalchemy.exc import DBAPIError

    # Schema probe: processing requires acceptance; the surrounding savepoint rolls it back.
    with world.place.uow() as uow:
        row = {
            "tenant_id": world.place.tenant_id,
            "id": new_id(),
            "release_validation_id": validation_id,
            "release_validation_attempt_id": frozen.id,
            "combination_group_id": booked.combination_group["id"],
            "entity_id": world.entity_id,
            "book_code": "ASC606",
            "seq": 1,
            "upgrade_state": "PENDING",
            "created_at": uow.now,
            "created_by": uow.principal.id,
            "created_by_kind": uow.principal.kind.value,
        }
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(insert(release_group_processing).values(**row))
        savepoint = uow.session.begin_nested()
        original_apply(
            uow.session,
            "release_validation",
            validation_id,
            to_status="PASSED",
            expected_status="RUNNING",
            set_values={"accepted_attempt_id": frozen.id, "enabled_at": uow.now},
        )
        from erev_api.db.tables import book
        from erev_api.domain.contracts.processing_seed import seed

        with world.place.uow() as concurrent:
            assert (
                concurrent.session.scalar(
                    text("SELECT pg_try_advisory_xact_lock(hashtextextended(:key, 0))"),
                    {
                        "key": f"erev:upgrade-group:{world.place.tenant_id}:"
                        f"{booked.combination_group['id']}"
                    },
                )
                is True
            )
            with pytest.raises(DBAPIError) as processing_busy, uow.session.begin_nested():
                seed(uow, validation_id=validation_id)
            assert getattr(processing_busy.value.orig, "sqlstate", None) == "55P03"
        assert uow.session.scalar(select(func.count()).select_from(release_group_processing)) == 0
        seed_probe = uow.session.begin_nested()
        enabled_books = set(uow.session.scalars(select(book.c.code).where(book.c.is_enabled)))
        assert seed(uow, validation_id=validation_id) == len(enabled_books)
        seeded = list(uow.session.execute(select(release_group_processing)).mappings())
        assert {item["book_code"] for item in seeded} == enabled_books
        assert len({item["combination_group_id"] for item in seeded}) == 1
        assert len({item["entity_id"] for item in seeded}) == 1
        assert all(item["upgrade_state"] == "PENDING" and item["seq"] == 1 for item in seeded)
        manifest = uow.session.scalar(
            select(release_validation.c.processing_seed).where(
                release_validation.c.id == validation_id
            )
        )
        from erev_api.db.tables import posting_attribution

        source_id = uow.session.scalar(
            select(combination_group.c.head_computation_id).where(
                combination_group.c.id == booked.combination_group["id"]
            )
        )
        posting_id = uow.session.scalar(
            select(subledger_posting.c.id).where(
                subledger_posting.c.contract_computation_id == source_id,
                subledger_posting.c.book_code == "ASC606",
            )
        )
        assert source_id is not None and posting_id is not None
        attribution = {
            "tenant_id": world.place.tenant_id,
            "id": new_id(),
            "subledger_posting_id": posting_id,
            "release_validation_attempt_id": frozen.id,
            "combination_group_id": booked.combination_group["id"],
            "contract_computation_id": source_id,
            "cause": "UPGRADE_CONTEXT",
            "sampled": False,
            "origin_periods": {},
            "created_at": uow.now,
            "created_by": uow.principal.id,
            "created_by_kind": uow.principal.kind.value,
        }
        attribution_probe = uow.session.begin_nested()
        uow.session.execute(insert(posting_attribution).values(**attribution))
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                insert(posting_attribution).values(**{**attribution, "id": new_id()})
            )
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(update(posting_attribution).values(sampled=True))
        attribution_probe.rollback()
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                insert(posting_attribution).values(**{**attribution, "cause": "UPGRADE_ONLY"})
            )
        terminal_probe = uow.session.begin_nested()
        uow.session.execute(
            insert(release_group_processing).values(
                **{
                    **row,
                    "id": new_id(),
                    "seq": 2,
                    "upgrade_state": "POSTED",
                    "contract_computation_id": source_id,
                    "subledger_posting_ids": [posting_id],
                }
            )
        )
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(insert(posting_attribution).values(**attribution))
        terminal_probe.rollback()
        audit_count = len(uow.audit_events)
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                update(release_validation)
                .where(release_validation.c.id == validation_id)
                .values(processing_seed={"version": 1, "count": 999})
            )
        # A later group cannot be silently added by a retry after acceptance.
        insert_contract_rows(uow.session, world.place.tenant_id)
        assert seed(uow, validation_id=validation_id) == len(seeded)
        assert len(uow.audit_events) == audit_count
        assert (
            uow.session.scalar(
                select(release_validation.c.processing_seed).where(
                    release_validation.c.id == validation_id
                )
            )
            == manifest
        )
        seed_probe.rollback()
        assert (
            uow.session.scalar(
                select(release_validation.c.processing_seed).where(
                    release_validation.c.id == validation_id
                )
            )
            is None
        )
        assert uow.session.scalar(select(func.count()).select_from(release_group_processing)) == 0
        uow.session.execute(insert(release_group_processing).values(**row))
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(insert(release_group_processing).values(**{**row, "id": new_id()}))
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                update(release_group_processing)
                .where(release_group_processing.c.id == row["id"])
                .values(reason="rewritten")
            )
        with pytest.raises(Problem) as unmanifested:
            seed(uow, validation_id=validation_id)
        assert unmanifested.value.slug == "validation-failed"
        source_computation = uow.session.scalar(
            select(combination_group.c.head_computation_id).where(
                combination_group.c.id == booked.combination_group["id"]
            )
        )
        assert source_computation is not None
        deferred = {
            **row,
            "id": new_id(),
            "seq": 2,
            "upgrade_state": "DEFERRED_NO_OPEN_PERIOD",
            "contract_computation_id": source_computation,
            "unresolved_origin_period_keys": ["FY2026-P01"],
        }
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(insert(release_group_processing).values(**{**deferred, "seq": 3}))
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                insert(release_group_processing).values(
                    **{**deferred, "subledger_posting_ids": [new_id()]}
                )
            )
        posting_ids = list(
            uow.session.scalars(
                select(subledger_posting.c.id).where(
                    subledger_posting.c.contract_computation_id == source_computation,
                    subledger_posting.c.book_code == "ASC606",
                )
            )
        )
        assert posting_ids
        posting_probe = uow.session.begin_nested()
        uow.session.execute(
            insert(release_group_processing).values(
                **{**deferred, "subledger_posting_ids": posting_ids}
            )
        )
        # Resolving origins must preserve postings recorded during an earlier deferral.
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                insert(release_group_processing).values(**{**deferred, "id": new_id(), "seq": 3})
            )
        uow.session.execute(
            insert(release_group_processing).values(
                **{
                    **deferred,
                    "id": new_id(),
                    "seq": 3,
                    "upgrade_state": "POSTED",
                    "subledger_posting_ids": posting_ids,
                    "unresolved_origin_period_keys": [],
                }
            )
        )
        posting_probe.rollback()
        uow.session.execute(insert(release_group_processing).values(**deferred))
        settled = {
            **deferred,
            "id": new_id(),
            "seq": 3,
            "upgrade_state": "SETTLED_NO_POSTING",
            "unresolved_origin_period_keys": [],
            "reason": "NO_POSTING_DELTA",
        }
        uow.session.execute(insert(release_group_processing).values(**settled))
        with pytest.raises(DBAPIError), uow.session.begin_nested():
            uow.session.execute(
                insert(release_group_processing).values(**{**settled, "id": new_id(), "seq": 4})
            )
        assert uow.session.scalar(select(func.count()).select_from(release_group_processing)) == 3
        savepoint.rollback()
        assert uow.session.scalar(select(func.count()).select_from(release_group_processing)) == 0

    from datetime import timedelta

    from erev_api.domain.contracts.upgrade_report import upgrade_validation_report
    from erev_api.domain.reports import validation_sources
    from erev_api.domain.reports.builders import RETAINED, ReportParams, SourceCollector
    from erev_api.domain.reports.builders import upgrade_validation as builder
    from erev_api.domain.reports.outputs import json_rows
    from erev_api.domain.reports.outputs.csv import render_csv

    entities = tuple(UUID(value) for value in payload["attempt"]["required_entity_ids"])
    with world.place.uow() as uow:
        source = validation_sources.read_for_report(
            uow,
            frozen.id,
            entity_ids=entities,
            known_at=uow.now,
            expected_file_id=dataset.file_id,
            expected_sha256=dataset.sha256,
        )
        report = upgrade_validation_report(source.content)
        collector = SourceCollector()
        params = ReportParams(
            builder.CODE,
            1,
            {"release_validation_attempt_id": str(frozen.id)},
            entities,
            uow.now,
            sources=collector,
            output_format="JSON",
        )
        built = builder.build(uow, params)
        assert built.control_totals["dataset_sha256"] == dataset.sha256
        assert built.row_count == len(report.rows)
        original_json = json_rows(built)
        original_csv = render_csv(built)
        binding = collector.binding(
            row_keys=[str(row["row_key"]) for row in built.rows],
            strategy=RETAINED,
            report_version=1,
        )
        requester = uow.ctx
        bound_params = replace(params, binding=binding, sources=None)
        assert report.control_totals["groups"] == 2
        assert report.control_totals["outcomes"] == {
            "DIFFERENCES" if different else "MATCH": 1,
            "NO_SOURCE": 1,
        }
        with pytest.raises(Problem):
            validation_sources.read_for_report(
                uow, frozen.id, entity_ids=entities, known_at=uow.now - timedelta(days=1)
            )
    flag_change(world, booked.contract["id"], 2)
    with world.place.uow() as uow, pytest.raises(Problem) as completed_drift:
        assess(uow, validation_id=validation_id, attempt_id=frozen.id)
    assert completed_drift.value.slug == "validation-failed"
    with world.place.uow() as uow:
        repeated = validation_sources.read_for_report(
            uow,
            frozen.id,
            entity_ids=entities,
            known_at=uow.now,
            expected_file_id=dataset.file_id,
            expected_sha256=dataset.sha256,
        )
        assert repeated == source
        assert upgrade_validation_report(repeated.content) == report
        rebuilt = builder.build(uow, bound_params)
        assert json_rows(rebuilt) == original_json
        assert render_csv(rebuilt) == original_csv
        assert rebuilt.control_totals == built.control_totals
        broken = replace(binding, evidence={})
        with pytest.raises(Problem):
            builder.build(uow, replace(bound_params, binding=broken))
        with pytest.raises(Problem):
            builder.build(uow, replace(bound_params, filters={"outcome": "MATCH"}))
    with world.place.uow(system_principal(world.place.tenant_id)) as uow:
        with pytest.raises(Problem):
            builder.build(uow, bound_params)
        assert builder.build(uow, replace(bound_params, requester=requester)) == built
        revoked = replace(
            requester,
            principal=replace(
                requester.principal, permissions=requester.principal.permissions - {"report.run"}
            ),
        )
        with pytest.raises(Problem):
            builder.build(uow, replace(bound_params, requester=revoked))
        wrong_tenant = replace(
            requester, principal=replace(requester.principal, tenant_id=UUID(int=999))
        )
        with pytest.raises(Problem):
            builder.build(uow, replace(bound_params, requester=wrong_tenant))

    from erev_api.db.tables import report_run
    from erev_api.domain.reports import framework
    from erev_api.schemas.reports import ReportRunCreateIn

    with world.place.uow() as uow:
        report_id, _ = framework.create_run(
            uow,
            ReportRunCreateIn(
                report_code="upgrade_validation",
                output_format="CSV" if different else "JSON",
                parameters={"release_validation_attempt_id": str(frozen.id)},
            ),
        )
        report_job = uow.session.scalar(
            select(report_run.c.job_id).where(report_run.c.id == report_id)
        )
        uow.commit()
    from erev_api.domain.contracts.validation_reports import attach

    original_report_id = report_id
    with world.place.uow() as uow, pytest.raises(Problem) as pending_report:
        attach(uow, validation_id=validation_id, attempt_id=frozen.id, report_id=report_id)
    assert pending_report.value.slug == "invalid-transition"
    outputs = []
    for iteration in range(2):
        assert report_job is not None
        with world.place.uow() as uow:
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == report_job)
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            report_job, world.place.tenant_id, attempt=1, runtime=runtime, delivered_task_id=task_id
        )
        with world.place.uow() as uow:
            saved = (
                uow.session.execute(select(report_run).where(report_run.c.id == report_id))
                .mappings()
                .one()
            )
            assert saved["status"] == "SUCCEEDED", dict(saved)
            outputs.append((saved["output_sha256"], saved["control_totals"]))
            if iteration == 0:
                report_id, _ = framework.rerun(uow, report_id)
                report_job = uow.session.scalar(
                    select(report_run.c.job_id).where(report_run.c.id == report_id)
                )
                uow.commit()
    assert outputs[0] == outputs[1]
    # A corrupted output cannot become the attempt's official report.
    import io

    from erev_api.domain.contracts import validation_reports
    from erev_api.files.store import open_file as read_report_file

    def corrupt_report_file(*args: Any, **kwargs: Any) -> Any:
        metadata, stream = read_report_file(*args, **kwargs)
        stream.close()
        return metadata, io.BytesIO(b"corrupted report output")

    with monkeypatch.context() as patch:
        patch.setattr(validation_reports, "open_file", corrupt_report_file)
        with world.place.uow() as uow, pytest.raises(Problem) as corrupted_report:
            attach(
                uow, validation_id=validation_id, attempt_id=frozen.id, report_id=original_report_id
            )
        assert corrupted_report.value.slug == "validation-failed"
    # Rollback leaves the attachment unset; publication commits once and retries are read-only.
    with world.place.uow() as uow:
        attach(uow, validation_id=validation_id, attempt_id=frozen.id, report_id=original_report_id)
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(
                select(release_validation_attempt.c.report_run_id).where(
                    release_validation_attempt.c.id == frozen.id
                )
            )
            is None
        )
        attach(uow, validation_id=validation_id, attempt_id=frozen.id, report_id=original_report_id)
        uow.commit()
    with world.place.uow() as uow:
        version = uow.session.scalar(
            select(release_validation_attempt.c.row_version).where(
                release_validation_attempt.c.id == frozen.id
            )
        )
        assert (
            attach(
                uow, validation_id=validation_id, attempt_id=frozen.id, report_id=original_report_id
            )
            == original_report_id
        )
        assert (
            uow.session.scalar(
                select(release_validation_attempt.c.row_version).where(
                    release_validation_attempt.c.id == frozen.id
                )
            )
            == version
        )
        assert (
            uow.session.scalar(
                select(release_validation.c.enabled_at).where(
                    release_validation.c.id == validation_id
                )
            )
            is None
        )
        uow.commit()
    with world.place.uow() as uow, pytest.raises(Problem) as replacement_report:
        attach(uow, validation_id=validation_id, attempt_id=frozen.id, report_id=report_id)
    assert replacement_report.value.slug == "invalid-transition"
    no_report_grant = replace(
        world.place.principal,
        permissions=world.place.principal.permissions - {"report.run"},
    )
    with world.place.uow(no_report_grant) as uow, pytest.raises(Problem) as unauthorized:
        attach(uow, validation_id=validation_id, attempt_id=frozen.id, report_id=original_report_id)
    assert unauthorized.value.slug == "forbidden"

    from erev_api.db.tables import api_client
    from erev_api.enums import PrincipalKind
    from support.rows import api_client_values

    client_row = api_client_values(
        world.place.tenant_id,
        status="ACTIVE",
        is_all_entities=True,
        scopes=["report.run", "report.export"],
    )
    with world.place.uow() as uow:
        uow.session.execute(insert(api_client).values(**client_row))
        uow.commit()
    client_principal = replace(
        world.place.principal,
        kind=PrincipalKind.API_CLIENT,
        id=client_row["id"],
        membership_id=None,
        permissions=frozenset({"report.run", "report.export"}),
        permission_scopes={"report.run": "*", "report.export": "*"},
        roles=(),
        auth_method="client_credentials",
        mfa_verified_at=None,
        session_id=None,
    )
    for client_revoked in (False, True):
        with world.place.uow(client_principal) as uow:
            client_report, _ = framework.create_run(
                uow,
                ReportRunCreateIn(
                    report_code="upgrade_validation",
                    output_format="CSV",
                    parameters={"release_validation_attempt_id": str(frozen.id)},
                ),
            )
            client_job = uow.session.scalar(
                select(report_run.c.job_id).where(report_run.c.id == client_report)
            )
            assert client_job is not None
            uow.commit()
        with world.place.uow() as uow:
            if client_revoked:
                uow.session.execute(
                    update(api_client)
                    .where(api_client.c.id == client_row["id"])
                    .values(status="REVOKED", updated_at=uow.now)
                )
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == client_job)
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            client_job, world.place.tenant_id, attempt=1, runtime=runtime, delivered_task_id=task_id
        )
        if client_revoked:
            with world.place.uow() as uow:
                uow.session.execute(
                    text("UPDATE procrastinate_jobs SET status = 'succeeded' WHERE id = :id"),
                    {"id": task_id},
                )
                task_id = uow.session.scalar(
                    select(job.c.procrastinate_job_id).where(job.c.id == client_job)
                )
                uow.session.execute(
                    text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                    {"id": task_id},
                )
                uow.commit()
            run_job(
                client_job,
                world.place.tenant_id,
                attempt=2,
                runtime=runtime,
                delivered_task_id=task_id,
            )
        with world.place.uow() as uow:
            saved = (
                uow.session.execute(select(report_run).where(report_run.c.id == client_report))
                .mappings()
                .one()
            )
            if client_revoked:
                assert saved["output_file_id"] is None
                problem = uow.session.scalar(select(job.c.problem).where(job.c.id == client_job))
                assert problem is not None and problem["type"].endswith("/forbidden")
            else:
                assert saved["status"] == "SUCCEEDED", dict(saved)

    principal = world.place.principal
    denied = [
        replace(principal, permissions=principal.permissions - {"report.run"}),
        replace(
            principal,
            permission_scopes={
                **principal.permission_scopes,
                "report.run": frozenset({UUID(int=999)}),
            },
        ),
        replace(principal, tenant_id=UUID(int=999)),
        system_principal(principal.tenant_id),
    ]
    original_open = open_file

    def unread(*args: Any, **kwargs: Any) -> Any:
        pytest.fail("Unauthorized report opened its dataset")

    monkeypatch.setattr(validation_sources, "open_file", unread)
    for actor in denied:
        with world.place.uow(actor) as uow, pytest.raises(Problem):
            validation_sources.read_for_report(
                uow, frozen.id, entity_ids=entities, known_at=uow.now
            )
    with world.place.uow() as uow:
        with pytest.raises(Problem):
            validation_sources.read_for_report(uow, frozen.id, entity_ids=(), known_at=uow.now)
        with pytest.raises(Problem):
            validation_sources.read_for_report(
                uow,
                frozen.id,
                entity_ids=entities,
                known_at=uow.now,
                expected_file_id=UUID(int=999),
                expected_sha256=dataset.sha256,
            )
    monkeypatch.setattr(validation_sources, "open_file", original_open)

    def tampered(*args: Any, **kwargs: Any) -> Any:
        meta, stream = original_open(*args, **kwargs)
        with stream:
            content = stream.read() + b" "
        return meta, io.BytesIO(content)

    monkeypatch.setattr(validation_sources, "open_file", tampered)
    with world.place.uow() as uow, pytest.raises(Problem):
        validation_sources.read_for_report(uow, frozen.id, entity_ids=entities, known_at=uow.now)


def test_validation_fanout_rolls_back_partial_binding(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    from erev_api.db.tables import (
        job,
        release_validation,
        release_validation_attempt,
        release_validation_group,
    )
    from erev_api.domain.contracts import validation_fanout
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.jobs.registry import defer
    from support.rows import insert_contract_rows

    k01(world)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        insert_contract_rows(uow.session, world.place.tenant_id)
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        uow.commit()
    calls = 0

    def fail_second(*args: Any, **kwargs: Any) -> Any:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("fanout interrupted")
        return defer(*args, **kwargs)

    monkeypatch.setattr(validation_fanout, "defer", fail_second)
    with world.place.uow() as uow, pytest.raises(RuntimeError, match="fanout interrupted"):
        validation_fanout.enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(
                select(func.count()).select_from(job).where(job.c.kind == "REPLAY_VERIFY")
            )
            == 0
        )
        assert set(uow.session.scalars(select(release_validation_group.c.child_job_id))) == {None}
        assert uow.session.scalar(select(release_validation.c.status)) == "PENDING"
        assert uow.session.scalar(select(release_validation_attempt.c.status)) == "SELECTED"


def test_validation_fanout_recovers_lost_dispatch_once(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    from erev_api.db.tables import job
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.domain.contracts.validation_fanout import enqueue
    from erev_api.jobs import registry
    from sqlalchemy import text

    k01(world)
    validation_id = pending_validation(world)
    original = registry._dispatch_hook
    monkeypatch.setattr(registry, "_dispatch_hook", lambda **kwargs: None)
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        first = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        uow.commit()
    monkeypatch.setattr(registry, "_dispatch_hook", original)
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == first.job_ids[0])
            )
            is None
        )
        for _ in range(2):
            retry = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
            assert (
                retry.job_ids == first.job_ids
                and retry.created == 0
                and retry.recovery_requested == 1
            )
        uow.commit()
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == first.job_ids[0])
            )
            is not None
        )
        assert (
            uow.session.scalar(
                text("SELECT count(*) FROM procrastinate_jobs WHERE args->>'job_id' = :id"),
                {"id": str(first.job_ids[0])},
            )
            == 1
        )
        retry = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        assert retry.created == retry.recovery_requested == 0
        uow.commit()


@pytest.mark.parametrize("failure", ["drift", "mode", "release"])
def test_validation_fanout_refuses_changed_basis(
    world: World, monkeypatch: pytest.MonkeyPatch, failure: str
) -> None:
    from erev_api.db.tables import job
    from erev_api.domain.contracts import validation_fanout
    from erev_api.domain.contracts.validation_attempts import freeze

    booked = k01(world)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        if failure == "mode":
            validation_fanout.enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        uow.commit()
    if failure == "drift":
        flag_change(world, booked.contract["id"], 2)
    if failure == "release":
        monkeypatch.setattr(validation_fanout, "current_release", lambda: None)
    with world.place.uow() as uow, pytest.raises(Problem):
        validation_fanout.enqueue(
            uow, validation_id=validation_id, attempt_id=frozen.id, trigger="REPLAY_VERIFY"
        )
    with world.place.uow() as uow:
        assert uow.session.scalar(
            select(func.count()).select_from(job).where(job.c.kind == "REPLAY_VERIFY")
        ) == (1 if failure == "mode" else 0)


def test_validation_fanout_serializes_concurrent_deliveries(world: World) -> None:
    import threading

    from erev_api.db.tables import job
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.domain.contracts.validation_fanout import Fanout, enqueue
    from sqlalchemy import text
    from support.interleave import await_lock_wait, backend_pid, observing_checkouts

    k01(world)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        uow.commit()
    results: list[Fanout] = []
    errors: list[Exception] = []

    def duplicate() -> None:
        try:
            with world.place.uow() as uow:
                results.append(enqueue(uow, validation_id=validation_id, attempt_id=frozen.id))
                uow.commit()
        except Exception as error:
            errors.append(error)

    contender = threading.Thread(target=duplicate, name="replay-fanout-duplicate")
    try:
        with observing_checkouts() as backends, world.place.uow() as holder:
            first = enqueue(holder, validation_id=validation_id, attempt_id=frozen.id)
            holder_pid = backend_pid(holder.session)
            contender.start()
            _, _ = await_lock_wait(
                holder.session,
                holder_pid=holder_pid,
                backends=backends,
                timeout=15,
                expect="release_validation",
            )
            assert contender.is_alive()
            holder.commit()
    finally:
        if contender.ident is not None:
            contender.join(timeout=30)
    assert not contender.is_alive() and not errors
    assert len(results) == 1 and results[0].job_ids == first.job_ids and results[0].created == 0
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(
                select(func.count()).select_from(job).where(job.c.kind == "REPLAY_VERIFY")
            )
            == 1
        )
        assert (
            uow.session.scalar(
                text("SELECT count(*) FROM procrastinate_jobs WHERE args->>'job_id' = :id"),
                {"id": str(first.job_ids[0])},
            )
            == 1
        )


@pytest.mark.parametrize(
    "scenario", ["queued", "cancelled", "drift", "superseded", "premature_complete"]
)
def test_validation_aggregation_stored_lifecycle(world: World, scenario: str) -> None:
    from erev_api.domain.contracts import replay_job
    from erev_api.domain.contracts.validation_aggregation import assess
    from erev_api.domain.contracts.validation_artifacts import retain
    from erev_api.domain.contracts.validation_attempts import freeze, supersede
    from erev_api.domain.contracts.validation_fanout import enqueue
    from erev_api.domain.platform.jobs import cancel_job

    assert replay_job.replay is not None
    booked = k01(world)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        fanout = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        uow.commit()
    if scenario == "drift":
        flag_change(world, booked.contract["id"], 2)
    with world.place.uow() as uow:
        if scenario == "cancelled":
            cancel_job(uow, fanout.job_ids[0])
        if scenario == "superseded":
            supersede(uow, validation_id=validation_id, attempt_id=frozen.id)
        if scenario == "premature_complete":
            from erev_api.db.transitions import apply

            apply(
                uow.session,
                "release_validation_attempt",
                frozen.id,
                to_status="COMPLETE",
                expected_status="RUNNING",
                set_values={},
            )
        if scenario in {"drift", "superseded", "premature_complete"}:
            with pytest.raises(Problem):
                assess(uow, validation_id=validation_id, attempt_id=frozen.id)
            with pytest.raises(Problem):
                retain(uow, validation_id=validation_id, attempt_id=frozen.id)
        else:
            result = assess(uow, validation_id=validation_id, attempt_id=frozen.id)
            assert result.status == ("INCOMPLETE" if scenario == "cancelled" else None)
            assert result.reasons == (("CANCELLED",) if scenario == "cancelled" else ())
            if scenario == "cancelled":
                assert retain(uow, validation_id=validation_id, attempt_id=frozen.id).file_id
                assert assess(uow, validation_id=validation_id, attempt_id=frozen.id) == result
            else:
                with pytest.raises(Problem):
                    retain(uow, validation_id=validation_id, attempt_id=frozen.id)


def test_validation_aggregation_waits_for_retries_then_reports_exhaustion(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    import json

    from erev_api.controls.release import current_release
    from erev_api.db.tables import job, release_validation_attempt
    from erev_api.domain.contracts import replay_job
    from erev_api.domain.contracts.replay_execution import execute_running
    from erev_api.domain.contracts.validation_aggregation import assess
    from erev_api.domain.contracts.validation_artifacts import retain
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.domain.contracts.validation_fanout import enqueue
    from erev_api.jobs.context import JobRuntime
    from erev_api.jobs.registry import run_job
    from sqlalchemy import text

    booked = k01(world)
    computed(world.place, booked.combination_group["id"])
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        child_id = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id).job_ids[0]
        runtime = JobRuntime(
            clock=uow.clock, keyring=uow.keyring, files=uow.files, engine_release=current_release()
        )
        uow.commit()

    def fail(bundle: Any) -> Any:
        raise EngineError("ENGINE_INVARIANT_VIOLATED", "aggregate retry probe")

    def execute(context: JobContext, **kwargs: Any) -> Any:
        return execute_running(context, **kwargs, engine=fail)

    monkeypatch.setattr(replay_job, "execute_running", execute)
    for attempt in (1, 2, 3):
        with world.place.uow() as uow:
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == child_id)
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            child_id,
            world.place.tenant_id,
            attempt=attempt,
            runtime=runtime,
            delivered_task_id=task_id,
        )
        with world.place.uow() as uow:
            result = assess(uow, validation_id=validation_id, attempt_id=frozen.id)
            assert result.status == ("INCOMPLETE" if attempt == 3 else None)
            assert result.reasons == (("FAILED_EXHAUSTED",) if attempt == 3 else ())
            if attempt == 3:
                dataset = retain(uow, validation_id=validation_id, attempt_id=frozen.id)
                _, stream = open_file(
                    uow.session, dataset.file_id, files=uow.files, keyring=uow.keyring
                )
                with stream:
                    payload = json.loads(stream.read())
                assert assess(uow, validation_id=validation_id, attempt_id=frozen.id) == result
                assert payload["groups"][0]["outcome"] == "ERROR"
                assert payload["groups"][0]["orchestration_state"] == "FAILED_EXHAUSTED"
                assert (
                    uow.session.scalar(select(release_validation_attempt.c.status)) == "INCOMPLETE"
                )
            else:
                with pytest.raises(Problem):
                    retain(uow, validation_id=validation_id, attempt_id=frozen.id)
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'succeeded' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()


def test_validation_dataset_is_encrypted_held_and_not_generically_downloadable(
    world: World,
) -> None:
    from erev_api.db.tables import release_validation_attempt
    from erev_api.db.transitions import apply
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.enums import FilePurpose

    validation_id = pending_validation(world)
    content = b'{"dataset_format":1,"groups":[]}'
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        stored = store_file(
            uow,
            purpose=FilePurpose.VALIDATION_DATASET,
            stream=io.BytesIO(content),
            original_filename="validation.json",
            media_type="application/json",
        )
        apply(
            uow.session,
            "release_validation_attempt",
            frozen.id,
            to_status=None,
            set_values={"dataset_file_id": stored["id"], "dataset_sha256": stored["sha256"]},
        )
        meta, stream = open_file(uow.session, stored["id"], files=uow.files, keyring=uow.keyring)
        with stream:
            assert stream.read() == content
        assert meta["sha256"] == hashlib.sha256(content).hexdigest()
        assert uow.files.exists(str(meta["storage_key"]) + ".dek")
        with uow.files.open(str(meta["storage_key"])) as raw:
            assert raw.read() != content
        holds = file_evidence.holds_in(uow.session, stored["id"])
        assert any(hold.reference.table == "release_validation_attempt" for hold in holds)
        assert file_evidence.first_refusing(holds) is not None
        assert (
            uow.session.scalar(select(release_validation_attempt.c.dataset_file_id)) == stored["id"]
        )
        uow.commit()
    for suffix in ("", "/content"):
        response = get(world.app, f"/api/v1/files/{stored['id']}{suffix}", world.place.author)
        assert response.status_code == 404, response.text


@pytest.mark.parametrize("changed", [False, True])
def test_replay_difference_artifact_is_atomic_and_reused(
    world: World, monkeypatch: pytest.MonkeyPatch, changed: bool
) -> None:
    import json

    from erev_api.db.tables import computation_verification
    from erev_api.domain.contracts import verification_store
    from erev_api.domain.contracts.replay_execution import execute_running
    from erev_api.domain.contracts.validation_artifacts import differences
    from erev_engine.bundle import Diagnostic

    case = verification_case(world)
    attempt_id = case.decision.release_validation_attempt_id
    assert attempt_id is not None

    def candidate(bundle: Any) -> Any:
        result = engine()(bundle)
        if changed:
            return replace(
                result,
                diagnostics=(
                    *result.diagnostics,
                    Diagnostic("DIFF_PROBE", "INFO", None, None, None, {}),
                ),
            )
        return result

    with world.place.uow() as uow:
        files_before = uow.session.scalar(select(func.count()).select_from(file_object))
        assert files_before is not None
    original = verification_store._Store.set_orchestration

    def interrupt(*args: Any, **kwargs: Any) -> None:
        raise RuntimeError("interrupt artifact acceptance")

    monkeypatch.setattr(verification_store._Store, "set_orchestration", interrupt)
    with pytest.raises(RuntimeError, match="interrupt artifact"):
        execute_running(
            case.context,
            attempt_id=attempt_id,
            group_id=case.group_id,
            trigger="UPGRADE_VALIDATE",
            engine=candidate,
        )
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(file_object)) == files_before
        assert uow.session.scalar(select(func.count()).select_from(computation_verification)) == 0
    monkeypatch.setattr(verification_store._Store, "set_orchestration", original)
    result = execute_running(
        case.context,
        attempt_id=attempt_id,
        group_id=case.group_id,
        trigger="UPGRADE_VALIDATE",
        engine=candidate,
    )
    assert result.outcome.value == ("DIFFERENCES" if changed else "MATCH")
    with world.place.uow() as uow:
        row = uow.session.execute(select(computation_verification)).mappings().one()
        file_id = row["difference_file_id"]
        assert file_id is not None
        meta, stream = open_file(uow.session, file_id, files=uow.files, keyring=uow.keyring)
        with stream:
            content = stream.read()
        payload = json.loads(content)
        assert meta["purpose"] == "VALIDATION_DATASET"
        assert meta["sha256"] == hashlib.sha256(content).hexdigest()
        assert uow.files.exists(str(meta["storage_key"]) + ".dek")
        assert payload["verification_id"] == str(result.row_id)
        assert payload["release_validation_attempt_id"] == str(attempt_id)
        assert payload["job_id"] == str(case.context.job_id)
        assert payload["task_id"] == case.context.task_id
        assert payload["summary"] == row["difference_summary"]
        assert bool(payload["differences"]) is changed
        assert list(differences(uow, dict(row))) == payload["differences"]
        with pytest.raises(Problem):
            differences(uow, {**row, "task_id": row["task_id"] + 1})
        with pytest.raises(Problem):
            differences(uow, {**row, "difference_file_id": None})
        assert (
            file_evidence.first_refusing(file_evidence.holds_in(uow.session, file_id)) is not None
        )
    again = execute_running(
        case.context,
        attempt_id=attempt_id,
        group_id=case.group_id,
        trigger="UPGRADE_VALIDATE",
        engine=candidate,
    )
    assert again == result
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(file_object)) == files_before + 1
    for suffix in ("", "/content"):
        response = get(world.app, f"/api/v1/files/{file_id}{suffix}", world.place.author)
        assert response.status_code == 404, response.text


def test_validation_population_fence_serializes_writers(world: World) -> None:
    from erev_api.domain.contracts.validation_fence import hold
    from erev_api.problems import from_db_error
    from sqlalchemy.exc import DBAPIError
    from support.rows import combination_group_values

    tenant = world.place.tenant_id
    with world.place.uow() as writer:
        writer.session.execute(insert(combination_group).values(**combination_group_values(tenant)))
        # A second writer shares the tenant gate; final acceptance cannot pass either writer.
        with world.place.uow() as second_writer:
            second_writer.session.execute(
                insert(combination_group).values(**combination_group_values(tenant))
            )
            with world.place.uow() as acceptance, pytest.raises(Problem) as refused:
                hold(acceptance)
            assert refused.value.slug == "lock-conflict"
            second_writer.commit()
        writer.commit()
    with world.place.uow() as acceptance:
        hold(acceptance)
        with pytest.raises(DBAPIError) as refused_write, world.place.uow() as writer:
            writer.session.execute(
                insert(combination_group).values(**combination_group_values(tenant))
            )
        problem = from_db_error(refused_write.value)
        assert problem is not None and problem.slug == "lock-conflict"
        # Reentrant calls retain the same transaction reservation.
        hold(acceptance)
        acceptance.commit()
    with world.place.uow() as writer:
        writer.session.execute(insert(combination_group).values(**combination_group_values(tenant)))
        writer.commit()
    with world.place.uow() as acceptance:
        hold(acceptance)
        # Rollback also releases the fence.
    with world.place.uow() as writer:
        writer.session.execute(insert(combination_group).values(**combination_group_values(tenant)))
        writer.commit()


def test_validation_population_fence_is_tenant_scoped(world: World, keyring: KeyRing) -> None:
    from erev_api.domain.contracts.validation_fence import hold
    from support.factories import tenant_factory, tenant_id_of
    from support.rows import combination_group_values

    other = tenant_id_of(tenant_factory(keyring=keyring))
    with world.place.uow() as acceptance:
        hold(acceptance)
        with world.place.uow(system_principal(other)) as writer:
            writer.session.execute(
                insert(combination_group).values(**combination_group_values(other))
            )
            writer.commit()
        with world.place.uow(system_principal(other)) as other_acceptance:
            hold(other_acceptance)


@pytest.mark.parametrize("group_count", [0, 21])
def test_processing_seed_includes_unsampled_groups(world: World, group_count: int) -> None:
    from erev_api.controls.release import current_release
    from erev_api.db.tables import (
        book,
        combination_group_member,
        job,
        release_group_processing,
        release_validation_group,
    )
    from erev_api.db.transitions import apply
    from erev_api.domain.contracts.processing_seed import seed
    from erev_api.domain.contracts.validation_artifacts import retain
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.domain.contracts.validation_fanout import enqueue
    from erev_api.jobs.context import JobRuntime
    from erev_api.jobs.registry import run_job
    from sqlalchemy import text
    from support.rows import (
        combination_group_member_values,
        contract_event_values,
        insert_contract_rows,
    )

    with world.place.uow() as uow:
        groups = set()
        for _ in range(group_count):
            rows = insert_contract_rows(uow.session, world.place.tenant_id, head_stream_version=1)
            event = contract_event_values(
                world.place.tenant_id,
                contract_id=rows.contract_id,
                contracting_entity_id=rows.entity_id,
                stream_version=1,
                effective_date=date(2025, 1, 1),
            )
            uow.session.execute(insert(contract_event).values(**event))
            uow.session.execute(
                insert(combination_group_member).values(
                    **combination_group_member_values(
                        world.place.tenant_id,
                        combination_group_id=rows.group_id,
                        contract_id=rows.contract_id,
                        join_event_id=event["id"],
                    )
                )
            )
            groups.add(rows.group_id)
        uow.commit()
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        fanout = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        assert len(fanout.job_ids) == min(group_count, 20)
        runtime = JobRuntime(
            clock=uow.clock, keyring=uow.keyring, files=uow.files, engine_release=current_release()
        )
        uow.commit()
    for job_id in fanout.job_ids:
        with world.place.uow() as uow:
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == job_id)
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            job_id, world.place.tenant_id, attempt=1, runtime=runtime, delivered_task_id=task_id
        )
    with world.place.uow() as uow:
        retain(uow, validation_id=validation_id, attempt_id=frozen.id)
        with pytest.raises(Problem) as premature:
            seed(uow, validation_id=validation_id)
        assert premature.value.slug == "invalid-transition"
        # This tests the seed boundary after acceptance, not the unfinished acceptance command.
        apply(
            uow.session,
            "release_validation",
            validation_id,
            to_status="PASSED",
            expected_status="RUNNING",
            set_values={"accepted_attempt_id": frozen.id, "enabled_at": uow.now},
        )
        books = set(uow.session.scalars(select(book.c.code).where(book.c.is_enabled)))
        assert seed(uow, validation_id=validation_id) == len(groups) * len(books)
        seeded_groups = set(
            uow.session.scalars(select(release_group_processing.c.combination_group_id))
        )
        selected_groups = set(
            uow.session.scalars(select(release_validation_group.c.combination_group_id))
        )
        assert seeded_groups == groups
        assert len(seeded_groups - selected_groups) == max(0, group_count - 20)
        uow.commit()
    with world.place.uow() as uow:
        assert seed(uow, validation_id=validation_id) == len(groups) * len(books)
        assert not uow.audit_events


@pytest.mark.parametrize("has_previous, bill", [(False, False), (True, False), (True, True)])
def test_computation_persists_upgrade_facts_atomically(
    world: World, monkeypatch: pytest.MonkeyPatch, has_previous: bool, bill: bool
) -> None:
    from erev_api.controls.release import current_release
    from erev_api.db.tables import job, posting_attribution, release_group_processing
    from erev_api.db.transitions import apply
    from erev_api.domain.contracts import replay_job, upgrade_persistence
    from erev_api.domain.contracts.processing_seed import seed
    from erev_api.domain.contracts.validation_artifacts import retain
    from erev_api.domain.contracts.validation_attempts import freeze
    from erev_api.domain.contracts.validation_fanout import enqueue
    from erev_api.jobs.context import JobRuntime
    from erev_api.jobs.registry import run_job
    from sqlalchemy import text

    assert replay_job.replay is not None
    if bill:
        from datetime import timedelta

        from erev_api.enums import RegistryCategory
        from support.rows import publish_registry_version

        with world.place.uow() as uow:
            publish_registry_version(
                uow.session,
                tenant_id=world.place.tenant_id,
                category=RegistryCategory.ACCOUNTING_POLICY,
                values={"billing.posting": "ENGINE"},
                at=uow.now - timedelta(days=365),
            )
            uow.commit()
    booked = k01(world)
    group_id = booked.combination_group["id"]
    if has_previous:
        computed(world.place, group_id)
    validation_id = pending_validation(world)
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        fanout = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        assert len(fanout.job_ids) == 1
        runtime = JobRuntime(
            clock=uow.clock, keyring=uow.keyring, files=uow.files, engine_release=current_release()
        )
        uow.commit()
    child_id = fanout.job_ids[0]
    with world.place.uow() as uow:
        task_id = uow.session.scalar(select(job.c.procrastinate_job_id).where(job.c.id == child_id))
        uow.session.execute(
            text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"), {"id": task_id}
        )
        uow.commit()
    run_job(child_id, world.place.tenant_id, attempt=1, runtime=runtime, delivered_task_id=task_id)
    with world.place.uow() as uow:
        retain(uow, validation_id=validation_id, attempt_id=frozen.id)
        # Exercise post-acceptance persistence; the acceptance command itself remains separate.
        apply(
            uow.session,
            "release_validation",
            validation_id,
            to_status="PASSED",
            expected_status="RUNNING",
            set_values={"accepted_attempt_id": frozen.id, "enabled_at": uow.now},
        )
        assert seed(uow, validation_id=validation_id) == 1
        uow.commit()
    from dataclasses import replace

    from erev_api.domain.reports.builders import RETAINED, ReportParams, SourceCollector
    from erev_api.domain.reports.builders import engine_trueup_register as register
    from erev_api.domain.reports.outputs import json_rows
    from erev_api.domain.reports.outputs.csv import render_csv

    with world.place.uow() as uow:
        collector = SourceCollector()
        register_params = ReportParams(
            register.CODE,
            1,
            {"release_validation_id": str(validation_id)},
            (world.entity_id,),
            uow.now,
            sources=collector,
            output_format="JSON",
        )
        pending_report = register.build(uow, register_params)
        assert pending_report.control_totals["states"] == {"PENDING": 1}
        assert pending_report.rows[0]["amount_functional"] is None
        pending_binding = collector.binding(
            row_keys=[str(r["row_key"]) for r in pending_report.rows],
            strategy=RETAINED,
            report_version=1,
        )

    if bill:
        from erev_api.events.payloads import BillingRecordedV1
        from erev_api.money import MoneyIn

        with world.place.uow() as uow:
            append_events(
                uow,
                contract_id=booked.contract["id"],
                expected_stream_version=2,
                events=[
                    EventIn(
                        event_type=ContractEventType.BILLING_RECORDED,
                        effective_date=date(2026, 2, 1),
                        payload=BillingRecordedV1(
                            invoice_number="UPGRADE-BILL-1",
                            line_external_id="UPGRADE-BILL-1-L1",
                            obligation_key="O1",
                            amount=MoneyIn(amount="120000.00", currency="USD"),
                            issue_date=date(2026, 2, 1),
                        ),
                    )
                ],
                origin="UI",
            )
            uow.commit()
    counted = (
        contract_computation,
        computation_evidence,
        file_object,
        subledger_posting,
        posting_attribution,
        release_group_processing,
    )
    with world.place.uow() as uow:
        before = [uow.session.scalar(select(func.count()).select_from(table)) for table in counted]
    if bill:
        with world.place.uow() as uow:
            probe = bundles.build(uow.session, group_id, uow.now, (), ComputationTrigger.COMMAND)
            assert any(event.event_type == "BILLING_RECORDED" for event in probe.events)
            assert engine()(probe).books[0].posting_intents, (
                "Billing must produce actual posting lines"
            )
        read_file = open_file

        def corrupted_input(*args: Any, **kwargs: Any) -> Any:
            metadata, stream = read_file(*args, **kwargs)
            stream.close()
            return metadata, io.BytesIO(b"corrupted previous input")

        with monkeypatch.context() as patch:
            patch.setattr(upgrade_persistence, "open_file", corrupted_input)
            with pytest.raises(Problem) as invalid_input:
                computed(world.place, group_id)
            assert invalid_input.value.slug == "validation-failed"
        with world.place.uow() as uow:
            assert [
                uow.session.scalar(select(func.count()).select_from(table)) for table in counted
            ] == before
    original = upgrade_persistence.record

    def interrupted(*args: Any, **kwargs: Any) -> Any:
        original(*args, **kwargs)
        raise RuntimeError("upgrade persistence interrupted")

    with monkeypatch.context() as patch:
        patch.setattr(upgrade_persistence, "record", interrupted)
        with pytest.raises(RuntimeError, match="upgrade persistence interrupted"):
            computed(world.place, group_id)
    with world.place.uow() as uow:
        assert [
            uow.session.scalar(select(func.count()).select_from(table)) for table in counted
        ] == before
    bundle, output, stored = computed(world.place, group_id)
    with world.place.uow() as uow:
        events = (
            uow.session.execute(
                select(release_group_processing).order_by(release_group_processing.c.seq)
            )
            .mappings()
            .all()
        )
        assert len(events) == 2
        assert events[-1]["contract_computation_id"] == stored["id"]
        assert events[-1]["upgrade_state"] == (
            "SETTLED_NO_POSTING" if has_previous and not bill else "POSTED"
        )
        attributed = uow.session.execute(select(posting_attribution)).mappings().all()
        assert {str(row["subledger_posting_id"]) for row in attributed} == set(
            stored["subledger_posting_ids"].values()
        )
        assert all(
            row["cause"] == ("UPGRADE_AND_OTHER_INPUTS" if bill else "UPGRADE_CONTEXT")
            for row in attributed
        )
        assert all(row["release_validation_attempt_id"] == frozen.id for row in attributed)
        assert all(row["origin_periods"]["periods"] for row in attributed)
        final_report = register.build(uow, replace(register_params, known_at=uow.now, sources=None))
        assert final_report.control_totals["states"] == {str(events[-1]["upgrade_state"]): 1}
        assert {r["posting_id"] for r in final_report.rows if r["posting_id"]} == {
            str(r["subledger_posting_id"]) for r in attributed
        }
        assert final_report.rows[0]["group_terminal"] is True
        assert render_csv(final_report)
        bound_report = register.build(
            uow,
            replace(
                register_params,
                sources=None,
                binding=pending_binding,
            ),
        )
        assert json_rows(bound_report) == json_rows(pending_report)
        assert render_csv(bound_report) == render_csv(pending_report)
        replayed = computation.persist(uow, bundle, output)
        assert replayed["replayed"] and replayed["id"] == stored["id"]
        assert uow.session.scalar(select(func.count()).select_from(release_group_processing)) == 2
        assert uow.session.scalar(select(func.count()).select_from(posting_attribution)) == len(
            attributed
        )
        uow.commit()

    # The published report uses the normal task/output path, including current grants.
    if bill:
        from erev_api.db.tables import report_run
        from erev_api.domain.reports import framework
        from erev_api.schemas.reports import ReportRunCreateIn

        with world.place.uow() as uow:
            report_id, _ = framework.create_run(
                uow,
                ReportRunCreateIn(
                    report_code=register.CODE,
                    output_format="CSV",
                    parameters={"release_validation_id": str(validation_id)},
                ),
            )
            report_job = uow.session.scalar(
                select(report_run.c.job_id).where(report_run.c.id == report_id)
            )
            uow.commit()
        outputs = []
        for iteration in range(2):
            assert report_job is not None
            with world.place.uow() as uow:
                task_id = uow.session.scalar(
                    select(job.c.procrastinate_job_id).where(job.c.id == report_job)
                )
                uow.session.execute(
                    text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                    {"id": task_id},
                )
                uow.commit()
            run_job(
                report_job,
                world.place.tenant_id,
                attempt=1,
                runtime=runtime,
                delivered_task_id=task_id,
            )
            with world.place.uow() as uow:
                saved = (
                    uow.session.execute(select(report_run).where(report_run.c.id == report_id))
                    .mappings()
                    .one()
                )
                assert saved["status"] == "SUCCEEDED", dict(saved)
                outputs.append((saved["output_sha256"], saved["control_totals"]))
                if iteration == 0:
                    report_id, _ = framework.rerun(uow, report_id)
                    report_job = uow.session.scalar(
                        select(report_run.c.job_id).where(report_run.c.id == report_id)
                    )
                    uow.commit()
        assert outputs[0] == outputs[1]


@pytest.mark.parametrize(
    "level,different,action",
    [
        ("MINOR", False, "accept"),
        ("MINOR", True, "accept"),
        ("MAJOR", False, "accept"),
        ("MAJOR", True, "approve"),
        ("MAJOR", True, "reject"),
        ("MAJOR", True, "drift"),
        ("MAJOR", True, "supersede"),
    ],
)
def test_release_finalization_lifecycle(
    world: World, monkeypatch: pytest.MonkeyPatch, level: str, different: bool, action: str
) -> None:
    from erev_api.controls.release import current_release
    from erev_api.db.tables import (
        approval_request,
        job,
        release_group_processing,
        release_validation,
        release_validation_attempt,
        report_run,
    )
    from erev_api.domain.contracts import replay_job, validation_acceptance
    from erev_api.domain.contracts.replay_execution import execute_running
    from erev_api.domain.contracts.validation_artifacts import retain
    from erev_api.domain.contracts.validation_attempts import freeze, supersede
    from erev_api.domain.contracts.validation_fanout import enqueue
    from erev_api.domain.contracts.validation_reports import attach
    from erev_api.domain.reports import framework
    from erev_api.jobs.context import JobRuntime
    from erev_api.jobs.registry import run_job
    from erev_api.schemas.reports import ReportRunCreateIn
    from erev_engine.bundle import Diagnostic
    from sqlalchemy import text
    from support.reference import approve, reject

    booked = k01(world)
    computed(world.place, booked.combination_group["id"])
    validation_id = pending_validation(world, level=level)
    with world.place.uow() as uow:
        frozen = freeze(uow, validation_id=validation_id, attempt_no=1)
        fanout = enqueue(uow, validation_id=validation_id, attempt_id=frozen.id)
        runtime = JobRuntime(
            clock=uow.clock, keyring=uow.keyring, files=uow.files, engine_release=current_release()
        )
        uow.commit()

    def candidate(bundle: Any) -> Any:
        output = engine()(bundle)
        return (
            replace(
                output,
                diagnostics=(
                    *output.diagnostics,
                    Diagnostic("APPROVAL_PROBE", "INFO", None, None, None, {}),
                ),
            )
            if different
            else output
        )

    monkeypatch.setattr(
        replay_job,
        "execute_running",
        lambda context, **kwargs: execute_running(context, **kwargs, engine=candidate),
    )

    def run(job_id: UUID) -> None:
        with world.place.uow() as uow:
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == job_id)
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            job_id, world.place.tenant_id, attempt=1, runtime=runtime, delivered_task_id=task_id
        )

    for job_id in fanout.job_ids:
        run(job_id)
    with world.place.uow() as uow:
        retain(uow, validation_id=validation_id, attempt_id=frozen.id)
        report_id, _ = framework.create_run(
            uow,
            ReportRunCreateIn(
                report_code="upgrade_validation",
                output_format="JSON",
                parameters={"release_validation_attempt_id": str(frozen.id)},
            ),
        )
        report_job = uow.session.scalar(
            select(report_run.c.job_id).where(report_run.c.id == report_id)
        )
        uow.commit()
    assert report_job is not None
    run(report_job)
    with world.place.uow() as uow:
        attach(uow, validation_id=validation_id, attempt_id=frozen.id, report_id=report_id)
        uow.commit()
    expected = "PASSED" if not different else "FAILED" if level == "MINOR" else "PENDING_APPROVAL"
    # An uncommitted acceptance leaves both enablement and the processing queue unchanged.
    with world.place.uow() as uow:
        assert (
            validation_acceptance.finalize(
                uow, validation_id=validation_id, attempt_id=frozen.id, report_id=report_id
            )
            == expected
        )
    with world.place.uow() as uow:
        assert uow.session.scalar(select(release_validation.c.status)) == "RUNNING"
        assert uow.session.scalar(select(func.count()).select_from(release_group_processing)) == 0
        assert (
            validation_acceptance.finalize(
                uow, validation_id=validation_id, attempt_id=frozen.id, report_id=report_id
            )
            == expected
        )
        uow.commit()
    with world.place.uow() as uow:
        assert (
            validation_acceptance.finalize(
                uow, validation_id=validation_id, attempt_id=frozen.id, report_id=report_id
            )
            == expected
        )
        assert not uow.audit_events
        request_id = uow.session.scalar(select(release_validation_attempt.c.approval_request_id))
        uow.commit()
    if expected == "PENDING_APPROVAL":
        assert request_id is not None
        with world.place.uow() as uow:
            request = (
                uow.session.execute(
                    select(approval_request).where(approval_request.c.id == request_id)
                )
                .mappings()
                .one()
            )
            assert request["preparer_kind"] == "SYSTEM" and request["status"] == "PENDING"
            assert uow.session.scalar(select(release_validation.c.enabled_at)) is None
        if action == "drift":
            flag_change(world, booked.contract["id"], 2)
        if action == "supersede":
            with world.place.uow() as uow:
                assert supersede(uow, validation_id=validation_id, attempt_id=frozen.id)
                replacement = freeze(uow, validation_id=validation_id, attempt_no=2)
                assert replacement.id != frozen.id
                uow.commit()
            expected = "INCOMPLETE"
        elif action == "reject":
            response = reject(world.app, str(request_id), world.marcus, "Review declined")
            assert response.status_code == 200, response.text
            expected = "REJECTED"
        else:
            response = approve(world.app, str(request_id), world.marcus)
            assert response.status_code == (409 if action == "drift" else 200), response.text
            expected = "INCOMPLETE" if action == "drift" else "APPROVED"
    with world.place.uow() as uow:
        saved = uow.session.execute(select(release_validation)).mappings().one()
        assert saved["status"] == expected
        enabled = expected in {"PASSED", "APPROVED"}
        assert (saved["enabled_at"] is not None) == enabled
        assert (saved["processing_seed"] is not None) == enabled
        assert uow.session.scalar(
            select(func.count()).select_from(release_group_processing)
        ) == int(enabled)
        if action in {"drift", "supersede"}:
            request = (
                uow.session.execute(
                    select(approval_request).where(approval_request.c.id == request_id)
                )
                .mappings()
                .one()
            )
            assert request["status"] == "VOIDED" and request["void_reason"] == "STALE_SUBJECT"


def previous_release(
    world: World, version: str, *, level: str = "PATCH", enabled: bool = True
) -> UUID:
    from datetime import timedelta
    from uuid import uuid4

    from erev_api.controls.release import current_release
    from erev_api.db.session import release_session
    from erev_api.db.tables import engine_release, release_validation

    current = current_release()
    assert current is not None
    release_id = uuid4()
    with release_session(request_id="release-history-probe") as session:
        session.execute(
            insert(engine_release).values(
                id=release_id,
                engine_version=version,
                build_sha=release_id.hex,
                schema_revision=current.schema_revision,
                validation_level=level,
            )
        )
    if enabled:
        with world.place.uow() as uow:
            uow.session.execute(
                insert(release_validation).values(
                    tenant_id=world.place.tenant_id,
                    id=uuid4(),
                    engine_release_id=release_id,
                    effective_level="BOOTSTRAP",
                    status="PASSED",
                    enabled_at=uow.now - timedelta(days=1),
                    created_at=uow.now,
                    created_by=uow.principal.id,
                    created_by_kind=uow.principal.kind.value,
                    updated_at=uow.now,
                    updated_by=uow.principal.id,
                    updated_by_kind=uow.principal.kind.value,
                )
            )
            uow.commit()
    return release_id


@pytest.mark.parametrize(
    "scenario", ["bootstrap", "historical", "minor", "skipped_major", "same_version", "downgrade"]
)
def test_tenant_validation_creation(world: World, scenario: str) -> None:
    from erev_api.controls.release import current_release
    from erev_api.db.tables import release_validation
    from erev_api.domain.contracts.validation_creation import ensure

    current = current_release()
    assert current is not None
    if scenario == "historical":
        booked = k01(world)
        computed(world.place, booked.combination_group["id"])
    if scenario in {"minor", "skipped_major"}:
        previous_release(world, "0.2.0")
    if scenario == "skipped_major":
        previous_release(world, "0.2.5", level="MAJOR", enabled=False)
        previous_release(world, "0.2.5", level="PATCH", enabled=False)
    if scenario == "same_version":
        previous_release(world, current.engine_version)
    if scenario == "downgrade":
        previous_release(world, "9.0.0")
        with world.place.uow() as uow, pytest.raises(Problem) as refused:
            ensure(uow)
        assert refused.value.slug == "release-mismatch"
        return
    with world.place.uow() as uow:
        first = ensure(uow)
        assert (
            first["effective_level"]
            == {
                "bootstrap": "BOOTSTRAP",
                "historical": "MAJOR",
                "minor": "MINOR",
                "skipped_major": "MAJOR",
                "same_version": "PATCH",
            }[scenario]
        )
        assert (first["enabled_at"] is not None) == (scenario == "bootstrap")
    with world.place.uow() as uow:
        assert (
            uow.session.scalar(
                select(release_validation.c.id).where(
                    release_validation.c.engine_release_id == current.id
                )
            )
            is None
        )
        saved = ensure(uow)
        uow.commit()
    with world.place.uow() as uow:
        assert ensure(uow) == saved
        assert not uow.audit_events


def test_release_orchestration_api_resumes_to_enabled(world: World) -> None:
    from erev_api.controls.release import current_release
    from erev_api.db.tables import job, release_validation_attempt, report_run
    from erev_api.domain.contracts import replay_job
    from erev_api.jobs.context import JobRuntime
    from erev_api.jobs.registry import run_job
    from sqlalchemy import text
    from support.reference import post

    assert replay_job.replay is not None
    # A tenant enabled on the preceding engine before its first computation has a NO_SOURCE
    # group to validate. No historical engine output is fabricated for this upgrade.
    previous_release(world, "0.2.0")
    k01(world)
    path = "/api/v1/release-validations/current/advance"

    def advance() -> dict[str, Any]:
        response = post(world.app, path, world.place.author, {})
        assert response.status_code == 200, response.text
        return dict(response.json())

    first = advance()
    assert first["phase"] == "WAITING_REPLAY" and len(first["job_ids"]) == 1
    assert advance() == first
    with world.place.uow() as uow:
        runtime = JobRuntime(
            clock=uow.clock, keyring=uow.keyring, files=uow.files, engine_release=current_release()
        )

    def run(job_id: str) -> None:
        with world.place.uow() as uow:
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == UUID(job_id))
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            UUID(job_id),
            world.place.tenant_id,
            attempt=1,
            runtime=runtime,
            delivered_task_id=task_id,
        )

    run(first["job_ids"][0])
    reporting = advance()
    assert reporting["phase"] == "WAITING_REPORT" and len(reporting["job_ids"]) == 1
    assert advance() == reporting
    run(reporting["job_ids"][0])
    attached = advance()
    assert attached["phase"] == "READY_FOR_ACCEPTANCE"
    finished = advance()
    assert finished["phase"] == "ENABLED" and finished["status"] == "PASSED"
    assert advance()["phase"] == "ENABLED"
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(release_validation_attempt)) == 1
        assert uow.session.scalar(select(func.count()).select_from(report_run)) == 1


def test_parallel_actual_reads_recorded_journals_and_balances(world: World) -> None:
    from dataclasses import replace

    from erev_api.domain.journals import views
    from erev_api.domain.reports import parallel_actual
    from erev_api.domain.reports.parallel_comparison import Balance, Period, Snapshot

    booked = k01(world)
    computed(world.place, booked.combination_group["id"])
    period = Period("AVM-US", "ASC606", "USD", "FY2026-P01", date(2026, 1, 1), date(2026, 1, 31))
    expected = Snapshot(
        "population-evidence",
        "a" * 64,
        frozenset({period}),
        (),
        (Balance(period, "SF-ORD-10001", Decimal(0), ("legacy-row:1",)),),
    )
    with world.place.uow() as uow:
        cutoff = uow.session.scalar(select(func.clock_timestamp()))
        assert cutoff is not None
        source = parallel_actual.read(
            uow, expected=expected, entity_ids=(world.entity_id,), known_at=cutoff
        )
        assert source.snapshot.periods == frozenset({period})
        assert source.snapshot.balances
        assert source.snapshot.journals
        assert all(row.source_rows for row in source.snapshot.journals)
        assert all(
            "subledger-line:" in ref for row in source.snapshot.journals for ref in row.source_rows
        )
        normal = views.legacy_view(
            uow,
            ["AVM-US"],
            "ASC606",
            period.start,
            period.end,
            "GROSS",
            known_at=cutoff,
            historical=True,
        )
        assert {row.account: (row.debit, row.credit) for row in source.snapshot.journals} == {
            row.account: (row.debit, row.credit) for row in normal.by_account
        }
        for row in source.snapshot.balances:
            assert row.liability is not None and row.asset_and_unbilled is not None
            assert row.net_position == row.liability - row.asset_and_unbilled
        repeated = parallel_actual.read(
            uow, expected=expected, entity_ids=(world.entity_id,), known_at=cutoff
        )
        assert repeated == source
        incomplete = replace(
            expected,
            balances=(
                *expected.balances,
                Balance(period, "not-yet-imported", Decimal(0), ("legacy-row:2",)),
            ),
        )
        missing = parallel_actual.read(
            uow, expected=incomplete, entity_ids=(world.entity_id,), known_at=cutoff
        )
        assert not missing.snapshot.periods and not missing.snapshot.journals
        assert missing.evidence["unavailable_periods"][0]["missing_contracts"] == [
            "not-yet-imported"
        ]
    denied = replace(world.place.principal, permissions=frozenset())
    with world.place.uow(principal=denied) as uow, pytest.raises(Problem) as refused:
        parallel_actual.read(uow, expected=expected, entity_ids=(world.entity_id,), known_at=cutoff)
    assert refused.value.slug == "forbidden"


def test_disclosure_child_plan_queues_bound_month_and_quarter_reports(world: World) -> None:
    from erev_api.db.tables import report_run
    from erev_api.domain.reports import disclosure_plan, tie_outs

    with world.place.uow() as uow:
        entities = tie_outs.entities(uow.session, (world.entity_id,))
        calendars = tie_outs.calendars(uow.session, entities)
        calendar = calendars[world.entity_id]
        selected = calendar.periods[1]
        plan = disclosure_plan.requests(
            entity_code=entities[0].code,
            book="ASC606",
            calendar=calendar,
            period_key=selected.key,
            known_at=uow.now,
            quarter=True,
        )
        children = disclosure_plan.queue(uow, plan)
        assert len(children) == 16
        assert len({child.run_id for child in children}) == 16
        stored = (
            uow.session.execute(
                select(report_run).where(report_run.c.id.in_([child.run_id for child in children]))
            )
            .mappings()
            .all()
        )
        assert len(stored) == 16
        assert all(row["entity_ids"] == [world.entity_id] for row in stored)
        assert all(row["known_at"] == uow.now for row in stored)
        assert all(row["output_format"] == "JSON" for row in stored)
        assert all(child.parameters_sha256 for child in children)
        from dataclasses import replace

        invalid = replace(
            plan[1],
            body=plan[1].body.model_copy(
                update={"parameters": {**plan[1].body.parameters, "unsupported_selector": True}}
            ),
        )
        count_before = uow.session.scalar(select(func.count()).select_from(report_run))
        with pytest.raises(Problem):
            disclosure_plan.queue(uow, (plan[0], invalid))
        assert uow.session.scalar(select(func.count()).select_from(report_run)) == count_before
        # No commit: queued child jobs belong to the eventual parent's atomic transaction.


def test_disclosure_children_attach_once_and_survive_retry(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    from dataclasses import replace

    from erev_api.db.tables import disclosure_snapshot, report_run
    from erev_api.domain.reports import disclosure_children, disclosure_readiness, tie_outs
    from support.principals import sign_in
    from support.principals import workspace as signed_workspace
    from support.reference import post

    with world.place.uow() as uow:
        observation_time = max(uow.now, uow.session.scalar(select(func.clock_timestamp())))
    world.place.clock.set(observation_time)
    with world.place.uow() as uow:
        entities = tie_outs.entities(uow.session, (world.entity_id,))
        calendar = tie_outs.calendars(uow.session, entities)[world.entity_id]
        period = calendar.periods[1]
    actor = signed_workspace(
        world.app,
        world.place.author.member,
        sign_in(world.app, world.place.author.member.email),
        world.place.author.secret,
    )
    for invalid in (
        {"entity_codes": [entities[0].code], "period_key": period.key},
        {"entity_codes": [entities[0].code], "book": "ASC606"},
        {"entity_codes": [entities[0].code], "book": "LEGACY", "period_key": period.key},
    ):
        refused = post(
            world.app,
            "/api/v1/report-runs",
            actor,
            {
                "report_code": "disclosure_pack",
                "output_format": "JSON",
                "parameters": invalid,
            },
        )
        assert refused.status_code == 422, refused.text
    response = post(
        world.app,
        "/api/v1/report-runs",
        actor,
        {
            "report_code": "disclosure_pack",
            "output_format": "JSON",
            "parameters": {
                "entity_codes": [entities[0].code],
                "book": "ASC606",
                "period_key": period.key,
                "quarter": True,
            },
        },
    )
    assert response.status_code == 202, response.text
    parent_id = UUID(response.headers["X-Erev-Report-Run-Id"])
    with world.place.uow() as uow:
        parent = dict(
            uow.session.execute(select(report_run).where(report_run.c.id == parent_id))
            .mappings()
            .one()
        )
        parent_params = {"report_run_id": str(parent_id)}
        first = disclosure_children.ensure(uow, parent["id"])
        count = uow.session.scalar(select(func.count()).select_from(report_run))
        assert len(first) == 16
        assert not disclosure_readiness.ready(uow.session, parent["job_id"], parent_params)
        with pytest.raises(Problem):
            disclosure_readiness.ready(uow.session, first[0].job_id, parent_params)
        assert disclosure_children.ensure(uow, parent["id"]) == first
        assert uow.session.scalar(select(func.count()).select_from(report_run)) == count
        stored = uow.session.scalar(
            select(report_run.c.child_report_run_ids).where(report_run.c.id == parent["id"])
        )
        assert stored == [child.run_id for child in first]
        uow.commit()
    with world.place.uow() as uow:
        assert disclosure_children.ensure(uow, parent["id"]) == first
        assert uow.session.scalar(select(func.count()).select_from(report_run)) == count
        # Fresh authorization still applies when all child tasks already exist.
        previous = uow.ctx
        uow.ctx = replace(
            previous,
            principal=replace(previous.principal, permissions=frozenset(), permission_scopes={}),
        )
        try:
            with pytest.raises(Problem) as refused:
                disclosure_children.ensure(uow, parent["id"])
            assert refused.value.slug == "forbidden"
        finally:
            uow.ctx = previous
        # Deliberately corrupt the attachment order: do not silently create replacement runs.
        uow.session.execute(
            update(report_run)
            .where(report_run.c.id == parent["id"])
            .values(child_report_run_ids=list(reversed(stored)))
        )
        with pytest.raises(Problem):
            disclosure_children.ensure(uow, parent["id"])
        with pytest.raises(Problem):
            disclosure_readiness.ready(uow.session, parent["job_id"], parent_params)
        assert uow.session.scalar(select(func.count()).select_from(report_run)) == count

    from erev_api.db.tables import job
    from erev_api.domain.reports import disclosure_results
    from erev_api.jobs.context import JobRuntime
    from erev_api.jobs.registry import run_job
    from sqlalchemy import text

    runtime = JobRuntime(
        clock=world.place.clock, keyring=world.place.keyring, files=world.place.files
    )
    with world.place.uow() as uow:
        pending = disclosure_results.collect(uow, parent["id"])
        assert set(pending.pending) == {child.run_id for child in first}
        assert pending.children == ()
    with world.place.uow() as uow:
        parent_task_id = uow.session.scalar(
            select(job.c.procrastinate_job_id).where(job.c.id == parent["job_id"])
        )
        uow.session.execute(
            text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
            {"id": parent_task_id},
        )
        uow.commit()
    run_job(
        parent["job_id"],
        world.place.tenant_id,
        attempt=1,
        runtime=runtime,
        delivered_task_id=parent_task_id,
    )
    with world.place.uow() as uow:
        waiting_job = (
            uow.session.execute(select(job).where(job.c.id == parent["job_id"])).mappings().one()
        )
        assert waiting_job["state"] == "QUEUED"
        assert waiting_job["started_at"] is None
        delivery = uow.session.execute(
            text("SELECT args FROM procrastinate_jobs WHERE id = :id"),
            {"id": waiting_job["procrastinate_job_id"]},
        ).scalar_one()
        assert delivery["attempt"] == 1
        assert waiting_job["procrastinate_job_id"] != parent_task_id
        assert (
            uow.session.scalar(select(report_run.c.status).where(report_run.c.id == parent["id"]))
            == "QUEUED"
        )
    for child in first:
        with world.place.uow() as uow:
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == child.job_id)
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            child.job_id,
            world.place.tenant_id,
            attempt=1,
            runtime=runtime,
            delivered_task_id=task_id,
        )
    with world.place.uow() as uow:
        assert disclosure_readiness.ready(uow.session, parent["job_id"], parent_params)
        completed = disclosure_results.collect(uow, parent["id"])
        assert completed.pending == ()
        assert len(completed.children) == 16
        assert [child.source for child in completed.children] == list(first)
        assert all(child.output_sha256 for child in completed.children)

    from erev_api.domain.reports import disclosure_expedients

    with world.place.uow() as uow:
        policy_cutoff = max(uow.now, uow.session.scalar(select(func.clock_timestamp())))
        policy_facts = disclosure_expedients.capture(
            uow,
            entity_id=world.entity_id,
            book="ASC606",
            known_at=policy_cutoff,
        )
        policy_data = disclosure_expedients.dataset(policy_facts)
        assert len(policy_data.rows) == 10
        cost = next(
            row for row in policy_data.rows if row["registry_key"] == "costs.obtain_expedient"
        )
        assert cost["value"] == "Applied"
        assert "340-40-25-4" in cost["disclosure"]
        assert policy_facts["resolved"]["costs.obtain_expedient"]["level"] in {
            "DEFAULT",
            "T",
            "E",
            "B",
        }
        ifrs = disclosure_expedients.capture(
            uow,
            entity_id=world.entity_id,
            book="IFRS15",
            known_at=policy_cutoff,
        )
        assert ifrs["resolved"]["rpo.exemption_royalty_vc"]["value"] == "DO_NOT_APPLY"
        assert all(
            "ASC " not in row["disclosure"] for row in disclosure_expedients.dataset(ifrs).rows
        )

    from erev_api.domain.reports import disclosure_assembly

    with world.place.uow() as uow:
        pack = disclosure_assembly.prepare(uow, parent["id"])
        assert isinstance(pack, disclosure_assembly.Pack)
        assert [section.number for section in pack.sections] == list(range(1, 8))
        assert all(section.tables for section in pack.sections)
        assert pack.policy_facts["known_at"] == parent["known_at"].isoformat()
        assert {s.number for s in pack.sections if s.interim_required} == {1, 2, 4, 5}
        assert len(pack.sections[-1].tables[0].data.rows) == 10
        assert len(pack.children) == 16

        import io

        import openpyxl
        from erev_api.domain.reports import disclosure_workbook, framework

        parent_row = dict(
            uow.session.execute(select(report_run).where(report_run.c.id == parent["id"]))
            .mappings()
            .one()
        )
        definition = framework.definition_row(uow.session, "disclosure_pack")
        assert definition is not None
        pack_stamp = framework._run_stamp(
            uow.session,
            parent_row,
            definition,
            output_sha256="c" * 64,
            style="PARENTHESES",
        )
        stamps = {}
        for child in pack.children:
            child_definition = framework.definition_row(uow.session, child.run["report_code"])
            assert child_definition is not None
            stamps[child.source.run_id] = framework._run_stamp(
                uow.session,
                child.run,
                child_definition,
                output_sha256=child.output_sha256,
                style="PARENTHESES",
            )
        workbook_bytes = disclosure_workbook.render(pack, pack_stamp, stamps)
        workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes))
        assert len(workbook.worksheets) == 8
        assert workbook.sheetnames[0] == "Contents"
        first_chapter = [cell.value for row in workbook.worksheets[1] for cell in row]
        assert "Period — Total" in first_chapter
        assert "Quarter to date — Total" in first_chapter
        for section, worksheet in zip(pack.sections, workbook.worksheets[1:], strict=True):
            values = [cell.value for row in worksheet for cell in row]
            assert section.title in values
            for table in section.tables:
                if table.source_run_id is not None:
                    assert stamps[table.source_run_id].report_run_no in values
                    assert table.source_sha256 in values
        bad_stamps = dict(stamps)
        first_id = pack.sections[0].tables[0].source_run_id
        assert first_id is not None
        bad_stamps[first_id] = replace(bad_stamps[first_id], output_sha256="d" * 64)
        with pytest.raises(Problem):
            disclosure_workbook.render(pack, pack_stamp, bad_stamps)

    from erev_api.domain.reports import disclosure_snapshots

    original_writer = disclosure_snapshots.write_draft

    def write_twice(uow, identity, binding):
        ids = original_writer(uow, identity, binding)
        assert original_writer(uow, identity, binding) == ids
        return ids

    monkeypatch.setattr(disclosure_snapshots, "write_draft", write_twice)

    # Publish the actual parent through the worker; the public catalogue gate remains closed.
    with world.place.uow() as uow:
        parent_task_id = uow.session.scalar(
            select(job.c.procrastinate_job_id).where(job.c.id == parent["job_id"])
        )
        uow.session.execute(
            text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
            {"id": parent_task_id},
        )
        uow.commit()
    run_job(
        parent["job_id"],
        world.place.tenant_id,
        attempt=1,
        runtime=runtime,
        delivered_task_id=parent_task_id,
    )
    with world.place.uow() as uow:
        published = (
            uow.session.execute(select(report_run).where(report_run.c.id == parent["id"]))
            .mappings()
            .one()
        )
        assert published["status"] == "SUCCEEDED", published["problem"]
        snapshots = list(
            uow.session.execute(
                select(disclosure_snapshot).where(
                    disclosure_snapshot.c.report_run_id == parent["id"]
                )
            ).mappings()
        )
        assert len(snapshots) == 7
        assert {row["id"] for row in snapshots} == set(published["disclosure_snapshot_ids"])
        assert all(
            row["period_lock_id"] is None and row["period_id"] == period.id for row in snapshots
        )
        assert sum(len(row["data"]["children"]) for row in snapshots) == 14
        assert (
            uow.session.scalar(select(job.c.state).where(job.c.id == parent["job_id"]))
            == "SUCCEEDED"
        )
        metadata, stream = open_file(
            uow.session, published["output_file_id"], files=uow.files, keyring=uow.keyring
        )
        with stream:
            output = stream.read()
        import json

        from erev_engine.canonical import sha256_hex

        assert all(sha256_hex(row["data"]) == row["data_sha256"] for row in snapshots)
        snapshot_hashes = {row["disclosure_kind"]: row["data_sha256"] for row in snapshots}
        document = json.loads(output)
        assert len(document["sections"]) == 7
        comparisons = [
            table
            for section in document["sections"]
            for table in section["tables"]
            if table["basis"] == "PERIOD_AND_QUARTER"
        ]
        assert len(comparisons) == 6
        assert all(sha256_hex(table["dataset"]) == table["source_sha256"] for table in comparisons)
        assert any(
            column["key"] == "quarter:total" for column in comparisons[0]["dataset"]["columns"]
        )
        assert document["row_count"] == published["row_count"]
        assert document["control_totals"]["children"] == 16
        assert (
            hashlib.sha256(output).hexdigest() == published["output_sha256"] == metadata["sha256"]
        )
        evidence = published["source_binding"]["evidence"]
        assert sha256_hex(evidence["disclosure_inputs"]) == evidence["disclosure_inputs_sha256"]
        assert len(evidence["disclosure_inputs"]["children"]) == 16
        assert (
            evidence["disclosure_inputs"]["policies"]["known_at"] == parent["known_at"].isoformat()
        )

    snapshots_page = get(
        world.app,
        "/api/v1/disclosure-snapshots",
        actor,
        {
            "report_run_id": str(parent["id"]),
            "limit": 3,
            "count": True,
        },
    )
    assert snapshots_page.status_code == 200, snapshots_page.text
    assert snapshots_page.headers["X-Erev-Total-Count"] == "7"
    assert snapshots_page.headers["Cache-Control"] == "no-store"
    first_items = snapshots_page.json()["items"]
    assert len(first_items) == 3
    assert all(item["source"] == "DRAFT" and "data" not in item for item in first_items)
    next_snapshots = get(
        world.app,
        "/api/v1/disclosure-snapshots",
        actor,
        {
            "report_run_id": str(parent["id"]),
            "limit": 3,
            "cursor": snapshots_page.json()["next_cursor"],
        },
    )
    assert next_snapshots.status_code == 200, next_snapshots.text
    assert {item["id"] for item in first_items}.isdisjoint(
        item["id"] for item in next_snapshots.json()["items"]
    )
    snapshot_detail = get(world.app, f"/api/v1/disclosure-snapshots/{first_items[0]['id']}", actor)
    assert snapshot_detail.status_code == 200, snapshot_detail.text
    assert sha256_hex(snapshot_detail.json()["data"]) == snapshot_detail.json()["data_sha256"]
    expedients = get(
        world.app,
        "/api/v1/disclosure-snapshots",
        actor,
        {
            "report_run_id": str(parent["id"]),
            "disclosure_kind": "EXPEDIENTS_USED",
            "entity": entities[0].code,
            "book": "ASC606",
            "period": period.key,
        },
    )
    assert expedients.status_code == 200 and len(expedients.json()["items"]) == 1
    from erev_api.domain.reports import disclosure_queries

    with world.place.uow() as uow:
        excluded = replace(uow.principal, permission_scopes={"report.run": frozenset()})
        assert not uow.session.execute(disclosure_queries.statement(excluded)).all()
        with pytest.raises(Problem) as hidden:
            disclosure_queries.get(uow.session, excluded, UUID(first_items[0]["id"]))
        assert hidden.value.slug == "not-found"

    detail = get(world.app, f"/api/v1/report-runs/{parent['id']}", actor)
    assert detail.status_code == 200, detail.text
    download = get(world.app, f"/api/v1/report-runs/{parent['id']}/output", actor)
    assert download.status_code == 200, download.text
    assert download.json() == document

    index = get(world.app, f"/api/v1/report-runs/{parent['id']}/data", actor, {"limit": 1})
    assert index.status_code == 200, index.text
    assert len(index.json()["sections"]) == 7
    policy_table = index.json()["sections"][-1]["tables"][0]
    assert policy_table["key"] == "expedients" and policy_table["row_count"] == 10
    table_page = get(
        world.app,
        f"/api/v1/report-runs/{parent['id']}/data",
        actor,
        {"section_table": "expedients", "limit": 1},
    )
    assert table_page.status_code == 200, table_page.text
    assert len(table_page.json()["items"]) == 1
    assert (
        table_page.json()["columns"] == document["sections"][-1]["tables"][0]["dataset"]["columns"]
    )
    cursor = table_page.json()["next_cursor"]
    assert cursor is not None
    next_page = get(
        world.app,
        f"/api/v1/report-runs/{parent['id']}/data",
        actor,
        {"section_table": "expedients", "limit": 1, "cursor": cursor},
    )
    assert next_page.status_code == 200
    assert next_page.json()["items"] != table_page.json()["items"]
    wrong_table = get(
        world.app,
        f"/api/v1/report-runs/{parent['id']}/data",
        actor,
        {"section_table": "disaggregation:PERIOD", "cursor": cursor},
    )
    assert wrong_table.status_code == 422
    unknown = get(
        world.app, f"/api/v1/report-runs/{parent['id']}/data", actor, {"section_table": "missing"}
    )
    assert unknown.status_code == 422

    from erev_api.domain.reports import disclosure_inputs

    with world.place.uow() as uow:
        original = dict(published)
        retained_pack, retained_stamps, retained_facts = disclosure_inputs.rebuild(original)
        assert len(retained_pack.children) == len(retained_stamps) == 16
        assert retained_facts == evidence["disclosure_inputs"]
        with pytest.raises(Problem):
            disclosure_inputs.rebuild({**original, "book_code": "IFRS15"})
        import copy

        damaged = copy.deepcopy(original)
        damaged["source_binding"]["evidence"]["disclosure_inputs"]["children"][0][
            "output_sha256"
        ] = "f" * 64
        with pytest.raises(Problem):
            disclosure_inputs.rebuild(damaged)
        count_before = uow.session.scalar(select(func.count()).select_from(report_run))
    response = post(world.app, f"/api/v1/report-runs/{original['id']}/rerun", actor, {})
    assert response.status_code == 202, response.text
    rerun_id = UUID(response.headers["X-Erev-Report-Run-Id"])
    rerun_job_id = UUID(response.json()["id"])
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(report_run)) == count_before + 1
    with world.place.uow() as uow:
        rerun_task_id = uow.session.scalar(
            select(job.c.procrastinate_job_id).where(job.c.id == rerun_job_id)
        )
        uow.session.execute(
            text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
            {"id": rerun_task_id},
        )
        uow.commit()
    monkeypatch.setattr(
        disclosure_assembly, "prepare", lambda *args: pytest.fail("live capture on rerun")
    )
    run_job(
        rerun_job_id,
        world.place.tenant_id,
        attempt=1,
        runtime=runtime,
        delivered_task_id=rerun_task_id,
    )
    with world.place.uow() as uow:
        repeated = (
            uow.session.execute(select(report_run).where(report_run.c.id == rerun_id))
            .mappings()
            .one()
        )
        assert repeated["status"] == "SUCCEEDED", repeated["problem"]
        assert repeated["output_sha256"] == original["output_sha256"]
        assert repeated["control_totals"] == original["control_totals"]
        assert repeated["tie_out_results"] == original["tie_out_results"]
        assert repeated["child_report_run_ids"] == original["child_report_run_ids"]
        assert repeated["source_binding"]["evidence"] == original["source_binding"]["evidence"]

        repeated_snapshots = list(
            uow.session.execute(
                select(disclosure_snapshot).where(disclosure_snapshot.c.report_run_id == rerun_id)
            ).mappings()
        )
        assert len(repeated_snapshots) == 7
        assert set(repeated["disclosure_snapshot_ids"]).isdisjoint(
            original["disclosure_snapshot_ids"]
        )
        assert {
            row["disclosure_kind"]: row["data_sha256"] for row in repeated_snapshots
        } == snapshot_hashes

    from erev_api.domain.close import freeze
    from erev_api.domain.reports import disclosure_freeze
    from erev_api.domain.reports.snapshots import SnapshotScope

    freeze_scope = SnapshotScope(
        entity_id=world.entity_id,
        book_code="ASC606",
        period_id=period.id,
        known_at=parent["known_at"],
        frozen_at=parent["known_at"],
    )
    with world.place.uow() as uow:
        before_capture = uow.session.scalar(select(func.count()).select_from(report_run))
        with pytest.raises(Problem) as refused_capture:
            disclosure_freeze.capture(uow, freeze_scope, quarter=True)
        assert refused_capture.value.slug == "forbidden"
        with freeze.system_unit(uow) as reader:
            frozen_inputs = disclosure_freeze.capture(reader, freeze_scope, quarter=True)
            with pytest.raises(Problem) as ordinary_policy:
                disclosure_expedients.capture(
                    reader, entity_id=world.entity_id, book="ASC606", known_at=parent["known_at"]
                )
            assert ordinary_policy.value.slug == "forbidden"
        assert uow.session.scalar(select(func.count()).select_from(report_run)) == before_capture
        assert len(frozen_inputs["children"]) == 16
        assert frozen_inputs["policy_sha256"] == sha256_hex(frozen_inputs["policies"])
        assert frozen_inputs["known_at"] == parent["known_at"].isoformat()
        expected_inputs = {
            child["key"]: child for child in evidence["disclosure_inputs"]["children"]
        }
        for captured in frozen_inputs["children"]:
            assert captured["dataset"] == expected_inputs[captured["key"]]["dataset"]
            assert captured["output_sha256"] == sha256_hex(captured["dataset"])
            assert "run_id" not in captured and "job_id" not in captured

        from erev_api.db.tables import period_lock, period_state_transition
        from erev_api.domain.reports import disclosure_lock, disclosure_queries
        from support.rows import CloseParts, insert_approval_request, period_lock_values

        transition = (
            uow.session.execute(
                select(period_state_transition)
                .where(
                    period_state_transition.c.entity_id == world.entity_id,
                    period_state_transition.c.period_id == period.id,
                    period_state_transition.c.book_code == "ASC606",
                )
                .limit(1)
            )
            .mappings()
            .one()
        )
        lock_values = period_lock_values(
            world.place.tenant_id,
            parts=CloseParts(
                calendar_id=UUID(int=0),
                entity_id=world.entity_id,
                period_id=period.id,
                period_state_transition_id=transition["id"],
                approval_request_id=insert_approval_request(
                    uow.session, tenant_id=world.place.tenant_id, entity_id=world.entity_id
                ),
                file_id=UUID(int=0),
            ),
            cutoff_known_at=parent["known_at"],
            created_at=uow.now,
            created_by=uow.principal.id,
            created_by_kind=uow.principal.kind.value,
        )
        uow.session.execute(insert(period_lock).values(**lock_values))
        # A failed publication must not leave a partial population of successful source runs.
        malformed = copy.deepcopy(frozen_inputs)
        malformed["children"][-1]["output_sha256"] = "f" * 64
        audit_before = len(uow.audit_events)
        with pytest.raises(Problem):
            disclosure_lock.publish(uow, lock_values["id"], malformed)
        assert len(uow.audit_events) == audit_before
        assert uow.session.scalar(select(func.count()).select_from(report_run)) == before_capture
        frozen_parent_id = disclosure_lock.publish(uow, lock_values["id"], frozen_inputs)
        frozen_parent = (
            uow.session.execute(select(report_run).where(report_run.c.id == frozen_parent_id))
            .mappings()
            .one()
        )
        assert frozen_parent["status"] == "SUCCEEDED"
        assert frozen_parent["job_id"] is None
        assert frozen_parent["period_lock_id"] == lock_values["id"]
        assert frozen_parent["source_binding"] is None
        assert len(frozen_parent["child_report_run_ids"]) == 16
        assert len(frozen_parent["disclosure_snapshot_ids"]) == 7
        assert (
            uow.session.scalar(select(func.count()).select_from(report_run)) == before_capture + 17
        )
        for identity in frozen_parent["disclosure_snapshot_ids"]:
            snapshot = disclosure_queries.get(
                uow.session, replace(uow.principal, permission_scopes={"report.run": "*"}), identity
            )
            assert snapshot.source == "LOCKED"
            assert snapshot.period_lock_id == lock_values["id"]
            assert snapshot.data["known_at"] == parent["known_at"].isoformat()
            assert sha256_hex(snapshot.data) == snapshot.data_sha256
            if snapshot.disclosure_kind.value == "EXPEDIENTS_USED":
                assert len(snapshot.data["supporting_children"]) == 2
                for child in snapshot.data["supporting_children"]:
                    assert child["dataset"] == expected_inputs[child["key"]]["dataset"]
            for child in snapshot.data["children"]:
                assert child["job_id"] is None
                assert child["dataset"] == expected_inputs[child["key"]]["dataset"]
                stored = (
                    uow.session.execute(
                        select(report_run).where(report_run.c.id == UUID(child["run_id"]))
                    )
                    .mappings()
                    .one()
                )
                assert stored["status"] == "SUCCEEDED" and stored["job_id"] is None
                assert stored["output_sha256"] == child["output_sha256"]
        with pytest.raises(Problem) as duplicate:
            disclosure_lock.publish(uow, lock_values["id"], frozen_inputs)
        assert duplicate.value.slug == "invalid-transition"

        uow.commit()
    frozen_detail = get(world.app, f"/api/v1/report-runs/{frozen_parent_id}", actor)
    assert frozen_detail.status_code == 200, frozen_detail.text
    frozen_data = get(
        world.app, f"/api/v1/report-runs/{frozen_parent_id}/data", actor, {"limit": 1}
    )
    assert frozen_data.status_code == 200, frozen_data.text
    assert len(frozen_data.json()["sections"]) == 7
    saved = get(
        world.app, "/api/v1/disclosure-snapshots", actor, {"report_run_id": str(frozen_parent_id)}
    )
    assert saved.status_code == 200, saved.text
    assert len(saved.json()["items"]) == 7
    assert all(row["source"] == "LOCKED" for row in saved.json()["items"])

    # Frozen exports rebuild from E-65. Child output files and current policy/builders are not read.
    from erev_api.domain.reports import disclosure_frozen

    monkeypatch.setattr(
        disclosure_expedients, "capture", lambda *args, **kwargs: pytest.fail("live policies")
    )
    with world.place.uow() as uow:
        frozen = disclosure_frozen.read(uow.session, dict(frozen_parent))
        assert len(frozen.pack.children) == 16
        assert tuple(frozen_parent["disclosure_snapshot_ids"]) == frozen.snapshot_ids
        damaged_run = {**frozen_parent, "known_at": uow.now}
        if uow.now != frozen_parent["known_at"]:
            with pytest.raises(Problem):
                disclosure_frozen.read(uow.session, damaged_run)
        bad_attached = {**frozen_parent, "disclosure_snapshot_ids": []}
        with pytest.raises(Problem):
            disclosure_frozen.read(uow.session, bad_attached)

    def finish_frozen(response: Any) -> Mapping[str, Any]:
        assert response.status_code == 202, response.text
        identity = UUID(response.headers["X-Erev-Report-Run-Id"])
        job_id = UUID(response.json()["id"])
        with world.place.uow() as uow:
            task_id = uow.session.scalar(
                select(job.c.procrastinate_job_id).where(job.c.id == job_id)
            )
            uow.session.execute(
                text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
                {"id": task_id},
            )
            uow.commit()
        run_job(
            job_id, world.place.tenant_id, attempt=1, runtime=runtime, delivered_task_id=task_id
        )
        with world.place.uow() as uow:
            built = dict(
                uow.session.execute(select(report_run).where(report_run.c.id == identity))
                .mappings()
                .one()
            )
            assert built["status"] == "SUCCEEDED", built["problem"]
            assert built["source_binding"] is None
            assert built["disclosure_snapshot_ids"] == frozen_parent["disclosure_snapshot_ids"]
        return built

    for output_format in ("JSON", "XLSX", "PDF"):
        built = finish_frozen(
            post(
                world.app,
                "/api/v1/report-runs",
                actor,
                {
                    "report_code": "disclosure_pack",
                    "output_format": output_format,
                    "parameters": {"period_lock_id": str(lock_values["id"]), "quarter": True},
                },
            )
        )
        assert built["output_sha256"] == frozen_parent["output_sha256"]
        exported = get(world.app, f"/api/v1/report-runs/{built['id']}/output", actor)
        assert exported.status_code == 200, exported.text
        if output_format == "XLSX":
            assert exported.content.startswith(b"PK")
        elif output_format == "PDF":
            assert exported.content.startswith(b"%PDF-")
        else:
            assert len(exported.json()["sections"]) == 7
        attached = get(
            world.app,
            "/api/v1/disclosure-snapshots",
            actor,
            {"report_run_id": str(built["id"]), "limit": 3, "count": True},
        )
        assert attached.status_code == 200, attached.text
        assert attached.headers["X-Erev-Total-Count"] == "7"

    monthly = finish_frozen(
        post(
            world.app,
            "/api/v1/report-runs",
            actor,
            {
                "report_code": "disclosure_pack",
                "output_format": "JSON",
                "parameters": {"period_lock_id": str(lock_values["id"])},
            },
        )
    )
    assert len(monthly["child_report_run_ids"]) == 9
    repeated_lock = finish_frozen(
        post(world.app, f"/api/v1/report-runs/{frozen_parent_id}/rerun", actor, {})
    )
    assert repeated_lock["output_sha256"] == frozen_parent["output_sha256"]
    with world.place.uow() as uow:
        assert uow.session.scalar(select(func.count()).select_from(disclosure_snapshot)) == 21


def test_forecast_event_capture_is_atomic_and_immutable(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Real scenario guards, fact validation, atomic capture and publication immutability."""
    from erev_api.db.tables import forecast_event_set
    from erev_api.domain.contracts import events
    from erev_api.domain.forecast import event_sets
    from erev_api.schemas.events import EventAppendIn

    booked = k01(world)
    identity = booked.contract["id"]
    principal = replace(
        world.place.principal,
        permissions=world.place.principal.permissions | {"scenario.use"},
        permission_scopes={**world.place.principal.permission_scopes, "scenario.use": "*"},
    )
    with world.place.uow(principal) as uow, pytest.raises(Problem, match="scenario sandbox"):
        event_sets.create(uow, {"code": "BASE", "name": "Forecast base"})
    from datetime import UTC, datetime, timedelta

    from erev_api.db.tables import job, tenant_membership
    from erev_api.domain.platform import snapshot_job
    from erev_api.enums import JobKind
    from erev_api.jobs import registry
    from erev_api.jobs.context import JobRuntime
    from sqlalchemy import text
    from support.http import call
    from support.principals import Member, cookie_headers, sign_in
    from support.principals import workspace as signed_workspace
    from support.snapshots import confirm_retention

    clock = world.place.clock
    clock.advance(max(timedelta(minutes=1), datetime.now(UTC) - clock.now() + timedelta(minutes=1)))
    source_actor = world.place.author
    source_actor = signed_workspace(
        world.app,
        source_actor.member,
        sign_in(world.app, source_actor.member.email),
        source_actor.secret,
    )
    with tenant_session(principal.db_context) as session:
        confirm_retention(session, principal.tenant_id, at=clock.now() - timedelta(days=1))
    created = call(
        world.app,
        "POST",
        "/api/v1/scenarios",
        headers=cookie_headers(source_actor.token, source_actor.csrf_token),
        json={"name": "Forecast fixture", "known_at": clock.now().isoformat()},
    )
    assert created.status_code == 202, created.text
    job_id = UUID(created.json()["id"])
    with tenant_session(principal.db_context) as session:
        pending = session.execute(select(job).where(job.c.id == job_id)).mappings().one()
        session.execute(
            text("UPDATE procrastinate_jobs SET status = 'doing' WHERE id = :id"),
            {"id": pending["procrastinate_job_id"]},
        )
    monkeypatch.setitem(
        registry.HANDLERS,
        JobKind.TENANT_SNAPSHOT,
        registry.HandlerSpec(
            handler=snapshot_job.tenant_snapshot_export,
            retry=snapshot_job.SNAPSHOT_RETRY,
            on_failure=snapshot_job.snapshot_failed,
        ),
    )
    registry.run_job(
        job_id,
        principal.tenant_id,
        attempt=1,
        runtime=JobRuntime(clock=clock, keyring=world.place.keyring, files=world.place.files),
    )
    with tenant_session(principal.db_context) as session:
        finished = session.execute(select(job).where(job.c.id == job_id)).mappings().one()
        assert finished["state"] == "SUCCEEDED", finished["problem"]
    listed = call(
        world.app, "GET", "/api/v1/scenarios", headers=cookie_headers(source_actor.token, key=False)
    )
    target_id = UUID(listed.json()["items"][0]["scenario_tenant_id"])
    with tenant_session(DbContext(tenant_id=target_id, user_id=None, entity_scope="*")) as session:
        membership_id = session.execute(
            select(tenant_membership.c.id).where(tenant_membership.c.user_id == principal.id)
        ).scalar_one()
    target_member = Member(
        user_id=source_actor.member.user_id,
        email=source_actor.member.email,
        tenant_id=target_id,
        membership_id=membership_id,
    )
    sandbox_actor = signed_workspace(
        world.app, target_member, sign_in(world.app, target_member.email), source_actor.secret
    )
    world = replace(
        world,
        place=workspace(world.app, clock, world.place.keyring, world.place.files, sandbox_actor),
    )
    principal = world.place.principal
    with world.place.uow(principal) as uow:
        draft = event_sets.create(uow, {"code": "BASE", "name": "Forecast base"})
        uow.commit()
    head = world.place.scalar(
        select(contract.c.head_stream_version).where(contract.c.id == identity)
    )
    body = EventAppendIn.model_validate(
        {
            "events": [
                {
                    "event_type": "PROGRESS_RECORDED",
                    "effective_date": "2026-01-31",
                    "payload": {
                        "obligation_key": "O2",
                        "cumulative_progress_ratio": "0.40",
                        "measure": "OUTPUT_PERCENT",
                    },
                }
            ]
        }
    )
    # A failure after actual event execution must also undo facts/computation/audit buffers.
    original_change = event_sets.change

    def fail_capture(*_args: Any, **_kwargs: Any) -> None:
        raise Problem("validation-failed", "capture failed")

    monkeypatch.setattr(event_sets, "change", fail_capture)
    with world.place.uow(principal) as uow:
        with pytest.raises(Problem, match="capture failed"):
            events.record_events(
                uow,
                contract_id=identity,
                expected_stream_version=head,
                body=body,
                forecast_event_set_id=draft["id"],
                engine=engine(),
            )
        assert (
            uow.session.execute(
                select(contract.c.head_stream_version).where(contract.c.id == identity)
            ).scalar_one()
            == head
        )
        uow.commit()
    monkeypatch.setattr(event_sets, "change", original_change)
    from support.http import call
    from support.principals import cookie_headers

    headers = cookie_headers(world.place.author.token, world.place.author.csrf_token)
    headers["If-Match"] = f'"s{head}"'
    path = f"/api/v1/contracts/{identity}/events"
    response = call(
        world.app,
        "POST",
        path,
        headers=headers,
        params={"forecast_event_set_id": str(draft["id"])},
        json=body.model_dump(mode="json"),
    )
    assert response.status_code == 201 and "events" in response.json(), response.text
    repeated = call(
        world.app,
        "POST",
        path,
        headers=headers,
        params={"forecast_event_set_id": str(draft["id"])},
        json=body.model_dump(mode="json"),
    )
    assert repeated.status_code == 201 and repeated.json() == response.json(), repeated.text
    captured = world.place.rows(
        select(forecast_event_set).where(forecast_event_set.c.id == draft["id"])
    )[0]
    assert captured["api_commands"] == [
        {
            "method": "POST",
            "path": f"/contracts/{identity}/events",
            "body": body.model_dump(mode="json"),
        }
    ]
    with world.place.uow(principal) as uow:
        event_sets.change(uow, draft["id"], expected_version=captured["row_version"], publish=True)
        uow.commit()
    with world.place.uow(principal) as uow, pytest.raises(Problem, match="draft version"):
        events.record_events(
            uow,
            contract_id=identity,
            expected_stream_version=head,
            body=body,
            forecast_event_set_id=draft["id"],
            engine=engine(),
        )

    from support.factories import ImportWorld
    from support.forecast_inputs import exercise_forecast_imports

    exercise_forecast_imports(
        ImportWorld(
            app=world.app,
            actor=sandbox_actor,
            clock=clock,
            runtime=JobRuntime(clock=clock, keyring=world.place.keyring, files=world.place.files),
        )
    )
