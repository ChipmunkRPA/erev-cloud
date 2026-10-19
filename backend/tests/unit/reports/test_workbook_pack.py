"""Pack worksheets retain distinct schemas, source identity and safe omission notes."""

import io
from dataclasses import replace
from datetime import UTC, datetime

import openpyxl
import pytest
from erev_api.domain.reports.outputs import Column, ReportData, RunStamp
from erev_api.domain.reports.outputs.xlsx import WorkbookSection, render_xlsx_pack

STAMP = RunStamp(
    "disclosure_pack",
    "Disclosure pack",
    1,
    "RPT-PACK",
    {},
    ("US",),
    "ASC606",
    "2026-09-30",
    "Frozen close",
    "1.0",
    "abc",
    "Maya",
    datetime(2026, 10, 1, tzinfo=UTC),
    "a" * 64,
)
DATA = ReportData(
    (Column("name", "Category", "text"), Column("amount", "Revenue", "money")),
    ({"row_key": "one", "name": "=1+1", "amount": {"amount": "123.45", "currency": "USD"}},),
    {"revenue": {"amount": "123.45", "currency": "USD"}},
    ({"code": "TO_REVENUE", "result": "PASS", "expected": "123.45", "actual": "123.45"},),
)


def test_seven_sections_keep_source_stamps_money_tieouts_and_omissions() -> None:
    sections = tuple(
        WorkbookSection(
            str(index),
            "A very long chapter/title which repeats'",
            DATA,
            replace(STAMP, report_run_no=f"CHILD-{index}", output_sha256=str(index) * 64),
        )
        for index in range(1, 7)
    ) + (
        WorkbookSection(
            "7",
            "Expedients",
            ReportData((), (), {}),
            replace(STAMP, report_run_no="CHILD-7"),
            notes=("Omitted under the elected relief.",),
            omitted=True,
        ),
    )
    raw = render_xlsx_pack(sections, STAMP)
    assert render_xlsx_pack(sections, STAMP) == raw
    book = openpyxl.load_workbook(io.BytesIO(raw))
    assert len(book.worksheets) == 8
    assert len(set(book.sheetnames)) == 8
    assert all(len(name) <= 31 and "/" not in name for name in book.sheetnames)
    index_values = [cell.value for row in book["Contents"] for cell in row]
    assert "RPT-PACK" in index_values and "Omitted" in index_values
    for index, sheet in enumerate(book.worksheets[1:7], start=1):
        cells = [cell for row in sheet for cell in row]
        values = [cell.value for cell in cells]
        assert f"CHILD-{index}" in values and str(index) * 64 in values
        assert "PASS" in values and 123.45 in values
        assert "'=1+1" in values
        assert all(cell.data_type != "f" and cell.hyperlink is None for cell in cells)
        assert sheet.freeze_panes is not None
    omitted = [cell.value for row in book.worksheets[-1] for cell in row]
    assert "Omitted under the elected relief." in omitted
    assert 123.45 not in omitted


def test_omitted_rows_or_totals_cannot_leak_into_workbook() -> None:
    section = WorkbookSection("s", "Section", DATA, STAMP, ("Omitted",), True)
    with pytest.raises(ValueError):
        render_xlsx_pack((section,), STAMP)
    with pytest.raises(ValueError):
        render_xlsx_pack((replace(section, data=replace(DATA, rows=())),), STAMP)
    with pytest.raises(ValueError):
        render_xlsx_pack((), STAMP)


def test_pdf_pack_has_source_stamps_index_and_safe_omissions() -> None:
    from erev_api.domain.reports.outputs.pdf import render_pdf_chapters
    from erev_api.domain.reports.outputs.xlsx import WorkbookChapter
    from tests.domain.reports.test_framework import pdf_pages

    populated = replace(
        DATA,
        rows=tuple({**DATA.rows[0], "row_key": str(n)} for n in range(130)),
    )
    chapters = (
        WorkbookChapter(
            "1", "Revenue", (WorkbookSection("revenue", "Revenue", populated, STAMP, ("Period",)),)
        ),
        WorkbookChapter(
            "2",
            "Expedients",
            (
                WorkbookSection(
                    "relief",
                    "Expedients",
                    ReportData((), (), {}),
                    STAMP,
                    ("Omitted under elected relief.",),
                    True,
                ),
            ),
        ),
    )
    raw = render_pdf_chapters(chapters, STAMP)
    assert render_pdf_chapters(chapters, STAMP) == raw
    pages = pdf_pages(raw)
    assert pages[0][:2] == ["Contents", "Run RPT-PACK"]
    assert "2" in pages[0]  # First chapter starts after the contents page.
    assert "Output SHA-256: " + "a" * 64 in pages[1]
    assert pages[2][:2] == ["Category", "Revenue (USD)"]
    assert "Omitted under elected relief." in pages[-1]
    assert "123.45" not in pages[-1]
    assert [page[-1] for page in pages] == [
        f"Run RPT-PACK · page {n} of {len(pages)}" for n in range(1, len(pages) + 1)
    ]
    with pytest.raises(ValueError):
        render_pdf_chapters(
            (
                WorkbookChapter(
                    "bad",
                    "Bad",
                    (WorkbookSection("leak", "Leak", DATA, STAMP, ("Omitted",), True),),
                ),
            ),
            STAMP,
        )
