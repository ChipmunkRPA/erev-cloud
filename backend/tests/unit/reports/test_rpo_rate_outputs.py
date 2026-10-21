"""RPO exports expose pinned FX legs without rounding their exact decimal strings."""

import io
from dataclasses import replace

import openpyxl
import pytest
from erev_api.domain.reports.outputs.pdf import render_pdf
from erev_api.domain.reports.outputs.xlsx import WorkbookSection, render_xlsx, render_xlsx_pack
from tests.unit.reports.test_workbook_pack import DATA, STAMP


@pytest.mark.parametrize("packed", [False, True])
def test_retained_rpo_rates_are_readable_and_exact(packed: bool) -> None:
    source = {
        "entity_id": "entity-1",
        "period_key": "FY2026-P02",
        "period_end": "2026-02-28",
        "base_currency": "EUR",
        "quote_currency": "USD",
        "rate_type": "average",
        "rate": "1.234567890123456789",
        "rate_id": "=original-rate",
        "version_id": "approved-version",
        "rate_cutoff": "2026-03-01T12:00:00+00:00",
    }
    rates = [
        source,
        {
            **source,
            "base_currency": "USD",
            "rate_type": "closing",
            "rate": "1",
            "rate_id": None,
            "version_id": None,
        },
    ]
    data = replace(DATA, control_totals={"translation_rates": rates})
    stamp = replace(STAMP, report_code="rpo_rollforward", report_name="RPO rollforward")
    raw = (
        render_xlsx_pack((WorkbookSection("rpo", "RPO", data, stamp),), STAMP)
        if packed
        else render_xlsx(data, stamp)
    )
    sheet = openpyxl.load_workbook(io.BytesIO(raw)).worksheets[-1]
    title = next(
        row
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row, 1).value == "Retained RPO exchange rates"
    )
    header = title + 2
    assert sheet.cell(header, 7).value == "Retained rate (exact)"
    assert sheet.cell(header + 1, 7).value == source["rate"]
    assert sheet.cell(header + 1, 7).data_type == "s"
    assert sheet.cell(header + 1, 8).value == "'=original-rate"
    assert sheet.cell(header + 1, 9).value == "approved-version"
    assert sheet.cell(header + 2, 7).value == "1"
    assert sheet.cell(header + 2, 8).value == ""
    pdf = render_pdf(data, stamp)
    assert b"Retained RPO exchange rates" in pdf
    assert source["rate"].encode() in pdf
    assert b"approved-version" in pdf
    assert data.control_totals["translation_rates"] == rates
