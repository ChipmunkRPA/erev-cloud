"""Real queued call-export builds retain the requester through promotion and revocation."""

import csv
import io
from collections.abc import Iterator, Mapping
from dataclasses import replace
from typing import Any
from uuid import UUID

import openpyxl
import pytest
from erev_api.auth.keyring import KeyRing
from erev_api.auth.principal import system_principal
from erev_api.clock import FrozenClock
from erev_api.db.session import DbContext, tenant_session
from erev_api.db.tables import ai_model_log, job, report_definition, report_run, role_assignment
from erev_api.domain.ai.export_visibility import CODE
from erev_api.domain.reports import framework
from erev_api.domain.reports.outputs import ReportData
from erev_api.enums import JobKind
from erev_api.jobs.context import JobAttempt, JobContext, JobRuntime
from erev_api.problems import Problem
from fastapi import FastAPI
from sqlalchemy import insert, select, update
from sqlalchemy.orm import Session
from support.db import TestDatabase, fresh_head
from support.factories import stamp_test_release
from support.principals import colleague, enrolled, member
from support.reference import get, post
from support.rows import RowContext, ai_model_log_row, insert_custom_role, insert_role_assignment
from tests.api.test_tenant_api import app as app  # noqa: F401
from tests.domain.reports.test_framework import run_deferred
from tests.domain.reports.test_framework import runtime as runtime  # noqa: F401

pytestmark = pytest.mark.pg


def _definition(database: TestDatabase) -> None:
    # Test-only registration. The public catalogue remains disabled until request/replay
    # authorization and its screen menu are integrated. Exercise the real report pipeline.
    with database.owner_engine.begin() as connection:
        connection.execute(
            insert(report_definition).values(
                code=CODE,
                version=1,
                name="AI call log export",
                kind="EXTRACT",
                description="Current authorized call evidence; reruns are new live evaluations.",
                parameters_schema={
                    "type": "object",
                    "properties": {},
                    "additionalProperties": False,
                },
                output_formats=["CSV", "XLSX"],
                ipe_logic={
                    "version": 1,
                    "source_tables": ["ai_model_log"],
                    "joins": [],
                    "filters": ["Original requester scope and current source/export permissions"],
                    "parameters": [],
                },
                tie_outs=[],
            )
        )


@pytest.fixture(autouse=True)
def export_definition(test_database: TestDatabase) -> Iterator[None]:
    # Global catalogue rows must not leak to other tests. This isolated test database uses
    # the same reset discipline as migration probes, with no change to real seed migrations.
    fresh_head()
    _definition(test_database)
    try:
        yield
    finally:
        fresh_head()
        test_database.app_engine.dispose()
        test_database.owner_engine.dispose()


@pytest.mark.parametrize("output_format", ["CSV", "XLSX"])
def test_call_export_worker_keeps_own_scope_and_rechecks_before_replay(
    app: FastAPI,
    keyring: KeyRing,
    clock: FrozenClock,
    runtime: JobRuntime,
    output_format: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    stamp_test_release()
    owner = member(keyring, clock)
    peer = colleague(owner.tenant_id, "worker_peer")
    context = DbContext(owner.tenant_id, owner.user_id, "*")
    with tenant_session(context) as session:
        insert_custom_role(
            session,
            tenant_id=owner.tenant_id,
            code="call_exporter",
            permissions=["ai.use", "report.export"],
        )
        assignment = insert_role_assignment(
            session,
            tenant_id=owner.tenant_id,
            membership_id=owner.membership_id,
            role_code="call_exporter",
        )
        ids = []
        for person in (owner, peer):
            row = ai_model_log_row(RowContext(person.tenant_id, person.user_id), session)
            row.update(created_by=person.user_id, created_by_kind="USER")
            session.execute(insert(ai_model_log).values(**row))
            ids.append(str(row["id"]))
    actor = enrolled(app, clock, owner)
    body = {"report_code": CODE, "parameters": {}, "output_format": output_format}
    started = post(app, "/api/v1/report-runs", actor, body)
    assert started.status_code == 202, started.text
    run_id = UUID(started.headers["x-erev-report-run-id"])
    job_id = UUID(started.json()["id"])
    with tenant_session(context) as session:
        binding = session.scalar(
            select(report_run.c.ai_export_access).where(report_run.c.id == run_id)
        )
        assert binding is not None and binding["scope"] == "OWN"
        promoted = insert_role_assignment(
            session,
            tenant_id=owner.tenant_id,
            membership_id=owner.membership_id,
            role_code="tenant_admin",
        )
    run_deferred(owner.tenant_id, job_id, runtime)
    completed = get(app, f"/api/v1/report-runs/{run_id}", actor)
    assert completed.status_code == 200, completed.text
    assert completed.json()["status"] == "SUCCEEDED", completed.text
    assert completed.json()["row_count"] == 1
    download = get(app, f"/api/v1/report-runs/{run_id}/output", actor)
    assert download.status_code == 200, download.text
    if output_format == "CSV":
        records = list(csv.DictReader(io.StringIO(download.content.decode("utf-8"))))
        assert [record["Call id"] for record in records] == ids[:1]
    else:
        workbook = openpyxl.load_workbook(io.BytesIO(download.content))
        values = [str(cell.value) for row in workbook.active for cell in row]
        assert ids[0] in values and ids[1] not in values
    pending_wide = post(app, "/api/v1/report-runs", actor, body)
    assert pending_wide.status_code == 202, pending_wide.text
    pending_id = UUID(pending_wide.headers["x-erev-report-run-id"])
    pending_job = UUID(pending_wide.json()["id"])
    with tenant_session(context, read_only=True) as session:
        assert (
            session.scalar(select(report_run.c.ledger_heads).where(report_run.c.id == run_id)) == {}
        )
        pending_binding = session.scalar(
            select(report_run.c.ai_export_access).where(report_run.c.id == pending_id)
        )
        assert pending_binding is not None and pending_binding["scope"] == "TENANT"
    rerun = post(app, f"/api/v1/report-runs/{run_id}/rerun", actor, {})
    assert rerun.status_code == 202, rerun.text
    rerun_id = UUID(rerun.headers["x-erev-report-run-id"])
    rerun_job = UUID(rerun.json()["id"])
    with tenant_session(context) as session:
        assert (
            session.scalar(select(report_run.c.ai_export_access).where(report_run.c.id == rerun_id))
            == binding
        )
        # Run the real handler under a persisted live attempt, leaving job settlement pending
        # to exercise its completed-output replay path after the report has committed.
        session.execute(
            update(job)
            .where(job.c.id == rerun_job)
            .values(state="RUNNING", updated_by_kind="SYSTEM")
        )
    jc = JobContext(
        job_id=rerun_job,
        tenant_id=owner.tenant_id,
        kind=JobKind.REPORT_RUN,
        principal=system_principal(owner.tenant_id, on_behalf_of_id=owner.user_id),
        runtime=replace(runtime, attempt=JobAttempt(rerun_job, owner.tenant_id, None)),
    )
    framework.run_report(jc, {"report_run_id": str(rerun_id)})
    framework.run_report(jc, {"report_run_id": str(rerun_id)})
    during_render = post(app, "/api/v1/report-runs", actor, body)
    assert during_render.status_code == 202, during_render.text
    during_id = UUID(during_render.headers["x-erev-report-run-id"])
    original_render = framework._render

    def revoke_during_render(
        session: Session,
        run: Mapping[str, Any],
        definition: Mapping[str, Any],
        data: ReportData,
        style: str,
    ) -> framework._Rendered:
        assert run["id"] == during_id and data.row_count == 2
        rendered = original_render(session, run, definition, data, style)
        with tenant_session(context) as revoke_session:
            revoke_session.execute(
                update(role_assignment)
                .where(role_assignment.c.id == promoted)
                .values(revoked_at=clock.now(), revoked_by_kind="SYSTEM")
            )
        return rendered

    with monkeypatch.context() as patch:
        patch.setattr(framework, "_render", revoke_during_render)
        run_deferred(owner.tenant_id, UUID(during_render.json()["id"]), runtime)
    with tenant_session(context, read_only=True) as session:
        refused_run = (
            session.execute(select(report_run).where(report_run.c.id == during_id)).mappings().one()
        )
        assert refused_run["status"] == "RUNNING"
        assert refused_run["output_file_id"] is None
        assert refused_run["manifest_file_id"] is None
        assert refused_run["row_count"] is None

    with tenant_session(context) as session:
        assert (
            session.scalar(select(report_run.c.row_count).where(report_run.c.id == rerun_id)) == 1
        )
        session.execute(
            update(role_assignment)
            .where(
                role_assignment.c.id.in_([assignment, promoted]),
                role_assignment.c.revoked_at.is_(None),
            )
            .values(revoked_at=clock.now(), revoked_by_kind="SYSTEM")
        )
    with pytest.raises(Problem) as refused:
        framework.run_report(jc, {"report_run_id": str(rerun_id)})
    assert refused.value.slug == "forbidden"
    assert get(app, f"/api/v1/report-runs/{run_id}/output", actor).status_code == 403
    run_deferred(owner.tenant_id, pending_job, runtime)
    with tenant_session(context, read_only=True) as session:
        pending = (
            session.execute(select(report_run).where(report_run.c.id == pending_id))
            .mappings()
            .one()
        )
        assert pending["status"] == "QUEUED"
        assert pending["output_file_id"] is None
        assert pending["manifest_file_id"] is None
        assert pending["row_count"] is None


def test_source_only_roles_and_cached_commands_recheck_retained_access(
    app: FastAPI,
    test_database: TestDatabase,
    keyring: KeyRing,
    clock: FrozenClock,
    runtime: JobRuntime,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from erev_api.db.tables import legal_entity, user_session
    from erev_api.idempotency import store
    from support.http import call
    from support.principals import cookie_headers
    from support.rows import legal_entity_row

    stamp_test_release()
    owner = member(keyring, clock)
    administrator = colleague(owner.tenant_id, "export_settings_only")
    context = DbContext(owner.tenant_id, owner.user_id, "*")
    with tenant_session(context) as session:
        for code, permissions, person in (
            ("calls_own", ["ai.use", "report.export"], owner),
            ("calls_admin", ["settings.manage", "report.export"], administrator),
        ):
            insert_custom_role(
                session, tenant_id=owner.tenant_id, code=code, permissions=permissions
            )
            insert_role_assignment(
                session,
                tenant_id=owner.tenant_id,
                membership_id=person.membership_id,
                role_code=code,
            )
    actor = enrolled(app, clock, owner)
    admin = enrolled(app, clock, administrator)
    body = {"report_code": CODE, "parameters": {}, "output_format": "CSV"}
    for reader in (actor, admin):
        definitions = get(app, "/api/v1/report-definitions", reader)
        assert definitions.status_code == 200, definitions.text
        assert [value["code"] for value in definitions.json()["items"]] == [CODE]
        denied = post(
            app, "/api/v1/report-runs", reader, {**body, "report_code": "revenue_waterfall"}
        )
        assert denied.status_code == 403, denied.text
    headers = cookie_headers(actor.token, actor.csrf_token)
    first = call(app, "POST", "/api/v1/report-runs", json=body, headers=headers)
    assert first.status_code == 202, first.text
    run_id = first.headers["x-erev-report-run-id"]
    repeated = call(app, "POST", "/api/v1/report-runs", json=body, headers=headers)
    assert repeated.status_code == 202 and repeated.json()["id"] == first.json()["id"]
    for injected in (
        {"actor_id": str(administrator.user_id)},
        {"ai_export_access": {"scope": "TENANT"}},
    ):
        refused = post(app, "/api/v1/report-runs", actor, {**body, "parameters": injected})
        assert refused.status_code == 422, refused.text
    run_deferred(owner.tenant_id, UUID(first.json()["id"]), runtime)
    for suffix in ("", "/output", "/output?part=manifest"):
        hidden = get(app, f"/api/v1/report-runs/{run_id}{suffix}", admin)
        assert hidden.status_code == 404, hidden.text
    assert get(app, f"/api/v1/jobs/{first.json()['id']}", admin).status_code == 404
    rerun_headers = cookie_headers(actor.token, actor.csrf_token)
    rerun_path = f"/api/v1/report-runs/{run_id}/rerun"
    rerun = call(app, "POST", rerun_path, json={}, headers=rerun_headers)
    assert rerun.status_code == 202, rerun.text
    assert (
        call(app, "POST", rerun_path, json={}, headers=rerun_headers).json()["id"]
        == rerun.json()["id"]
    )

    # Promotion must not reinterpret an old OWN request as a new TENANT request. The current
    # named-entity export grant still permits OWN, while a new workspace export must refuse.
    with tenant_session(context) as session:
        entity = legal_entity_row(RowContext(owner.tenant_id, owner.user_id), session)
        session.execute(insert(legal_entity).values(**entity))
        session.execute(
            update(role_assignment)
            .where(
                role_assignment.c.membership_id == owner.membership_id,
                role_assignment.c.revoked_at.is_(None),
            )
            .values(revoked_at=clock.now(), revoked_by_kind="SYSTEM")
        )
        insert_custom_role(
            session,
            tenant_id=owner.tenant_id,
            code="settings_only",
            permissions=["settings.manage"],
        )
        limited = insert_role_assignment(
            session,
            tenant_id=owner.tenant_id,
            membership_id=owner.membership_id,
            role_code="calls_own",
            entity_ids=[entity["id"]],
        )
        settings = insert_role_assignment(
            session,
            tenant_id=owner.tenant_id,
            membership_id=owner.membership_id,
            role_code="settings_only",
        )
    own_again = call(app, "POST", "/api/v1/report-runs", json=body, headers=headers)
    assert own_again.status_code == 202 and own_again.json()["id"] == first.json()["id"]
    assert call(app, "POST", rerun_path, json={}, headers=rerun_headers).status_code == 202
    assert get(app, f"/api/v1/report-runs/{run_id}/output", actor).status_code == 200
    assert post(app, "/api/v1/report-runs", actor, body).status_code == 403
    with tenant_session(context) as session:
        session.execute(
            update(role_assignment)
            .where(role_assignment.c.id.in_([limited, settings]))
            .values(revoked_at=clock.now(), revoked_by_kind="SYSTEM")
        )
        insert_role_assignment(
            session,
            tenant_id=owner.tenant_id,
            membership_id=owner.membership_id,
            role_code="calls_own",
        )

    # Capture a TENANT request while the same actor has settings.manage. After demotion the
    # request body would now qualify for OWN; replay must still authorize its retained TENANT run.
    with tenant_session(context) as session:
        promotion = insert_role_assignment(
            session,
            tenant_id=owner.tenant_id,
            membership_id=owner.membership_id,
            role_code="calls_admin",
        )
    wide_headers = cookie_headers(actor.token, actor.csrf_token)
    wide = call(app, "POST", "/api/v1/report-runs", json=body, headers=wide_headers)
    assert wide.status_code == 202, wide.text
    with tenant_session(context) as session:
        session.execute(
            update(role_assignment)
            .where(role_assignment.c.id == promotion)
            .values(revoked_at=clock.now(), revoked_by_kind="SYSTEM")
        )
    refused_wide = call(app, "POST", "/api/v1/report-runs", json=body, headers=wide_headers)
    assert refused_wide.status_code == 404, refused_wide.text
    # The handler is bypassed on replay. Observe an actual stored Replay and prove that
    # authorization still refuses its retained TENANT scope before it reaches the response.
    original_begin = store.begin
    observed = []

    def observe_replay(*args: Any, **kwargs: Any) -> Any:
        result = original_begin(*args, **kwargs)
        if isinstance(result, store.Replay):
            observed.append(True)
        return result

    with monkeypatch.context() as patch:
        patch.setattr(store, "begin", observe_replay)
        late_refused = call(app, "POST", "/api/v1/report-runs", json=body, headers=wide_headers)
    assert observed == [True]
    assert late_refused.status_code == 404, late_refused.text
    listed = get(app, "/api/v1/report-runs", actor)
    assert listed.status_code == 200, listed.text
    assert wide.headers["x-erev-report-run-id"] not in {row["id"] for row in listed.json()["items"]}

    # Grants can also change while store.begin waits. Revoke only after it has produced the
    # stored response: cached request authentication must not survive that wait boundary.
    revoked_during_lookup = []

    def revoke_after_lookup(*args: Any, **kwargs: Any) -> Any:
        result = original_begin(*args, **kwargs)
        if isinstance(result, store.Replay):
            with tenant_session(context) as session:
                session.execute(
                    update(role_assignment)
                    .where(
                        role_assignment.c.membership_id == owner.membership_id,
                        role_assignment.c.revoked_at.is_(None),
                    )
                    .values(revoked_at=clock.now(), revoked_by_kind="SYSTEM")
                )
            revoked_during_lookup.append(True)
        return result

    with monkeypatch.context() as patch:
        patch.setattr(store, "begin", revoke_after_lookup)
        refreshed_refusal = call(app, "POST", "/api/v1/report-runs", json=body, headers=headers)
    assert revoked_during_lookup == [True]
    assert refreshed_refusal.status_code == 404, refreshed_refusal.text

    # Both cached commands refuse after the source role is gone, even though their replies exist.
    with tenant_session(context) as session:
        session.execute(
            update(role_assignment)
            .where(
                role_assignment.c.membership_id == owner.membership_id,
                role_assignment.c.revoked_at.is_(None),
            )
            .values(revoked_at=clock.now(), revoked_by_kind="SYSTEM")
        )
    assert call(app, "POST", "/api/v1/report-runs", json=body, headers=headers).status_code == 404
    assert call(app, "POST", rerun_path, json={}, headers=rerun_headers).status_code == 404

    # The normal authenticated-session MFA requirement also precedes cached administrator replies.
    admin_headers = cookie_headers(admin.token, admin.csrf_token)
    admin_run = call(app, "POST", "/api/v1/report-runs", json=body, headers=admin_headers)
    assert admin_run.status_code == 202, admin_run.text
    with test_database.owner_engine.begin() as connection:
        connection.execute(
            update(user_session)
            .where(user_session.c.user_id == administrator.user_id)
            .values(mfa_verified_at=None)
        )
    mfa_refused = call(app, "POST", "/api/v1/report-runs", json=body, headers=admin_headers)
    assert mfa_refused.status_code == 403, mfa_refused.text
    assert mfa_refused.json()["type"].endswith("/mfa-required")
