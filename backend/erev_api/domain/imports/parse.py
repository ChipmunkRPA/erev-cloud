"""IPL-02 parsing of import files (05 §5.5 IPL-02, §6.13 UPL-05 to UPL-08; 03 REQ-DAT-007; 04
T-IMP-01 ``sheet_rule``, table 15.4-B ``FORMULA_NO_CACHED_VALUE``; BUILD_SPEC DIN-1).

``read_xlsx`` opens the workbook with openpyxl ``read_only=True, data_only=True, keep_links=False``
(defusedxml installed, UPL-05) and reads the first visible worksheet (``FIRST_SHEET``). Formulas are
never evaluated. openpyxl reads a formula without a cached value as an empty cell that is present in
the sheet, so only when such cells exist is the sheet read a second time without ``data_only`` to
name the cells that hold formulas (UPL-06). Dimensions are reset, so a file whose stored dimension
is wrong is still read to its last row.

``read_csv`` decodes UTF-8 with an optional BOM through the ``csv`` module (RFC 4180; UPL-08).

``cell_text`` renders a cell as text: numbers through their shortest decimal string (REQ-DAT-007),
dates in ISO form, text as given.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Any, BinaryIO, Final

import openpyxl
from openpyxl.cell.read_only import EmptyCell

__all__ = ["CSV_SHEET", "ParseError", "SheetRows", "cell_text", "read_csv", "read_xlsx"]

# UPL-08: the per-field limit of the csv module.
CSV_FIELD_LIMIT: Final = 65_536
# [J] L4-1-Q-29: T-IMP-03 ``sheet_name`` of a CSV file.
CSV_SHEET: Final = "CSV"
VISIBLE: Final = "visible"


class ParseError(Exception):
    """The file cannot be read as a workbook or a CSV file."""


@dataclass(frozen=True, slots=True)
class SheetRows:
    """One sheet as read: the header cells, every later row with its number, and the cells that
    hold a formula without a cached value as ``(row number, column index)``."""

    sheet_name: str
    headers: tuple[Any, ...]
    rows: tuple[tuple[int, tuple[Any, ...]], ...]
    formula_cells: frozenset[tuple[int, int]]


def shortest_decimal(value: float) -> str:
    """The shortest decimal string of a spreadsheet number, without exponent or trailing zeros."""
    try:
        text = format(Decimal(repr(value)), "f")
    except InvalidOperation:
        return repr(value)
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return "0" if text in ("-0", "") else text


def cell_text(value: Any) -> str | None:
    """The text of a cell (T-IMP-03 ``raw``); None for an empty cell."""
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return shortest_decimal(value)
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


def _first_visible(workbook: Any) -> Any:
    for sheet in workbook.worksheets:
        if getattr(sheet, "sheet_state", VISIBLE) == VISIBLE:
            return sheet
    raise ParseError("the workbook has no visible worksheet")


def _cells(workbook: Any) -> tuple[str, Iterator[tuple[int, Any]]]:
    sheet = _first_visible(workbook)
    sheet.reset_dimensions()
    return str(sheet.title), enumerate(sheet.iter_rows(), start=1)


def read_xlsx(stream: BinaryIO) -> SheetRows:
    """The first visible worksheet of an XLSX file (IPL-02; UPL-06)."""
    headers: tuple[Any, ...] = ()
    rows: list[tuple[int, tuple[Any, ...]]] = []
    candidates: set[tuple[int, int]] = set()
    try:
        workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True, keep_links=False)
    except Exception as error:  # a corrupt workbook is a finding, not a crash
        raise ParseError("the workbook cannot be read") from error
    try:
        sheet_name, numbered = _cells(workbook)
        for number, cells in numbered:
            values: list[Any] = []
            for index, cell in enumerate(cells):
                empty = isinstance(cell, EmptyCell)
                value = None if empty else cell.value
                if value is None and not empty:
                    candidates.add((number, index))
                values.append(value)
            if number == 1:
                headers = tuple(values)
            else:
                rows.append((number, tuple(values)))
    finally:
        workbook.close()
    formulas: set[tuple[int, int]] = set()
    if candidates:
        stream.seek(0)
        formula_book = openpyxl.load_workbook(
            stream, read_only=True, data_only=False, keep_links=False
        )
        try:
            _, numbered = _cells(formula_book)
            for number, cells in numbered:
                for index, cell in enumerate(cells):
                    if (number, index) in candidates and getattr(cell, "data_type", None) == "f":
                        formulas.add((number, index))
        finally:
            formula_book.close()
    return SheetRows(
        sheet_name=sheet_name,
        headers=headers,
        rows=tuple(rows),
        formula_cells=frozenset(formulas),
    )


def read_csv(stream: BinaryIO) -> SheetRows:
    """A CSV file: UTF-8 with an optional BOM, RFC 4180 records; empty fields are empty cells."""
    csv.field_size_limit(CSV_FIELD_LIMIT)
    headers: tuple[Any, ...] = ()
    rows: list[tuple[int, tuple[Any, ...]]] = []
    text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    try:
        for number, record in enumerate(csv.reader(text), start=1):
            values = tuple(None if value == "" else value for value in record)
            if number == 1:
                headers = values
            else:
                rows.append((number, values))
    except (UnicodeDecodeError, csv.Error) as error:
        raise ParseError("the CSV file cannot be read") from error
    finally:
        text.detach()
    return SheetRows(
        sheet_name=CSV_SHEET, headers=headers, rows=tuple(rows), formula_cells=frozenset()
    )
