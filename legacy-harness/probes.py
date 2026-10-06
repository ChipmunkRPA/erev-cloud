#!/usr/bin/env python3
"""Targeted probes of legacy behaviours that the UAT scenario does not exercise.

Runs the ORIGINAL eRev.py with the same stubs as replay.py (imported, not re-implemented).
Each probe starts from the shipped DB, clicks the real Reset button, then replays a short
sequence. Generated inputs are derived from the shipped UAT workbooks with openpyxl and are
saved under <out-dir>/probes/inputs/ for inspection.

Outputs: <golden-dir>/probes/<probe-id>/{contract_live.csv, popups.csv, events.csv,
         je_*.csv, probe.json}
"""
from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

import replay as H

UAT = H.UAT
RESET = ("button8", "!Reset Database Completely!", None, None)
SSP = ("button1", "Load SSPs", UAT / "SSP Upload UAT/SKU SSP Template.xlsx", None)
CS1 = ("button2", "Load Contracts", UAT / "Contract Setup UAT/Contract Setup Template 1.1.2023.xlsx", None)
CS2 = ("button2", "Load Contracts", UAT / "Contract Setup UAT/Contract Setup Template 2.1.2023.xlsx", None)
DB131_FILE = UAT / "Delivery and Billing UAT/Contract Progress Tracking Template 1.31.2023.xlsx"
DB131 = ("button3", "Load Delivery and Billing", DB131_FILE, "2023-01-31")
DB228 = ("button3", "Load Delivery and Billing",
         UAT / "Delivery and Billing UAT/Contract Progress Tracking Template 2.28.2023.xlsx", "2023-02-28")
MOD0615_FILE = UAT / "Contract Mod UAT/Contract Modification Template 06.15.2023 - prospective reduction.xlsx"


def _edit_workbook(src: Path, dst: Path, edit):
    from openpyxl import load_workbook
    wb = load_workbook(src)
    ws = wb.active
    header = [c.value for c in ws[1]]
    edit(ws, {name: i + 1 for i, name in enumerate(header)})
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return dst


def gen_blank_memo(inputs: Path):
    def edit(ws, col):
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, col["Contract Unique Name"]).value == "Contract 1" and \
                    ws.cell(r, col["POB Unique ID"]).value in ("POB #2", "POB #3"):
                ws.cell(r, col["Memo 3"]).value = None
    return _edit_workbook(H.COPY_DIR / DB131_FILE,
                          inputs / "P1 progress 1.31.2023 - Memo 3 blank on Contract 1 POB #2 and POB #3.xlsx", edit)


def gen_noop_mod(inputs: Path):
    def edit(ws, col):
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col["Mod Billing"]).value = 0
            ws.cell(r, col["Mod Qty"]).value = 0
            for m in ("Memo 1", "Memo 2", "Memo 3"):
                ws.cell(r, col[m]).value = f"Probe no-op mod 02.15.23 ({m})"
    return _edit_workbook(H.COPY_DIR / MOD0615_FILE,
                          inputs / "P2 no-op prospective mod - Contract 1 POB #1 qty 0 billing 0.xlsx", edit)


def gen_over_delivery(inputs: Path):
    def edit(ws, col):
        ws.cell(2, col["Current Delivery"]).value = 6  # Contract 1 POB #1: 6 + duplicate row 1 = 7 vs 5 remaining
    return _edit_workbook(H.COPY_DIR / DB131_FILE,
                          inputs / "P3 progress 1.31.2023 - Contract 1 POB #1 over-delivered.xlsx", edit)


PROBES = [
    dict(id="P1-blank-memo-drops-progress-rows",
         question="Does a blank memo cell silently drop a delivery/billing row?",
         code_ref="eRev.py l.909-913: df_progress.groupby([... 'Memo 1', 'Memo 2', 'Memo 3']).sum() uses pandas dropna=True",
         sequence=[RESET, SSP, CS1, ("button3", "Load Delivery and Billing", gen_blank_memo, "2023-01-31")],
         reports=[("je_gross", "CustomDialog", "Gross Revenue JEs", ["2023-01-01", "2023-01-31"])],
         watch=[("Contract 1", "POB #1"), ("Contract 1", "POB #2"), ("Contract 1", "POB #3")],
         baseline="UAT replay step 04: Contract 1 POB #2 delivered 1.0 / billed 100 / revenue 118.5332; POB #3 delivered 0.5 / billed 100 / revenue 48.3152"),
    dict(id="P2-mod-reposts-pre-asc606-in-delta-je",
         question="Does a modification version re-post the previous Current Pre-ASC606 Revenue in the Revenue Adjustment JE?",
         code_ref="Mods never reset 'Current Pre-ASC606 Revenue (Net Design Only)' (l.1469-1473, 1965-1968, 2383-2386); "
                  "journal_entries_delta posts that field for every version in the range (l.2966-2979)",
         sequence=[RESET, SSP, CS1, DB131, ("button4", "Prospective Contract Mod", gen_noop_mod, "2023-02-15")],
         reports=[("je_gross_feb", "CustomDialog", "Gross Revenue JEs", ["2023-02-01", "2023-02-28"]),
                  ("je_delta_feb", "CustomDialog", "Revenue Adjustment JEs", ["2023-02-01", "2023-02-28"]),
                  ("je_delta_jan", "CustomDialog", "Revenue Adjustment JEs", ["2023-01-01", "2023-01-31"])],
         watch=[("Contract 1", "POB #2"), ("Contract 1", "POB #3")],
         baseline="No Pre-ASC606 revenue was loaded in February; the correct February delta JE is the gross JE (no Pre-ASC606 lines)"),
    dict(id="P3-over-delivery-validation",
         question="What does the loader do when a delivery exceeds the remaining quantity?",
         code_ref="eRev.py l.959-974 (Flag_Qty / Flag_Billing) and the outer except at l.1163-1168",
         sequence=[RESET, SSP, CS1, ("button3", "Load Delivery and Billing", gen_over_delivery, "2023-01-31")],
         reports=[], watch=[("Contract 1", "POB #1")],
         baseline="Contract 1 POB #1 remaining qty 5; upload asks for 7"),
    dict(id="P4-duplicate-upload-double-counts",
         question="Is re-uploading the same delivery/billing file idempotent?",
         code_ref="eRev.py l.1112-1154: every upload appends a new version; no file, event or period de-duplication",
         sequence=[RESET, SSP, CS1, CS2, DB131, DB228, DB228],
         reports=[("je_gross_feb", "CustomDialog", "Gross Revenue JEs", ["2023-02-01", "2023-02-28"])],
         watch=[("Contract 1", "POB #1"), ("Contract 3", "POB #3"), ("Contract 4", "POB #1")],
         baseline="UAT replay (single 2.28 upload): February gross JE totals Dr/Cr 224.03; Contract 1 POB #1 delivered cum 3.0"),
]


def latest_rows(pd, live):
    if live is None or live.empty:
        return live
    idx = live.groupby("Record Unique ID without time")["Processing Time Log"].transform("max") == live["Processing Time Log"]
    return live[idx]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--golden-dir", default=str(H.DEFAULT_GOLDEN))
    ap.add_argument("--out-dir", default=str(H.DEFAULT_OUT))
    ap.add_argument("--allow-modified-copy", action="store_true")
    args = ap.parse_args()
    golden = Path(args.golden_dir).expanduser().resolve() / "probes"
    out = Path(args.out_dir).expanduser().resolve() / "probes"
    H.preflight(args.allow_modified_copy)
    if golden.exists():
        shutil.rmtree(golden)
    if out.exists():
        shutil.rmtree(out)
    golden.mkdir(parents=True)
    out.mkdir(parents=True)
    inputs = out / "inputs"

    import pandas as pd
    shutil.copyfile(H.FIXTURE_DB, H.COPY_DIR / H.DB_REL)
    H.collect_app_exports(out / "_stray_exports_before_run")
    L = H.Legacy()
    live_cols = [c for c, _ in H.table_schema(H.FIXTURE_DB, "Contract_Live")]
    watch_cols = ["Current Delivery - Cumulative", "Current Billing - Cumulative", "Current Rev Rec - Cumulative",
                  "Current Pre-ASC606 Revenue (Net Design Only)", "Current Pre-ASC606 Revenue (Net Design Only) - Cumulative",
                  "Current Remaining Qty", "Current Remaining Allocation", "Current Period"]

    for p in PROBES:
        pdir = golden / p["id"]
        pdir.mkdir(parents=True)
        H.EVENTS.step = p["id"]
        tokens = H.TimeTokens()
        seq_results = []
        for i, (button, label, file, date_input) in enumerate(p["sequence"]):
            path = file(inputs) if callable(file) else (H.COPY_DIR / file if file else None)
            before = H.read_table(pd, L.db_path, "Contract_Live")
            L.click(button, label, path, date_input, f"{p['id']}:{i:02d}:{label}")
            after = H.read_table(pd, L.db_path, "Contract_Live")
            if after is not None:
                tokens.register(after["Processing Time Log"].unique(), f"{p['id'].split('-')[0]}-{i:02d}")
            handler_popups = [dict(kind=e["kind"], title=e["title"], text=e["text"])
                              for e in H.EVENTS.records
                              if e["step"] == p["id"] and e["phase"].startswith(f"{p['id']}:{i:02d}:") and e["category"] == "popup"]
            seq_results.append(dict(i=i, button_label=label, file=Path(path).name if path else None, date_input=date_input,
                                    rows_before=0 if before is None else len(before),
                                    rows_after=0 if after is None else len(after), popups=handler_popups))
        live = H.read_table(pd, L.db_path, "Contract_Live")
        H.write_csv(tokens.normalize(live), pdir / "contract_live.csv", live_cols)
        reports = {}
        for name, dialog, button_text, rng in p["reports"]:
            frame, _ = L.report(dialog, button_text, rng, f"{p['id']}:report:{name}")
            H.write_csv(frame, pdir / f"{name}.csv", ["Record Unique ID without time", "Account", "Amount"])
            reports[name] = dict(window=rng, totals=H.je_totals(frame))
        H.collect_app_exports(out / p["id"] / "app_exports")
        lt = latest_rows(pd, live)
        watched = []
        for k, pob in p["watch"]:
            r = lt[(lt["Contract Unique Name"] == k) & (lt["POB Unique ID"] == pob)]
            if len(r):
                row = r.iloc[0]
                watched.append(dict(contract=k, pob=pob, **{c: (None if pd.isna(row[c]) else
                                                                (row[c] if isinstance(row[c], str) else round(float(row[c]), 4)))
                                                            for c in watch_cols}))
        events = H.EVENTS.for_step(p["id"])
        H.write_csv(pd.DataFrame([e for e in events if e["category"] == "popup"], columns=H.EVENT_COLUMNS), pdir / "popups.csv")
        H.write_csv(pd.DataFrame(events, columns=H.EVENT_COLUMNS), pdir / "events.csv")
        rec = dict(id=p["id"], question=p["question"], code_ref=p["code_ref"], baseline=p["baseline"],
                   sequence=seq_results, reports=reports, watched_latest=watched)
        (pdir / "probe.json").write_text(json.dumps(rec, indent=2, default=str) + "\n")
        last = seq_results[-1]
        print(f"[{p['id']}] last step rows {last['rows_before']} -> {last['rows_after']} popups={[x['title'] for x in last['popups']]}")
    L.close()
    print(f"probes -> {golden}")


if __name__ == "__main__":
    main()
